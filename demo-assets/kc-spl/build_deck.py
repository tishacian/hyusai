#!/usr/bin/env python3
"""Build the Knowledge Capture demo deck (Andritz / BBA120 spunlace)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))

# Palette
NAVY = RGBColor(0x0F, 0x2A, 0x43)
TEAL = RGBColor(0x12, 0x8C, 0x7D)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
LIGHT = RGBColor(0xF2, 0xF5, 0xF7)
GREY = RGBColor(0x55, 0x5B, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1B, 0x1F, 0x24)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def rect(slide, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = "Calibri"
    return tb


def header(slide, kicker, title):
    rect(slide, 0, 0, SW, Inches(1.15), NAVY)
    rect(slide, 0, Inches(1.15), SW, Pt(4), TEAL)
    textbox(slide, Inches(0.55), Inches(0.12), Inches(12), Inches(0.4),
            [(kicker, 13, TEAL, True)])
    textbox(slide, Inches(0.55), Inches(0.42), Inches(12.3), Inches(0.7),
            [(title, 26, WHITE, True)])


def add_image_fit(slide, path, x, y, max_w, max_h, border=True):
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    box_ar = max_w / max_h
    if ar > box_ar:
        w = max_w
        h = Emu(int(max_w / ar))
    else:
        h = max_h
        w = Emu(int(max_h * ar))
    px = x + Emu(int((max_w - w) / 2))
    py = y + Emu(int((max_h - h) / 2))
    if border:
        rect(slide, px - Pt(2), py - Pt(2), w + Pt(4), h + Pt(4), RGBColor(0xD0, 0xD7, 0xDD))
    slide.shapes.add_picture(path, px, py, width=w, height=h)


def bullets(slide, x, y, w, h, items, size=15, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            txt, lvl, bold, color = it
        else:
            txt, lvl, bold, color = it, 0, False, DARK
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        p.space_after = Pt(gap)
        r = p.add_run()
        prefix = "" if lvl == 0 else ("•  " if lvl == 1 else "–  ")
        r.text = prefix + txt
        r.font.size = Pt(size - lvl)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb


# ----------------------------------------------------------------- Slide 1 : Title
s = add_slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(4.0), SW, Pt(5), TEAL)
textbox(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.0),
        [("KNOWLEDGE CAPTURE — DÉMO", 18, TEAL, True)])
textbox(s, Inches(0.8), Inches(2.55), Inches(11.7), Inches(1.4),
        [("Capture de savoir expert — Ligne spunlace BBA120 (JETLACE)", 34, WHITE, True)])
textbox(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(1.6),
        [("Co-construction de plan assistée par l'IA, oracle de relance et", 17, RGBColor(0xC9, 0xD6, 0xDF), False),
         ("remontée de contexte en direct, ancrée sur les notices techniques SPL.", 17, RGBColor(0xC9, 0xD6, 0xDF), False),
         ("", 8, WHITE, False),
         ("Workspace Andritz  ·  Collection : andritz-notices-techniques-spl-pilot", 14, TEAL, True)])

# ----------------------------------------------------------------- Slide 2 : Parcours
s = add_slide()
header(s, "VUE D'ENSEMBLE", "Le parcours de démo en 5 étapes")
steps = [
    ("1", "Préparation", "Titre, sujet, interlocuteur, durée estimée."),
    ("2", "Mode de capture", "« Session avec plan » — l'IA co-construit le plan."),
    ("3", "Co-construction", "L'expert décrit ; l'IA produit un plan à 3 niveaux."),
    ("4", "Lancement", "Plan validé · banque de questions en préparation."),
    ("5", "Échange live", "Voix + remontée de contexte + questions de l'oracle."),
]
y = Inches(1.6)
for num, t, d in steps:
    rect(s, Inches(0.55), y, Inches(0.7), Inches(0.7), TEAL)
    textbox(s, Inches(0.55), y + Inches(0.07), Inches(0.7), Inches(0.6),
            [(num, 24, WHITE, True)], align=PP_ALIGN.CENTER)
    textbox(s, Inches(1.45), y - Inches(0.02), Inches(11), Inches(0.45),
            [(t, 18, NAVY, True)])
    textbox(s, Inches(1.45), y + Inches(0.38), Inches(11), Inches(0.4),
            [(d, 14, GREY, False)])
    y += Inches(1.0)
rect(s, Inches(0.55), Inches(6.75), Inches(12.2), Inches(0.5), LIGHT)
textbox(s, Inches(0.7), Inches(6.8), Inches(12), Inches(0.4),
        [("Pré-requis démo : autoriser le micro du navigateur — les panneaux live ne se peuplent qu'avec la boucle voix.", 12.5, GREY, True)])

# ----------------------------------------------------------------- Slide 3 : Étape 1 Préparation
s = add_slide()
header(s, "ÉTAPE 1 · PRÉPARATION", "« Nouvelle capture » — valeurs à saisir")
fields = [
    ("Titre de session *", "Exploitation et optimisation de la ligne spunlace BBA120 (JETLACE)"),
    ("Sujet (optionnel)", "Hydroliage HP, maintenance injecteurs/pompes, réglage capteurs, optimisation énergie/cadence/homogénéité"),
    ("Personne interrogée", "Jean Mercier — Expert procédé non-tissés"),
    ("Durée estimée (min)", "25"),
]
y = Inches(1.5)
for k, v in fields:
    rect(s, Inches(0.55), y, Inches(2.7), Inches(0.85), NAVY)
    textbox(s, Inches(0.62), y + Inches(0.08), Inches(2.55), Inches(0.7),
            [(k, 12.5, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(3.25), y, Inches(3.4), Inches(0.85), LIGHT)
    textbox(s, Inches(3.38), y + Inches(0.05), Inches(3.15), Inches(0.78),
            [(v, 11.5, DARK, False)], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.95)
textbox(s, Inches(0.55), y + Inches(0.05), Inches(6), Inches(0.4),
        [("→ Cliquer « Choisir le mode de capture ».", 13, TEAL, True)])
add_image_fit(s, os.path.join(HERE, "01_preparation.png"), Inches(6.95), Inches(1.45), Inches(6.0), Inches(5.6))

# ----------------------------------------------------------------- Slide 4 : Étape 2 Mode
s = add_slide()
header(s, "ÉTAPE 2 · MODE DE CAPTURE", "« Session avec plan »")
bullets(s, Inches(0.55), Inches(1.55), Inches(5.9), Inches(5.0), [
    ("Choisir la carte « Session avec plan ».", 0, True, NAVY),
    ("L'IA co-construit le plan à partir de vos propos.", 1, False, DARK),
    ("Laisser le champ « Plan document » vide", 0, True, NAVY),
    ("On co-construit le plan en direct à l'étape suivante.", 1, False, DARK),
    ("Cliquer « Co-construire le plan ».", 0, True, NAVY),
], size=16, gap=12)
add_image_fit(s, os.path.join(HERE, "02_mode.png"), Inches(6.7), Inches(1.45), Inches(6.25), Inches(5.6))

# ----------------------------------------------------------------- Slide 5 : Étape 3 Co-construction (texte)
s = add_slide()
header(s, "ÉTAPE 3 · CO-CONSTRUCTION", "Texte expert à coller, puis « Envoyer »")
rect(s, Inches(0.55), Inches(1.5), Inches(12.2), Inches(2.55), LIGHT)
expert = ("Je suis expert procédé sur la ligne BBA120. Je veux transmettre mon savoir sur l'exploitation et "
          "l'optimisation de la ligne spunlace JETLACE : carde → nappe → hydroliage par injecteurs haute "
          "pression → séchage → enroulement. Points clés : gestion des pompes haute pression URACA/KD724, "
          "réglage des injecteurs (prewetting et autoclamped), entretien des garnitures de carde, qualité de "
          "l'eau du circuit d'hydroliage, réglage des capteurs XS1/XS2/ZCT. Objectifs : optimisation "
          "énergétique, augmentation de cadence, et homogénéité du non-tissé.")
textbox(s, Inches(0.75), Inches(1.65), Inches(11.8), Inches(2.3),
        [(expert, 14.5, DARK, False)])
textbox(s, Inches(0.55), Inches(4.2), Inches(12), Inches(0.4),
        [("→ L'IA produit un plan à 3 niveaux. Quand il convient : « Valider le plan » (la banque de questions se prépare en arrière-plan).", 13, TEAL, True)])
# plan généré
rect(s, Inches(0.55), Inches(4.75), Inches(12.2), Inches(0.42), NAVY)
textbox(s, Inches(0.7), Inches(4.79), Inches(12), Inches(0.4),
        [("Plan généré (3 niveaux, 10 sujets) — hiérarchie préservée :", 13, WHITE, True)])
plan_items = [
    ("1. Introduction à la ligne spunlace BBA120 (JETLACE)", 0, True, NAVY),
    ("Présentation générale · Importance dans l'industrie", 2, False, GREY),
    ("2. Composants et équipements clés", 0, True, NAVY),
    ("Pompes HP URACA/KD724 · Injecteurs HP · Capteurs XS1/XS2/ZCT", 2, False, GREY),
    ("3. Optimisation de la ligne", 0, True, NAVY),
    ("Énergie · Cadence · Homogénéité du non-tissé", 2, False, GREY),
    ("4. Maintenance de la ligne", 0, True, NAVY),
    ("Garnitures de carde · Qualité de l'eau du circuit d'hydroliage", 2, False, GREY),
]
bullets(s, Inches(0.6), Inches(5.25), Inches(12), Inches(2.1), plan_items, size=13.5, gap=2)

# ----------------------------------------------------------------- Slide 6 : screenshot plan
s = add_slide()
header(s, "ÉTAPE 3 · CO-CONSTRUCTION", "Plan hiérarchique généré par l'IA")
add_image_fit(s, os.path.join(HERE, "03_plan.png"), Inches(0.55), Inches(1.45), Inches(12.2), Inches(5.7))

# ----------------------------------------------------------------- Slide 7 : Étape 4 Lancement
s = add_slide()
header(s, "ÉTAPE 4 · LANCEMENT", "Plan validé — prêt pour l'échange")
bullets(s, Inches(0.55), Inches(1.55), Inches(5.9), Inches(5.0), [
    ("« 10 sujets / 25 min » · « Plan validé ».", 0, True, NAVY),
    ("La banque de questions se prépare en arrière-plan.", 1, False, DARK),
    ("Un seul appel à l'action : « Démarrer la conversation ».", 0, True, NAVY),
    ("Le plan devient un rappel non-bloquant — il n'impose aucun format.", 1, False, DARK),
    ("Cliquer « Démarrer la conversation » pour ouvrir la session live.", 0, True, TEAL),
], size=16, gap=12)
add_image_fit(s, os.path.join(HERE, "04_launch.png"), Inches(6.7), Inches(1.45), Inches(6.25), Inches(5.6))

# ----------------------------------------------------------------- Slide 8 : Étape 5 Session live
s = add_slide()
header(s, "ÉTAPE 5 · ÉCHANGE LIVE", "Session live — contexte & oracle en direct")
bullets(s, Inches(0.55), Inches(1.5), Inches(5.0), Inches(5.4), [
    ("Parlez librement (micro).", 0, True, NAVY),
    ("Transcription live par chunk pendant que vous parlez.", 1, False, DARK),
    ("« Contexte retrouvé »", 0, True, TEAL),
    ("Passages des notices SPL remontés en ~0,3–1,2 s.", 1, False, DARK),
    ("« Questions de l'oracle »", 0, True, TEAL),
    ("Angles à clarifier : rationale, signaux terrain, provenance.", 1, False, DARK),
    ("Relance non-bloquante : « Avez-vous terminé sur ce point ? »", 0, True, NAVY),
    ("Terminer par « Créer la proposition » (relecture HITL).", 0, True, NAVY),
], size=14.5, gap=8)
add_image_fit(s, os.path.join(HERE, "05_session.png"), Inches(5.75), Inches(1.45), Inches(7.15), Inches(5.6))

# ----------------------------------------------------------------- Slide 9 : Script parlé
s = add_slide()
header(s, "SCRIPT PARLÉ", "4 tours d'expert prêts à dire")
turns = [
    ("Tour 1 — Architecture & hydroliage",
     "« Sur la BBA120, le cœur du procédé c'est l'hydroliage : la nappe issue des cardes passe sous une série "
     "d'injecteurs haute pression alimentés par les pompes URACA KD724, autour de 250 bars pour un grammage de 50 g/m². »"),
    ("Tour 2 — Maintenance injecteurs",
     "« Pour la maintenance des injecteurs on surveille l'usure des buses et l'état des strip-carriers ; un joint "
     "O-ring défaillant provoque des fuites et une chute de pression. Un prewetting bien réglé évite les défauts d'homogénéité. »"),
    ("Tour 3 — Capteurs",
     "« Côté instrumentation, les XS1 et XS2 sont des détecteurs de proximité ; un ZCT mal réglé déclenche des arrêts intempestifs en production. »"),
    ("Tour 4 — Optimisation cadence / énergie",
     "« Pour augmenter la cadence, le goulot se situe au séchage et sur la pression d'hydroliage ; si on pousse trop "
     "la vitesse sans ajuster, l'homogénéité chute. Côté énergie, les pompes HP sont le premier poste de consommation. »"),
]
y = Inches(1.45)
for t, q in turns:
    rect(s, Inches(0.55), y, Pt(5), Inches(1.25), TEAL)
    textbox(s, Inches(0.8), y, Inches(12), Inches(0.4),
            [(t, 14.5, NAVY, True)])
    textbox(s, Inches(0.8), y + Inches(0.36), Inches(12.0), Inches(0.95),
            [(q, 12.5, DARK, False)])
    y += Inches(1.42)

# ----------------------------------------------------------------- Slide 10 : Preuve d'ancrage
s = add_slide()
header(s, "PREUVE D'ANCRAGE", "Le contexte remonte bien des notices techniques SPL")
textbox(s, Inches(0.55), Inches(1.35), Inches(12), Inches(0.5),
        [("Requête de test (hydroliage / pompes URACA) sur la collection andritz-notices-techniques-spl-pilot — extraits réels remontés :", 13.5, GREY, True)])
proofs = [
    ("DESCRIPTION OF THE INJECTOR", "« …the strip is auto-clamped by the water pressure on the strip-carrier insuring the watertightness and the maintain in position. »"),
    ("URACA GmbH & Co. KG — pompes HP", "« URACA GmbH & Co. KG, Bad Urach, Germany… » — documentation des pompes haute pression."),
    ("Réglage de pression", "« …system pressure adjusting at the minimum number of revolutions depends on the nozzle bar characteristic. »"),
    ("Perfojet — pièces de rechange", "« Spare parts recommended for preventive maintenance… » (injecteurs / strip-carriers)."),
]
y = Inches(1.95)
for t, q in proofs:
    rect(s, Inches(0.55), y, Inches(3.6), Inches(1.0), LIGHT)
    textbox(s, Inches(0.68), y + Inches(0.08), Inches(3.4), Inches(0.85),
            [(t, 12.5, NAVY, True)], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, Inches(4.35), y + Inches(0.02), Inches(8.4), Inches(0.98),
            [(q, 12, DARK, False)], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.1)
rect(s, Inches(0.55), Inches(6.5), Inches(12.2), Inches(0.55), TEAL)
textbox(s, Inches(0.7), Inches(6.56), Inches(12), Inches(0.45),
        [("Le retrieval de capture est désormais branché sur les notices techniques SPL transverses au workspace — le contexte colle au discours BBA120.", 12.5, WHITE, True)])

# ----------------------------------------------------------------- Slide 11 : Conseils présentateur
s = add_slide()
header(s, "CONSEILS PRÉSENTATEUR", "Pour une démo nickel")
bullets(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(5.4), [
    ("Autoriser le micro du navigateur AVANT de démarrer la session live.", 0, True, NAVY),
    ("C'est le seul moyen de peupler « Contexte retrouvé » et « Questions de l'oracle » (boucle voix).", 1, False, DARK),
    ("Coller le texte expert d'un coup à l'étape 3, puis « Envoyer ».", 0, True, NAVY),
    ("Laisser l'IA générer le plan à 3 niveaux ; insister sur la hiérarchie préservée.", 1, False, DARK),
    ("Pendant l'échange, pointer du doigt les deux panneaux latéraux à chaque tour.", 0, True, NAVY),
    ("Montrer la remontée de contexte (~0,3–1,2 s) et les questions anticipées de l'oracle.", 1, False, DARK),
    ("Insister : le plan est un rappel non-bloquant — l'expert n'est jamais enfermé dans un format.", 0, True, NAVY),
    ("Conclure par « Créer la proposition » → connaissance structurée à relire (HITL).", 0, True, TEAL),
], size=15, gap=10)

out = os.path.join(HERE, "Knowledge_Capture_Demo_Andritz_BBA120.pptx")
prs.save(out)
print("Saved:", out)
print("Slides:", len(prs.slides._sldIdLst))
