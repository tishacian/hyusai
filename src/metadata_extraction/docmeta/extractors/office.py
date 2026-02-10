"""
Microsoft Office Metadata Extractor

This module provides functionality to extract metadata from Microsoft Office files.
Supports DOCX, XLSX, and PPTX formats with comprehensive metadata extraction.
"""

import logging
from typing import Any

# Import required libraries for Office document processing
import docx
from openpyxl import load_workbook
from pptx import Presentation

from docmeta.core.defaults import (
    create_office_metadata,  # create_office_metadata returns dict[str, Any]
)
from docmeta.utils.text_processing import (
    count_text_statistics,
    extract_common_metadata_fields,
    extract_keywords,
    get_file_extension,
)

# Configure logger
logger = logging.getLogger(__name__)


def _extract_with_docx(path: str) -> tuple[dict[str, Any], str | None]:
    """
    Extract metadata and text from a DOCX file.

    Args:
        path: Path to the DOCX file

    Returns:
        tuple[dict[str, Any], str | None]: Specific DOCX metadata and text content.
    """
    metadata = create_office_metadata()

    doc = docx.Document(path)
    core_properties = doc.core_properties

    field_mapping = {
        "author": "author",
        "title": "title",
        "subject": "subject",
        "last_modified_by": "last_modified_by",
    }
    extract_common_metadata_fields(core_properties, metadata, field_mapping)
    extract_keywords(core_properties, metadata, "keywords")

    metadata["application"] = "Microsoft Word"
    metadata["paragraph_count"] = len(doc.paragraphs)

    text_elements = [para.text for para in doc.paragraphs]
    stats = count_text_statistics(text_elements)
    metadata["word_count"] = stats["word_count"]
    metadata["character_count"] = stats["char_count"]

    if stats["word_count"] > 0:
        metadata["num_pages"] = max(1, stats["word_count"] // 500)

    text_content = "\n".join(stats["all_text"]) if stats["all_text"] else None
    return metadata, text_content


def _extract_with_xlsx(path: str) -> tuple[dict[str, Any], str | None]:
    """
    Extract metadata and text from an XLSX file.

    Args:
        path: Path to the XLSX file

    Returns:
        tuple[dict[str, Any], str | None]: Specific XLSX metadata and text content.
    """
    metadata = create_office_metadata()

    workbook = load_workbook(path, read_only=True)
    properties = workbook.properties

    field_mapping = {
        "author": "creator",
        "title": "title",
        "subject": "subject",
        "last_modified_by": "lastModifiedBy",
    }
    extract_common_metadata_fields(properties, metadata, field_mapping)
    extract_keywords(properties, metadata, "keywords")

    metadata["application"] = "Microsoft Excel"
    metadata["num_sheets"] = len(workbook.worksheets)

    total_cells = 0
    text_elements = []
    char_count = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    total_cells += 1
                    cell_text = str(cell.value)
                    if isinstance(cell.value, str):
                        char_count += len(cell_text)
                        if cell_text.strip():
                            text_elements.append(cell_text)

    metadata["character_count"] = char_count
    metadata["word_count"] = total_cells  # Using cell count as word equivalent

    text_content = "\n".join(text_elements) if text_elements else None
    workbook.close()
    return metadata, text_content


def _extract_with_pptx(path: str) -> tuple[dict[str, Any], str | None]:
    """
    Extract metadata and text from a PPTX file.

    Args:
        path: Path to the PPTX file

    Returns:
        tuple[dict[str, Any], str | None]: Specific PPTX metadata and text content.
    """
    metadata = create_office_metadata()

    presentation = Presentation(path)
    core_properties = presentation.core_properties

    field_mapping = {
        "author": "author",
        "title": "title",
        "subject": "subject",
        "last_modified_by": "last_modified_by",
    }
    extract_common_metadata_fields(core_properties, metadata, field_mapping)
    extract_keywords(core_properties, metadata, "keywords")

    metadata["application"] = "Microsoft PowerPoint"
    metadata["num_slides"] = len(presentation.slides)
    metadata["num_pages"] = len(presentation.slides)

    text_elements = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text_elements.append(shape.text)

    stats = count_text_statistics(text_elements)
    metadata["word_count"] = stats["word_count"]
    metadata["character_count"] = stats["char_count"]

    text_content = "\n".join(stats["all_text"]) if stats["all_text"] else None
    return metadata, text_content


def extract_office_metadata(path: str) -> tuple[dict[str, Any], str | None]:
    """
    Extract metadata and text from a Microsoft Office file.

    Args:
        path: Path to the Office file

    Returns:
        tuple[dict[str, Any], str | None]: Specific Office metadata and text content.
    """
    # File existence validation is handled by the factory
    extension = get_file_extension(path)

    if extension == ".docx":
        return _extract_with_docx(path)
    elif extension == ".xlsx":
        return _extract_with_xlsx(path)
    elif extension == ".pptx":
        return _extract_with_pptx(path)
    else:
        # Should not happen if called via registered extensions
        logger.warning(f"Unsupported Office file type for explicit extraction: {path}")
        return create_office_metadata(), None
