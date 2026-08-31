#!/usr/bin/env python3
"""Build the illustrated Fayçal PR→PO playbook Word from in-repo figures."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "nawa-pr-to-po-live-demo-playbook.docx"
ART = Path("/opt/cursor/artifacts/nawa_pr_to_po_live_demo_playbook.docx")
FIGS = HERE / "nawa-pr-to-po-playbook-figures"

SHA = "cd40fe337285058120ec7236d587b1853ca72c42"
URL = "https://agentium.papai.ai/work/pr-to-po?workspace=nawa&lang=en"


def _set_run_font(run, *, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def para(doc, text, *, italic=False, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        _set_run_font(run)
    return p


def table(doc, rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            tbl.rows[i].cells[j].text = ""
            p = tbl.rows[i].cells[j].paragraphs[0]
            run = p.add_run(cell)
            _set_run_font(run, size=10, bold=(i == 0))
    doc.add_paragraph()
    return tbl


def shot(doc, name, caption):
    path = FIGS / name
    if not path.exists():
        para(doc, f"[Screenshot missing: {name}]", italic=True)
        return
    pic = doc.add_picture(str(path), width=Inches(6.4))
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    run = cap.add_run(caption)
    _set_run_font(run, size=9, italic=True, color=RGBColor(0x55, 0x55, 0x55))
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    heading(doc, "NAWA PR → PO — live desk playbook")
    para(doc, "For Fayçal and the presenter · Desk language: English")
    para(doc, f"URL: {URL}")
    para(doc, f"Live revision: {SHA} (hard-reload the tab). Verified 31 August 2026.")
    para(
        doc,
        "This is an operator desk, not a numbered demo script. "
        "The four proofs live in this document. On screen you only see work.",
    )

    heading(doc, "Who gets what", 2)
    table(
        doc,
        [
            ["Give this", "To whom"],
            ["This Word (or the markdown sibling)", "Fayçal — proofs and walkthrough"],
            [
                "nawa-pr-to-po-live-demo-presenter-fr.md",
                "Presenter — one French page to hold during the talk",
            ],
        ],
    )

    heading(doc, "Open this URL", 2)
    para(doc, URL, bold=True)
    bullet(doc, "Hard-reload once (Ctrl+Shift+R / Cmd+Shift+R).")
    bullet(doc, "Workspace nawa. Title on screen: PR to PO.")
    bullet(doc, "Presenter login: thibaud.ishacian@datategy.net.")
    bullet(doc, "Live Hikma MCP. No HANA. Write stays Sealed.")
    para(
        doc,
        "Subtitle: System → Flow → Run → Decision. The factory reads Hikma, "
        "compiles a dossier, and writes nothing to SAP until you decide.",
        italic=True,
    )

    heading(doc, "What Fayçal must see (four proofs)", 2)
    para(
        doc,
        "The screen does not number these. You point at the matching card.",
    )
    table(
        doc,
        [
            ["#", "Proof", "Where you point", "What must be true"],
            [
                "1",
                "Live SAP MCP",
                "KPI row, stations, Read selected requisition, dossier facts",
                "get_A_PurchaseRequisitionItem and _by_key return a real PR. "
                "listTaskCollection feeds the approval. fi_Validate / "
                "get_A_PurReqnAcctAssgmt feed budget. get_A_PurchaseOrder* "
                "feed orders, supplier, type. Compile shows python_recipe_v1.",
            ],
            [
                "2",
                "Small agentic task",
                "Summarise in the compiled dossier",
                "azure_llm_v1 POST /chat/completion prompt_type: factual. "
                "Two clear sentences. Invent nothing.",
            ],
            [
                "3",
                "PO POST package",
                "Draft / Purchase order composed for SAP",
                "hikma post_A_PurchaseOrder, type NB, Send Sealed. "
                "JSON sealed: true, called: false. Client 300 has no NB range; "
                "ZAPO is refused. Success = composed POST, not a live create.",
            ],
            [
                "4",
                "Nice-to-have",
                "Ask the factory / Open the portal",
                "Read-only. Chips: Majority supplier, Orders read, Next step, "
                "Write status. The write stays sealed.",
            ],
        ],
    )
    para(doc, "If the Write station says Sealed, that is correct.", bold=True)

    heading(doc, "Walkthrough (about three minutes)", 2)
    bullet(doc, "Hard-reload. Header reads PR to PO. Four counts. Five stations. Write = Sealed.")
    bullet(
        doc,
        "Click Run the factory. Wait until Compile is Done and the hero reads "
        "something like “Dossier compiled on STICKER WHITE — majority supplier 1000001737.”",
    )
    bullet(
        doc,
        "Proof 1 — Point at the counts and the tools under the stations. "
        "Click Read selected requisition if justification is empty.",
    )
    bullet(doc, "Proof 2 — Click Summarise. Two sentences appear via azure_llm_v1.")
    bullet(
        doc,
        "Proof 3 — Scroll to Draft / Purchase order composed for SAP. "
        "Row: POST · hikma · post_A_PurchaseOrder · Send Sealed. "
        "Say: this is the POST we would send; it is not sent.",
    )
    bullet(doc, "Proof 4 — Chip Write status. Answer: the write stays sealed.")
    para(doc, "There is no button that posts a purchase order.")
    para(
        doc,
        "Talk track: “This is the PR to PO factory. It reads Hikma through MCP, "
        "compiles a dossier, and writes nothing until a human decides.”",
        italic=True,
    )

    heading(doc, "Screen map", 2)
    heading(doc, "Top of the desk", 3)
    para(
        doc,
        "Lineage, briefing, counts, stations, compiled dossier. "
        "Actions in the dossier header: Read selected requisition, Summarise.",
    )
    shot(doc, "desk-top.png", "Figure 1 — Operator desk after a live read. Write stays Sealed.")

    heading(doc, "Compiled dossier (proofs 1 and 2)", 3)
    para(
        doc,
        "Facts name their tools. Justification via get_A_PurchaseRequisitionItem_by_key. "
        "Summary via azure_llm_v1.",
    )
    shot(doc, "dossier.png", "Figure 2 — Live counts, stations, and the compiled dossier.")

    heading(doc, "Draft purchase order (proof 3)", 3)
    bullet(doc, "Kicker Draft. Title Purchase order composed for SAP.")
    bullet(doc, "Hint: Draft purchase order. It is not sent; the write stays sealed.")
    bullet(doc, "Row: Method POST · Server hikma · Tool post_A_PurchaseOrder · Send Sealed.")
    bullet(doc, "Footnote: Client 300 has no NB number range. The API refuses ZAPO.")
    shot(doc, "post.png", "Figure 3 — Draft hikma POST: type NB, sealed, not sent.")

    heading(doc, "Ask the factory (proof 4)", 3)
    para(doc, "Hint: Questions about the live reads. The write stays sealed.")
    shot(
        doc,
        "ask.png",
        "Figure 4 — Write status answers No; the write station stays sealed.",
    )

    heading(doc, "Portal", 3)
    para(
        doc,
        "Hint: Answers from the requisitions and orders just read. The write stays sealed. "
        "Stay on get-PR / get-PO questions. Older summarise prompts in the portal thread "
        "are session history, not a second write.",
    )
    shot(doc, "portal.png", "Figure 5 — Factory portal on the live reads. Write stays sealed.")

    heading(doc, "Live example (31 August 2026, SHA cd40fe33)", 2)
    table(
        doc,
        [
            ["Field", "Value"],
            [
                "Headline",
                "Dossier compiled on STICKER WHITE — majority supplier 1000001737.",
            ],
            ["Counts", "Requisitions 50 · Tasks 2 · Orders 5 · Receipts 0"],
            [
                "Stations",
                "Connect Done · Read Done (get_A_PurchaseRequisitionItem) · "
                "Compile Done (python_recipe_v1) · Decide Ready · Write Sealed",
            ],
            ["Selected requisition", "STICKER WHITE via get_A_PurchaseRequisitionItem"],
            [
                "Approval task",
                "Release TR transaction 1000008 56B 200 via listTaskCollection",
            ],
            ["PR in the POST", "2000276449 item 10"],
            [
                "Justification",
                "STICKER WHITE via get_A_PurchaseRequisitionItem_by_key",
            ],
            [
                "Summary",
                "The purchase requisition is for white stickers. "
                "No additional justification details are provided in the text. via azure_llm_v1",
            ],
            ["Supplier / type (read)", "1000001737 / ZAPO via get_A_PurchaseOrder"],
            ["Budget", "ok via fi_Validate"],
            [
                "POST",
                "hikma post_A_PurchaseOrder, PurchaseOrderType NB, "
                "PurchasingOrganization CPO, sealed true, called false",
            ],
        ],
    )
    para(
        doc,
        "Ask → Write status: No. The write station stays sealed. "
        "A cycle can compile a dossier; SAP is not written until you decide.",
        italic=True,
    )
    para(doc, "The exact PR can move on a later read. The tools and the sealed POST must not.")

    heading(doc, "Exact labels on the desk", 2)
    table(
        doc,
        [
            ["Place", "English"],
            ["Title", "PR to PO"],
            ["Primary / secondary", "Run the factory · Start a cycle"],
            ["Dossier", "Compiled dossier / Compiled from the live SAP reads."],
            ["Dossier actions", "Read selected requisition · Summarise"],
            ["Draft card", "Draft · Purchase order composed for SAP"],
            ["Ask chips", "Majority supplier · Orders read · Next step · Write status"],
            ["Portal", "Open the portal · Factory portal"],
            ["Write station", "Sealed"],
        ],
    )
    para(
        doc,
        "Banned on this desk: pipeline, workflow, job. "
        "Removed: Success criteria, numbered beats 1–4, Shown / Play / Waiting, "
        "“Compose the POST is the success of this demo”.",
    )

    heading(doc, "Seven-minute script", 2)
    bullet(doc, "0:00 — URL, nawa, English, hard-reload. “Operator desk. Not a slide.”")
    bullet(doc, "0:30 — Counts + stations. “Live SAP, through MCP. No HANA.”")
    bullet(
        doc,
        "1:30 — Dossier + Read selected requisition + Summarise. "
        "“One short model call. Two sentences. Nothing invented.”",
    )
    bullet(
        doc,
        "3:00 — Draft card. “Type NB, sealed, not sent. Client 300 has no NB range.”",
    )
    bullet(doc, "5:00 — Write status, then Open the portal if there is time.")
    bullet(doc, "6:30 — Stop. Write still Sealed. Offer Start a cycle only if they ask.")

    heading(doc, "If something looks wrong", 2)
    bullet(doc, "Old tiles Shown / Play / Success criteria → hard-reload. Removed on cd40fe33.")
    bullet(doc, "Blank desk / login loop → login, then hard-reload this desk URL. Do not open / first.")
    bullet(doc, "token expired or HTTP 407 → restart the Hikma Cloud Foundry MCP apps, not the Agentium password.")
    bullet(doc, "Counts at 0 → Run the factory once and wait.")
    bullet(doc, "Empty justification → Read selected requisition once, then Summarise once.")
    bullet(doc, "Write still Sealed after Run → correct.")
    bullet(doc, "Receipts 0 / Limited read → expected. sap_gr is often 403.")
    bullet(doc, "Someone asks to create the PO live → refuse. The composed POST is the artefact.")

    heading(doc, "Do not", 2)
    bullet(doc, "Do not post a purchase order. No tools/call on post_A_PurchaseOrder.")
    bullet(doc, "Do not call fi_EnableForPurchasing or fi_Discard.")
    bullet(doc, "Do not invent MCP aliases (A_*, YY1_*). Do not turn HANA on.")
    bullet(doc, "Do not seed nawa on this VM. Do not move the demo-agentic tag.")
    bullet(doc, "Do not put the four proofs back on the screen as numbered tiles.")
    bullet(doc, "Do not put a password or an OAuth secret on a slide.")

    heading(doc, "Technical anchors", 2)
    bullet(doc, "Desk /work/pr-to-po · Binding procurement.pr_to_po.run · System 28345b5a-…")
    bullet(doc, "Images agentium-{backend,worker,frontend}:cd40fe337285")
    bullet(doc, "Rollback of this desk slice: AGENTIUM_IMAGE_TAG=3eb72f95f6b6")
    bullet(doc, "Writes via /read → HTTP 400 “not a read”.")
    bullet(doc, "Markdown source: docs/ops/nawa-pr-to-po-live-demo-playbook.md")
    bullet(doc, "Presenter card (FR): docs/ops/nawa-pr-to-po-live-demo-presenter-fr.md")

    end = doc.add_paragraph()
    run = end.add_run("End of playbook. Walk the desk. Leave SAP sealed.")
    _set_run_font(run, bold=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    if ART.parent.is_dir():
        doc.save(ART)
    return OUT


if __name__ == "__main__":
    path = build()
    print("wrote", path, "bytes", path.stat().st_size)
