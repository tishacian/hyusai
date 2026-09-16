#!/usr/bin/env python3
"""Build the Nawa Limited Access Word in the Hermann assessment-note style.

Visual contract (cloned from the 7 October talking-points source):
Aptos, title 25 pt, navy subtitle 14 pt (#17365D), Heading 1 15 pt, body 10.5 pt,
navy callouts, two-column tables with navy header, #F4F7FA zebra, #D9D9D9 borders.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "nawa-limited-access-release-template.docx"
OUT = HERE / "nawa-limited-access-release-7-october-2026.docx"
ART = Path("/opt/cursor/artifacts/nawa_limited_access_release_7_october_2026.docx")

NAVY = RGBColor(0x17, 0x36, 0x5D)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ZEBRA = "F4F7FA"
HEADER_FILL = "17365D"
BORDER = "D9D9D9"
FONT = "Aptos"

# Printable width at 0.72" side margins on US Letter ≈ 7.06"
USABLE_DXA = 10160


def _set_run(run, *, size: float, bold: bool = False, color: RGBColor = BLACK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:cs"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)


def _keep_with_next(p) -> None:
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(OxmlElement("w:keepNext"))


def _cant_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def _clear_body(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _strip_title_border(doc: Document) -> None:
    pPr = doc.styles["Title"].element.find(qn("w:pPr"))
    if pPr is None:
        return
    border = pPr.find(qn("w:pBdr"))
    if border is not None:
        pPr.remove(border)


def _spacing(p, *, before=0, after=0, line=269) -> None:
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        pPr.append(sp)
    sp.set(qn("w:before"), str(before))
    sp.set(qn("w:after"), str(after))
    sp.set(qn("w:line"), str(line))
    sp.set(qn("w:lineRule"), "auto")


def title(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p, before=0, after=160, line=240)
    run = p.add_run(text)
    _set_run(run, size=25, bold=True)


def subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _spacing(p, before=0, after=360, line=269)
    run = p.add_run(text)
    _set_run(run, size=14, bold=True, color=NAVY)


def h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    _spacing(p, before=180, after=100, line=269)
    # heading() already inserted the text; restyle the runs
    if p.runs:
        p.runs[0].text = text
        _set_run(p.runs[0], size=15, bold=True)
    else:
        run = p.add_run(text)
        _set_run(run, size=15, bold=True)
    _keep_with_next(p)


def body(doc: Document, text: str, *, size: float = 10.5, keep=False) -> None:
    p = doc.add_paragraph()
    _spacing(p, before=0, after=160, line=276)
    run = p.add_run(text)
    _set_run(run, size=size, bold=False)
    if keep:
        _keep_with_next(p)


def callout(doc: Document, text: str, *, size: float = 10.5) -> None:
    p = doc.add_paragraph()
    _spacing(p, before=80, after=200, line=276)
    run = p.add_run(text)
    _set_run(run, size=size, bold=True, color=NAVY)


def blank(doc: Document) -> None:
    p = doc.add_paragraph()
    _spacing(p, before=0, after=0, line=240)
    p.add_run("")


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    _spacing(p, before=40, after=40, line=276)
    if p.runs:
        p.runs[0].text = text
        _set_run(p.runs[0], size=10)
    else:
        run = p.add_run(text)
        _set_run(run, size=10)


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def _borders(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)


def _cell_margin(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for edge, tag in (("top", "top"), ("left", "start"), ("bottom", "bottom"), ("right", "end")):
        el = tcMar.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tcMar.append(el)
        el.set(qn("w:w"), "110")
        el.set(qn("w:type"), "dxa")
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), "center")


def _set_width(cell, dxa: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:type"), "dxa")
    tcW.set(qn("w:w"), str(dxa))


def table(doc: Document, headers: tuple[str, str], rows: list[tuple[str, str]], *, col0_dxa: int) -> None:
    col1_dxa = USABLE_DXA - col0_dxa
    tbl = doc.add_table(rows=1 + len(rows), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    tblPr = tbl._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl._tbl.insert(0, tblPr)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(USABLE_DXA))
    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblPr.append(jc)
    jc.set(qn("w:val"), "center")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tbl._tbl.insert(1, grid)
    for w in (col0_dxa, col1_dxa):
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)

    widths = (col0_dxa, col1_dxa)
    data = [headers, *rows]
    for ri, row in enumerate(data):
        tr = tbl.rows[ri]
        if ri == 0:
            trPr = tr._tr.get_or_add_trPr()
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            trPr.append(hdr)
        for ci, text in enumerate(row):
            cell = tr.cells[ci]
            _set_width(cell, widths[ci])
            _cell_margin(cell)
            _borders(cell)
            if ri == 0:
                _shade(cell, HEADER_FILL)
                fill_color, bold, size = WHITE, True, 9
            else:
                if ri % 2 == 0:
                    _shade(cell, ZEBRA)
                fill_color, bold, size = BLACK, False, 9
            p = cell.paragraphs[0]
            p.clear()
            _spacing(p, before=0, after=0, line=240)
            run = p.add_run(text)
            _set_run(run, size=size, bold=bold, color=fill_color)
        _cant_split(tr)
    blank(doc)


def build() -> Path:
    if not TEMPLATE.exists():
        raise SystemExit(f"missing style template: {TEMPLATE}")
    doc = Document(TEMPLATE)
    _clear_body(doc)
    _strip_title_border(doc)

    title(doc, "Nawa Limited Access Release")
    subtitle(doc, "Client assessment baseline available 7 October 2026")

    body(
        doc,
        "Nawa is the white-label of Agentium. From 7 October 2026 it will be "
        "available for a limited client assessment in a controlled, versioned "
        "environment. The assessment is for the intended end users — data "
        "scientists and AI engineers — so they can operate the product "
        "unattended across the declared functional envelope. Nawa does not "
        "receive a parallel feature set. Every Agentium capability included in "
        "the release is inherited by Nawa by transitivity, under the same "
        "contracts. This protects production systems, proprietary assets and "
        "unrelated environments, and it gives both parties a stable and "
        "reproducible basis for review.",
        size=11,
    )

    h1(doc, "Purpose of the controlled release")
    body(
        doc,
        "The objective is not a single representative agent and not a guided "
        "replay of one System. The 7 October baseline is reached when a data "
        "scientist or AI engineer can use Nawa in complete autonomy across the "
        "Agentium functional scope included in the release: Work, Create / "
        "Studio, Systems, Flows, Runs, Decisions, Data & Models, evaluation, "
        "identity and audit.",
    )
    body(
        doc,
        "The release brings those existing product families into one consistent "
        "configuration. It does not expand the agreed functional scope. It does "
        "not open Marketplace listing, cluster portability or factory-scale "
        "rollout.",
    )

    h1(doc, "Why access before 7 October is not relevant")
    table(
        doc,
        ("Reason", "Assessment impact"),
        [
            (
                "Moving product baseline",
                "Agentium is still aligning the families that Nawa will inherit. "
                "Observations made now may not correspond to the released configuration.",
            ),
            (
                "Incomplete autonomy context",
                "End users would still need a guided session. Findings would describe "
                "a presenter-led path, not unattended use of the envelope.",
            ),
            (
                "Incomplete inheritance",
                "A capability that is not yet consistently enforced on Agentium is "
                "not yet a Nawa capability, even if a seeded example exists.",
            ),
            (
                "Limited reproducibility",
                "Changes made during consolidation could prevent either party from "
                "reproducing an observation raised against an earlier build.",
            ),
            (
                "Disproportionate review effort",
                "Reviewing an intermediate configuration would require the client to "
                "reassess the same areas after the controlled release is established.",
            ),
        ],
        col0_dxa=2800,
    )
    callout(
        doc,
        "The proposed timing does not reduce the client’s ability to perform due "
        "diligence. It ensures that the review applies to the version intended to "
        "serve as the assessment baseline.",
        size=10,
    )

    h1(doc, "Product roadmap — Agentium capabilities inherited by Nawa")
    body(
        doc,
        "The roadmap below is the Agentium product plan for 16 September–7 October "
        "2026. Each row is a generic platform capability. Nawa receives it by "
        "transitivity when that capability is in the controlled release. It is "
        "not a Nawa-only build and it is not scoped to one System.",
        keep=True,
    )
    table(
        doc,
        ("Agentium capability", "Inherited Nawa outcome"),
        [
            (
                "Work",
                "Published Experiences are launched from Work. The operator does not "
                "need a presenter URL.",
            ),
            (
                "Create / Studio",
                "A published System can be bound, released and deployed through the "
                "same authoring path.",
            ),
            (
                "System / Flow",
                "The published Flow is the executable contract. Drafts are not "
                "selectable; versions do not float in silence.",
            ),
            (
                "Run",
                "Execution is a canonical Run with recoverable evidence, not a chat "
                "transcript.",
            ),
            (
                "Decision",
                "Accept / reject is an explicit human act. It records an opinion; it "
                "does not by itself apply a write.",
            ),
            (
                "Skills, connectors and patterns",
                "The governed catalogue is the one presented in the environment. "
                "Sealed writes stay sealed.",
            ),
            (
                "Model routing",
                "Routing and fallback are workspace policy, applied to every "
                "in-envelope System.",
            ),
            (
                "Data & Models",
                "Dataset → model version → evaluation → published Skill is a product "
                "loop, not a seeded demo only.",
            ),
            (
                "Evaluation and observability",
                "Evaluation and prompt-regression evidence refer to the same released "
                "configuration as the Run.",
            ),
            (
                "Identity, mandates and approvals",
                "Identity, permitted actions and the approval path are the same "
                "contracts on every inherited System.",
            ),
            (
                "Audit and operational telemetry",
                "Execution events, traces and operating signals belong to the Run "
                "ledger.",
            ),
            (
                "Impact / Hypervisor",
                "Displayed measures are measured, declared or missing. Missing is not "
                "shown as zero. Unattested financial totals are not the baseline.",
            ),
            (
                "Deployment artefacts",
                "The handover pack matches the versioned Compose configuration of the "
                "release. Kubernetes / Marketplace packaging is outside this envelope.",
            ),
        ],
        col0_dxa=3200,
    )
    callout(
        doc,
        "Release principle: one Agentium product envelope, inherited by Nawa, "
        "operable without a guided session.",
    )

    h1(doc, "Weekly product sequence — 16 September to 7 October 2026")
    body(
        doc,
        "The following activities consolidate, standardise and consistently enforce "
        "existing Agentium capabilities so that Nawa can inherit them. They do not "
        "introduce additional functional scope.",
        keep=True,
    )
    table(
        doc,
        ("Week", "Agentium product work (inherited by Nawa)"),
        [
            (
                "16–21 September — Envelope",
                "Freeze the 7 October capability map. Apply the honesty contract "
                "(measured / declared / missing / sealed) across Work, Create, "
                "Systems, Runs, Decisions and Impact. Align lexicon and navigation "
                "so the same families are discoverable without a presenter.",
            ),
            (
                "22–28 September — Operate",
                "Make the generic operate loop autonomous: launch an Experience, "
                "execute its System, open the Run evidence from the Work result, "
                "record a Decision. Align evaluation evidence with that Run. Any "
                "in-envelope System must follow this loop, not only a hero path.",
            ),
            (
                "29 September–5 October — Build and Data & Models",
                "Make the generic build loop autonomous: bind a published System in "
                "Studio, pass ready-check, cut an immutable Release, deploy Pilot or "
                "In-service. Make the data loop autonomous: dataset, model version, "
                "evaluation, published Skill, model-routing policy.",
            ),
            (
                "6–7 October — Autonomy cut",
                "Unattended dry-run of the declared envelope by data-scientist / "
                "AI-engineer profiles. Open limited access only for the families that "
                "passed. A family that still requires Datategy is labelled as such or "
                "left out of the envelope.",
            ),
        ],
        col0_dxa=3400,
    )

    h1(doc, "Release consolidation in progress")
    body(
        doc,
        "The following activities concern consolidation, standardisation and "
        "consistent enforcement of existing Agentium capabilities. They are part of "
        "establishing the controlled release baseline and do not introduce "
        "additional functional scope. Nawa inherits the result by transitivity.",
        keep=True,
    )
    table(
        doc,
        ("Consolidation area", "Release activity"),
        [
            (
                "Work and Experience runtime",
                "Consolidate the launcher, certified renderer pin and deployment "
                "records so published Experiences are operable from Work.",
            ),
            (
                "Create / Studio and System binding",
                "Standardise bind → ready-check → immutable Release → Pilot / In-service.",
            ),
            (
                "Orchestration and execution lifecycle",
                "Consolidate the supported build, execution, recovery and completion "
                "paths in one release configuration.",
            ),
            (
                "Versioning, replay and traceability",
                "Standardise version identification, execution snapshots, replay "
                "behaviour and Run lineage.",
            ),
            (
                "Connector, skill and pattern catalogues",
                "Curate the approved components presented in the controlled environment.",
            ),
            (
                "Model routing policies",
                "Apply the supported routing and fallback policies as workspace "
                "policy, not per demo.",
            ),
            (
                "Data & Models",
                "Align datasets, model versions, champion designation and published "
                "Skills with the released configuration.",
            ),
            (
                "Evaluation and prompt regression",
                "Align evaluation results and regression evidence with the released "
                "configuration.",
            ),
            (
                "Identity, mandates and approvals",
                "Standardise the identity, permitted actions and approval path used "
                "across the envelope.",
            ),
            (
                "Audit and operational telemetry",
                "Consolidate execution events, traces and operating signals for the "
                "inherited Systems.",
            ),
            (
                "Governance records",
                "Align governance information with the configuration included in the "
                "release.",
            ),
            (
                "Deployment and handover artefacts",
                "Harmonise the configuration and handover artefacts associated with "
                "the reference environment.",
            ),
        ],
        col0_dxa=3380,
    )

    h1(doc, "Capabilities presented on 7 October")
    body(
        doc,
        "The session is designed so that a data scientist or AI engineer can "
        "operate the inherited Agentium families together, rather than watching "
        "multiple disconnected demonstrations. The walkthrough is not limited to "
        "one System.",
        keep=True,
    )
    table(
        doc,
        ("Demonstrated capability", "What will be shown"),
        [
            (
                "Work",
                "Launch of a published Experience chosen by the operator, not only a "
                "presenter bookmark.",
            ),
            (
                "System / Flow configuration",
                "The published Flow, approved components and released configuration.",
            ),
            (
                "Orchestration and execution",
                "Execution of the System and visibility of its main processing steps "
                "as a Run.",
            ),
            (
                "Version, snapshot and replay",
                "Identification of the executed version, Run history and controlled "
                "replay.",
            ),
            (
                "Create / Studio",
                "Binding a published System and the Release / deploy path used for "
                "the envelope.",
            ),
            (
                "Data & Models",
                "Dataset, model versions, evaluation evidence and a Skill consumed by "
                "a System.",
            ),
            (
                "Connectors, skills and patterns",
                "The governed catalogue components used in the environment, including "
                "sealed writes.",
            ),
            (
                "Model routing",
                "Application of the configured routing and fallback policy.",
            ),
            (
                "Evaluation evidence",
                "Evaluation results and prompt-regression evidence bound to the same "
                "configuration.",
            ),
            (
                "Identity and approvals",
                "Identity, permitted actions, mandate and the applicable approval point.",
            ),
            (
                "Traceability and audit",
                "Execution trace, relevant operational telemetry and audit events.",
            ),
            (
                "Decision",
                "An explicit human accept / reject, recorded as opinion, not as an "
                "automatic apply.",
            ),
            (
                "Impact",
                "Measures shown as measured, declared or missing. No unattested "
                "financial total.",
            ),
            (
                "Deployment and handover",
                "The configuration and handover artefacts available for the released "
                "envelope.",
            ),
        ],
        col0_dxa=3240,
    )

    h1(doc, "Outside the 7 October envelope")
    body(
        doc,
        "These items remain Agentium product work. They are not inherited by Nawa "
        "in this limited access release.",
        keep=True,
    )
    table(
        doc,
        ("Out of envelope", "Why it is not in the 7 October baseline"),
        [
            (
                "Azure Marketplace listing",
                "A commercial listing is not a condition of unattended product use on "
                "the current environment.",
            ),
            (
                "Kubernetes / Helm portability",
                "The reference environment is the versioned Compose configuration. "
                "Cluster packaging is a later product increment.",
            ),
            (
                "Unsealed external writes",
                "Connectors that are sealed in the release stay sealed. A write is "
                "not implied by a Decision.",
            ),
            (
                "Hypervisor as attested financial ROI",
                "Impact may show measured, declared or missing values. It does not "
                "certify economic totals.",
            ),
            (
                "Adoption journey for business users",
                "The 7 October end users are data scientists and AI engineers. "
                "Business-user adoption chrome is a separate acceptance.",
            ),
            (
                "Factory-scale agent count",
                "The release is the product envelope, not a volume of agents.",
            ),
        ],
        col0_dxa=3400,
    )

    h1(doc, "Limited environment")
    bullet(
        doc,
        "The declared Agentium envelope on the Nawa workspace, in an isolated "
        "assessment environment.",
    )
    bullet(doc, "Controlled test scenarios and non-production data.")
    bullet(
        doc,
        "Role-restricted, time-limited access for agreed data-scientist and "
        "AI-engineer participants.",
    )
    bullet(
        doc,
        "No access to production systems, unrelated customer environments or "
        "unrestricted administrative functions.",
    )
    bullet(
        doc,
        "No penetration testing, source-code review or destructive testing unless "
        "separately agreed in writing.",
    )
    callout(
        doc,
        "The assessment scope can be agreed before 7 October so that the session "
        "addresses the client’s priority questions while remaining anchored to the "
        "same controlled release.",
    )

    doc.core_properties.title = "Nawa Limited Access Release"
    doc.core_properties.subject = (
        "Agentium product envelope inherited by Nawa — 7 October 2026"
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    ART.parent.mkdir(parents=True, exist_ok=True)
    ART.write_bytes(OUT.read_bytes())
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
    print(ART)
