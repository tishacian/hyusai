# -*- coding: utf-8 -*-
"""Note client Andritz — Proposition d'avenant / offre de poursuite (Word brandé Datategy)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "logo_datategy.png")
OUT_DIR = os.path.join(HERE, "out")
OUT = os.path.join(OUT_DIR, "Datategy-Andritz-Proposition-avenant-poursuite-2026-07.docx")
OUT_CLIENT = os.path.join(
    os.path.expanduser("~"),
    "Developer/DATATEGY/papAI/POC/ONE/ANDRITZ/project-andritz",
    "Datategy-Andritz-Proposition-avenant-poursuite-2026-07.docx",
)

# Grille poursuite (HT / mois)
HOSTING = 6_500
T1_RUN = 1_500
# Net capacité Datategy (inchangé vs baseline ETP)
LOW_NET, MED_NET, HIGH_NET = 6_600, 14_200, 18_900
# Assistance Datategy avec Éric O’Neill absorbée dans le forfait (j/mois)
LOW_ASSIST_J, MED_ASSIST_J, HIGH_ASSIST_J = 3, 4, 5
ASSIST_DAY = 1_000  # valorisation interne indicative
LOW_CAP = LOW_NET + LOW_ASSIST_J * ASSIST_DAY  # 9 600
MED_CAP = MED_NET + MED_ASSIST_J * ASSIST_DAY  # 18 200
HIGH_CAP = HIGH_NET + HIGH_ASSIST_J * ASSIST_DAY  # 23 900
LOW_ALL = LOW_CAP + HOSTING  # 16 100
MED_ALL = MED_CAP + HOSTING  # 24 700
HIGH_ALL = HIGH_CAP + HOSTING  # 30 400
RUN_ALL = HOSTING + T1_RUN  # 8 000


def C(h):
    return RGBColor.from_string(h)


NAVY = C("0B2545")
BLUE = C("2D6CDF")
GREEN = C("0E8F62")
AMBER = C("B45309")
INK = C("151A23")
MUT = C("596371")
WHITE = C("FFFFFF")
LINE = "D7DCE3"
BODY = "Arial"


def _runfmt(run, size, color, bold=False, italic=False, mono=False):
    run.font.name = "Consolas" if mono else BODY
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _shade(cell, hexfill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(sh)


def _cell_margins(cell, t=80, b=80, l=140, r=140):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, v in (
        ("top", t),
        ("bottom", b),
        ("start", l),
        ("end", r),
        ("left", l),
        ("right", r),
    ):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def _borders(table, color=LINE, sz=4):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    table._tbl.tblPr.append(b)


def _left_accent(cell, hexcolor, sz=30):
    tcPr = cell._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    e = OxmlElement("w:left")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), hexcolor)
    bd.append(e)
    tcPr.append(bd)


def _pbottom(p, color="C9D2DE", sz=6):
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "4")
    e.set(qn("w:color"), color)
    pb.append(e)
    pPr.append(pb)


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY
        r.font.name = BODY
    _pbottom(p)
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = BLUE
        r.font.name = BODY
        r.font.size = Pt(12)
    return p


def para(doc, text, size=10.5, color=INK, bold=False, italic=False, before=3, after=5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.15
    r = p.add_run(text)
    _runfmt(r, size, color, bold=bold, italic=italic)
    return p


def bullet(doc, parts, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.12
    if isinstance(parts, str):
        parts = [(parts, False, INK)]
    for text, bold, color in parts:
        _runfmt(p.add_run(text), size, color, bold=bold)
    return p


def callout(doc, title, body, accent=BLUE, accent_hex="2D6CDF", fill="E5EDFB"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Cm(17)
    _shade(cell, fill)
    _left_accent(cell, accent_hex, sz=34)
    _cell_margins(cell, t=100, b=100, l=180, r=150)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    _runfmt(p.add_run(title.upper()), 9, accent, bold=True, mono=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.14
    _runfmt(p2.add_run(body), 10.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def table(doc, headers, rows, col_widths=None, header_fill="0B2545"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        _shade(cell, header_fill)
        _cell_margins(cell, t=60, b=60, l=100, r=100)
        cell.text = ""
        p = cell.paragraphs[0]
        _runfmt(p.add_run(h), 9, WHITE, bold=True)
    for r_i, row in enumerate(rows):
        fill = "F7F9FC" if r_i % 2 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            _shade(cell, fill)
            _cell_margins(cell, t=55, b=55, l=100, r=100)
            cell.text = ""
            p = cell.paragraphs[0]
            _runfmt(p.add_run(str(val)), 9.5, INK)
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def package_card(doc, name, price, subtitle, accent_hex, fill, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Cm(17)
    _shade(cell, fill)
    _left_accent(cell, accent_hex, sz=36)
    _cell_margins(cell, t=90, b=90, l=170, r=140)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    _runfmt(p.add_run(name), 12, NAVY, bold=True)
    _runfmt(p.add_run(f"   ·   {price}"), 12, C(accent_hex), bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(4)
    _runfmt(p2.add_run(subtitle), 9.5, MUT, italic=True)
    for label, text in lines:
        bp = cell.add_paragraph()
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(1)
        _runfmt(bp.add_run(f"{label}  "), 9.5, C(accent_hex), bold=True)
        _runfmt(bp.add_run(text), 9.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t


def build():
    os.makedirs(OUT_DIR, exist_ok=True)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    st = doc.styles["Heading 1"]
    st.font.name = BODY
    st.font.size = Pt(15)
    st.font.bold = True
    st.font.color.rgb = NAVY
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

    st2 = doc.styles["Heading 2"]
    st2.font.name = BODY
    st2.font.size = Pt(12)
    st2.font.bold = True
    st2.font.color.rgb = BLUE
    st2.paragraph_format.space_before = Pt(10)
    st2.paragraph_format.space_after = Pt(3)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    sec.left_margin = sec.right_margin = Cm(2.0)

    f = sec.footer
    f.is_linked_to_previous = False
    fp = f.paragraphs[0]
    _runfmt(
        fp.add_run(
            "Confidentiel · Datategy × ANDRITZ France · Proposition d’avenant — juillet 2026"
        ),
        8,
        MUT,
    )

    # ---- En-tête ----
    tb = doc.add_table(rows=1, cols=2)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    lc, rc = tb.cell(0, 0), tb.cell(0, 1)
    lc.width = Cm(13.5)
    rc.width = Cm(3.5)
    lc.text = ""
    pe = lc.paragraphs[0]
    _runfmt(pe.add_run("PROPOSITION COMMERCIALE · AVENANT"), 10, BLUE, bold=True, mono=True)
    pt = lc.add_paragraph()
    pt.paragraph_format.space_before = Pt(2)
    _runfmt(pt.add_run("Poursuite du dispositif capacitaire IA"), 18, NAVY, bold=True)
    ps = lc.add_paragraph()
    ps.paragraph_format.space_before = Pt(3)
    _runfmt(
        ps.add_run("ANDRITZ France SAS  ·  Datategy  ·  Juillet 2026  ·  Montants HT"),
        9.5,
        MUT,
    )
    rp = rc.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(LOGO):
        rp.add_run().add_picture(LOGO, width=Cm(2.6))
    p = doc.add_paragraph()
    _pbottom(p, "C9D2DE", sz=6)
    p.paragraph_format.space_after = Pt(8)

    para(
        doc,
        "La phase initiale du dispositif capacitaire est arrivée à son terme. "
        "Quatre applications métier sont disponibles sur la plateforme papAI / Agentium. "
        "La présente note propose les conditions de poursuite : un hébergement dimensionné "
        "à l’usage réel, et une capacité de service Datategy modulable — incluant, selon le "
        "forfait, une assistance Datategy avec Éric O’Neill (cadrage, ownership, adoption), "
        "dans la continuité de la collaboration déjà engagée.",
        after=8,
    )

    callout(
        doc,
        "En une phrase",
        "Deux postes Datategy uniquement : (1) hébergement plateforme / applications, "
        "(2) capacité de service. L’assistance avec Éric O’Neill est intégrée dans le forfait "
        "capacité Datategy — pas une prestation séparée facturée au Client.",
    )

    # ---- 1. Périmètre livré ----
    h1(doc, "1.  Périmètre déjà mis à disposition")
    para(
        doc,
        "Les applications suivantes sont opérationnelles dans l’environnement Andritz :",
    )
    bullet(doc, [("Recherche documentaire", True, BLUE), (" — interrogation des manuels et corpus techniques, avec corrections possibles sur les informations retrouvées.", False, INK)])
    bullet(doc, [("Capture de connaissances", True, BLUE), (" — formalisation et publication du savoir expert.", False, INK)])
    bullet(doc, [("Client 360", True, BLUE), (" — vision client consolidée (priorisation et compléments selon roadmap).", False, INK)])
    bullet(doc, [("Rapports site FSE", True, BLUE), (" — génération assistée disponible ; ajustements attendus après premiers usages.", False, INK)])
    para(
        doc,
        "S’y ajoutent le socle plateforme (indexation, dépôt sécurisé / SFTP, orchestration, "
        "environnements de run) nécessaire au fonctionnement multi-utilisateurs.",
        color=MUT,
        after=8,
    )

    # ---- 2. Structure de l'offre ----
    h1(doc, "2.  Structure de l’offre")
    para(doc, "Chaque mensualité se compose de deux postes Datategy :")
    table(
        doc,
        ["Poste", "Rôle", "Caractère"],
        [
            [
                "Hébergement",
                "Infrastructure, plateforme, applications en production, volume data",
                f"Montant fixe — {HOSTING:,} € HT / mois".replace(",", " "),
            ],
            [
                "Capacité / service",
                "Delivery Datategy +, selon forfait, assistance Datategy avec Éric O’Neill "
                "(cadrage, ownership métier, adoption)",
                "Au choix : Run, Low, Medium ou High",
            ],
        ],
        col_widths=[3.5, 8.0, 5.5],
    )
    para(
        doc,
        "Le support de base (Tier 1 : disponibilité, incidents, maintien en condition "
        "opérationnelle) est inclus dans les forfaits Low, Medium et High. Il n’apparaît "
        "comme ligne distincte que dans le forfait Run (hébergement sans capacité de build).",
    )
    callout(
        doc,
        "Assistance avec Éric O’Neill",
        "Dans la continuité du pilotage, du cadrage et de l’ownership déjà assurés pendant "
        "la phase initiale, les forfaits Low, Medium et High incluent une assistance "
        "Datategy avec Éric O’Neill (respectivement 3, 4 et 5 jours par mois). "
        "Cette assistance est portée et facturée par Datategy dans le forfait capacité — "
        "elle ne constitue pas une prestation séparée du Client.",
        accent=GREEN,
        accent_hex="0E8F62",
        fill="E2F2EC",
    )

    # ---- 3. Trajectoire ----
    h1(doc, "3.  Trajectoire proposée")
    h2(doc, "3.1  Août 2026 — ancrage de l’hébergement")
    para(
        doc,
        "Mois de transition (période estivale, préparation des démonstrations de rentrée). "
        "Datategy propose d’ancrer dès août le nouveau poste d’hébergement, avec la capacité "
        "de service offerte pour ce mois.",
    )
    table(
        doc,
        ["Poste", "Montant HT", "Commentaire"],
        [
            ["Hébergement", f"{HOSTING:,} €".replace(",", " "), "Tarif cible — dès août"],
            ["Capacité / service", "0 € (offerte)", "Geste de transition"],
            ["Total août", f"{HOSTING:,} €".replace(",", " "), "Ancrage de l’infrastructure"],
        ],
        col_widths=[5.0, 4.5, 7.5],
    )

    h2(doc, "3.2  À partir de septembre 2026 — régime choisi")
    para(
        doc,
        "À compter de septembre, l’hébergement reste à "
        f"{HOSTING:,} € HT / mois. ".replace(",", " ")
        + "ANDRITZ choisit le forfait de capacité adapté à sa feuille de route.",
    )
    table(
        doc,
        ["Forfait", "Capacité / service", "Hébergement", "Total mensuel HT"],
        [
            ["Run", f"{T1_RUN:,} € (Tier 1 seul)".replace(",", " "), f"{HOSTING:,} €".replace(",", " "), f"{RUN_ALL:,} €".replace(",", " ")],
            [
                "Low",
                f"{LOW_CAP:,} € (dont {LOW_ASSIST_J} j assistance É. O’Neill)".replace(",", " "),
                f"{HOSTING:,} €".replace(",", " "),
                f"{LOW_ALL:,} €".replace(",", " "),
            ],
            [
                "Medium",
                f"{MED_CAP:,} € (dont {MED_ASSIST_J} j assistance É. O’Neill)".replace(",", " "),
                f"{HOSTING:,} €".replace(",", " "),
                f"{MED_ALL:,} €".replace(",", " "),
            ],
            [
                "High",
                f"{HIGH_CAP:,} € (dont {HIGH_ASSIST_J} j assistance É. O’Neill)".replace(",", " "),
                f"{HOSTING:,} €".replace(",", " "),
                f"{HIGH_ALL:,} €".replace(",", " "),
            ],
        ],
        col_widths=[2.5, 7.5, 3.2, 3.8],
    )
    para(
        doc,
        "Rappel phase initiale : 17 000 € HT / mois (capacité 1,5 ETP et hébergement "
        "exploratoire inclus dans un forfait unique).",
        color=MUT,
        italic=True,
        after=8,
    )

    # ---- 4. Hébergement ----
    h1(doc, f"4.  Poste Hébergement — {HOSTING:,} € HT / mois".replace(",", " "))
    para(
        doc,
        "Pendant la phase initiale, l’enveloppe d’hébergement (~2 800 € HT / mois) "
        "correspondait à un usage exploratoire et séquentiel : les fonctions "
        "consommatrices de ressources étaient utilisées les unes après les autres.",
    )
    para(
        doc,
        "L’usage actuel et cible est différent : quatre applications multi-utilisateurs, "
        "avec des fonctions qui tournent en parallèle (voix, raisonnement, recherche, "
        "OCR / indexation), sur un volume documentaire plus important et plus complexe "
        "que prévu. Un canal SFTP et une capacité de volume ont été mis en place pour "
        "absorber les données nécessaires.",
    )

    h2(doc, "4.1  Ce que couvre le poste")
    table(
        doc,
        ["Composant", "Contenu"],
        [
            ["Compute applicatif", "Backend, interface, workers d’ingestion et traitements asynchrones"],
            [
                "Compute IA (chemins isolés)",
                "Raisonnement (grand modèle) · voix (STT/TTS) · OCR / embeddings / rerank · routeur / modèles rapides",
            ],
            ["Data", "Base applicative, index vectoriel, stockage documentaire, rétention"],
            ["Ingest", "Dépôt sécurisé / SFTP et volume dimensionné pour le corpus"],
            ["Exploitation", "Supervision, sauvegardes, snapshots, bande passante"],
        ],
        col_widths=[4.5, 12.5],
    )

    h2(doc, "4.2  Topologie — de l’exploratoire au concurrent")
    table(
        doc,
        ["", "Phase initiale (~2 800 €)", "Poursuite (~6 500 €)"],
        [
            [
                "Organisation compute",
                "Ressources partagées, file séquentielle",
                "Chemins isolés pour éviter les files d’attente",
            ],
            [
                "Applications",
                "Exploration, usage limité en parallèle",
                "4 applications multi-utilisateurs",
            ],
            [
                "Données",
                "Volume et complexité contenus",
                "OCR / indexation / retrieval dimensionnés au corpus réel",
            ],
            [
                "Ingest",
                "Flux ad hoc",
                "SFTP + volume provisionné",
            ],
        ],
        col_widths=[3.5, 6.5, 7.0],
    )
    callout(
        doc,
        "Point important",
        (
            f"Le montant de {HOSTING:,} € HT / mois est le même quel que soit le forfait "
            f"de capacité choisi (Run, Low, Medium ou High). Il finance exclusivement "
            f"l’infrastructure nécessaire au fonctionnement concurrent des applications. "
            f"Il ne porte pas l’assistance métier / ownership."
        ).replace(",", " "),
        accent=AMBER,
        accent_hex="B45309",
        fill="F6ECE0",
    )

    # ---- 5. Forfaits ----
    h1(doc, "5.  Forfaits de capacité / service")
    para(
        doc,
        "Les montants ci-dessous sont les totaux mensuels HT (hébergement + capacité). "
        "Pour Low / Medium / High, le forfait capacité inclut à la fois la capacité "
        "delivery Datategy et l’assistance Datategy avec Éric O’Neill.",
    )

    h2(doc, "5.1  Composition du forfait capacité (Low / Medium / High)")
    table(
        doc,
        ["Forfait", "Capacité delivery", "Assistance É. O’Neill", "Capacité facturée", "Total avec héberg."],
        [
            [
                "Low",
                f"{LOW_NET:,} €".replace(",", " "),
                f"{LOW_ASSIST_J} j / mois",
                f"{LOW_CAP:,} €".replace(",", " "),
                f"{LOW_ALL:,} €".replace(",", " "),
            ],
            [
                "Medium",
                f"{MED_NET:,} €".replace(",", " "),
                f"{MED_ASSIST_J} j / mois",
                f"{MED_CAP:,} €".replace(",", " "),
                f"{MED_ALL:,} €".replace(",", " "),
            ],
            [
                "High",
                f"{HIGH_NET:,} €".replace(",", " "),
                f"{HIGH_ASSIST_J} j / mois",
                f"{HIGH_CAP:,} €".replace(",", " "),
                f"{HIGH_ALL:,} €".replace(",", " "),
            ],
        ],
        col_widths=[2.5, 3.5, 3.5, 3.8, 3.7],
    )
    para(
        doc,
        "L’assistance couvre notamment le pilotage métier, le cadrage et l’expression "
        "de besoin, le conseil technologique, et l’ownership des applications — "
        "dans la continuité du dispositif déjà expérimenté avec succès.",
        color=MUT,
        after=6,
    )

    h2(doc, "5.2  Run — hébergement et maintien opérationnel")
    package_card(
        doc,
        "Run",
        f"{RUN_ALL:,} € HT / mois".replace(",", " "),
        "Les applications tournent telles quelles — pas de backlog de développement",
        "0E8F62",
        "E2F2EC",
        [
            ("Hébergement", f"{HOSTING:,} € — infrastructure et plateforme".replace(",", " ")),
            ("Support Tier 1 / SLA", f"{T1_RUN:,} € — disponibilité, incidents, redémarrages, health-checks".replace(",", " ")),
            ("Assistance É. O’Neill", "Non incluse (forfait sans capacité de build)"),
            ("Hors scope", "Évolutions fonctionnelles, nouvelles applications, tuning produit, formations étendues"),
        ],
    )

    h2(doc, "5.3  Low — cadence réduite")
    package_card(
        doc,
        "Low",
        f"{LOW_ALL:,} € HT / mois".replace(",", " "),
        f"Capacité {LOW_CAP:,} € + hébergement {HOSTING:,} €".replace(",", " "),
        "2D6CDF",
        "E5EDFB",
        [
            ("Delivery Datategy", f"{LOW_NET:,} € — itérations, corrections, accompagnement léger".replace(",", " ")),
            ("Assistance É. O’Neill", f"{LOW_ASSIST_J} jours / mois — inclus dans le forfait capacité"),
            ("Hébergement", f"{HOSTING:,} €".replace(",", " ")),
            ("Support Tier 1", "Inclus (diffusé — pas de ligne séparée)"),
            ("Typiquement", "Formations / kick-off, évolutions guidées par les retours, finitions d’applications proches"),
        ],
    )

    h2(doc, "5.4  Medium — rythme de référence")
    package_card(
        doc,
        "Medium",
        f"{MED_ALL:,} € HT / mois".replace(",", " "),
        f"Capacité {MED_CAP:,} € + hébergement {HOSTING:,} €".replace(",", " "),
        "0B2545",
        "E7ECF4",
        [
            ("Delivery Datategy", f"{MED_NET:,} € — rythme comparable à la phase initiale (1,5 ETP)".replace(",", " ")),
            ("Assistance É. O’Neill", f"{MED_ASSIST_J} jours / mois — inclus dans le forfait capacité"),
            ("Hébergement", f"{HOSTING:,} €".replace(",", " ")),
            ("Support Tier 1", "Inclus (diffusé)"),
            ("Typiquement", "Plusieurs chantiers en parallèle (FSE, campagnes, Client 360, sync data, etc.)"),
        ],
    )

    h2(doc, "5.5  High — accélération")
    package_card(
        doc,
        "High",
        f"{HIGH_ALL:,} € HT / mois".replace(",", " "),
        f"Capacité {HIGH_CAP:,} € + hébergement {HOSTING:,} €".replace(",", " "),
        "B45309",
        "F6ECE0",
        [
            ("Delivery Datategy", f"{HIGH_NET:,} € — accélération et multi-chantiers (~2 ETP)".replace(",", " ")),
            ("Assistance É. O’Neill", f"{HIGH_ASSIST_J} jours / mois — inclus dans le forfait capacité"),
            ("Hébergement", f"{HOSTING:,} €".replace(",", " ")),
            ("Support Tier 1", "Inclus (diffusé)"),
            ("Typiquement", "Roadmap ambitieuse avec plusieurs flux de livraison simultanés"),
        ],
    )

    # ---- 6. Récap ----
    h1(doc, "6.  Récapitulatif des mensualités")
    table(
        doc,
        ["Période / forfait", "Capacité", "Hébergement", "Total HT / mois"],
        [
            ["Phase initiale (référence)", "14 200 € (1,5 ETP)", "~2 800 € (inclus bundle)", "17 000 €"],
            ["Août 2026", "Offerte", f"{HOSTING:,} €".replace(",", " "), f"{HOSTING:,} €".replace(",", " ")],
            ["Sept.+ · Run", f"{T1_RUN:,} € (T1)".replace(",", " "), f"{HOSTING:,} €".replace(",", " "), f"{RUN_ALL:,} €".replace(",", " ")],
            [
                "Sept.+ · Low",
                f"{LOW_CAP:,} € (dont {LOW_ASSIST_J} j É. O’Neill)".replace(",", " "),
                f"{HOSTING:,} €".replace(",", " "),
                f"{LOW_ALL:,} €".replace(",", " "),
            ],
            [
                "Sept.+ · Medium",
                f"{MED_CAP:,} € (dont {MED_ASSIST_J} j É. O’Neill)".replace(",", " "),
                f"{HOSTING:,} €".replace(",", " "),
                f"{MED_ALL:,} €".replace(",", " "),
            ],
            [
                "Sept.+ · High",
                f"{HIGH_CAP:,} € (dont {HIGH_ASSIST_J} j É. O’Neill)".replace(",", " "),
                f"{HOSTING:,} €".replace(",", " "),
                f"{HIGH_ALL:,} €".replace(",", " "),
            ],
        ],
        col_widths=[4.0, 5.5, 3.5, 4.0],
    )

    # ---- 7. Modalités ----
    h1(doc, "7.  Modalités")
    bullet(doc, [("Cadre", True, NAVY), (" — avenant au contrat de prestation de services en vigueur (dispositif capacitaire IA).", False, INK)])
    bullet(doc, [("Montants", True, NAVY), (" — exprimés hors taxes ; TVA en sus au taux en vigueur.", False, INK)])
    bullet(doc, [("Contractant", True, NAVY), (" — Datategy est l’unique prestataire facturé ; l’assistance avec Éric O’Neill est incluse dans le forfait capacité Datategy.", False, INK)])
    bullet(doc, [("Hébergement", True, NAVY), (" — poste fixe mensuel, commun à tous les forfaits de poursuite ; distinct de la capacité.", False, INK)])
    bullet(doc, [("Capacité", True, NAVY), (" — forfait mensuel global et indivisible pour le niveau choisi ; ajustable par avenant selon la feuille de route.", False, INK)])
    bullet(doc, [("Support Tier 1", True, NAVY), (" — inclus dans Low / Medium / High ; facturé explicitement uniquement dans Run.", False, INK)])
    bullet(doc, [("Facturation", True, NAVY), (" — mensuelle, selon les modalités du contrat (échéance et envoi électronique).", False, INK)])

    callout(
        doc,
        "Prochaine étape",
        "Valider le forfait de septembre (Run, Low, Medium ou High) et confirmer "
        "l’ancrage de l’hébergement dès août. Datategy rédigera ensuite l’avenant "
        "correspondant pour signature.",
        accent=GREEN,
        accent_hex="0E8F62",
        fill="E2F2EC",
    )

    para(
        doc,
        "Datategy reste à votre disposition pour détailler la feuille de route associée "
        "à chaque forfait et pour caler le calendrier de rentrée.",
        color=MUT,
        italic=True,
        before=8,
    )

    doc.save(OUT)
    client_dir = os.path.dirname(OUT_CLIENT)
    if os.path.isdir(client_dir):
        doc.save(OUT_CLIENT)
        print("Wrote", OUT_CLIENT)
    print("Wrote", OUT)
    return OUT


if __name__ == "__main__":
    build()
