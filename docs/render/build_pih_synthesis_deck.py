# -*- coding: utf-8 -*-
"""Synthesis deck (16:9, EN) summarising the Agentium Product Overview for PIH/PowerMind.

Does NOT modify the product document. Reuses the visual system of build_pih_deck.py
and embeds the diagrams rendered by pih_diagrams.py.
Output: out/Agentium-Synthesis-PIH.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import pih_diagrams as dg

HERE = os.path.dirname(os.path.abspath(__file__))
DIAG = os.path.join(HERE, "assets", "diagrams")
LOGO = os.path.join(HERE, "assets", "logo_datategy.png")
OUT = os.path.join(HERE, "out", "Agentium-Synthesis-PIH.pptx")

# ---------------------------------------------------------------- tokens
SW, SH = 13.333, 7.5
MX = 0.62
INK    = RGBColor(0x15, 0x1A, 0x23)
MUT    = RGBColor(0x59, 0x63, 0x71)
LINE   = RGBColor(0xD7, 0xDC, 0xE3)
PANEL  = RGBColor(0xF3, 0xF5, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
NAVY   = RGBColor(0x0B, 0x25, 0x45)
BLUE   = RGBColor(0x2D, 0x6C, 0xDF)
PURPLE = RGBColor(0x6D, 0x28, 0xD9)
GREEN  = RGBColor(0x0E, 0x8F, 0x62)
AMBER  = RGBColor(0xB4, 0x53, 0x09)
LIGHTBLUE = RGBColor(0xBE, 0xD0, 0xE8)
TINTBLUE  = RGBColor(0xEA, 0xF1, 0xFD)
F      = "Arial"
FM     = "Courier New"
FOOTER = "Confidential — Datategy · Agentium · June 2026"

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]
PAGE = [0]


# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)


def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=12, color=INK, bold=False, mono=False, first=False,
         before=4, after=0, align=PP_ALIGN.LEFT, spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(0 if first else before)
    p.space_after = Pt(after)
    try:
        p.line_spacing = spacing
    except Exception:
        pass
    r = p.add_run()
    r.text = text
    r.font.name = FM if mono else F
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return p


def add_run(p, text, size=12, color=INK, bold=False, mono=False):
    r = p.add_run()
    r.text = text
    r.font.name = FM if mono else F
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return r


def rect(s, x, y, w, h, fill, line_color=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(0.75)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    sp.shadow.inherit = False
    return sp


def chrome(s, kicker, title, accent=NAVY):
    PAGE[0] += 1
    rect(s, MX, 0.52, 0.30, 0.045, accent)
    tf = box(s, MX + 0.42, 0.40, SW - 2 * MX - 0.42, 0.3)
    para(tf, kicker.upper(), size=10.5, color=accent, bold=True, mono=True, first=True)
    tf = box(s, MX, 0.72, SW - 2 * MX, 0.62)
    para(tf, title, size=25, color=INK, bold=True, first=True)
    rect(s, MX, 1.38, SW - 2 * MX, 0.012, LINE)
    tf = box(s, MX, SH - 0.42, 9.0, 0.3)
    para(tf, FOOTER, size=8.5, color=MUT, first=True)
    tf = box(s, SW - MX - 1.0, SH - 0.42, 1.0, 0.3)
    para(tf, f"{PAGE[0]:02d}", size=9, color=MUT, mono=True, first=True, align=PP_ALIGN.RIGHT)


def bullets(tf, items, size=12, color=INK, first=True, before=6, spacing=1.05):
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            head, body = it
            p = para(tf, "▪ ", size=size, color=color, first=(first and i == 0), before=before, spacing=spacing)
            add_run(p, head + " — ", size=size, color=INK, bold=True)
            add_run(p, body, size=size, color=color)
        else:
            para(tf, "▪ " + it, size=size, color=color, first=(first and i == 0), before=before, spacing=spacing)


def panel_card(s, x, y, w, h, accent, title, title_size=13):
    rect(s, x, y, w, h, WHITE, line_color=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.045)
    rect(s, x, y + 0.12, 0.05, h - 0.24, accent)
    tf = box(s, x + 0.22, y + 0.14, w - 0.4, 0.35)
    para(tf, title, size=title_size, color=accent, bold=True, first=True)
    return box(s, x + 0.22, y + 0.54, w - 0.4, h - 0.7)


def style_cell(cell, text, size=10.5, color=INK, bold=False, fill=None, mono=False,
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill if fill is not None else WHITE
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.10)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FM if mono else F
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def pic(s, name, w, y, x=None, caption=None):
    """Embed a diagram PNG (width-driven), horizontally centred unless x given."""
    if x is None:
        x = (SW - w) / 2
    p = s.shapes.add_picture(os.path.join(DIAG, name), Inches(x), Inches(y), width=Inches(w))
    if caption:
        cy = y + Emu(p.height).inches + 0.06
        tf = box(s, MX, cy, SW - 2 * MX, 0.3)
        para(tf, caption, size=9, color=MUT, align=PP_ALIGN.CENTER, first=True)
    return p


def band(s, x, y, w, h, fill, title, body, title_color=WHITE, body_color=LIGHTBLUE):
    rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    tf = box(s, x + 0.28, y + 0.16, w - 0.56, h - 0.32)
    para(tf, title, size=12.5, color=title_color, bold=True, mono=True, first=True)
    para(tf, body, size=12, color=body_color, before=5, spacing=1.06)


# ================================================================ S1 — Cover
def s1_cover():
    s = slide()
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, 0.06, BLUE)
    tf = box(s, MX, 1.7, SW - 2 * MX, 0.4)
    para(tf, "TECHNICAL PRODUCT OVERVIEW   ·   SYNTHESIS", size=12, color=BLUE, bold=True, mono=True, first=True)
    tf = box(s, MX, 2.25, SW - 2 * MX, 1.2)
    para(tf, "Agentium", size=46, color=WHITE, bold=True, first=True)
    tf = box(s, MX, 3.5, SW - 2 * MX, 0.6)
    para(tf, "Sovereign Agent Orchestration, Building & Evaluation", size=20, color=LIGHTBLUE, first=True)
    rect(s, MX, 4.35, 2.0, 0.03, BLUE)
    tf = box(s, MX, 4.6, SW - 2 * MX, 0.9)
    para(tf, "Prepared for PowerMind", size=14, color=WHITE, bold=True, first=True)
    para(tf, "Sovereign AI for PIH, Qatar and the region", size=12.5, color=RGBColor(0x8F, 0xA8, 0xC7), before=3)
    tf = box(s, MX, SH - 0.7, SW - 2 * MX, 0.3)
    para(tf, "Datategy   ·   June 2026", size=11, color=RGBColor(0x8F, 0xA8, 0xC7), mono=True, first=True)


# ================================================================ S2 — Executive summary
def s2_exec():
    s = slide()
    chrome(s, "Executive summary", "An operating system for intelligent systems", NAVY)
    tf = box(s, MX, 1.62, SW - 2 * MX, 0.7)
    p = para(tf, "Agentium is a single platform to ", size=13, color=INK, first=True, spacing=1.1)
    add_run(p, "orchestrate, build and evaluate", size=13, color=NAVY, bold=True)
    add_run(p, " autonomous agents — on infrastructure the customer fully controls.", size=13, color=INK)
    cw, gap = (SW - 2 * MX - 2 * 0.3) / 3, 0.3
    y, h = 2.45, 2.25
    cards = [
        (BLUE, "Orchestrate", ["Stateful, replayable runs", "Checkpoints & resubmission", "Governed tool-calling"]),
        (PURPLE, "Build", ["Typed skills + visual flow", "Versioning & rollback", "Reusable capabilities"]),
        (GREEN, "Evaluate", ["Automatic scoring", "Thresholds & review queue", "Human feedback loop"]),
    ]
    for i, (c, t, items) in enumerate(cards):
        x = MX + i * (cw + gap)
        tf = panel_card(s, x, y, cw, h, c, t)
        bullets(tf, items, size=11.5, before=6)
    band(s, MX, 5.0, SW - 2 * MX, 1.0, NAVY, "THE SOVEREIGN THESIS",
         "Agentium controls the whole stack — models (on-prem), the action ledger and the governance "
         "plane. Sovereignty is enforced rather than promised, and auditability becomes proof rather than logs.")


# ================================================================ S3 — One platform
def s3_platform():
    s = slide()
    chrome(s, "One platform", "One unified, sovereign platform", BLUE)
    pic(s, "d8_portfolio.png", 9.6, 1.66,
        caption="Data & ML, knowledge and agentic capabilities on one platform — one deployment, one sovereign boundary.")


# ================================================================ S4 — Agentium model
def s4_model():
    s = slide()
    chrome(s, "The model", "From data to governed action", PURPLE)
    tf = box(s, MX, 1.6, SW - 2 * MX, 0.45)
    p = para(tf, "System  →  Run  →  Evaluation  →  Decision  →  Action", size=15, color=NAVY, bold=True, mono=True, first=True, align=PP_ALIGN.CENTER)
    pic(s, "d2_chain.png", 8.6, 2.05,
        caption="A System runs, each run is measured and evaluated, producing a governed decision and a traced action.")


# ================================================================ S5 — Orchestration
def s5_orchestration():
    s = slide()
    chrome(s, "Orchestration", "Stateful, replayable, governed execution", BLUE)
    cw = (SW - 2 * MX - 0.3) / 2
    y, h = 1.7, 2.45
    tf = panel_card(s, MX, y, cw, h, BLUE, "Execution lifecycle")
    bullets(tf, ["Initialise → Plan → Execute", "Observe → Evaluate → Adapt",
                 "Adaptive feedback between phases"], size=11.5)
    tf = panel_card(s, MX + cw + 0.3, y, cw, h, PURPLE, "Robust by design")
    bullets(tf, ["Immutable flow snapshots & checkpoints",
                 "Replay & resubmission with overrides",
                 "Parent-child lineage · idempotent skills"], size=11.5)
    band(s, MX, 4.45, SW - 2 * MX, 1.55, PANEL, "EXECUTION MODES",
         "Real-time decision · batch processing · event-driven automation · continuous monitoring · "
         "human-augmented — the same governed runtime, triggered five ways.",
         title_color=NAVY, body_color=INK)


# ================================================================ S6 — Engineering foundation
def s6_foundation():
    s = slide()
    chrome(s, "Engineering foundation", "Neither a black box nor a from-scratch reinvention", GREEN)
    cw = (SW - 2 * MX - 0.3) / 2
    y, h = 1.66, 2.5
    tf = panel_card(s, MX, y, cw, h, GREEN, "The cream of open source")
    bullets(tf, ["FastAPI · PostgreSQL · Qdrant", "Keycloak · Celery · SQLAlchemy",
                 "Ollama / vLLM / TGI · LiveKit", "BM25 · RRF · cross-encoder"], size=11.5)
    tf = panel_card(s, MX + cw + 0.3, y, cw, h, PURPLE, "Where Agentium engineers beyond")
    bullets(tf, ["In-house orchestration kernel", "Deterministic token budget",
                 "Query-adaptive retrieval fusion", "Evaluation loop · tamper-evident audit"], size=11.5)
    band(s, MX, 4.5, SW - 2 * MX, 1.0, NAVY, "BOUND AND OPTIMISED",
         "The orchestration kernel is built in-house — not LangGraph, CrewAI or AutoGen — and binds proven "
         "open-source engines into one governed, sovereign runtime.")


# ================================================================ S7 — Feature catalog
def s7_catalog():
    s = slide()
    chrome(s, "Feature catalog", "Capabilities at a glance", NAVY)
    groups = [
        (BLUE, "Agent Builder", "Six-step builder, visual flows, versioning"),
        (PURPLE, "Orchestration", "Modes, checkpoints, replay, skill registry"),
        (GREEN, "Agent Evaluation", "Scoring, thresholds, review, feedback"),
        (AMBER, "Security & Identity", "RBAC + ABAC, agent identity, guardrails"),
        (NAVY, "Observability & Audit", "Tool-call ledger, tracing, decision trail"),
        (BLUE, "Model & Token", "On-prem serving, budget, adaptive fusion"),
        (GREEN, "Document Intelligence", "Ingestion, OCR, hybrid retrieval, isolation"),
        (PURPLE, "Connectors & Knowledge", "SharePoint, secure deposit, capture"),
        (AMBER, "Cockpit & Roles", "Hypervisor, steering, four operating roles"),
    ]
    cols, gap = 3, 0.28
    cw = (SW - 2 * MX - (cols - 1) * gap) / cols
    ch, vgap = 1.32, 0.22
    y0 = 1.66
    for i, (c, t, body) in enumerate(groups):
        r, col = divmod(i, cols)
        x = MX + col * (cw + gap)
        y = y0 + r * (ch + vgap)
        rect(s, x, y, cw, ch, WHITE, line_color=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        rect(s, x, y + 0.1, 0.05, ch - 0.2, c)
        tf = box(s, x + 0.2, y + 0.16, cw - 0.34, ch - 0.3)
        para(tf, t, size=12, color=c, bold=True, first=True)
        para(tf, body, size=10, color=MUT, before=4, spacing=1.05)


# ================================================================ S8 — Sovereign deployment
def s8_deployment():
    s = slide()
    chrome(s, "Sovereign deployment", "The entire stack inside the boundary — models included", AMBER)
    pic(s, "d5_deployment.png", 7.7, 1.7, x=MX)
    bx = MX + 7.7 + 0.4
    tf = panel_card(s, bx, 1.7, SW - MX - bx, 3.1, AMBER, "Reference baseline")
    bullets(tf, [("On-prem", "7-9 GPUs · ~320-384 GB VRAM"),
                 ("Compute", "~272 cores · ~1.85 TB RAM"),
                 ("Footprint", "~7 servers · 25 GbE"),
                 ("Air-gapped", "signed offline metering, no phone-home")], size=11)


# ================================================================ S9 — Alignment to PIH needs
def s9_alignment():
    s = slide()
    chrome(s, "Alignment to PIH", "Five needs, mapped to capability", NAVY)
    rows = [
        ("Sovereignty", NAVY, "On-prem serving & local data plane", "Inference inside the boundary; workspace-scoped vector, relational & object stores."),
        ("Robustness", BLUE, "Stateful runs · replay · resubmission", "Immutable snapshots, checkpoints, parent-child lineage, replay with overrides, idempotent skills."),
        ("Security", PURPLE, "RBAC + ABAC · agent identity · guardrails", "Keycloak OIDC, five roles, attribute conditions, distinct agent principals, tenant isolation."),
        ("Observability & Scale", GREEN, "Tool-call ledger · audit · async scale", "Skill-invocation ledger, pipeline tracing, exportable audit, decision trail; async workers."),
        ("Token Performance", AMBER, "Deterministic budget & adaptive fusion", "Ratio-based budget with caps, dense/sparse fusion, budgeted reranking, multi-provider routing."),
    ]
    x, y = MX, 1.62
    w = SW - 2 * MX
    tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(x), Inches(y), Inches(w), Inches(4.9)).table
    tbl.columns[0].width = Inches(2.45)
    tbl.columns[1].width = Inches(3.7)
    tbl.columns[2].width = Inches(w - 2.45 - 3.7)
    for j, htxt in enumerate(["PIH need", "Agentium capability", "How it is delivered"]):
        style_cell(tbl.cell(0, j), htxt, size=11, color=WHITE, bold=True, fill=NAVY)
    for i, (need, c, cap, how) in enumerate(rows, start=1):
        style_cell(tbl.cell(i, 0), need, size=11, color=c, bold=True, fill=TINTBLUE if i % 2 else WHITE)
        style_cell(tbl.cell(i, 1), cap, size=10, color=INK, fill=PANEL if i % 2 == 0 else WHITE)
        style_cell(tbl.cell(i, 2), how, size=9.5, color=MUT, fill=PANEL if i % 2 == 0 else WHITE)


# ================================================================ S10 — Positioning
def s10_positioning():
    s = slide()
    chrome(s, "Positioning", "Owned capability, not rented intelligence", PURPLE)
    cw = (SW - 2 * MX - 0.3) / 2
    y, h = 1.7, 3.0
    tf = panel_card(s, MX, y, cw, h, NAVY, "Agentium — owned capability")
    bullets(tf, ["Owns the intelligence: models on-prem",
                 "Sovereign, replicable deployment",
                 "Proof — tamper-evident audit & evaluation",
                 "Deterministic token budget",
                 "Arabic / MENA-ready"], size=11.5)
    tf = panel_card(s, MX + cw + 0.3, y, cw, h, MUT, "Rented-intelligence platforms")
    bullets(tf, ["Depend on third-party hosted models",
                 "Sovereignty promised, not enforced",
                 "Logs, not provable lineage",
                 "Opaque, usage-priced token spend",
                 "Generic language coverage"], size=11.5, color=MUT)
    band(s, MX, 4.95, SW - 2 * MX, 1.05, PANEL,
         "QUESTION WORTH PUTTING TO ANY AGENT PLATFORM",
         "“Who owns the model, the action ledger and the governance plane — and can you prove it offline?”",
         title_color=PURPLE, body_color=INK)


# ================================================================ S11 — Value model
def s11_value():
    s = slide()
    chrome(s, "Value model", "Five dimensions of value", GREEN)
    pic(s, "d11_value_model.png", 7.4, 1.72, x=MX)
    bx = MX + 7.4 + 0.4
    tf = panel_card(s, bx, 1.72, SW - MX - bx, 3.05, GREEN, "What compounds over time")
    bullets(tf, [("Platform", "the sovereign operating system"),
                 ("System capacity", "agents & suites in production"),
                 ("Knowledge capacity", "governed enterprise memory"),
                 ("Center of excellence", "local skills & ownership"),
                 ("Sovereign infra", "the controlled boundary")], size=10.5)


# ================================================================ S12 — Close
def s12_close():
    s = slide()
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, 0.06, BLUE)
    tf = box(s, MX, 1.9, SW - 2 * MX, 0.4)
    para(tf, "WHY AGENTIUM FOR PIH", size=12, color=BLUE, bold=True, mono=True, first=True)
    tf = box(s, MX, 2.45, SW - 2 * MX, 1.0)
    para(tf, "The constraint becomes the advantage", size=32, color=WHITE, bold=True, first=True)
    tf = box(s, MX, 3.8, SW - 2 * MX, 2.2)
    for head, body in [("Whole-stack control", "models, action ledger and governance — inside your boundary"),
                       ("Proof, not logs", "tamper-evident audit, replayable runs, scored evaluations"),
                       ("Sovereignty enforced", "deployed, replicable and operated by your own teams")]:
        p = para(tf, "▪ ", size=14, color=BLUE, before=10, spacing=1.1)
        add_run(p, head + " — ", size=14, color=WHITE, bold=True)
        add_run(p, body, size=14, color=LIGHTBLUE)
    tf = box(s, MX, SH - 0.7, SW - 2 * MX, 0.3)
    para(tf, "Datategy   ·   Agentium   ·   June 2026", size=11, color=RGBColor(0x8F, 0xA8, 0xC7), mono=True, first=True)


def build():
    dg.render_all()  # ensure diagrams exist in assets/diagrams
    s1_cover(); s2_exec(); s3_platform(); s4_model(); s5_orchestration()
    s6_foundation(); s7_catalog(); s8_deployment(); s9_alignment()
    s10_positioning(); s11_value(); s12_close()
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print("saved:", os.path.relpath(OUT, HERE), "·", len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
