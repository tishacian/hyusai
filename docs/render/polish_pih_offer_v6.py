# -*- coding: utf-8 -*-
"""Polish the PIH offer v4 -> v6 (rebuilt clean from v4; v5 was wrong, discarded).

v5 bug reports (2026-07-08) and root causes:
  1. "Double puce" — the "plain list item" paragraphs (Authentication handling,
     Agent definitions., etc.) already carry NATIVE Word/Google-Docs list
     numbering (<w:numPr><w:numId>) which auto-renders a bullet/number glyph.
     This is invisible via python-docx's `paragraph_format` (only exposes direct
     <w:ind>, not <w:numPr>), so v5 wrongly assumed these had no marker and
     prepended a literal "•"/"1." text — creating a duplicate glyph. FIX: do not
     touch these paragraphs at all; the platform already renders their bullets.
  2. "Texte hors tableau" — pre-existing bug from the original iterate_pih_offer.py
     (v2): `table.add_row()` for the ManageEngine SDP connector row produced a
     bare row with none of the sibling rows' cell borders/margins/vAlign/font
     size/row height, so it renders unstyled, floating below the bordered table.
     FIX: copy the tcPr/rPr/pPr/trPr of a sibling row onto that row.
  3. "Mauvaise gestion des pages" — orphaned headings: 15 of the original
     document's Heading paragraphs carry an explicit keepNext=False override
     (a Google-Docs export quirk), so nothing stops a heading from being
     stranded at the bottom of a page, separated from its own following
     paragraph. This was latent in v4 and became visible once the (correct)
     spacing fix below changed pagination. FIX: force keepNext + keepLines on
     every heading in the document.

Fixes applied here (formatting only — no text is changed, removed or reworded):
  A. Same universal spacing fix as before: body ("normal"-style) paragraphs
     missing direct spacing get 12pt before, 1.15 line spacing, and 12pt after
     (0pt after + hanging indent for "•"-prefixed bullet paragraphs already
     written as literal bullets, e.g. "Our proposed solution" bullets).
  B. Table row fix for the ManageEngine SDP connector row.
  C. keepNext/keepLines forced true on every heading, to stop orphaned headings.

Input : docs/pih/Service Help Desk - Technical and Commercial offer - v4 Datategy.docx
Output: docs/pih/Service Help Desk - Technical and Commercial offer - v6 Datategy.docx
"""
import copy
import os
from docx import Document
from docx.shared import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
SRC = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v4 Datategy.docx")
OUT = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v6 Datategy.docx")

SPACE = Emu(152400)      # 12pt, matches the original document's own paragraph spacing
ZERO = Emu(0)
LEFT_INDENT = Emu(457200)
FIRST_LINE = Emu(-228600)
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

doc = Document(SRC)
BEFORE_TEXT = [p.text for p in doc.paragraphs]

# ============================================================ A. universal spacing fix
fixed = 0
for p in doc.paragraphs:
    if not p.style or p.style.name != "normal":
        continue
    if not p.text.strip():
        continue
    pf = p.paragraph_format
    if pf.space_before is not None:
        continue  # already explicitly formatted (original content) — leave untouched
    # native-numbered paragraphs (numPr) already have their own indent/spacing baked
    # in by the platform; do not touch them even if paragraph_format looked "empty".
    if p._p.find(f"{W_NS}pPr/{W_NS}numPr") is not None:
        continue
    is_literal_bullet = p.text.lstrip().startswith("•")
    pf.space_before = SPACE
    pf.line_spacing = 1.15
    if is_literal_bullet:
        pf.space_after = ZERO
        pf.left_indent = LEFT_INDENT
        pf.first_line_indent = FIRST_LINE
    else:
        pf.space_after = SPACE
    fixed += 1
print(f"spacing fixed on {fixed} paragraphs")

# ============================================================ B. fix the orphan table row
t2 = None
for t in doc.tables:
    if t.cell(0, 0).text.strip() == "Connector":
        t2 = t
        break
assert t2 is not None, "connector table not found"
last_text = t2.rows[-1].cells[0].text.strip()
assert last_text == "ManageEngine SDP connector", f"unexpected last row: {last_text!r}"

good_row = t2.rows[-2]._tr    # a correctly-styled sibling row
bad_row = t2.rows[-1]._tr

good_trpr = good_row.find(f"{W_NS}trPr")
if good_trpr is not None and bad_row.find(f"{W_NS}trPr") is None:
    bad_row.insert(0, copy.deepcopy(good_trpr))

good_tcs = good_row.findall(f"{W_NS}tc")
bad_tcs = bad_row.findall(f"{W_NS}tc")
for good_tc, bad_tc in zip(good_tcs, bad_tcs):
    good_tcpr = good_tc.find(f"{W_NS}tcPr")
    bad_tcpr = bad_tc.find(f"{W_NS}tcPr")
    if good_tcpr is None:
        continue
    new_tcpr = copy.deepcopy(good_tcpr)
    # keep the bad cell's own column width if it had one, else inherit the sibling's
    old_w = bad_tcpr.find(f"{W_NS}tcW") if bad_tcpr is not None else None
    if old_w is not None:
        w_in_new = new_tcpr.find(f"{W_NS}tcW")
        if w_in_new is not None:
            new_tcpr.replace(w_in_new, old_w)
    if bad_tcpr is not None:
        bad_tc.replace(bad_tcpr, new_tcpr)
    else:
        bad_tc.insert(0, new_tcpr)
    # match paragraph-level formatting (font size, line spacing, widow control)
    good_p = good_tc.find(f"{W_NS}p")
    bad_p = bad_tc.find(f"{W_NS}p")
    good_ppr = good_p.find(f"{W_NS}pPr") if good_p is not None else None
    if good_ppr is not None and bad_p is not None:
        bad_ppr = bad_p.find(f"{W_NS}pPr")
        new_ppr = copy.deepcopy(good_ppr)
        if bad_ppr is not None:
            bad_p.replace(bad_ppr, new_ppr)
        else:
            bad_p.insert(0, new_ppr)
        for r in bad_p.findall(f"{W_NS}r"):
            good_r = good_p.find(f"{W_NS}r")
            good_rpr = good_r.find(f"{W_NS}rPr") if good_r is not None else None
            if good_rpr is not None:
                r_rpr = r.find(f"{W_NS}rPr")
                new_rpr = copy.deepcopy(good_rpr)
                if r_rpr is not None:
                    r.replace(r_rpr, new_rpr)
                else:
                    r.insert(0, new_rpr)
print("table row fixed: ManageEngine SDP connector")

# ============================================================ C. no orphaned headings
n_headings = 0
for p in doc.paragraphs:
    if p.style and p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        n_headings += 1
print(f"keepNext/keepLines forced on {n_headings} headings")

doc.save(OUT)

# ============================================================ sanity: no wording lost
d2 = Document(OUT)
AFTER_TEXT = [p.text for p in d2.paragraphs]
assert len(BEFORE_TEXT) == len(AFTER_TEXT), "paragraph count changed!"
mismatches = [i for i, (b, a) in enumerate(zip(BEFORE_TEXT, AFTER_TEXT)) if b != a]
print(f"wording mismatches: {len(mismatches)} (must be 0)")
for i in mismatches[:10]:
    print(" ", i, BEFORE_TEXT[i][:60], "->", AFTER_TEXT[i][:60])

t2b = None
for t in d2.tables:
    if t.cell(0, 0).text.strip() == "Connector":
        t2b = t
        break
print("connector table rows:", len(t2b.rows), "last row:", [c.text[:30] for c in t2b.rows[-1].cells])
print("saved:", OUT)
