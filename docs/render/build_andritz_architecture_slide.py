#!/usr/bin/env python3
"""Executive architecture deck — workspace Andritz (Agentium / OmniRAG).

Outputs:
  - docs/render/out/ANDRITZ-architecture-executive.pptx
    · Slide 1 — Vue applicative (technos & flux métier)
    · Slide 2 — Vue plateforme on-prem complète (infra OVH + technos)
"""

from __future__ import annotations

import os
from dataclasses import dataclass

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SW, SH = 13.333, 7.5
MX = 0.45
GAP = 0.22
F = "Arial"

INK = RGBColor(0x15, 0x1A, 0x23)
MUT = RGBColor(0x59, 0x63, 0x71)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x2D, 0x6C, 0xDF)
TEAL = RGBColor(0x0E, 0x8F, 0x62)
PURPLE = RGBColor(0x6D, 0x28, 0xD9)
AMBER = RGBColor(0xB4, 0x53, 0x09)
SLATE = RGBColor(0xE8, 0xEC, 0xF1)
LINE = RGBColor(0xC5, 0xCD, 0xD8)
AI_FILL = RGBColor(0xEE, 0xF4, 0xFF)
INFRA_FILL = RGBColor(0xF0, 0xF7, 0xF4)
OVH_GREEN = RGBColor(0x0E, 0x6B, 0x4F)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "out")
OUT_PATH = os.path.join(OUT_DIR, "ANDRITZ-architecture-executive.pptx")

COL_W = (SW - 2 * MX - 2 * GAP) / 3
HDR_H = 0.36
SUB_H = 0.24
ITEM_H = 0.28
PAD = 0.10
ROW_GAP = 0.18


@dataclass
class Block:
    x: float
    y: float
    w: float
    h: float

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def right(self) -> float:
        return self.x + self.w

    @property
    def bottom(self) -> float:
        return self.y + self.h


def block_height(n_items: int) -> float:
    return HDR_H + SUB_H + PAD + n_items * ITEM_H + PAD


def rect(slide, x, y, w, h, fill, line_color=None, dashed: bool = False):
    sp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line_color:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(1.0)
        if dashed:
            sp.line.dash_style = 2
    else:
        sp.line.fill.background()
    try:
        sp.adjustments[0] = 0.06
    except Exception:
        pass
    sp.shadow.inherit = False
    return sp


def arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = LINE
    conn.line.width = Pt(1.4)
    return conn


def arrow_h(slide, a: Block, b: Block):
    arrow(slide, a.right, a.cy, b.x, b.cy)


def arrow_pt(slide, x1, y1, x2, y2):
    arrow(slide, x1, y1, x2, y2)


def slide_header(slide, title: str, subtitle: str):
    rect(slide, 0, 0, SW, 0.12, NAVY)
    hdr = slide.shapes.add_textbox(Inches(MX), Inches(0.18), Inches(SW - 2 * MX), Inches(0.48))
    hp = hdr.text_frame.paragraphs[0]
    hr = hp.add_run()
    hr.text = title
    hr.font.name = F
    hr.font.size = Pt(20)
    hr.font.bold = True
    hr.font.color.rgb = INK

    sub = slide.shapes.add_textbox(Inches(MX), Inches(0.64), Inches(SW - 2 * MX), Inches(0.30))
    sp = sub.text_frame.paragraphs[0]
    sr = sp.add_run()
    sr.text = subtitle
    sr.font.name = F
    sr.font.size = Pt(9.5)
    sr.font.color.rgb = MUT


def slide_footer(slide, y: float, lines: list[str]):
    foot = slide.shapes.add_textbox(Inches(MX), Inches(y), Inches(SW - 2 * MX - 1.9), Inches(0.58))
    ftf = foot.text_frame
    ftf.word_wrap = True
    for i, line in enumerate(lines):
        p = ftf.paragraphs[0] if i == 0 else ftf.add_paragraph()
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line
        r.font.name = F
        r.font.size = Pt(8.5)
        r.font.color.rgb = MUT

    brand = slide.shapes.add_textbox(Inches(SW - MX - 1.8), Inches(y), Inches(1.8), Inches(0.25))
    bp = brand.text_frame.paragraphs[0]
    bp.alignment = PP_ALIGN.RIGHT
    br = bp.add_run()
    br.text = "Datategy · Agentium"
    br.font.name = F
    br.font.size = Pt(8)
    br.font.bold = True
    br.font.color.rgb = NAVY


def superblock(slide, x, y, w, title, subtitle, fill, accent, items) -> Block:
    h = block_height(len(items))
    rect(slide, x, y, w, h, fill, line_color=accent)
    rect(slide, x, y, w, HDR_H, accent)

    title_tb = slide.shapes.add_textbox(
        Inches(x + 0.10), Inches(y + 0.07), Inches(w - 0.20), Inches(HDR_H - 0.10)
    )
    ttf = title_tb.text_frame
    ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tr = ttf.paragraphs[0].add_run()
    tr.text = title
    tr.font.name = F
    tr.font.size = Pt(10.5)
    tr.font.bold = True
    tr.font.color.rgb = WHITE

    body_tb = slide.shapes.add_textbox(
        Inches(x + 0.10), Inches(y + HDR_H + 0.04), Inches(w - 0.20), Inches(h - HDR_H - 0.08)
    )
    btf = body_tb.text_frame
    btf.word_wrap = True
    btf.vertical_anchor = MSO_ANCHOR.TOP
    btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = Pt(0)

    sp = btf.paragraphs[0]
    sp.space_after = Pt(3)
    sr = sp.add_run()
    sr.text = subtitle
    sr.font.name = F
    sr.font.size = Pt(7.5)
    sr.font.color.rgb = MUT

    for item in items:
        bp = btf.add_paragraph()
        bp.space_before = Pt(1)
        bp.space_after = Pt(1)
        bp.line_spacing = 1.05
        br = bp.add_run()
        br.text = f"• {item}"
        br.font.name = F
        br.font.size = Pt(8)
        br.font.color.rgb = INK

    return Block(x, y, w, h)


def build_application_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(
        slide,
        "Architecture Agentium — workspace Andritz",
        "Slide 1/2 · Vue applicative · technos & flux métier · inférence Qwen 3.5 on-prem (zéro API cloud IA)",
    )

    y1 = 1.00
    b_access = superblock(
        slide, MX, y1, COL_W,
        "Accès & sécurité", "Périmètre utilisateur", SLATE, NAVY,
        ["Keycloak (OIDC / SSO)", "nginx + TLS", "IAM Andritz activé"],
    )
    b_ui = superblock(
        slide, MX + COL_W + GAP, y1, COL_W,
        "Interface métier", "Profil business Andritz", SLATE, BLUE,
        ["Angular 20 (Chat Recherche)", "Knowledge Capture (voix)", "Admin : gouvernance & cockpit"],
    )
    b_orch = superblock(
        slide, MX + 2 * (COL_W + GAP), y1, COL_W,
        "Orchestration", "Services applicatifs", SLATE, TEAL,
        ["FastAPI + Uvicorn (API, SSE chat)", "Celery + RabbitMQ (jobs async)", "PostgreSQL (sessions, métier)"],
    )

    y2 = max(b_access.bottom, b_ui.bottom, b_orch.bottom) + ROW_GAP
    ai_w = 2 * COL_W + GAP
    b_ai = superblock(
        slide, MX, y2, ai_w,
        "Inférence IA — 100 % on-prem", "vLLM / TEI intra-cluster · zéro API externe", AI_FILL, PURPLE,
        [
            "Chat RAG · Qwen3.5-122B-A10B (vLLM)",
            "Capture oracle · Qwen3.5-35B-A3B (vLLM)",
            "Embeddings · BGE-M3 (TEI) · Voix · Whisper",
            "Rerank · cross-encoder ms-marco",
        ],
    )
    b_rag = superblock(
        slide, MX + ai_w + GAP, y2, COL_W,
        "Intelligence documentaire", "RAG industriel SPL", SLATE, AMBER,
        [
            "Recherche hybride Qdrant (dense + sparse)",
            "Scope « Contexte Andritz SPL »",
            "Fiches expertes auto-publiées",
            "Grounding & traçabilité retrieval",
        ],
    )

    y3 = max(b_ai.bottom, b_rag.bottom) + ROW_GAP
    b_store = superblock(
        slide, MX, y3, COL_W,
        "Stockage & données", "Couche persistance", SLATE, NAVY,
        ["PostgreSQL (métier, ledger, sessions)", "Qdrant (vecteurs SPL)", "MinIO (objets & fichiers)"],
    )
    b_ing = superblock(
        slide, MX + COL_W + GAP, y3, COL_W,
        "Ingestion documentaire", "Entrée corpus SPL", SLATE, TEAL,
        ["SFTP Secure Deposit (Andritz)", "OCR Tesseract · PDF / Excel / DOCX", "Celery → indexation Qdrant"],
    )
    b_obs = superblock(
        slide, MX + 2 * (COL_W + GAP), y3, COL_W,
        "Observabilité", "Pilotage & qualité", SLATE, BLUE,
        ["structlog · métriques API", "Historique chat / capture (admin)", "Golden tests SPL · réconciliation PG↔Qdrant"],
    )

    arrow_h(slide, b_access, b_ui)
    arrow_h(slide, b_ui, b_orch)
    arrow_pt(slide, b_orch.cx - 0.35, b_orch.bottom, b_ai.cx, b_ai.y)
    arrow_pt(slide, b_orch.cx + 0.35, b_orch.bottom, b_rag.cx, b_rag.y)
    arrow_h(slide, b_ai, b_rag)
    arrow_h(slide, b_ing, b_store)
    arrow_pt(slide, b_ing.cx, b_ing.y, b_rag.cx - 0.25, b_rag.bottom)
    arrow_pt(slide, b_store.cx, b_store.y, b_rag.cx - 0.55, b_rag.bottom)

    foot_y = max(b_store.bottom, b_ing.bottom, b_obs.bottom) + 0.16
    slide_footer(
        slide,
        foot_y,
        [
            "Flux : SFTP → ingestion → MinIO/PostgreSQL → BGE-M3 → Qdrant → RAG ← Qwen3.5-122B · capture ← Qwen3.5-35B-A3B",
            "Utilisateur → Angular → FastAPI → chat / capture · voix → Whisper → oracle Qwen",
        ],
    )


def build_platform_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_header(
        slide,
        "Plateforme on-prem complète — Andritz",
        "Slide 2/2 · Infrastructure OVH France (gauche) + stack technos Agentium (droite) · Qwen 3.5 · zéro API IA externe",
    )

    y0 = 1.00
    half_w = (SW - 2 * MX - GAP) / 2
    lx, rx = MX, MX + half_w + GAP

    # Enveloppe OVH (colonne gauche)
    rect(slide, lx - 0.05, y0 - 0.05, half_w + 0.10, 4.05, RGBColor(0xFA, 0xFC, 0xFB), OVH_GREEN, dashed=True)
    env_lbl = slide.shapes.add_textbox(Inches(lx), Inches(y0 - 0.04), Inches(half_w), Inches(0.20))
    el = env_lbl.text_frame.paragraphs[0].add_run()
    el.text = "Infrastructure on-prem · OVHcloud France"
    el.font.name = F
    el.font.size = Pt(8)
    el.font.bold = True
    el.font.color.rgb = OVH_GREEN

    app_lbl = slide.shapes.add_textbox(Inches(rx), Inches(y0 - 0.04), Inches(half_w), Inches(0.20))
    al = app_lbl.text_frame.paragraphs[0].add_run()
    al.text = "Stack applicatif & données · workspace Andritz"
    al.font.name = F
    al.font.size = Pt(8)
    al.font.bold = True
    al.font.color.rgb = BLUE

    # ── Colonne gauche : infra ──
    ly = y0 + 0.12
    b_infra = superblock(
        slide, lx, ly, half_w,
        "Cloud & plateforme OVH", "Réseau · K8s · sécurité · France UE",
        INFRA_FILL, OVH_GREEN,
        [
            "OVH LB · Istio Ingress · cert-manager TLS · VPC privé",
            "OVH MKS · workers CPU (FastAPI, Celery, Angular, RabbitMQ)",
            "Cinder PV · MinIO registre · Vault · Keycloak · snapshots",
        ],
    )
    ly = b_infra.bottom + ROW_GAP
    b_gpu = superblock(
        slide, lx, ly, half_w,
        "Tier GPU — inférence souveraine", "LLMaaS / EMBaaS · vLLM / TEI · Qwen 3.5",
        AI_FILL, PURPLE,
        [
            "2× L40S 48G · Qwen3.5-122B-A10B · vLLM TP=2 (chat RAG)",
            "H100 80G · Qwen3.5-35B-A3B · vLLM (capture oracle)",
            "4× L4 24G · BGE-M3 (TEI) · Whisper · ms-marco rerank",
        ],
    )

    # ── Colonne droite : applicatif ──
    ry = y0 + 0.12
    b_app = superblock(
        slide, rx, ry, half_w,
        "Stack Agentium Andritz", "Interface · orchestration · RAG",
        SLATE, BLUE,
        [
            "Keycloak · IAM · Angular 20 · Chat / Capture / Admin",
            "FastAPI + Uvicorn (SSE) · Celery · recherche hybride Qdrant",
            "Scope SPL · fiches expertes · grounding & traçabilité",
        ],
    )
    ry = b_app.bottom + ROW_GAP
    b_data = superblock(
        slide, rx, ry, half_w,
        "Données & pipeline", "Persistance · ingestion · observabilité",
        SLATE, AMBER,
        [
            "PostgreSQL · Qdrant · MinIO · SFTP Secure Deposit",
            "Tesseract · PDF/Excel/DOCX → BGE-M3 → indexation Qdrant",
            "structlog · golden tests · réconciliation PG↔Qdrant",
        ],
    )

    left_bottom = b_gpu.bottom
    right_bottom = b_data.bottom

    arrow_pt(slide, b_infra.right, b_infra.cy, rx, b_app.cy)
    arrow_pt(slide, b_gpu.right, b_gpu.cy, rx, b_data.cy)
    arrow_pt(slide, b_app.cx, b_app.bottom, b_data.cx, b_data.y)

    foot_y = max(left_bottom, right_bottom) + 0.16
    slide_footer(
        slide,
        foot_y,
        [
            "Gauche : hébergement & inférence 100 % intra-cluster OVH · Droite : technos Agentium Andritz · aucun prompt/completion vers API externe",
            "Production cible OVH MKS (cf. fiche hosting Andritz) · démo actuelle : VM Docker Compose — même stack logique, sizing réduit",
        ],
    )


def build():
    os.makedirs(OUT_DIR, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    build_application_slide(prs)
    build_platform_slide(prs)
    prs.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} (2 slides)")


if __name__ == "__main__":
    build()
