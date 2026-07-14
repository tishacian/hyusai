# -*- coding: utf-8 -*-
"""Iterate on the NEGOTIATED PIH offer v6 -> v7: align prose on the negotiated table.

The client-validated v6 (from negotiation) has an updated pricing TABLE
(10k / 26k / 14k / 10k · 50k / 60k · run 25k / 80k · 8 months) but the PROSE
still carries the pre-negotiation figures (25k initial, 80k waves, 105k/120k
totals, run 20k). This script aligns every prose figure on the negotiated
numbers, and adjusts the two value claims that the new run levels no longer
support ("roughly half the current IT licence", "TCO below current combined
spend") to keep the offer factually defensible.

Input : docs/pih/Service Help Desk - Technical and Commercial offer - v6 NEGOTIATED Datategy.docx
Output: docs/pih/Service Help Desk - Technical and Commercial offer - v7 Datategy.docx
"""
import os
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
SRC = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v6 NEGOTIATED Datategy.docx")
OUT = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v7 Datategy.docx")

doc = Document(SRC)


def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise SystemExit(f"ANCHOR NOT FOUND: {prefix!r}")


def set_text(p, text, bold_lead=None):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


# ---- exec summary: staged commitment figure
p_dv = find("We propose a phased engagement")
set_text(p_dv,
    "We propose a phased engagement: a foundation phase and a pilot wave with client-confirmed "
    "baselines, then wave-based migration and enhancement of all 39 use cases and scale-out across "
    "entities — each wave closing with quality gates and ≥30 days of hypercare. The commercial "
    "commitment is staged the same way: an initial engagement of approximately USD 10k puts the "
    "pilot in production, and each subsequent wave is engaged only on the measured results of the "
    "previous one. At your volumes, a 30–40% deflection/automation rate on the residual ~70,000 "
    "tickets represents 5,000–7,000 analyst-hours per year returned to higher-value work — measured "
    "live in the platform's value cockpit, so each wave justifies the next. The run model is a flat "
    "editor licence with operations internalised in your team.")

# ---- commercial principles
p_pri = find("The commercial model follows three principles")
set_text(p_pri,
    "The commercial model follows three principles: (1) no per-employee licensing — costs do not "
    "scale with your 12,200 users but with the value delivered; (2) a staged commitment — the only "
    "amount engaged at signature is the initial engagement (≈ USD 10k), which puts the pilot in "
    "production; the following waves are pre-priced and engaged wave by wave on the pilot's "
    "measured results; (3) a flat, predictable run. Two commercial options are proposed; amounts "
    "in this section are finalised at contract in the separate Price Schedule.")

# ---- Option 1 build & run
p_build = find("Build (staged commitment) —")
set_text(p_build,
    "an initial engagement of approximately USD 10k covers Phase 0-lite (environments, security "
    "model, ManageEngine + Entra ID + Teams connectivity, knowledge ingestion, golden sets) and a "
    "pilot wave proving one exemplar of each of the four foundational patterns that cover the 39 "
    "use cases: conversational deflection live on a representative population, proactive "
    "operations & RCA over ServiceDesk Plus history, identity & access lifecycle in "
    "parallel-run/simulate, and one approval-orchestrated flow end-to-end. Waves 2–3 — migration "
    "and enhancement of all 39 use cases, privileged writes under mandates, the remaining "
    "connectors (Egnyte, SolarWinds ARM, firewalls) and multi-entity scale-out — are pre-priced "
    "(approximately USD 40k) and engaged wave by wave on measured results, bringing the full build "
    "to approximately USD 50k over ≈ 8 months. The per-use-case day build-up is provided in the "
    "Build Effort Detail annex.",
    bold_lead="Build (staged commitment) — ")

p_run = find("Run —")
set_text(p_run,
    "annual platform licence and editor support (product updates, security patches, level-3 "
    "support) — USD 25k per year for the IT scope, below the current conversational platform's IT "
    "licence. Day-to-day operations are internalised in PIH's IT operations team, trained through "
    "the Datategy Academy programme; infrastructure (Azure tenant or on-premise) remains under "
    "PIH's control and billing.",
    bold_lead="Run — ")

# ---- Option 2
p_opt2 = find("Option 2 adds an Employee Services module")
set_text(p_opt2,
    "Option 2 adds an Employee Services module: employee Q&A and self-service over HR policies and "
    "services (leave, documents, onboarding questions), delivered through the same channels, the "
    "same knowledge machinery and the same SuccessFactors connector already required for the IT "
    "identity-lifecycle flows. HR systems remain the systems of record — the module is a deflection "
    "and self-service layer, not HR process execution — so the marginal build cost is limited "
    "(approximately USD 10k on top of Option 1, delivered with Waves 2–3, i.e. a full build of "
    "approximately USD 60k). The initial engagement is identical (≈ USD 10k).")

p_runfee = find("The run fee is unchanged")
set_text(p_runfee,
    "The Option 2 run fee is USD 80k per year and covers the full IT + Employee Services perimeter: "
    "platform licence, editor support and golden-set regression across both scopes, with day-to-day "
    "operations internalised in PIH's team for IT and employee services alike. Option 2 remains the "
    "recommended path: one front door for employees across IT and HR questions, the full "
    "replacement of the current conversational stack, and the conversion of rented workflows into "
    "an owned, governed, agentic platform whose returned analyst-hours are measured live in the "
    "value cockpit.")

doc.save(OUT)

# sanity
d2 = Document(OUT)
text = "\n".join(p.text for p in d2.paragraphs)
for pat, want in [("USD 105k", 0), ("USD 120k", 0), ("USD 25k covers", 0), ("USD 20k per year", 0),
                  ("USD 10k", None), ("USD 50k", None), ("USD 60k", None), ("USD 25k per year", None),
                  ("USD 80k per year", None)]:
    n = text.count(pat)
    status = "OK" if (want is None and n > 0) or (want == 0 and n == 0) else "FAIL"
    print(f"{status}  {pat!r}: {n}")
print("saved:", OUT)
