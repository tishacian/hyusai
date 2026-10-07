#!/usr/bin/env python3
"""Build FR/EN Word briefings and PowerPoint decks from the Hypervisor docs.

    python3 docs/render/build_hypervisor_decision_pack.py

Reads the decision briefing, the adoption deck, the Hypervisor user
stories and the COMEX MVP note (FR + EN).

Writes Office twins next to those notes (not under docs/render/out,
which is gitignored).
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt as Ppt

DOCS = Path(__file__).resolve().parents[1]

INK = RGBColor(0x15, 0x1A, 0x23)
NAVY = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x59, 0x63, 0x71)
TEAL = RGBColor(0x0E, 0x6B, 0x6A)

P_INK = PptColor(0x15, 0x1A, 0x23)
P_NAVY = PptColor(0x0B, 0x25, 0x45)
P_MUTE = PptColor(0x59, 0x63, 0x71)
P_LINE = PptColor(0xD7, 0xDC, 0xE3)
P_PANEL = PptColor(0xF3, 0xF5, 0xF8)
P_WHITE = PptColor(0xFF, 0xFF, 0xFF)
P_TEAL = PptColor(0x0E, 0x6B, 0x6A)

SW, SH, MX = 13.333, 7.5, 0.58
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def strip_md(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"(?<!\w)\*(?!\*)(.+?)(?<!\*)\*(?!\w)", r"\1", text)
    return text.replace("`", "")


def iter_inlines(text: str):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            yield "plain", text[pos : m.start()]
        token = m.group(0)
        if token.startswith("**"):
            yield "bold", token[2:-2]
        elif token.startswith("`"):
            yield "code", token[1:-1]
        elif token.startswith("["):
            yield "plain", token[1 : token.index("]")]
        else:
            yield "italic", token[1:-1]
        pos = m.end()
    if pos < len(text):
        yield "plain", text[pos:]


def set_run_font(run, *, size=11, bold=False, italic=False, color=None, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inlines(paragraph, text, *, size=11, color=INK):
    for kind, chunk in iter_inlines(text):
        run = paragraph.add_run(chunk)
        set_run_font(
            run,
            size=size,
            bold=(kind == "bold"),
            italic=(kind == "italic"),
            color=TEAL if kind == "code" else color,
            name="Consolas" if kind == "code" else "Calibri",
        )


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in raw):
            rows.append(raw)
        i += 1
    return rows, i


def parse_code(lines: list[str], start: int) -> tuple[str, int]:
    i = start + 1
    body = []
    while i < len(lines) and not lines[i].startswith("```"):
        body.append(lines[i])
        i += 1
    return "\n".join(body), i + 1


def parse_blocks(md: str) -> list[tuple]:
    lines = md.splitlines()
    if lines and lines[0].strip() == "---":
        i = 1
        while i < len(lines) and lines[i].strip() != "---":
            i += 1
        lines = lines[i + 1 :]
    blocks: list[tuple] = []
    i = 0
    para: list[str] = []

    def flush():
        text = " ".join(x.strip() for x in para if x.strip())
        para.clear()
        if text:
            blocks.append(("p", text))

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush()
            body, i = parse_code(lines, i)
            blocks.append(("code", body))
            continue
        if line.startswith("|"):
            flush()
            rows, i = parse_table(lines, i)
            if rows:
                blocks.append(("table", rows))
            continue
        if line.startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("h", level, line[level:].strip()))
            i += 1
            continue
        if line.startswith("> "):
            flush()
            quote = [line[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("> "):
                quote.append(lines[i][2:])
                i += 1
            blocks.append(("quote", " ".join(q.strip() for q in quote)))
            continue
        if re.match(r"^[-*] ", line):
            flush()
            while i < len(lines) and re.match(r"^[-*] ", lines[i]):
                blocks.append(("li", lines[i][2:].strip()))
                i += 1
            continue
        if re.match(r"^\d+\. ", line):
            flush()
            while i < len(lines) and re.match(r"^\d+\. ", lines[i]):
                blocks.append(("ol", re.sub(r"^\d+\. ", "", lines[i]).strip()))
                i += 1
            continue
        if line.strip() == "":
            flush()
            i += 1
            continue
        if line.strip() == "<br>":
            i += 1
            continue
        para.append(line)
        i += 1
    flush()
    return blocks


def build_docx(md_path: Path, out_path: Path, *, title: str, lang: str) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for key in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if key in doc.styles:
            doc.styles[key].font.color.rgb = NAVY
            doc.styles[key].font.name = "Calibri"

    core = doc.core_properties
    core.title = title
    core.author = "Agentium"
    core.language = "fr-FR" if lang == "fr" else "en-US"

    for block in parse_blocks(md_path.read_text()):
        kind = block[0]
        if kind == "h":
            level, text = block[1], strip_md(block[2])
            heading = doc.add_heading(text, level=min(level, 3))
            for run in heading.runs:
                set_run_font(run, size={1: 22, 2: 16, 3: 13}.get(level, 12), bold=True, color=NAVY)
        elif kind == "p":
            p = doc.add_paragraph()
            add_inlines(p, block[1])
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inlines(p, block[1], color=MUTED)
            for run in p.runs:
                run.italic = True
        elif kind == "li":
            p = doc.add_paragraph(style="List Bullet")
            for run in list(p.runs):
                run._element.getparent().remove(run._element)
            add_inlines(p, block[1])
        elif kind == "ol":
            p = doc.add_paragraph(style="List Number")
            for run in list(p.runs):
                run._element.getparent().remove(run._element)
            add_inlines(p, block[1])
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            run = p.add_run(block[1])
            set_run_font(run, size=9, name="Consolas", color=TEAL)
        elif kind == "table":
            rows = block[1]
            tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
            tbl.style = "Table Grid"
            for r_i, row in enumerate(rows):
                for c_i, cell in enumerate(row):
                    if c_i >= len(tbl.rows[r_i].cells):
                        continue
                    cell_p = tbl.rows[r_i].cells[c_i].paragraphs[0]
                    cell_p.clear()
                    add_inlines(cell_p, cell, size=9, color=INK)
                    if r_i == 0:
                        for run in cell_p.runs:
                            run.bold = True
                            run.font.color.rgb = NAVY
            doc.add_paragraph()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def ppt_rect(slide, x, y, w, h, fill, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Ppt(0.75)
    sp.shadow.inherit = False
    return sp


def ppt_box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


def ppt_para(tf, text, *, size=16, color=P_INK, bold=False, first=False, before=4, align=PP_ALIGN.LEFT, mono=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Ppt(0 if first else before)
    p.space_after = Ppt(2)
    r = p.add_run()
    r.text = text
    r.font.name = "Consolas" if mono else "Calibri"
    r.font.size = Ppt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return p


def add_ppt_inlines(p, text, *, size=16, color=P_INK):
    for kind, chunk in iter_inlines(text):
        r = p.add_run()
        r.text = chunk
        r.font.name = "Consolas" if kind == "code" else "Calibri"
        r.font.size = Ppt(size)
        r.font.bold = kind == "bold"
        r.font.italic = kind == "italic"
        r.font.color.rgb = P_TEAL if kind == "code" else color


def split_marp_slides(md: str) -> list[str]:
    if md.startswith("---"):
        end = md.find("\n---", 3)
        if end != -1:
            md = md[end + 4 :]
    parts = re.split(r"\n---\n", md)
    return [p.strip() for p in parts if p.strip()]


def chrome(slide, kicker: str, title: str, page: int, footer: str):
    ppt_rect(slide, 0, 0, SW, 0.08, P_NAVY)
    ppt_rect(slide, MX, 0.42, 0.28, 0.045, P_TEAL)
    tf = ppt_box(slide, MX + 0.40, 0.30, SW - 2 * MX - 0.4, 0.28)
    ppt_para(tf, kicker.upper(), size=11, color=P_TEAL, bold=True, first=True, mono=True)
    tf = ppt_box(slide, MX, 0.58, SW - 2 * MX, 0.70)
    ppt_para(tf, title, size=26, color=P_INK, bold=True, first=True)
    ppt_rect(slide, MX, 1.32, SW - 2 * MX, 0.012, P_LINE)
    tf = ppt_box(slide, MX, SH - 0.40, 10.4, 0.26)
    ppt_para(tf, footer, size=10, color=P_MUTE, first=True)
    tf = ppt_box(slide, SW - MX - 0.8, SH - 0.40, 0.8, 0.26)
    ppt_para(tf, f"{page:02d}", size=10, color=P_MUTE, first=True, align=PP_ALIGN.RIGHT, mono=True)


def add_table(slide, rows: list[list[str]], x, y, w, h):
    cols = max(len(r) for r in rows)
    table_shape = slide.shapes.add_table(len(rows), cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for c in range(cols):
        table.columns[c].width = Emu(int(Inches(w) / cols))
    for r_i, row in enumerate(rows):
        for c_i in range(cols):
            cell = table.cell(r_i, c_i)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = P_NAVY if r_i == 0 else (P_PANEL if r_i % 2 else P_WHITE)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = strip_md(row[c_i] if c_i < len(row) else "")
            r.font.name = "Calibri"
            r.font.size = Ppt(11 if cols <= 3 else 10)
            r.font.bold = r_i == 0
            r.font.color.rgb = P_WHITE if r_i == 0 else P_INK
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return table_shape


def build_pptx(md_path: Path, out_path: Path, *, kicker: str, footer: str) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    blank = prs.slide_layouts[6]
    slides = split_marp_slides(md_path.read_text())

    for idx, raw in enumerate(slides, start=1):
        blocks = parse_blocks(raw)
        slide = prs.slides.add_slide(blank)
        headings = [b for b in blocks if b[0] == "h"]
        rest = [b for b in blocks if b[0] != "h"]
        is_title = bool(headings) and headings[0][1] == 1

        if is_title:
            ppt_rect(slide, 0, 0, 0.18, SH, P_NAVY)
            ppt_rect(slide, 0, SH - 0.18, SW, 0.18, P_TEAL)
            tf = ppt_box(slide, 0.9, 2.05, 11.6, 0.4)
            ppt_para(tf, kicker.upper(), size=13, color=P_TEAL, bold=True, first=True, mono=True)
            tf = ppt_box(slide, 0.9, 2.5, 11.6, 1.1)
            ppt_para(tf, strip_md(headings[0][2]), size=40, color=P_INK, bold=True, first=True)
            y = 3.75
            for block in rest:
                if block[0] == "p":
                    tf = ppt_box(slide, 0.9, y, 11.6, 0.55)
                    ppt_para(tf, strip_md(block[1]), size=18, color=P_MUTE, first=True)
                    y += 0.55
            tf = ppt_box(slide, 0.9, SH - 0.70, 11.6, 0.3)
            ppt_para(tf, footer, size=11, color=P_MUTE, first=True)
            continue

        title = strip_md(headings[0][2]) if headings else kicker
        chrome(slide, kicker, title, idx, footer)
        y = 1.52
        tables = [b for b in rest if b[0] == "table"]
        others = [b for b in rest if b[0] != "table"]
        ol_n = 0

        for block in others:
            kind = block[0]
            if y > 6.55:
                break
            if kind == "p":
                tf = ppt_box(slide, MX, y, SW - 2 * MX, 0.42)
                p = tf.paragraphs[0]
                p.space_before = Ppt(0)
                add_ppt_inlines(p, block[1], size=16)
                y += 0.40
            elif kind == "quote":
                ppt_rect(slide, MX, y, 0.08, 0.55, P_TEAL)
                tf = ppt_box(slide, MX + 0.22, y, SW - 2 * MX - 0.22, 0.55)
                ppt_para(tf, strip_md(block[1]), size=16, color=P_NAVY, first=True, bold=True)
                y += 0.62
            elif kind in {"li", "ol"}:
                tf = ppt_box(slide, MX, y, SW - 2 * MX, 0.36)
                p = tf.paragraphs[0]
                if kind == "ol":
                    ol_n += 1
                    prefix = f"{ol_n}.  "
                else:
                    prefix = "•  "
                add_ppt_inlines(p, prefix + block[1], size=16)
                y += 0.34
            elif kind == "code":
                height = min(1.55, 0.26 * (block[1].count("\n") + 2))
                ppt_rect(slide, MX, y, SW - 2 * MX, height, P_PANEL)
                tf = ppt_box(slide, MX + 0.16, y + 0.08, SW - 2 * MX - 0.32, height - 0.1)
                first = True
                for line in block[1].splitlines() or [""]:
                    ppt_para(tf, line, size=13, color=P_TEAL, first=first, mono=True, before=1)
                    first = False
                y += height + 0.12

        if tables:
            rows = tables[0][1]
            remaining = max(1.4, SH - 0.55 - y)
            add_table(slide, rows, MX, y, SW - 2 * MX, min(remaining, 0.38 * len(rows) + 0.15))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def main() -> int:
    jobs = [
        (
            "fr",
            DOCS / "agentium-hypervisor-decision-strategy.md",
            DOCS / "agentium-hypervisor-decision-strategy.fr.docx",
            "Agentium — Hyperviseur, décision et industrialisation de l'IA",
            DOCS / "deck-agentium-decision-adoption.md",
            DOCS / "deck-agentium-decision-adoption.fr.pptx",
            "Agentium · Décision",
            "Datategy — Agentium — Confidentiel — demo/agentic",
        ),
        (
            "en",
            DOCS / "agentium-hypervisor-decision-strategy.en.md",
            DOCS / "agentium-hypervisor-decision-strategy.en.docx",
            "Agentium — Hypervisor, decision and industrializing AI",
            DOCS / "deck-agentium-decision-adoption.en.md",
            DOCS / "deck-agentium-decision-adoption.en.pptx",
            "Agentium · Decision",
            "Datategy — Agentium — Confidential — demo/agentic",
        ),
    ]
    written = []
    for lang, brief_md, brief_docx, title, deck_md, deck_pptx, kicker, footer in jobs:
        if not brief_md.exists() or not deck_md.exists():
            print(f"missing source for {lang}", file=sys.stderr)
            return 1
        written.append(build_docx(brief_md, brief_docx, title=title, lang=lang))
        written.append(build_pptx(deck_md, deck_pptx, kicker=kicker, footer=footer))
    extra_docx = [
        (
            "fr",
            DOCS / "agentium-hypervisor-user-stories.md",
            DOCS / "agentium-hypervisor-user-stories.fr.docx",
            "User stories : Hyperviseur Agentium",
        ),
        (
            "en",
            DOCS / "agentium-hypervisor-user-stories.en.md",
            DOCS / "agentium-hypervisor-user-stories.en.docx",
            "User stories: Agentium Hypervisor",
        ),
        (
            "fr",
            DOCS / "agentium-hypervisor-mvp-comex.md",
            DOCS / "agentium-hypervisor-mvp-comex.fr.docx",
            "Hyperviseur Agentium — MVP COMEX",
        ),
        (
            "en",
            DOCS / "agentium-hypervisor-mvp-comex.en.md",
            DOCS / "agentium-hypervisor-mvp-comex.en.docx",
            "Agentium Hypervisor — COMEX MVP",
        ),
    ]
    for lang, src, dest, title in extra_docx:
        if not src.exists():
            print(f"missing {src}", file=sys.stderr)
            return 1
        written.append(build_docx(src, dest, title=title, lang=lang))
    for path in written:
        print(f"{path}  ({path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
