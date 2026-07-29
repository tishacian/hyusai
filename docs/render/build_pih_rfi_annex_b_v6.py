#!/usr/bin/env python3
"""Annex B v6 — from v5: AgentOps overage tiers raised to $120/$140/$160
(recurring target ~680k$/year at 400 live agents, baseline unchanged).
Also fixes the stale $50/$60/$70 tier values left in Commercial Terms."""
import os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(HERE, "..", "pih")
SRC = os.path.join(PIH, "RFI-PIH-AI-Factory-Annex-B-Pricing-Schedule-v5.xlsx")
OUT = os.path.join(PIH, "RFI-PIH-AI-Factory-Annex-B-Pricing-Schedule-v6.xlsx")

TIERS = {"Low": 120, "Medium": 140, "High": 160}

wb = load_workbook(SRC)

# ------------------------------------------------ Pricing Summary
ps = wb["Pricing Summary"]
ps["A1"] = "Annex B — Pricing Schedule (v6, indicative, RFI stage)"
for row in range(1, ps.max_row + 1):
    label = str(ps.cell(row, 1).value or "")
    for tier, price in TIERS.items():
        if label == f"AgentOps overage — {tier} impact":
            ps.cell(row, 3).value = f"${price}"

# ------------------------------------------------ Illustrative Envelope
ie = wb["Illustrative Programme Envelope"]
ie["A1"] = "Illustrative programme envelope — 400 agents (v6, indicative)"
# Rows 27-29: Low/Medium/High overage agents; C holds display price, D the formula.
ie["C27"] = "$120/month"; ie["D27"] = "=B27*120"
ie["C28"] = "$140/month"; ie["D28"] = "=B28*140"
ie["C29"] = "$160/month"; ie["D29"] = "=B29*160"
mix = {27: 140, 28: 140, 29: 70}
weighted = (mix[27] * 120 + mix[28] * 140 + mix[29] * 160) / sum(mix.values())
ie["C30"] = f"Illustrative weighted average: ${weighted:.0f}/month"

# ------------------------------------------------ Commercial Terms & Capacity
ct = wb["Commercial Terms & Capacity"]
for row in range(1, ct.max_row + 1):
    if str(ct.cell(row, 1).value or "").startswith("Post-build baseline"):
        # Reconcile with the $8,900/month figure used everywhere else (8,900 x 12 = 106.8k).
        ct.cell(row, 1).value = "Post-build baseline — $8,900/month (≈ $107k/year)"
    if str(ct.cell(row, 1).value or "") == "Agents above 50":
        ct.cell(row, 2).value = (
            "AgentOps overage applies per accepted live agent and month: Low $120; Medium $140; "
            "High $160. High-impact scope includes reinforced drift watch and quarterly "
            "adversarial/red-team regression."
        )

wb.save(OUT)
print("saved:", OUT)

# ------------------------------------------------ verification
wb2 = load_workbook(OUT)
ps2, ie2, ct2 = wb2["Pricing Summary"], wb2["Illustrative Programme Envelope"], wb2["Commercial Terms & Capacity"]
ok = True
found = {t: False for t in TIERS}
for row in range(1, ps2.max_row + 1):
    label = str(ps2.cell(row, 1).value or "")
    for tier, price in TIERS.items():
        if label == f"AgentOps overage — {tier} impact":
            found[tier] = ps2.cell(row, 3).value == f"${price}"
assert all(found.values()), f"summary tiers: {found}"
assert ie2["D27"].value == "=B27*120" and ie2["D29"].value == "=B29*160"
baseline_annual = 8900 * 12
overage_monthly = 140 * 120 + 140 * 140 + 70 * 160
annual = baseline_annual + overage_monthly * 12
print(f"recurring check: baseline {baseline_annual:,}/yr + overage {overage_monthly:,}/mo -> {annual:,}/yr")
assert 650_000 <= annual <= 700_000, annual
stale = [
    r for r in range(1, ct2.max_row + 1)
    if "Low $50" in str(ct2.cell(r, 2).value or "")
    or "100k$/year" in str(ct2.cell(r, 1).value or "")
]
assert not stale, f"stale tiers remain rows {stale}"
assert any(
    "$107k/year" in str(ct2.cell(r, 1).value or "") for r in range(1, ct2.max_row + 1)
), "baseline label not reconciled"
print("all checks passed")
