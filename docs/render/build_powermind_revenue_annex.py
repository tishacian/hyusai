#!/usr/bin/env python3
"""Internal annex for the business introducer (PowerMind): price composition
of the PIH AI Factory offer and definition of the revenue-share basis.
Client-facing Annex B (all-inclusive) is unchanged; this document only maps
each client line to its economic components and states where the 50/50 applies."""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(HERE, "..", "pih")
OUT = os.path.join(PIH, "Datategy-PowerMind-Revenue-Share-Annex-DRAFT.docx")

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text, bold=False, italic=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    return par

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htxt in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = htxt
        for par in cell.paragraphs:
            for run in par.runs:
                run.bold = True
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = str(val)
    return t

# ---------------------------------------------------------------- title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Annex — Price Composition & Revenue-Share Basis\nPIH AI Factory Programme")
run.bold = True
run.font.size = Pt(16)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
srun = sub.add_run(
    "Datategy SAS × PowerMind — INTERNAL & CONFIDENTIAL — DRAFT for discussion\n"
    "Supplements the business-introduction agreement. Not for distribution to PIH."
)
srun.italic = True
srun.font.color.rgb = RGBColor(0x88, 0x00, 0x00)

# ---------------------------------------------------------------- 1
h("1. Purpose and principles")
p(
    "The PIH RFI mandates an all-inclusive unit price per delivered agent (team, tooling, testing, "
    "deployment, knowledge transfer) with no separate licence line during the Factory programme. "
    "Datategy's client-facing pricing (Annex B) therefore does not expose a software-licence price. "
    "This annex defines, for the sole purpose of the revenue-share agreement between Datategy and "
    "PowerMind, the economic composition of each client-facing pricing line and the basis on which "
    "the agreed revenue share applies."
)
p("Principles:", bold=True)
for txt in [
    "The client-facing pricing of Annex B is the sole commercial reference towards PIH and is not modified by this annex.",
    "Revenue share applies to the editor (software/licence-equivalent) component of collected revenues only; production and managed-services labour is excluded.",
    "The share is computed on amounts actually collected from PIH (net of taxes, travel at cost and third-party pass-through), and settled quarterly with a line-by-line statement.",
    "Any future change to Annex B pricing lines triggers an update of the mapping table in Section 3 under the same principles.",
]:
    doc.add_paragraph(txt, style="List Bullet")

# ---------------------------------------------------------------- 2
h("2. Client-facing pricing lines (Annex B v7 recap)")
table(
    ["Client line", "Unit", "Indicative price"],
    [
        ("Agent build — catalogue (reuse instance)", "per delivered agent", "Simple ≈ $3.5–5k · Medium ≈ $9–13k · Complex ≈ $20–30k"),
        ("First-of-type premium", "per identified first-of-type agent", "+50% / +50–75% / +75–100%"),
        ("UCC Factory Validation Pilot", "13 weeks, fixed capacity", "≈ $90–110k"),
        ("T&M rate card", "per day", "$149 / $349 / $440"),
        ("Enterprise Platform & Managed Operations Baseline (post-build)", "first 50 live agents / month", "$8,900"),
        ("AgentOps overage — Low / Medium / High impact", "per live agent above 50 / month", "$120 / $140 / $160"),
        ("Enhanced End-User Support (optional)", "per live agent / month", "$50"),
        ("SaaS inference (PIH-approved)", "—", "$0 Datategy fee (direct PIH–cloud billing)"),
        ("Sovereign inference — cloud-hosted (Azure Qatar) or PIH-provisioned on-premise", "—", "$0 Datategy infrastructure fee"),
        ("Sovereign inference — partner-provided (optional managed service)", "—", "Price on request"),
        ("Sovereign model-ops subscription (PIH-provisioned option)", "per year", "Scoped at contract"),
    ],
)

# ---------------------------------------------------------------- 3
h("3. Economic composition and revenue-share basis")
p(
    "Each client line decomposes into an editor component (Agentium platform: runtime, updates, "
    "security patches, connector/model compatibility, evaluation tooling, telemetry) and a delivery "
    "component (CoE production, managed-operations labour, support). The revenue share applies to "
    "the editor component as defined below."
)
table(
    ["Client line", "Editor component (share basis)", "Delivery component (excluded)", "Share basis"],
    [
        (
            "Agent build (catalogue + first-of-type premiums)",
            "None exposed — runtime is contractually included at $0 during the programme (competitive positioning required by the RFI).",
            "CoE production: design, authoring, integration, testing, security, deployment, hypercare.",
            "0% of line",
        ),
        (
            "UCC pilot & T&M",
            "None — pure delivery capacity.",
            "Staff augmentation, discovery, change requests.",
            "0% of line",
        ),
        (
            "Platform & Managed Operations Baseline ($8,900/mo)",
            "Platform Subscription: runtime continuity, updates, patches, compatibility, golden-set regression tooling, N3 product support.",
            "Standard AgentOps labour and knowledge service for the first 50 agents.",
            "50% of line = editor component",
        ),
        (
            "AgentOps overage ($120/$140/$160 per agent/mo)",
            "Platform telemetry, prompt-regression and evaluation tooling: $20 per agent/month.",
            "Ops labour: monitoring, exception management, reviews, red-team regression.",
            "$20 per agent/month",
        ),
        (
            "Enhanced End-User Support ($50 per agent/mo)",
            "None — direct L1 functional user assistance.",
            "Support labour during business hours.",
            "0% of line",
        ),
        (
            "Post-programme maintenance / Platform Subscription (successor contract, if any)",
            "Fully editor by nature.",
            "Any bundled managed services to be carved out at signature.",
            "100% of editor portion",
        ),
        (
            "Inference (SaaS, cloud-hosted sovereign or PIH-provisioned)",
            "None — $0 Datategy fee, pass-through or PIH infrastructure.",
            "—",
            "n/a",
        ),
        (
            "Partner-provided managed inference and sovereign model-ops subscription (if selected)",
            "None by default — infrastructure and operations service; any embedded platform-software portion to be carved out at signature.",
            "GPU capacity operations, model serving, monitoring.",
            "0% of line (carve-out reviewable at signature)",
        ),
    ],
)
p(
    "The $20 per agent/month platform component matches the AgentOps split already defined in "
    "Datategy's commercial model (editor telemetry vs transferable ops labour) and remains valid "
    "if PIH later appoints a third-party operator: the editor component stays with Datategy and "
    "remains in the share basis.",
    italic=True,
)

# ---------------------------------------------------------------- 4
h("4. Illustrative computation — steady state at 400 live agents")
table(
    ["Item", "Computation", "Monthly", "Annual"],
    [
        ("Baseline — editor component (50%)", "$8,900 × 50%", "$4,450", "$53,400"),
        ("Overage platform component", "350 agents × $20", "$7,000", "$84,000"),
        ("Editor basis — total", "", "$11,450", "$137,400"),
        ("PowerMind share (50%)", "", "$5,725", "$68,700 / year"),
    ],
)
p(
    "Build revenues (≈ $4.3M for 400 agents at the official 40/40/20 mix, before premiums and volume "
    "reduction) are production revenues carried by the Datategy CoE at cost and are excluded from the "
    "50/50 basis under Section 3.",
)

# ---------------------------------------------------------------- 5
h("5. Alternative structure (for discussion)")
p(
    "Should the parties prefer a basis less sensitive to the editor/delivery qualification of "
    "individual lines, the following alternative may replace Section 3: a reduced introduction fee "
    "on a broader basis — e.g. 10–15% of collected build revenues plus 50% of the editor components "
    "defined above. This aligns both parties on total programme win rather than on line "
    "classification. Selection of either structure must be recorded before submission of the "
    "Datategy response to PIH."
)

# ---------------------------------------------------------------- 6
h("6. Governance")
for txt in [
    "Quarterly settlement within 30 days of quarter end, on collected cash, with a line-by-line statement reconciling PIH invoices to the Section 3 mapping.",
    "Audit right: once per year, on reasonable notice, limited to PIH-programme invoicing records.",
    "Exclusions from all bases: taxes, travel and accommodation at cost, cloud/inference pass-through, penalties and credit notes.",
    "Duration: aligned with the business-introduction agreement; the editor-component share survives for revenues collected under contracts signed during its term (including the post-programme subscription).",
    "Confidentiality: this annex and the composition it defines are strictly confidential between Datategy and PowerMind and shall not be disclosed to PIH or any third party.",
]:
    doc.add_paragraph(txt, style="List Bullet")

p("")
p("DRAFT — figures and percentages to be validated by both parties before signature.", italic=True)

doc.save(OUT)
print("saved:", OUT)
