# -*- coding: utf-8 -*-
"""Produce Annex B v2 from the RFI-stage first draft.

v2 makes the commercial model internally coherent without publishing the
shortlist-only commercial values:
  - rebuilds the 400-agent envelope with first-of-type premiums;
  - defines incremental, non-retroactive volume discounts;
  - defines tier classification / design-card price-lock rules;
  - changes the delivery wording to remote-first;
  - separates AgentOps (transferable managed service) from the future Platform
    Subscription (editor rights after the build programme);
  - adds explicit platform-rights and capacity/payment-gate sheets;
  - makes knowledge capacity scoped by source type rather than promising an
    unbounded page allowance.

The precise subscription fee, factory throughput, knowledge allowance and pack
prices intentionally remain "proposal-stage / shortlist" parameters. They are
active decisions in the pricing canvas, not RFI-stage promises.
"""
import os
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
SRC = os.path.join(PIH, "RFI-PIH-AI-Factory-Annex-B-Pricing-Schedule.xlsx")
OUT = os.path.join(PIH, "RFI-PIH-AI-Factory-Annex-B-Pricing-Schedule-v2.xlsx")

NAVY = "0B2545"; PANEL = "F3F5F8"; AMBER = "FFF8E7"; GREEN = "E2F2EC"
H_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
B_FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
TITLE = Font(name="Arial", bold=True, size=14, color=NAVY)
NOTE = Font(name="Arial", italic=True, size=9, color="596371")
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_P = PatternFill("solid", fgColor=PANEL)
FILL_A = PatternFill("solid", fgColor=AMBER)
FILL_G = PatternFill("solid", fgColor=GREEN)
THIN = Border(*[Side(style="thin", color="D7DCE3")] * 4)
MONEY = "#,##0"


def style_row(ws, row, ncols, header=False, fill=None):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = H_FONT if header else B_FONT
        cell.border = THIN
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if header:
            cell.fill = FILL_H
        elif fill:
            cell.fill = fill


def clean_sheet(ws):
    ws.delete_rows(1, ws.max_row)


wb = load_workbook(SRC)

# ---------------------------------------------------------------- Reading Guide
rg = wb["Reading Guide"]
clean_sheet(rg)
rg["A1"] = "Annex B — Pricing Schedule (v2, indicative, RFI stage)"
rg["A1"].font = TITLE
guide_rows = [
    ("Scope", "PIH AI Factory Programme — response of Datategy. All figures USD, exclusive of taxes, "
              "indicative at RFI stage; firm pricing follows PIH clarifications, design cards and the "
              "confirmed agent mix at proposal / shortlist stage."),
    ("RFI pricing format", "The RFI requires an all-inclusive unit price per agent by Simple / Medium / "
                          "Complex tier, T&M day rates, an AgentOps monthly fee per live agent and "
                          "transparent inference treatment. This annex maps directly to that format."),
    ("Principle 1", "No per-user licence and no inference markup. SaaS inference is passed through at "
                    "cloud tariff; sovereign inference runs on PIH infrastructure."),
    ("Principle 2", "The catalogue build unit is a reuse instance: a new agent configured from patterns, "
                    "skills and connectors already delivered and validated. Unit price is fixed at "
                    "design-card approval; first-of-type and reuse status are determined before build."),
    ("Definition — reuse instance", "A new agent that reuses validated patterns, skills and connectors. "
                                    "Its business rules, parameters, rights and acceptance criteria are "
                                    "adapted, but its technical mechanism is not rebuilt from zero."),
    ("Definition — first-of-type", "The first agent creating or materially extending a pattern, skill or "
                                   "connector. It carries a separately scoped premium because it creates "
                                   "the reusable asset used by later reuse instances."),
    ("Principle 3", "The Agentium runtime is included for programme agents while an active factory "
                    "throughput commitment is in place. Post-programme rights are defined in the "
                    "'Platform Rights' sheet."),
    ("Principle 4", "Knowledge capacity is shared memory serving all agents. It is scoped by source "
                    "type and volume — never charged per agent — and priced separately from agent build."),
    ("Principle 5", "Delivery is remote-first through Datategy's AI & Data Center of Excellence. "
                    "Targeted Doha presence is limited to jointly agreed mobilisation, security/UAT and "
                    "critical go-live gates."),
    ("Contents", "Build Unit Prices · T&M Rate Card · AgentOps & Recurring · Illustrative Envelope · "
                 "Platform Rights · Capacity & Commercial Gates"),
    ("Validity", "90 days from submission. Rates assume a ≥3-month engagement commitment."),
]
row = 3
for i, (key, value) in enumerate(guide_rows):
    rg.cell(row=row, column=1, value=key).font = BOLD
    rg.cell(row=row, column=2, value=value)
    style_row(rg, row, 2, fill=FILL_P if i % 2 else None)
    rg.cell(row=row, column=1).font = BOLD
    row += 1
rg.column_dimensions["A"].width = 20
rg.column_dimensions["B"].width = 112

# ----------------------------------------------------------- Build Unit Prices
bu = wb["Build Unit Prices"]
clean_sheet(bu)
bu["A1"] = "Agent build — unit prices by complexity tier (factory execution)"
bu["A1"].font = TITLE
bu["A2"] = ("Catalogue price applies to a reuse instance: a new agent configured from patterns, skills "
            "and connectors already delivered and validated. It covers detailed design (design card), "
            "flow authoring, connector mapping, integration, unit/integration/UAT testing including "
            "golden-set creation, security validation, deployment and one knowledge-transfer session. "
            "First-of-type work is separately identified and priced before build, never retrospectively.")
bu["A2"].font = NOTE
headers = ["Tier", "Eligibility / objective classification", "RFI build time", "Catalogue price — reuse instance", "First-of-type premium"]
for col, header in enumerate(headers, 1):
    bu.cell(row=4, column=col, value=header)
style_row(bu, 4, 5, header=True)
tiers = [
    ("Simple", "One system; existing pattern and skill; deterministic or single-step generative; "
               "read or reversible write; maximum one standard approval.", "Up to 4 weeks",
     "≈ $3,500 – 5,000", "+50%"),
    ("Medium", "Two to three systems; pattern adaptation; exception handling and/or standard "
               "approval flow; no new critical connector; governed write may be included.", "Up to 8 weeks",
     "≈ $9,000 – 13,000", "+50–75%"),
    ("Complex", "New pattern or skill; three or more systems, privileged write, regulatory/high-impact "
                "controls, multi-agent coordination or reinforced parallel-run.", "Up to 12 weeks",
     "≈ $20,000 – 30,000", "+75–100%"),
]
row = 5
for i, values in enumerate(tiers):
    for col, value in enumerate(values, 1):
        bu.cell(row=row, column=col, value=value)
    style_row(bu, row, 5, fill=FILL_P if i % 2 else None)
    row += 1
row += 1
notes = [
    ("Price lock", "PIH and Datategy approve tier, first-of-type/reuse status, systems and acceptance "
                   "criteria in the design card before build. A reuse instance reuses delivered and "
                   "validated technical mechanisms; the resulting unit price is then fixed."),
    ("Volume degressivity", "−10% applies only to incremental reuse-instance units 101–250; −20% only "
                            "to incremental reuse-instance units 251+. Count: accepted agents live in "
                            "production. No retroactive rebate; first-of-type premiums excluded."),
    ("Reuse commitment", "30% pattern reuse / 15% component standardisation targets are measured "
                         "against the first-of-type baseline per pattern and reported monthly from the "
                         "factory ledger."),
]
for i, (label, value) in enumerate(notes):
    bu.cell(row=row, column=1, value=label).font = BOLD
    bu.cell(row=row, column=2, value=value)
    bu.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    style_row(bu, row, 5, fill=FILL_G if i == 1 else (FILL_P if i % 2 else None))
    bu.cell(row=row, column=1).font = BOLD
    row += 1
bu.column_dimensions["A"].width = 20
bu.column_dimensions["B"].width = 62
bu.column_dimensions["C"].width = 16
bu.column_dimensions["D"].width = 25
bu.column_dimensions["E"].width = 20

# --------------------------------------------------------------- T&M Rate Card
tm = wb["T&M Rate Card"]
tm["A2"] = ("Delivery model: remote-first — nearshore build capacity from Datategy's AI & Data Center "
            "of Excellence, with senior architecture and QA from France. Targeted Doha presence is "
            "limited to jointly agreed mobilisation, security/UAT and critical go-live gates. Travel "
            "& accommodation at cost. Valid for a ≥3-month commitment.")
tm["A2"].font = NOTE

# ------------------------------------------------------ AgentOps & Recurring
ao = wb["AgentOps & Recurring"]
clean_sheet(ao)
ao["A1"] = "AgentOps, platform rights & recurring components (indicative)"
ao["A1"].font = TITLE
for col, header in enumerate(["Component", "Basis", "Indicative price", "Notes"], 1):
    ao.cell(row=3, column=col, value=header)
style_row(ao, 3, 4, header=True)
recurring_rows = [
    ("AgentOps — Low-impact agents", "per live agent / month", "$50",
     "L1 monitoring, exception management, prompt regression, reporting. Transferable to PIH or a "
     "third party if PIH elects an alternative operating model."),
    ("AgentOps — Medium-impact agents", "per live agent / month", "$60",
     "As above plus approval-flow supervision and monthly quality review."),
    ("AgentOps — High-impact agents", "per live agent / month", "$70",
     "As above plus reinforced drift watch and quarterly red-team regression."),
    ("Agentium runtime — active programme", "included", "$0",
     "Included for programme agents while PIH maintains an active factory throughput commitment. "
     "No separate user licence and no inference markup."),
    ("Platform Subscription — post-programme", "annual / capacity tier", "on proposal",
     "Required for build/modify rights, updates, security patches, compatibility, golden-set "
     "regression and N3 support after the programme. Right to run delivered versions remains portable."),
    ("Knowledge capacity — initial allowance", "by source type / volume", "scoped at proposal",
     "Allowance specifies native digital pages, scanned/OCR pages, languages, metadata and number of "
     "knowledge bases. It is shared memory, never priced per agent."),
    ("Knowledge ingestion / capacity packs", "per defined pack", "on proposal",
     "Additional source onboarding, OCR, cleaning/metadata, storage and knowledge-base packs are "
     "priced after the source inventory; no unbounded page commitment at RFI stage."),
    ("Inference — SaaS endpoints", "pass-through", "at cloud tariff",
     "No markup; per-agent token telemetry and cost exportable to Finance."),
    ("Inference — sovereign serving", "PIH infrastructure envelope", "on quote",
     "Open-weight models on PIH GPU capacity (Azure Qatar region, private cloud or on-premise)."),
    ("Inference optimisation trajectory", "jointly tracked KPI", "−30–50% / 12 months",
     "Caching, right-sized routing and prompt optimisation; quarterly reviews."),
]
row = 4
for i, values in enumerate(recurring_rows):
    for col, value in enumerate(values, 1):
        ao.cell(row=row, column=col, value=value)
    style_row(ao, row, 4, fill=FILL_P if i % 2 else None)
    row += 1
ao.column_dimensions["A"].width = 38
ao.column_dimensions["B"].width = 28
ao.column_dimensions["C"].width = 20
ao.column_dimensions["D"].width = 74

# -------------------------------------------------------- Illustrative Envelope
il = wb["Illustrative Envelope"]
clean_sheet(il)
il["A1"] = "Illustrative programme envelope — ~400 agents (indicative, like-for-like comparison)"
il["A1"].font = TITLE
il["A2"] = ("Purely illustrative: mix, reuse-instance unit prices, first-of-type count/premium and "
            "volume reduction are editable assumptions. The unit-price tier and first-of-type status "
            "are fixed per agent at design-card approval; this sheet is not a commitment to the mix.")
il["A2"].font = NOTE
for col, header in enumerate(["Tier / component", "# / input (editable)", "Unit price / input (editable)", "Subtotal"], 1):
    il.cell(row=4, column=col, value=header)
style_row(il, 4, 4, header=True)
mix = [("Simple — reuse instances", 180, 4200), ("Medium — reuse instances", 160, 10500), ("Complex — reuse instances", 60, 24000)]
row = 5
for tier, count, price in mix:
    il.cell(row=row, column=1, value=tier)
    il.cell(row=row, column=2, value=count)
    il.cell(row=row, column=3, value=price).number_format = MONEY
    il.cell(row=row, column=4, value=f"=B{row}*C{row}").number_format = MONEY
    style_row(il, row, 4)
    il.cell(row=row, column=2).fill = FILL_A
    il.cell(row=row, column=3).fill = FILL_A
    row += 1
il.cell(row=row, column=1, value="Reuse-instance build subtotal").font = BOLD
il.cell(row=row, column=4, value=f"=SUM(D5:D{row - 1})").number_format = MONEY
style_row(il, row, 4, fill=FILL_P); reuse_subtotal = row; row += 1
il.cell(row=row, column=1, value="First-of-type premiums").font = BOLD
il.cell(row=row, column=2, value=40)
il.cell(row=row, column=3, value=6000).number_format = MONEY
il.cell(row=row, column=4, value=f"=B{row}*C{row}").number_format = MONEY
style_row(il, row, 4)
il.cell(row=row, column=2).fill = FILL_A
il.cell(row=row, column=3).fill = FILL_A
first_type = row; row += 1
il.cell(row=row, column=1, value="Volume degressivity — reuse instances only (blended)").font = BOLD
il.cell(row=row, column=2, value=0.12).number_format = "0%"
il.cell(row=row, column=4, value=f"=-D{reuse_subtotal}*B{row}").number_format = MONEY
style_row(il, row, 4)
il.cell(row=row, column=2).fill = FILL_A
discount = row; row += 1
il.cell(row=row, column=1, value="INDICATIVE BUILD ENVELOPE").font = Font(name="Arial", bold=True, size=11, color=NAVY)
il.cell(row=row, column=4, value=f"=D{reuse_subtotal}+D{first_type}+D{discount}").number_format = MONEY
il.cell(row=row, column=4).font = Font(name="Arial", bold=True, size=11, color=NAVY)
style_row(il, row, 4, fill=FILL_G); build_total = row; row += 2
il.cell(row=row, column=1, value="AgentOps at steady state (400 agents, blended $58/month)").font = BOLD
il.cell(row=row, column=4, value="=400*58*12").number_format = MONEY
style_row(il, row, 4); row += 1
il.cell(row=row, column=1, value="Excluded / scoped separately").font = BOLD
il.cell(row=row, column=2, value="Knowledge source inventory and packs; sovereign-serving infrastructure; "
                               "post-programme Platform Subscription; travel at cost.").font = NOTE
il.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
style_row(il, row, 4, fill=FILL_P)
il.column_dimensions["A"].width = 53
for col in ("B", "C", "D"):
    il.column_dimensions[col].width = 24

# --------------------------------------------------------------- Platform Rights
if "Platform Rights" in wb.sheetnames:
    del wb["Platform Rights"]
pr = wb.create_sheet("Platform Rights")
pr["A1"] = "Platform rights during and after the factory programme"
pr["A1"].font = TITLE
pr["A2"] = ("The model preserves PIH portability while distinguishing the perpetual right to run "
            "delivered versions from the subscription rights needed to evolve and maintain the factory.")
pr["A2"].font = NOTE
for col, header in enumerate(["Right", "During active factory programme", "After programme / inactive throughput", "Commercial treatment"], 1):
    pr.cell(row=4, column=col, value=header)
style_row(pr, 4, 4, header=True)
rights = [
    ("Run delivered agent versions", "Included", "Perpetual on PIH-controlled infrastructure", "Included in delivered agent price"),
    ("Inspect / export delivered flows, skills & evidence", "Included", "Included", "Supports portability and no supplier-specific dependency"),
    ("Create or materially modify agents", "Included", "Requires active Platform Subscription", "Subscription terms at proposal / shortlist"),
    ("Platform updates, security patches & compatibility", "Included", "Requires active Platform Subscription", "Subscription terms at proposal / shortlist"),
    ("Golden-set regression, N3 support & product roadmap", "Included", "Requires active Platform Subscription", "Subscription terms at proposal / shortlist"),
    ("AgentOps operational service", "Optional managed service", "Optional; transferable to PIH or third party", "Per live agent / month if selected"),
]
row = 5
for i, values in enumerate(rights):
    for col, value in enumerate(values, 1):
        pr.cell(row=row, column=col, value=value)
    style_row(pr, row, 4, fill=FILL_P if i % 2 else None)
    row += 1
row += 1
pr.cell(row=row, column=1, value="Subscription trigger").font = BOLD
pr.cell(row=row, column=2, value="To be confirmed at proposal stage: final programme-wave acceptance or 90 days without a committed factory throughput plan.")
pr.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
style_row(pr, row, 4, fill=FILL_A)
pr.column_dimensions["A"].width = 36
pr.column_dimensions["B"].width = 34
pr.column_dimensions["C"].width = 46
pr.column_dimensions["D"].width = 42

# ---------------------------------------------------- Capacity & Commercial Gates
if "Capacity & Commercial Gates" in wb.sheetnames:
    del wb["Capacity & Commercial Gates"]
cg = wb.create_sheet("Capacity & Commercial Gates")
cg["A1"] = "Factory capacity, acceptance & commercial gates"
cg["A1"].font = TITLE
cg["A2"] = ("A 400-agent volume cannot be responsibly committed without the PIH agent mix, access to "
            "environments and an agreed delivery-capacity plan. Capacity is therefore expressed in "
            "factory points, not in an undifferentiated number of agents.")
cg["A2"].font = NOTE
for col, header in enumerate(["Topic", "RFI-stage position", "Proposal-stage commercial gate"], 1):
    cg.cell(row=4, column=col, value=header)
style_row(cg, 4, 3, header=True)
gates = [
    ("Factory-point model", "Simple reuse = 1 point · Medium reuse = 2.5 points · Complex reuse = 7.5 points. First-of-type work is separately identified in the design card.",
     "Base points/month, ramp-up and extension capacity are agreed once PIH confirms mix, phases and priority domains."),
    ("Capacity extension", "Datategy scales through internal recruitment and allocation of its AI & Data Center of Excellence capacity.",
     "Extension requires agreed notice, named Datategy delivery resources, PIH environment/API access and an agreed ramp-up plan."),
    ("Design-card gate", "Tier, first-of-type/reuse status, systems, success metrics, security class and acceptance evidence are approved before build.",
     "Price and planned delivery slot are fixed at design-card approval."),
    ("Acceptance", "An agent is counted live only after agreed UAT, security gate and production acceptance evidence.",
     "Volume reductions are calculated only on accepted reuse-instance agents."),
    ("Pilot cash flow", "The 13-week pilot / staff augmentation phase is invoiced monthly.",
     "Mobilisation payment and monthly invoicing cadence are agreed before resources are reserved."),
    ("Industrial-wave cash flow", "Build proceeds in planned waves with transparent factory-ledger reporting.",
     "Mobilisation / capacity reservation plus short wave or monthly milestones; no deferred single milestone at 200 agents."),
    ("Travel", "Remote-first delivery. Doha attendance only for jointly agreed mobilisation, security/UAT or critical go-live gates.",
     "Travel and accommodation at cost; named attendance and purpose agreed before travel."),
]
row = 5
for i, values in enumerate(gates):
    for col, value in enumerate(values, 1):
        cg.cell(row=row, column=col, value=value)
    style_row(cg, row, 3, fill=FILL_P if i % 2 else None)
    row += 1
cg.column_dimensions["A"].width = 28
cg.column_dimensions["B"].width = 70
cg.column_dimensions["C"].width = 70

wb.save(OUT)
print("saved:", OUT)
