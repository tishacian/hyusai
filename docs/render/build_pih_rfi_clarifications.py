# -*- coding: utf-8 -*-
"""RFI PIH AI Factory — Clarification Questions (formal submission letter, EN).

Output: docs/pih/RFI-PIH-AI-Factory-Clarification-Questions.docx
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "logo_datategy.png")
REPO_ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(REPO_ROOT, "docs", "pih", "RFI-PIH-AI-Factory-Clarification-Questions.docx")


def C(h):
    return RGBColor.from_string(h)


NAVY = C("0B2545"); BLUE = C("2D6CDF"); INK = C("151A23"); MUT = C("596371")
LINE = "D7DCE3"
BODY = "Arial"


def _runfmt(run, size, color, bold=False, italic=False, mono=False):
    run.font.name = "Consolas" if mono else BODY
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = color


def _pbottom(p, color=LINE, sz=6):
    pPr = p._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "4")
    e.set(qn("w:color"), color); pb.append(e); pPr.append(pb)


def _pagefield(run, code="PAGE"):
    r = run._r
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "begin"); r.append(e)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code; r.append(it)
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); r.append(sep)
    t = OxmlElement("w:t"); t.text = "1"; r.append(t)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); r.append(end)


def para(doc, text, size=10.5, color=INK, bold=False, italic=False, before=4, after=6, spacing=1.18):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = spacing
    _runfmt(p.add_run(text), size, color, bold=bold, italic=italic)
    return p


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY; r.font.name = BODY; r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(4)
    return p


def q_item(doc, num, ref, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(7); pf.space_after = Pt(2); pf.line_spacing = 1.2
    pf.left_indent = Cm(0.6); pf.first_line_indent = Cm(-0.6)
    _runfmt(p.add_run(f"{num}. "), 10.5, NAVY, bold=True)
    _runfmt(p.add_run(f"({ref}) "), 9, BLUE, bold=True, mono=True)
    _runfmt(p.add_run(text), 10.5, INK)
    return p


def build():
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = BODY; normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK; normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.18
    st = doc.styles["Heading 1"]; st.font.name = BODY; st.font.bold = True; st.font.color.rgb = NAVY
    st.paragraph_format.keep_with_next = True

    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(1.9)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.1); sec.footer_distance = Cm(1.1)

    # header: logo left + doc label right
    hdr = sec.header; hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hp.add_run().add_picture(LOGO, height=Cm(0.6))
    hp.add_run("    ")
    _runfmt(hp.add_run("RFI Clarification Questions — PIH AI Factory"), 8.5, MUT, italic=True)
    _pbottom(hp)

    # footer: confidential left, page right
    f = sec.footer; f.is_linked_to_previous = False
    ft = f.add_table(rows=1, cols=2, width=Cm(16.6)); ft.allow_autofit = False
    lc = ft.cell(0, 0); rc = ft.cell(0, 1); lc.width = Cm(12.1); rc.width = Cm(4.5)
    _runfmt(lc.paragraphs[0].add_run("Datategy — Confidential"), 8, MUT)
    rp = rc.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _runfmt(rp.add_run("Page "), 8, MUT); _pagefield(rp.add_run())
    _empty = f.paragraphs[0]; _empty._element.getparent().remove(_empty._element)

    # ---- letterhead block
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    _runfmt(p.add_run("Datategy"), 11, NAVY, bold=True)
    para(doc, "58 rue de Monceau, 75008 Paris — France  ·  contact@datategy.net  ·  datategy.net",
         size=8.5, color=MUT, before=0, after=14)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _runfmt(p.add_run("Paris, ‹ date ›"), 9.5, MUT, italic=True)

    para(doc, "To: Power International Holding — Group Data & AI Office", bold=True, before=10, after=1)
    para(doc, "Subject: PIH AI Factory — Request for Information — Clarification Questions",
         bold=True, color=NAVY, before=1, after=4)
    para(doc, "Reference: PIH AI Factory RFI, issued 1 July 2026  ·  Clarification deadline: 7 July 2026, 17:00 AST",
         size=9.5, italic=True, color=MUT, before=1, after=14)

    para(doc, "Dear PIH Group Data & AI Office,")
    para(doc, "Please find below our clarification questions on the AI Factory Request for Information, "
              "referenced to the relevant sections of the document. We thank you in advance for your consideration.",
         after=10)

    # ============================================================ A
    h1(doc, "A — Programme scope & sizing")
    q_item(doc, 1, "§2, §2C",
           "To help respondents size their response, could PIH share an indicative breakdown of the "
           "~400 agents by build tier (Simple / Medium / Complex), agent type (Transactional / Assistive / "
           "Analytical / Autonomous), and domain or phase — even at a coarse grain?")
    q_item(doc, 2, "§2, §2C",
           "Section 2 targets ~400 agents “within 6–7 months from mobilisation”, while Section 2C "
           "sequences three phases from Q3 2026 to Q1 2027. Could PIH clarify the agent volume expected "
           "within the initial 6–7-month window versus across the three phases?")
    q_item(doc, 3, "§2",
           "Ahead of the full specifications at shortlist stage, could PIH indicate the number and "
           "indicative complexity of the priority SAP agents (FI / MM / SD / HCM / FI-AA) targeted for the "
           "Phase 1 pilot?")
    q_item(doc, 4, "§3",
           "Is Tier 3 (multi-agent coordination / A2A delegation) in scope for this engagement, or is it "
           "roadmap-only at this stage?")

    # ============================================================ B
    h1(doc, "B — Delivery structure")
    q_item(doc, 5, "§2, §5",
           "Will PIH accept a consortium or lead-partner-plus-specialists structure (for example, an "
           "AI-agentic, evaluation and sovereignty specialist teamed with a systems integrator for SAP / "
           "RPA / Databricks delivery scale), or is a single prime contractor required? Are there any "
           "constraints on subcontracting?")
    q_item(doc, 6, "§5",
           "What level of on-site presence in Qatar is expected during the pilot and the programme "
           "(mandatory on-site, hybrid, or remote-permissible), and what data-residency or security "
           "constraints apply to remote delivery and system access?")
    q_item(doc, 7, "§5, §8",
           "Section 8 indicates the operate-and-monitor commercial model will be confirmed at shortlist "
           "stage. Should respondents provide indicative AgentOps pricing now, or describe the capability "
           "only?")

    # ============================================================ C
    h1(doc, "C — Factory platform & infrastructure")
    q_item(doc, 8, "§3, §4",
           "To plan integration effort and the Phase 1 critical path, could PIH indicate the current "
           "implementation status of the Abstraction Layer components (LLM gateway, model router, "
           "connector registry, prompt store) and the AI Control Tower — in production, in build, or in "
           "design? Will their API specifications and the connector-registry / MCP standard be available "
           "at mobilisation? Where components are still in build or design, would PIH consider a "
           "partner-provided reference implementation operating under PIH’s governance and branding?")
    q_item(doc, 9, "§3, §3B",
           "Where existing Hikmah components provide overlapping capability (for example conversational "
           "analytics or an agent registry), should respondents plan to build within those components, "
           "extend them, or contribute complementary runtime capability integrated through the "
           "Abstraction Layer?")
    q_item(doc, 10, "§3",
           "Section 3 states that vendor names are category examples rather than mandated choices. To "
           "what extent may a respondent propose alternative category tools — or bring its own "
           "agentic-orchestration, evaluation or governance layer — integrated via the Abstraction Layer "
           "and MCP-compliant interfaces, while SAP, Hikmah/Databricks, Azure and Microsoft 365 remain the "
           "fixed strategic platforms?")
    q_item(doc, 11, "§3B, §4E",
           "For the UCC pilot domains, what is the current coverage of Gold data products? Given the "
           "master-data-quality risk identified for SAP-posting agents, who owns data-quality remediation "
           "(PIH or the delivery partner), and is data-quality sign-off a PIH-retained gate?")
    q_item(doc, 12, "§2A.3",
           "Does PIH mandate specific AI coding-agent tools, or will the partner’s own tooling be "
           "accepted? How does PIH govern the security and intellectual-property aspects of AI-generated "
           "code on its side?")

    # ============================================================ D
    h1(doc, "D — Model strategy & sovereignty")
    q_item(doc, 13, "§3A",
           "For which agent classes and data-sensitivity tiers does PIH require private, sovereign or "
           "on-premise inference, as opposed to SaaS model endpoints (e.g. Azure OpenAI)? Is there a "
           "target proportion of agents that must remain on sovereign infrastructure?")
    q_item(doc, 14, "§3A",
           "Is open-weight, on-premise model serving in scope alongside Azure OpenAI, and does PIH "
           "currently have — or plan to provision — sovereign GPU capacity, or is this expected from "
           "the partner?")

    # ============================================================ E
    h1(doc, "E — Governance, security & environments")
    q_item(doc, 15, "§4C",
           "For High Impact agents, who performs the pre-go-live red-teaming (the delivery partner, PIH, "
           "or an independent third party)? Are any standards or certifications required of the delivery "
           "partner (e.g. ISO/IEC 42001)?")
    q_item(doc, 16, "§2, §6",
           "Will shortlisted respondents be granted access to a non-production SAP environment (S/4HANA "
           "+ BTP) and the priority-agent module specifications before or at contract, to enable a "
           "representative pilot?")
    q_item(doc, 17, "§4A",
           "Are the ground-truth datasets used for evaluation provided and owned by PIH, or is the "
           "delivery partner expected to build and maintain them per agent type?")

    # ============================================================ F
    h1(doc, "F — Commercial & intellectual property")
    q_item(doc, 18, "§2A",
           "Could PIH clarify the definition and measurement of the 30% / 15% reuse cost-reduction "
           "targets — reduction against which baseline, measured how, and whether these are contractual "
           "commitments or aspirational targets?")
    q_item(doc, 19, "§5, §5B",
           "For outcome- and adoption-linked payments: what portion of fees is expected to be at risk, "
           "and how is Finance-confirmed value attributed between the delivery partner and PIH’s own "
           "change-management structure (BU Champions) that drives adoption?")
    q_item(doc, 20, "§3, §6",
           "Regarding IP and portability: who owns the patterns, skills, prompts and connectors handed "
           "over to PIH at the end of the engagement? Where a partner brings pre-existing IP (for example "
           "an agentic framework, connectors, or evaluation tooling), how should it be licensed so as to "
           "satisfy the “no supplier-specific dependency” and portability requirements?")
    q_item(doc, 21, "§5",
           "For the mobilisation-to-Phase-1 window, is any mobilisation or interim payment foreseen, or "
           "should respondents assume milestone-only cash flow across the full 13-week pilot?")

    para(doc, "We thank you in advance for these clarifications.", before=16, after=10)
    para(doc, "Kind regards,", after=2)
    para(doc, "‹ Name › — Datategy", bold=True)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    build()
