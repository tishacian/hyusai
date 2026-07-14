# -*- coding: utf-8 -*-
"""Iterate on the PIH offer v2 → v3 (v2 untouched).

Fixes applied (review of 2026-07-08):
  #4 — References named: ANDRITZ (knowledge & retrieval), Mobilize Financial
       Services (AIOps), Renault (multi-agent document pipelines) — in the
       Experience & References section and the exec-summary differentiator.
  #6 — Platform availability (uptime) target added to Support & Maintenance
       (RFP §6.5 explicitly asks for uptime SLAs).
  #7 — Executive summary condensed to fit the 2-page limit (RFP §9.1).
  #5 — TOC: w:updateFields flag set so Word/Docs regenerates the table of
       contents on open (new sections currently missing from the static TOC).

Input : docs/pih/Service Help Desk - Technical and Commercial offer - v2 Datategy.docx
Output: docs/pih/Service Help Desk - Technical and Commercial offer - v3 Datategy.docx
"""
import os
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
SRC = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v2 Datategy.docx")
OUT = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v3 Datategy.docx")

doc = Document(SRC)

BODY_STYLE = None
for p in doc.paragraphs:
    if p.style is not None and p.style.name == "normal":
        BODY_STYLE = p.style
        break


def find(prefix, style=None):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix) and (style is None or (p.style and p.style.name == style)):
            return p
    raise SystemExit(f"ANCHOR NOT FOUND: {prefix!r}")


def delete(p):
    p._p.getparent().remove(p._p)


def set_text(p, text, bold_lead=None):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def insert_after(anchor, text="", style=None, bold_lead=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if style is not None:
        para.style = style
    elif BODY_STYLE is not None:
        para.style = BODY_STYLE
    if bold_lead:
        r = para.add_run(bold_lead)
        r.bold = True
    if text:
        para.add_run(text)
    return para


def insert_block(anchor, items):
    cur = anchor
    for kind, text in items:
        if kind == "h2":
            cur = insert_after(cur, text, style=doc.styles["Heading 2"])
        elif kind.startswith("b:"):
            cur = insert_after(cur, text, bold_lead=kind[2:])
        else:
            cur = insert_after(cur, text)
    return cur


def paras_between(start, end_prefixes):
    out, seen = [], False
    for p in doc.paragraphs:
        if p._p is start._p:
            seen = True
            continue
        if seen:
            t = p.text.strip()
            if any(t.startswith(e) for e in end_prefixes):
                break
            out.append(p)
    return out


# ============================================================ #7 — EXEC SUMMARY ≤ 2 PAGES
h_exec = find("Executive Summary", "Heading 1")
for p in paras_between(h_exec, ["Scope Of Work Response"]):
    delete(p)

insert_block(h_exec, [
    ("h2", "Our understanding of your need"),
    ("p", "PIH operates one of the region's most demanding IT service environments: approximately "
          "12,200 active users across 145+ companies, 289+ locations and 25 countries, generating "
          "around 5,800 support tickets per month (~70,000 a year) — on top of the volume already "
          "absorbed by the 39 automation use cases in production. Automation is proven at PIH; what "
          "limits it today is the platform: rule-based rather than agentic, without a credible "
          "roadmap toward agentic AI and LLM technologies, and with SLA performance below "
          "expectations. You are asking for a next-generation, AI-driven service layer that "
          "consolidates all 39 use cases onto one scalable platform, deflects and automates the "
          "residual ticket load, integrates deeply with your landscape, delivers real-time KPIs and "
          "ROI, adds proactive operations and root-cause analysis, and leaves your IT team trained, "
          "autonomous and free of vendor lock-in — with ManageEngine ServiceDesk Plus remaining the "
          "ITSM system of record throughout."),
    ("h2", "Our proposed solution"),
    ("p", "Datategy proposes Agentium, our agent orchestration, building and evaluation platform — a "
          "software product configured for PIH, not a bespoke development. Where the current "
          "platform executes fixed workflows, Agentium runs governed agents: they understand intent "
          "in natural language, retrieve grounded answers from your knowledge with source citations, "
          "decide within explicit business rules, and act on your systems through governed "
          "connectors — every action mandated, logged and replayable. Your requirements map onto "
          "four systems built from one shared library of skills and connectors:"),
    ("p", "•  Conversational Service Desk & Knowledge Deflection — grounded, cited answers on Teams, "
          "portal and email that deflect tickets before they are created; intelligent triage and "
          "routing for the rest."),
    ("p", "•  Proactive Operations & Root-Cause Analysis — trend analysis over your ticket history, "
          "recurring-incident clustering, and RCA powered by Datategy's published research — from "
          "reactive to preventive."),
    ("p", "•  Identity & Access Lifecycle — onboarding/offboarding and access management across "
          "AD/Entra, Microsoft 365, SuccessFactors, Oracle and SolarWinds ARM; every privileged "
          "action runs under a signed, revocable mandate."),
    ("p", "•  Approval-Orchestrated Requests — groups, mailboxes, announcements, firewall rules: "
          "request → multi-level approval → fulfilment, with rejection, reminder and escalation "
          "branches built in."),
    ("p", "Migration of the 39 use cases is an authoring exercise, not 39 rebuilds: your existing "
          "automations decompose into ~10 reusable patterns on our Flow Builder. We migrate, "
          "validate in parallel run, and enhance — with a phased plan, UAT and hypercare."),
    ("h2", "What makes Datategy different"),
    ("b:Agentic-native, with a real roadmap — ",
     "Agentium is built for agents (stateful, replayable runs; a tool-calling ledger; evaluation of "
     "every output), not workflows retrofitted with AI; our product roadmap is published, dated and "
     "demonstrable."),
    ("b:Deployment freedom and sovereignty — ",
     "your Azure tenant (Qatar region), private cloud or fully on-premise, migratable without "
     "rework; Arabic-capable open-weight models and/or Azure OpenAI, routed per data-sensitivity "
     "tier. No per-employee licence billing your 12,200 users."),
    ("b:Governance you can prove — ",
     "distinct agent identity, RBAC+ABAC, signed mandates with caps and scopes, tamper-evident "
     "audit, and full replay of any past decision."),
    ("b:Quality is measured, not promised — ",
     "golden test sets including refusal cases, grounded-answer and hallucination rates, drift "
     "monitoring, and KPI dashboards reconciled against the run ledger; methodology aligned with "
     "ISO/IEC 42001 — Datategy sits on the AFNOR committee (CN IA, mirror of ISO/IEC JTC 1/SC 42) "
     "that contributes to drafting these standards."),
    ("b:Proven at enterprise scale — ",
     "grounded retrieval over multi-million-chunk technical corpora for ANDRITZ; predictive IT "
     "operations (AIOps) for Mobilize Financial Services (Renault Group); governed multi-agent "
     "document pipelines with normed quality (SAE J2450) for Renault — see Experience & References."),
    ("b:You own it at the end — ",
     "structured enablement, knowledge transfer, and a platform your team extends without us. We "
     "are a software editor, not a body shop: our success is your autonomy."),
    ("h2", "Delivery and value"),
    ("p", "We propose a phased engagement: a foundation phase, then a pilot wave with "
          "client-confirmed baselines, then wave-based migration and enhancement of all 39 use "
          "cases and scale-out across entities — each wave closing with quality gates and ≥30 days "
          "of hypercare. At your volumes, a 30–40% deflection/automation rate on the residual "
          "~70,000 tickets represents 5,000–7,000 analyst-hours per year returned to higher-value "
          "work — measured live in the platform's value cockpit, so each wave justifies the next. "
          "The run model is a flat editor licence with operations internalised in your team, "
          "designed to remain below your current conversational platform's annual cost."),
    ("p", "Datategy is a French AI software editor, EU Trusted AI-labelled, UGAP-referenced, and a "
          "contributor to European AI standards. We would be proud to make PIH's service desk the "
          "region's reference for sovereign, governed, agentic IT operations."),
])

# ============================================================ #6 — UPTIME SLA
p_sup = find("Editor support is included in the run fee")
insert_after(p_sup,
    "Platform availability: indicative target of 99.5% monthly on the platform services (up to "
    "99.9% in a high-availability topology with redundant workers and database), measured on the "
    "software layer and reported in the KPI cockpit. The underlying infrastructure (PIH Azure "
    "tenant or on-premise) remains operated under PIH's own infrastructure SLAs; committed "
    "availability figures are finalised at contract once the target topology is selected.")

# ============================================================ #4 — NAMED REFERENCES
p_refintro = find("Datategy is a French AI software editor founded in 2016")
set_text(p_refintro,
         "Datategy is a French AI software editor founded in 2016 — EU Trusted AI-labelled, "
         "UGAP-referenced for the French public market, and a member of the AFNOR AI "
         "standardisation committee (CN IA / ISO/IEC JTC 1 SC 42). Three references representative "
         "of this engagement (reference interviews available on request):")

REF_FIXES = [
    ("•  Enterprise knowledge & retrieval at industrial scale",
     "•  ANDRITZ — enterprise knowledge & retrieval at industrial scale: a grounded, cited "
     "assistant over a multi-million-chunk technical corpus for the global industrial engineering "
     "group: hybrid retrieval, evaluation harness, expert knowledge capture."),
    ("•  Predictive IT operations (AIOps)",
     "•  Mobilize Financial Services (Renault Group) — predictive IT operations (AIOps): proactive "
     "prevention of critical IT incidents: large-scale log and metric ingestion, time-series "
     "forecasting, explainable alerts adopted by the IT teams."),
    ("•  Governed multi-agent document pipelines",
     "•  Renault — governed multi-agent document pipelines: industrialised translation and QA "
     "chains with normed quality measurement (SAE J2450) and human-in-the-loop review: "
     "agent-per-error-class design, golden-set gating, measured ROI."),
]
for prefix, new in REF_FIXES:
    set_text(find(prefix), new)

doc.save(OUT)

# ============================================================ #5 — TOC refresh on open
# The TOC is a field whose cached entries are stale (new H1/H2 sections missing).
# Setting w:updateFields makes Word regenerate all fields — including the TOC —
# the next time the document is opened.
import zipfile
import shutil

tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/settings.xml":
            xml = data.decode("utf-8")
            if "w:updateFields" not in xml:
                xml = xml.replace("<w:settings ", "<w:settings ", 1)
                # insert right after the opening <w:settings ...> tag
                i = xml.index(">", xml.index("<w:settings")) + 1
                xml = xml[:i] + '<w:updateFields w:val="true"/>' + xml[i:]
            data = xml.encode("utf-8")
        zout.writestr(item, data)
shutil.move(tmp, OUT)

# ---- sanity: exec summary length
words = 0
started = False
d2 = Document(OUT)
for p in d2.paragraphs:
    t = p.text.strip()
    if t == "Executive Summary":
        started = True
        continue
    if started:
        if t.startswith("Scope Of Work Response"):
            break
        words += len(t.split())
print(f"exec summary: {words} words (~{words / 500:.1f} pages at ~500 w/page)")
print("saved:", OUT)
