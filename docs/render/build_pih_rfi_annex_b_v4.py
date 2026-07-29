# -*- coding: utf-8 -*-
"""PIH AI Factory — Annex B v4, client-facing pricing schedule.

Three sheets only:
  1. Pricing Summary
  2. Illustrative Programme Envelope
  3. Commercial Terms & Delivery Capacity

The programme build price is all-inclusive. Post-build is a single Enterprise
Platform & Managed Operations Baseline: USD 100k/year for the first 50 agents,
then AgentOps overage by risk class.
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
OUT = os.path.join(PIH, "RFI-PIH-AI-Factory-Annex-B-Pricing-Schedule-v4.xlsx")

NAVY = "0B2545"; PANEL = "F3F5F8"; AMBER = "FFF8E7"; GREEN = "E2F2EC"
H_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
B_FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
TITLE = Font(name="Arial", bold=True, size=14, color=NAVY)
NOTE = Font(name="Arial", italic=True, size=9, color="596371")
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_P = PatternFill("solid", fgColor=PANEL)
FILL_A = PatternFill("solid", fgColor=AMBER)
FILL_G = PatternFill("solid", fgColor=GREEN)
THIN = Border(*[Side(style="thin", color="D7DCE3")] * 4)
MONEY = "#,##0"


def style_row(ws, row, ncols, *, header=False, fill=None):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = H_FONT if header else B_FONT
        cell.border = THIN
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if header:
            cell.fill = FILL_H
        elif fill:
            cell.fill = fill


def two_column(ws, row, label, body, *, fill=None):
    ws.cell(row=row, column=1, value=label).font = BOLD
    ws.cell(row=row, column=2, value=body)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    style_row(ws, row, 4, fill=fill)
    ws.cell(row=row, column=1).font = BOLD


wb = Workbook()

# ---------------------------------------------------------------- Pricing Summary
ps = wb.active
ps.title = "Pricing Summary"
ps["A1"] = "Annex B — Pricing Schedule (v4, indicative, RFI stage)"
ps["A1"].font = TITLE
ps["A2"] = (
    "PIH AI Factory Programme — Datategy response. USD, exclusive of taxes. "
    "Firm pricing follows PIH clarifications, approved design cards and the confirmed agent mix."
)
ps["A2"].font = NOTE

row = 4
for label, body in [
    ("Commercial principle",
     "The build price is all-inclusive per delivered agent: agreed knowledge scope, Agentium runtime "
     "during the Factory programme, testing, security, deployment, hypercare and one knowledge-transfer "
     "session. No per-user licence and no Datategy per-execution consumption charge."),
    ("Reuse instance",
     "A new agent configured from patterns, skills and connectors already delivered and validated. "
     "Business rules, rights, parameters and acceptance criteria are adapted; the technical mechanism "
     "is not rebuilt from zero."),
    ("First-of-type",
     "The first agent creating or materially extending a pattern, skill or connector. It is a subset "
     "of programme volume and carries a premium identified and fixed in the design card before build."),
    ("Design card",
     "The approved record that fixes tier, systems, reuse/first-of-type status, security class, "
     "success metrics, acceptance evidence, delivery slot and price before build."),
]:
    two_column(ps, row, label, body, fill=FILL_P if row % 2 else None)
    row += 1

row += 1
headers = ["Build tier", "Eligibility", "RFI elapsed build time*", "Catalogue price — reuse instance", "First-of-type premium"]
for col, header in enumerate(headers, 1):
    ps.cell(row=row, column=col, value=header)
style_row(ps, row, 5, header=True)
row += 1
tiers = [
    ("Simple",
     "One system; delivered pattern/skill; deterministic or single-step generative; read or reversible write; max. one standard approval.",
     "Up to 4 weeks", "≈ $3,500 – 5,000", "+50%"),
    ("Medium",
     "Two to three systems; delivered pattern adapted for business rules or exceptions; standard approval; no new material connector/skill.",
     "Up to 8 weeks", "≈ $9,000 – 13,000", "+50–75%"),
    ("Complex",
     "New/materially extended pattern, skill or connector; 3+ systems; privileged write; High-Impact controls; multi-agent or reinforced parallel-run.",
     "Up to 12 weeks", "≈ $20,000 – 30,000", "+75–100%"),
]
for index, values in enumerate(tiers):
    for col, value in enumerate(values, 1):
        ps.cell(row=row, column=col, value=value)
    style_row(ps, row, 5, fill=FILL_P if index % 2 else None)
    row += 1

for label, body in [
    ("* Build time",
     "Elapsed calendar time, assuming PIH timely access to environments, SMEs, data and approval/UAT participants. It is not a representation of Datategy internal effort."),
    ("First-of-type formula",
     "Premium is applied to the catalogue price selected in the approved design card, before any volume reduction. It is non-degressive and applies only to the explicitly identified first-of-type scope."),
    ("Volume reduction",
     "−10% applies only to incremental accepted reuse instances numbered 101–250; −20% applies only to incremental accepted reuse instances numbered 251+. No retroactive rebate; first-of-type agents are excluded."),
    ("RFI reuse targets",
     "Datategy proposes a 30% pattern-reuse / 15% component-standardisation measurement method, subject to PIH clarification of baseline, measurement and contractual status."),
]:
    two_column(ps, row, label, body, fill=FILL_G if label == "Volume reduction" else (FILL_P if row % 2 else None))
    row += 1

row += 1
for col, header in enumerate(["T&M / inference item", "Basis", "Indicative price", "Clarification"], 1):
    ps.cell(row=row, column=col, value=header)
style_row(ps, row, 4, header=True)
row += 1
commercial_rows = [
    ("AI / Data Engineer · Developer · QA/Evaluation", "remote/hybrid, per day", "$149",
     "Pilot staff augmentation, discovery, change requests or scope outside catalogue unit pricing only."),
    ("CoE Director / Delivery leadership", "remote/hybrid, per day", "$349",
     "Part-time programme leadership."),
    ("Solution Architect / Expert", "remote/hybrid, per day", "$440",
     "Per named person; two concurrently assigned experts represent $880/day in aggregate."),
    ("Targeted Doha attendance", "travel/accommodation", "at cost",
     "Remote-first delivery; attendance only for jointly agreed mobilisation, security/UAT or critical go-live gates."),
    ("Enterprise Platform & Managed Operations Baseline — post-build", "first 50 live agents", "$100k/year",
     "Single post-build baseline covering platform continuity, standard AgentOps and knowledge service for delivered agents."),
    ("AgentOps overage — Low impact", "per live agent above first 50 / month", "$50",
     "L1 monitoring, exception management, prompt-regression checks and reporting."),
    ("AgentOps overage — Medium impact", "per live agent above first 50 / month", "$60",
     "Low-impact scope plus approval-flow supervision and monthly quality review."),
    ("AgentOps overage — High impact", "per live agent above first 50 / month", "$70",
     "Medium-impact scope plus reinforced drift watch and quarterly adversarial/red-team regression."),
    ("Azure / SaaS inference, where PIH approved", "Datategy fee", "$0",
     "Azure/OpenAI consumption remains contracted and billed directly between PIH and its cloud provider; Datategy adds no markup."),
    ("Private / sovereign inference", "PIH-controlled GPU capacity", "$0 Datategy infrastructure fee",
     "Hosted on PIH-controlled Azure Qatar, private-cloud or on-premise capacity; model sizing is a PIH infrastructure decision."),
    ("Inference optimisation", "joint KPI", "target −30–50% / 12 months",
     "Measured against a jointly agreed baseline; target only, not a guaranteed reduction."),
]
for index, values in enumerate(commercial_rows):
    for col, value in enumerate(values, 1):
        ps.cell(row=row, column=col, value=value)
    style_row(ps, row, 4, fill=FILL_P if index % 2 else None)
    row += 1

ps.column_dimensions["A"].width = 32
ps.column_dimensions["B"].width = 66
ps.column_dimensions["C"].width = 24
ps.column_dimensions["D"].width = 64
ps.column_dimensions["E"].width = 22

# ---------------------------------------------------- Illustrative Programme Envelope
ie = wb.create_sheet("Illustrative Programme Envelope")
ie["A1"] = "Illustrative programme envelope — 400 agents (indicative)"
ie["A1"].font = TITLE
ie["A2"] = (
    "The 400-agent mix includes reuse and first-of-type agents. First-of-type premiums apply to a "
    "stated subset of the same 400 agents. Amber cells are editable assumptions."
)
ie["A2"].font = NOTE
for col, header in enumerate(["Tier", "Total agents incl. first-of-type", "Catalogue base price", "Base subtotal"], 1):
    ie.cell(row=4, column=col, value=header)
style_row(ie, 4, 4, header=True)

mix = [("Simple", 180, 4200), ("Medium", 160, 10500), ("Complex", 60, 24000)]
row = 5
for index, (tier, count, price) in enumerate(mix):
    ie.cell(row=row, column=1, value=tier)
    ie.cell(row=row, column=2, value=count)
    ie.cell(row=row, column=3, value=price).number_format = MONEY
    ie.cell(row=row, column=4, value=f"=B{row}*C{row}").number_format = MONEY
    style_row(ie, row, 4, fill=FILL_P if index % 2 else None)
    ie.cell(row=row, column=2).fill = FILL_A
    ie.cell(row=row, column=3).fill = FILL_A
    row += 1
ie.cell(row=row, column=1, value="Base unit-price subtotal — all 400 agents").font = BOLD
ie.cell(row=row, column=4, value="=SUM(D5:D7)").number_format = MONEY
style_row(ie, row, 4, fill=FILL_P)
base_total = row
row += 2

for col, header in enumerate(["First-of-type subset (included above)", "# agents", "Premium rate", "Premium subtotal"], 1):
    ie.cell(row=row, column=col, value=header)
style_row(ie, row, 4, header=True)
row += 1
first = [("Simple first-of-type", 20, 0.50, "$C$5"), ("Medium first-of-type", 15, 0.60, "$C$6"), ("Complex first-of-type", 5, 0.80, "$C$7")]
first_start = row
for index, (label, count, premium, price_ref) in enumerate(first):
    ie.cell(row=row, column=1, value=label)
    ie.cell(row=row, column=2, value=count)
    ie.cell(row=row, column=3, value=premium).number_format = "0%"
    ie.cell(row=row, column=4, value=f"=B{row}*{price_ref}*C{row}").number_format = MONEY
    style_row(ie, row, 4, fill=FILL_P if index % 2 else None)
    ie.cell(row=row, column=2).fill = FILL_A
    ie.cell(row=row, column=3).fill = FILL_A
    row += 1
ie.cell(row=row, column=1, value="First-of-type premium subtotal").font = BOLD
ie.cell(row=row, column=2, value=f"=SUM(B{first_start}:B{row - 1})")
ie.cell(row=row, column=4, value=f"=SUM(D{first_start}:D{row - 1})").number_format = MONEY
style_row(ie, row, 4, fill=FILL_P)
first_total = row
row += 2

for col, header in enumerate(["Volume reduction — accepted reuse only", "Formula / input", "Result", "Amount"], 1):
    ie.cell(row=row, column=col, value=header)
style_row(ie, row, 4, header=True)
row += 1
ie.cell(row=row, column=1, value="Eligible reuse-instance count")
ie.cell(row=row, column=2, value=f"=SUM(B5:B7)-B{first_total}")
ie.cell(row=row, column=3, value="All agents less first-of-type subset")
style_row(ie, row, 4)
reuse_count = row
row += 1
ie.cell(row=row, column=1, value="Units 101–250 — 10%")
ie.cell(row=row, column=2, value=f"=MAX(0,MIN(B{reuse_count},250)-100)")
ie.cell(row=row, column=3, value=f"=B{row}*10%")
style_row(ie, row, 4)
band_10 = row
row += 1
ie.cell(row=row, column=1, value="Units 251+ — 20%")
ie.cell(row=row, column=2, value=f"=MAX(0,B{reuse_count}-250)")
ie.cell(row=row, column=3, value=f"=B{row}*20%")
style_row(ie, row, 4)
band_20 = row
row += 1
ie.cell(row=row, column=1, value="Illustrative effective reuse reduction")
ie.cell(row=row, column=2, value=f"=(C{band_10}+C{band_20})/B{reuse_count}").number_format = "0.0%"
ie.cell(row=row, column=3, value="Pro-rata presentation of contractual incremental rule")
style_row(ie, row, 4, fill=FILL_P)
effective_rate = row
row += 1
ie.cell(row=row, column=1, value="Reuse base-price subtotal")
ie.cell(row=row, column=2, value=f"=D{base_total}-SUMPRODUCT(B{first_start}:B{first_total-1},$C$5:$C$7)").number_format = MONEY
ie.cell(row=row, column=3, value="Base price of reuse agents eligible for reduction")
style_row(ie, row, 4)
reuse_base = row
row += 1
ie.cell(row=row, column=1, value="Volume reduction amount")
ie.cell(row=row, column=4, value=f"=-B{reuse_base}*B{effective_rate}").number_format = MONEY
style_row(ie, row, 4)
volume_discount = row
row += 1

ie.cell(row=row, column=1, value="INDICATIVE BUILD ENVELOPE").font = Font(name="Arial", bold=True, size=11, color=NAVY)
ie.cell(row=row, column=4, value=f"=D{base_total}+D{first_total}+D{volume_discount}").number_format = MONEY
ie.cell(row=row, column=4).font = Font(name="Arial", bold=True, size=11, color=NAVY)
style_row(ie, row, 4, fill=FILL_G)
row += 2

for col, header in enumerate(["Post-build recurring regime", "Formula / input", "Result", "Annual amount"], 1):
    ie.cell(row=row, column=col, value=header)
style_row(ie, row, 4, header=True)
row += 1
ie.cell(row=row, column=1, value="Enterprise Platform & Managed Operations Baseline")
ie.cell(row=row, column=2, value="First 50 live agents")
ie.cell(row=row, column=3, value="$100k/year")
ie.cell(row=row, column=4, value=100000).number_format = MONEY
style_row(ie, row, 4, fill=FILL_P)
baseline = row
row += 1
ie.cell(row=row, column=1, value="Low-impact overage agents")
ie.cell(row=row, column=2, value=140)
ie.cell(row=row, column=3, value="$50/month")
ie.cell(row=row, column=4, value=f"=B{row}*50*12").number_format = MONEY
style_row(ie, row, 4)
ie.cell(row=row, column=2).fill = FILL_A
low_overage = row
row += 1
ie.cell(row=row, column=1, value="Medium-impact overage agents")
ie.cell(row=row, column=2, value=140)
ie.cell(row=row, column=3, value="$60/month")
ie.cell(row=row, column=4, value=f"=B{row}*60*12").number_format = MONEY
style_row(ie, row, 4, fill=FILL_P)
ie.cell(row=row, column=2).fill = FILL_A
medium_overage = row
row += 1
ie.cell(row=row, column=1, value="High-impact overage agents")
ie.cell(row=row, column=2, value=70)
ie.cell(row=row, column=3, value="$70/month")
ie.cell(row=row, column=4, value=f"=B{row}*70*12").number_format = MONEY
style_row(ie, row, 4)
ie.cell(row=row, column=2).fill = FILL_A
high_overage = row
row += 1
ie.cell(row=row, column=1, value="AgentOps overage subtotal — 350 agents")
ie.cell(row=row, column=2, value=f"=SUM(B{low_overage}:B{high_overage})")
ie.cell(row=row, column=3, value="Illustrative weighted average: $58/month")
ie.cell(row=row, column=4, value=f"=SUM(D{low_overage}:D{high_overage})").number_format = MONEY
style_row(ie, row, 4, fill=FILL_P)
overage = row
row += 1
ie.cell(row=row, column=1, value="ILLUSTRATIVE POST-BUILD RECURRING").font = Font(name="Arial", bold=True, size=11, color=NAVY)
ie.cell(row=row, column=4, value=f"=D{baseline}+D{overage}").number_format = MONEY
ie.cell(row=row, column=4).font = Font(name="Arial", bold=True, size=11, color=NAVY)
style_row(ie, row, 4, fill=FILL_G)
row += 1
ie.cell(row=row, column=1, value="Scope note").font = BOLD
ie.cell(row=row, column=2, value=(
    "Baseline starts post-build: final acceptance of the contracted programme, or 90 days after the "
    "last accepted delivery if no successor contracted factory plan is active. It includes platform "
    "continuity, standard AgentOps and knowledge service for the first 50 live agents. The overage "
    "is priced by Low/Medium/High impact in the Summary sheet. Material new corpus/OCR/language/source, "
    "connector, pattern or regulatory scope is classified before commitment as first-of-type, Complex "
    "or change request."
))
ie.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
style_row(ie, row, 4, fill=FILL_P)
ie.column_dimensions["A"].width = 56
ie.column_dimensions["B"].width = 30
ie.column_dimensions["C"].width = 31
ie.column_dimensions["D"].width = 24

# ------------------------------------------- Commercial Terms & Delivery Capacity
ct = wb.create_sheet("Commercial Terms & Capacity")
ct["A1"] = "Commercial terms & delivery capacity"
ct["A1"].font = TITLE
ct["A2"] = (
    "The build price remains the primary commercial unit. These terms define the post-build continuity "
    "regime, acceptance and delivery controls; they do not introduce undisclosed add-on pricing."
)
ct["A2"].font = NOTE
row = 4
terms = [
    ("Run delivered agent versions",
     "PIH receives a non-exclusive right to execute accepted delivered agent versions on PIH-controlled infrastructure. This does not include future product releases or a source-code transfer."),
    ("During active Factory programme",
     "Runtime, agreed knowledge scope, build, testing, security, deployment, hypercare and transfer are included in the all-inclusive delivered-agent price."),
    ("Post-build baseline — 100k$/year",
     "For the first 50 live agents: Platform Subscription, updates, security patches, connector/model compatibility, golden-set regression, N3 support, standard AgentOps and knowledge service necessary to operate delivered agents."),
    ("Agents above 50",
     "AgentOps overage applies per accepted live agent and month: Low $50; Medium $60; High $70. High-impact scope includes reinforced drift watch and quarterly adversarial/red-team regression."),
    ("PIH / third-party operating model",
     "PIH may operate agents itself or appoint another operator. Delivered flows, design cards, evaluation evidence and documented connector mappings are exportable. The operating model does not transfer Agentium product IP, future releases or source code."),
    ("Knowledge scope",
     "Normal knowledge scope is agreed in each design card and included in build/baseline. A material expansion — massive corpus, industrial OCR, new language, source-critical integration or new regulatory scope — is scoped and priced before commitment as first-of-type, Complex or change request."),
    ("SaaS / sovereign inference",
     "Where PIH approves Azure/SaaS endpoints, Datategy charges $0 and PIH pays cloud consumption directly. Private/sovereign inference runs on PIH-controlled GPU capacity; infrastructure is not a Datategy recurring fee."),
    ("Factory capacity unit",
     "Simple reuse = 1 unit; Medium reuse = 2.5 units; Complex reuse = 7.5 units. This is a delivery-planning unit only, not a price unit."),
    ("Internal scale-up",
     "Datategy scales autonomously through internal recruitment and allocation of its AI & Data Center of Excellence. Monthly delivery capacity is confirmed against the PIH mix, named Datategy resources and environment/API access."),
    ("Design card & acceptance",
     "Tier, systems, first-of-type/reuse status, security class, success metrics, acceptance evidence, delivery slot and price are approved before build. An agent counts live after UAT, security gate and production acceptance."),
    ("Cash flow",
     "Mobilisation payment and monthly pilot invoicing are agreed before resources are reserved. Industrial delivery is paid through short monthly or wave milestones; no deferred single milestone at 200 agents."),
    ("Delivery model",
     "Remote-first. Doha attendance only for jointly agreed mobilisation, security/UAT and critical go-live gates; travel/accommodation at cost."),
]
for index, (label, body) in enumerate(terms):
    two_column(ct, row, label, body, fill=FILL_P if index % 2 else None)
    row += 1
ct.column_dimensions["A"].width = 33
ct.column_dimensions["B"].width = 120

wb.save(OUT)
print("saved:", OUT)
