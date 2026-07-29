# -*- coding: utf-8 -*-
"""Two generic product slides (16:9, EN) — what Datategy's products do.

Standalone, reusable deck (no client-specific content): requested by the
intermediary to present Datategy's product line on his own.
Slide 1 — one sovereign platform, three capabilities (papAI / Document Center / Agentium).
Slide 2 — how it runs in production: governance loop, differentiators, deployment, proof.

Visual system and Datategy branding reused from build_pih_synthesis_deck.py.
Output: out/Datategy-Products-Overview-2slides.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "logo_datategy.png")
SCREENS = os.path.join(HERE, "assets", "screens")
SHOTS = os.path.join(HERE, "assets", "screenshots")
OUT = os.path.join(HERE, "out", "Datategy-Products-Overview-2slides.pptx")

SW, SH = 13.333, 7.5
MX = 0.62
INK = RGBColor(0x15, 0x1A, 0x23)
MUT = RGBColor(0x59, 0x63, 0x71)
LINE = RGBColor(0xD7, 0xDC, 0xE3)
PANEL = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x2D, 0x6C, 0xDF)
PURPLE = RGBColor(0x6D, 0x28, 0xD9)
GREEN = RGBColor(0x0E, 0x8F, 0x62)
AMBER = RGBColor(0xB4, 0x53, 0x09)
TINTBLUE = RGBColor(0xEA, 0xF1, 0xFD)
LIGHTBLUE = RGBColor(0xBE, 0xD0, 0xE8)
F = "Arial"
FM = "Courier New"
FOOTER = "Datategy — Confidential · July 2026"

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=12, color=INK, bold=False, mono=False, first=False,
         before=4, after=0, align=PP_ALIGN.LEFT, spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(0 if first else before)
    p.space_after = Pt(after)
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.name = FM if mono else F
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return p


def add_run(p, text, size=12, color=INK, bold=False):
    r = p.add_run()
    r.text = text
    r.font.name = F
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return r


def rect(s, x, y, w, h, fill, line_color=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(0.75)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    sp.shadow.inherit = False
    return sp


def logo(s):
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, Inches(SW - MX - 1.45), Inches(0.42), height=Inches(0.34))


def chrome(s, kicker, title, page, accent=NAVY):
    rect(s, MX, 0.52, 0.30, 0.045, accent)
    tf = box(s, MX + 0.42, 0.40, SW - 2 * MX - 2.0, 0.3)
    para(tf, kicker.upper(), size=10.5, color=accent, bold=True, mono=True, first=True)
    tf = box(s, MX, 0.72, SW - 2 * MX - 1.6, 0.62)
    para(tf, title, size=24, color=INK, bold=True, first=True)
    rect(s, MX, 1.38, SW - 2 * MX, 0.012, LINE)
    tf = box(s, MX, SH - 0.42, 9.0, 0.3)
    para(tf, FOOTER, size=8.5, color=MUT, first=True)
    tf = box(s, SW - MX - 1.0, SH - 0.42, 1.0, 0.3)
    para(tf, f"{page:02d}", size=9, color=MUT, mono=True, first=True, align=PP_ALIGN.RIGHT)
    logo(s)


def panel_card(s, x, y, w, h, accent, title, subtitle=None, title_size=13):
    rect(s, x, y, w, h, WHITE, line_color=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.045)
    rect(s, x, y + 0.12, 0.05, h - 0.24, accent)
    tf = box(s, x + 0.22, y + 0.13, w - 0.4, 0.55 if subtitle else 0.35)
    para(tf, title, size=title_size, color=accent, bold=True, first=True)
    if subtitle:
        para(tf, subtitle, size=9.5, color=MUT, mono=True, before=1)
    return box(s, x + 0.22, y + (0.66 if subtitle else 0.52), w - 0.4, h - (0.82 if subtitle else 0.68))


def screenshot(s, path, x, w, y=None, bottom=None):
    """Insert a full, uncropped screenshot fitted to width at its native
    aspect ratio. Anchor either by top (`y`) or by bottom edge (`bottom`)."""
    pic = s.shapes.add_picture(path, Inches(x), Inches(y or 0), width=Inches(w))
    if bottom is not None:
        pic.top = Inches(bottom) - pic.height
    pic.line.color.rgb = LINE
    pic.line.width = Pt(1.0)
    pic.shadow.inherit = False
    return pic


def bullets(tf, items, size=10.5, before=5, spacing=1.02):
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            head, body = it
            p = para(tf, "▪ ", size=size, color=INK, first=(i == 0), before=before, spacing=spacing)
            add_run(p, head + " — ", size=size, color=INK, bold=True)
            add_run(p, body, size=size, color=MUT)
        else:
            para(tf, "▪ " + it, size=size, color=MUT, first=(i == 0), before=before, spacing=spacing)


# ============================================================ SLIDE 1
s = slide()
chrome(s, "Datategy product line", "One sovereign AI platform, three capabilities", 1, accent=NAVY)

tf = box(s, MX, 1.52, SW - 2 * MX, 0.4)
p = para(tf, "papAI is the common sovereign foundation. ", size=11.5, color=INK, bold=True, first=True)
add_run(p, "Document Center and Agentium are not separate products: one platform, one deployment, one governance and security model — capabilities activated as needs grow.", size=11.5, color=MUT)

CW = (SW - 2 * MX - 2 * 0.24) / 3
CY, CH = 2.02, 4.36
cards = [
    (BLUE, "papAI", "DATA INTELLIGENCE",
     os.path.join(SCREENS, "papai-features.png"), "papAI — AutoML experiment: features, data quality, training",
     [
        ("What it does", "turns data into industrialised AI assets: DataOps, AutoML, ML/deep learning, fine-tuning, MLOps"),
        ("Built in", "native orchestration, embedded ML/Spark/RAG engines, no-code + pro-code, drift detection and retraining"),
        ("Typical uses", "predictive maintenance, AIOps, forecasting, risk scoring, fraud detection, computer vision"),
    ]),
    (PURPLE, "Document Center", "KNOWLEDGE INTELLIGENCE",
     # anonymised variant: client workspace data (Andritz) pixelated, UI structure kept
     os.path.join(SCREENS, "guide-recherche-anon.png"), "Document Center — grounded, source-cited answers on a technical corpus",
     [
        ("What it does", "turns documents and expert knowledge into governed, searchable enterprise memory"),
        ("Built in", "OmniRAG sovereign retrieval (peer-reviewed research), cited answers, grounding checks, per-collection access rights"),
        ("Typical uses", "technical documentation assistants, regulatory analysis, expert knowledge capture, document compliance"),
    ]),
    (GREEN, "Agentium", "DECISION INTELLIGENCE",
     os.path.join(SCREENS, "ui-flow-builder.png"), "Agentium — Flow Builder: a governed multi-agent pipeline",
     [
        ("What it does", "turns data and knowledge into governed intelligent systems: agents that decide and act under control"),
        ("Built in", "visual Flow Builder over typed skills, run engine with replay, golden-set evaluation, agent identity and mandates, audit ledger, model router"),
        ("Typical uses", "enterprise assistants, transactional process agents, approval workflows, multi-agent pipelines"),
    ]),
]
for i, (accent, title, sub, img, cap, items) in enumerate(cards):
    x = MX + i * (CW + 0.24)
    tf = panel_card(s, x, CY, CW, CH, accent, title, subtitle=sub)
    bullets(tf, items, size=9.0, before=3, spacing=1.0)
    screenshot(s, img, x + 0.22, CW - 0.44, bottom=CY + CH - 0.14)

rect(s, MX, 6.50, SW - 2 * MX, 0.52, TINTBLUE, line_color=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
tf = box(s, MX + 0.28, 6.56, SW - 2 * MX - 0.56, 0.42)
p = para(tf, "ONE GRAMMAR, REUSED EVERYWHERE.  ", size=10, color=NAVY, bold=True, mono=True, first=True)
add_run(p, "Skill → Capability → System → Suite: reusable primitives assemble into business systems and vertical Suites. ", size=10.5, color=INK)
add_run(p, "Each delivery enriches the catalogue — the next one starts from assets, not from a blank page.", size=10.5, color=INK, bold=True)

# ============================================================ SLIDE 2
s = slide()
chrome(s, "How it runs in production", "Governed, sovereign, measurable — by design", 2, accent=PURPLE)

tf = box(s, MX, 1.52, SW - 2 * MX, 0.4)
p = para(tf, "Every intelligent system follows one canonical chain: ", size=11.5, color=INK, bold=True, first=True)
add_run(p, "Objective → System → Run → Evaluation → Decision → Action. What runs is measured; what is measured is governed.", size=11.5, color=MUT)

CW2 = (SW - 2 * MX - 3 * 0.22) / 4
CY2, CH2 = 2.04, 2.42
pillars = [
    (BLUE, "Evaluated by default", [
        "Golden sets and quality scoring per system; industry-normed where it exists (e.g. SAE J2450 for translation)",
        "Drift detection in production; no prompt or model change without passing regression",
    ]),
    (PURPLE, "Governed & auditable", [
        "Agents hold their own identity with scoped, time-boxed mandates",
        "Full audit ledger: every run, tool call, cost and human decision traced and replayable",
    ]),
    (GREEN, "Human in the loop", [
        "The AI proposes, the expert decides: confidence thresholds, escalation, calibrated alerts",
        "Corrections feed back into knowledge bases, golden sets and models",
    ]),
    (AMBER, "Sovereign by design", [
        "SaaS, private cloud, on-premise or air-gap — data and models stay in the client's perimeter",
        "Open-weight model serving, edge AI and federated learning proven in production",
    ]),
]
for i, (accent, title, items) in enumerate(pillars):
    tf = panel_card(s, MX + i * (CW2 + 0.22), CY2, CW2, CH2, accent, title, title_size=12)
    bullets(tf, items, size=9.5, before=5)

y3 = 4.60
PH = 1.82
rect(s, MX, y3, SW - 2 * MX, PH, PANEL, line_color=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
shot_w = 2.48
shot_h = shot_w * 0.625
shot_x = SW - MX - 0.22 - shot_w
screenshot(s, os.path.join(SHOTS, "09_evaluation.png"), shot_x, shot_w, y=y3 + 0.08)
tf = box(s, shot_x, y3 + 0.08 + shot_h + 0.03, shot_w, 0.2)
para(tf, "Hypervisor — live cost, value & ROI", size=7.5, color=MUT, mono=True, first=True, align=PP_ALIGN.RIGHT)
tf = box(s, MX + 0.28, y3 + 0.16, shot_x - MX - 0.56, PH - 0.3)
p = para(tf, "PROVEN IN PRODUCTION.  ", size=10, color=NAVY, bold=True, mono=True, first=True)
add_run(p,
        "Knowledge & retrieval over multi-million-document industrial corpora · SAP S/4HANA-integrated agents for aftermarket sales campaigns · "
        "governed multi-agent translation with normed QA for a global automaker · AIOps for an automotive financial-services group · "
        "computer vision for rail and highway maintenance.", size=9.8, color=MUT)
p = para(tf, "Trust marks — ", size=9.8, color=INK, bold=True, before=6)
add_run(p, "EU Trusted AI start-up · UGAP-referenced · AI Act contributor · peer-reviewed R&D (OmniRAG, root-cause analysis).", size=9.8, color=MUT)

rect(s, MX, y3 + PH + 0.10, SW - 2 * MX, 0.54, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
tf = box(s, MX + 0.28, y3 + PH + 0.20, SW - 2 * MX - 0.56, 0.38)
p = para(tf, "The client does not buy isolated agents or licences per user — ", size=11.5, color=WHITE, bold=True, first=True)
add_run(p, "it deploys complete, governed systems ready to produce measurable value, on infrastructure it controls.", size=11.5, color=LIGHTBLUE)

prs.save(OUT)
print("saved:", OUT)
