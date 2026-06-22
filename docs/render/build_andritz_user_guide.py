# -*- coding: utf-8 -*-
"""Guide utilisateur métier (FR) — chat « Recherche » et Capture de connaissances.

Note succincte pour les utilisateurs métier du workspace (navigation simplifiée).
Les copies d'écran sont pré-câblées : déposez un PNG nommé comme attendu dans
assets/screens/ et il s'embarque automatiquement ; sinon un cadre indicatif
s'affiche.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SCREENS = os.path.join(HERE, "assets", "screens")
LOGO = os.path.join(HERE, "assets", "logo_datategy.png")
OUT = os.path.join(HERE, "out", "Guide-Utilisateur-Recherche-Capture.docx")


def C(h):
    return RGBColor.from_string(h)

NAVY = C("0B2545"); BLUE = C("2D6CDF"); GREEN = C("0E8F62"); AMBER = C("B45309")
PURPLE = C("6D28D9"); INK = C("151A23"); MUT = C("596371")
WHITE = C("FFFFFF"); LINE = "D7DCE3"; PANEL = "F3F5F8"
TINT = {"BLUE": "E5EDFB", "GREEN": "E2F2EC", "AMBER": "F6ECE0", "NAVY": "E7ECF4"}
BODY = "Arial"


def _runfmt(run, size, color, bold=False, italic=False, mono=False):
    run.font.name = "Consolas" if mono else BODY
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = color


def _shade(cell, hexfill):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(sh)


def _cell_margins(cell, t=80, b=80, l=140, r=140):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for tag, v in (("top", t), ("bottom", b), ("start", l), ("end", r), ("left", l), ("right", r)):
        e = OxmlElement(f"w:{tag}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def _borders(table, color=LINE, sz=4, style="single"):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), style); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color); b.append(e)
    table._tbl.tblPr.append(b)


def _left_accent(cell, hexcolor, sz=30):
    tcPr = cell._tc.get_or_add_tcPr(); bd = OxmlElement("w:tcBorders")
    e = OxmlElement("w:left"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0"); e.set(qn("w:color"), hexcolor); bd.append(e); tcPr.append(bd)


def _pbottom(p, color="C9D2DE", sz=6):
    pPr = p._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "4")
    e.set(qn("w:color"), color); pb.append(e); pPr.append(pb)


def h1(doc, text, color=NAVY):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = color; r.font.name = BODY
    _pbottom(p); return p


def para(doc, text, size=10.5, color=INK, bold=False, italic=False, before=3, after=4, spacing=1.13):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = spacing
    r = p.add_run(text); _runfmt(r, size, color, bold=bold, italic=italic); return p


def runs_para(doc, parts, before=3, after=4, spacing=1.13, size=10.5):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = spacing
    for text, bold, color in parts:
        _runfmt(p.add_run(text), size, color, bold=bold)
    return p


def bullet(doc, parts, size=10.5, before=1, after=1):
    p = doc.add_paragraph(style="List Bullet"); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.1
    if isinstance(parts, str):
        parts = [(parts, False, INK)]
    for text, bold, color in parts:
        _runfmt(p.add_run(text), size, color, bold=bold)
    return p


def step(doc, n, title, accent, accent_hex, lines):
    """A numbered step block: coloured number chip + title, then bullet lines."""
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0); cell.width = Cm(17); _shade(cell, "FFFFFF")
    _left_accent(cell, accent_hex, sz=30); _cell_margins(cell, t=70, b=70, l=170, r=140)
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    _runfmt(p.add_run(f"{n}  "), 12.5, accent, bold=True)
    _runfmt(p.add_run(title), 12, INK, bold=True)
    for parts in lines:
        bp = cell.add_paragraph(); bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(1); bp.paragraph_format.line_spacing = 1.1
        _runfmt(bp.add_run("•  "), 10.5, accent)
        if isinstance(parts, str):
            parts = [(parts, False, MUT)]
        for text, bold, color in parts:
            _runfmt(bp.add_run(text), 10.5, color, bold=bold)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def callout(doc, title, body, accent, accent_hex, fill_hex):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0); cell.width = Cm(17); _shade(cell, fill_hex)
    _left_accent(cell, accent_hex, sz=34); _cell_margins(cell, t=110, b=110, l=190, r=150)
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    _runfmt(p.add_run(title.upper()), 9, accent, bold=True, mono=True)
    p2 = cell.add_paragraph(); p2.paragraph_format.line_spacing = 1.14
    _runfmt(p2.add_run(body), 10.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def figure(doc, caption, fname, width=16.0):
    """Embed a screenshot if present in assets/screens, else a labelled frame."""
    path = os.path.join(SCREENS, fname)
    if os.path.exists(path):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Cm(width))
    else:
        t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, color="C2CAD6", sz=6, style="dashed")
        cell = t.cell(0, 0); cell.width = Cm(width); _shade(cell, "FAFBFD"); _cell_margins(cell, t=420, b=420)
        cell.text = ""; pp = cell.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runfmt(pp.add_run("Copie d'écran à insérer"), 10, MUT, bold=True, mono=True)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    _runfmt(c.add_run(caption), 8.5, MUT, italic=True)


def build():
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = BODY; normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK; normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.13
    st = doc.styles["Heading 1"]; st.font.name = BODY; st.font.size = Pt(15); st.font.bold = True
    st.font.color.rgb = NAVY; st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True

    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.8); sec.left_margin = sec.right_margin = Cm(2.0)

    # footer
    f = sec.footer; f.is_linked_to_previous = False
    fp = f.paragraphs[0]; _runfmt(fp.add_run("Confidentiel · Datategy · Agentium — Guide utilisateur"), 8, MUT)

    # ---- title block (compact, not a full cover) ----
    tb = doc.add_table(rows=1, cols=2); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    lc, rc = tb.cell(0, 0), tb.cell(0, 1); lc.width = Cm(13.5); rc.width = Cm(3.5)
    lc.text = ""
    pe = lc.paragraphs[0]; _runfmt(pe.add_run("GUIDE UTILISATEUR · WORKSPACE MÉTIER"), 10, BLUE, bold=True, mono=True)
    pt = lc.add_paragraph(); pt.paragraph_format.space_before = Pt(2)
    _runfmt(pt.add_run("Recherche & Capture de connaissances"), 20, NAVY, bold=True)
    rp = rc.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(LOGO):
        rp.add_run().add_picture(LOGO, width=Cm(2.6))
    p = doc.add_paragraph(); _pbottom(p, "C9D2DE", sz=6); p.paragraph_format.space_after = Pt(6)

    runs_para(doc, [
        ("Vous accédez directement au système. ", True, INK),
        ("Ce guide couvre les deux outils de votre espace : la ", False, INK),
        ("Recherche", True, BLUE),
        (" (poser des questions sur les connaissances du workspace) et la ", False, INK),
        ("Capture de connaissances", True, GREEN),
        (" (transformer le savoir d'un expert en une fiche publiée). C'est tout ce dont vous avez "
         "besoin au quotidien.", False, INK),
    ], after=6)

    # ============================================ 1. Accès & navigation
    h1(doc, "1.  Accès & navigation")
    para(doc,
         "Après connexion, vous arrivez directement dans votre espace de travail. En haut de l'écran, "
         "deux onglets seulement :")
    bullet(doc, [("Recherche", True, BLUE), (" — interroger les connaissances du workspace.", False, MUT)])
    bullet(doc, [("Capture de connaissances", True, GREEN), (" — formaliser et publier un savoir.", False, MUT)])
    para(doc,
         "À droite, le sélecteur d'espace de travail et votre compte. Pas de menu technique à "
         "configurer : l'interface est volontairement épurée.", color=MUT)
    figure(doc, "Barre de navigation métier — onglets « Recherche » et « Capture de connaissances ».",
           "guide-nav.png")

    # ============================================ 2. Recherche
    h1(doc, "2.  Recherche — interroger vos connaissances", color=BLUE)
    para(doc,
         "La Recherche répond à vos questions en s'appuyant uniquement sur les sources de votre "
         "workspace, et cite ses sources. Si l'information n'y figure pas, l'assistant le dit plutôt "
         "que d'inventer.")
    step(doc, "1", "Posez votre question", BLUE, "2D6CDF", [
        [("Onglet ", False, MUT), ("Recherche", True, INK), (". Saisissez votre question dans le champ « ", False, MUT),
         ("Posez votre question…", False, INK), (" » puis « ", False, MUT), ("Interroger", True, INK),
         (" » — ou dictez-la avec le micro.", False, MUT)],
        [("Soyez précis : mentionnez le ", False, MUT), ("code projet, la machine ou la référence", True, INK),
         (" si vous les connaissez — la réponse n'en sera que plus juste.", False, MUT)],
        [("Vos échanges restent à gauche (", False, MUT), ("Conversations", True, INK),
         (") ; le bouton « + » démarre une nouvelle recherche.", False, MUT)],
    ])
    step(doc, "2", "Lisez la réponse et ses sources", BLUE, "2D6CDF", [
        [("La réponse va droit au fait. Les renvois ", False, MUT), ("[1] [2]", True, INK),
         (" pointent les documents utilisés ; cliquez-les pour vérifier la source. Le bandeau ", False, MUT),
         ("Sources du workspace", True, INK), (" indique le périmètre interrogé.", False, MUT)],
        [("Besoin d'aller plus loin ? « ", False, MUT), ("Deep search", True, INK),
         (" » lance une recherche approfondie quand le sujet le mérite.", False, MUT)],
    ])
    figure(doc, "Recherche — réponse sourcée avec renvois numérotés, et accès « Deep search » / « Corriger · Compléter ».",
           "guide-recherche.png")

    step(doc, "3", "Corriger ou compléter une réponse", BLUE, "2D6CDF", [
        [("Une réponse incomplète ou imprécise (parfois signalée « ", False, MUT), ("Réponse à vérifier", True, INK),
         (" ») ? Sous la réponse, cliquez « ", False, MUT), ("Corriger / Compléter", True, INK),
         (" » : l'encart « ", False, MUT), ("Correction experte", True, INK),
         (" » s'ouvre et rappelle la question. (Réservé aux droits de relecture.)", False, MUT)],
        [("Saisissez votre correction (« ", False, MUT),
         ("Corrigez ou complétez la réponse. Vous pouvez aussi dicter au micro.", False, INK),
         (" »), ou ", False, MUT), ("dictez-la au micro", True, INK),
         (" : la transcription reste modifiable avant l'envoi.", False, MUT)],
        [("Validez : votre correction est ", False, MUT), ("publiée immédiatement", True, INK),
         (" comme connaissance experte (", False, MUT), ("sans revue", True, INK),
         ("), et l'assistant en accuse réception dans le fil.", False, MUT)],
        [("Dès lors, cette connaissance est ", False, MUT),
         ("priorisée dans les recherches suivantes", True, INK), (", devant les documents ingérés.", False, MUT)],
    ])
    figure(doc, "Recherche — encart « Correction experte » : corriger la réponse, au clavier ou à la dictée.",
           "guide-recherche-correction.png")

    # ============================================ 3. Capture de connaissances
    h1(doc, "3.  Capture de connaissances — formaliser le savoir", color=GREEN)
    para(doc,
         "La Capture transforme le savoir d'un expert (oral ou écrit) en une fiche publiée dans la "
         "base — donc immédiatement réutilisable par la Recherche. Le parcours suit cinq étapes, "
         "indiquées en haut de l'écran.")
    para(doc,
         "Dans cet espace, la revue est désactivée : une fois validée, la fiche est publiée "
         "immédiatement et priorisée dans les résultats de Recherche.", color=MUT, italic=True)
    step(doc, "1", "Préparer", GREEN, "0E8F62", [
        [("Donnez un ", False, MUT), ("Titre de session", True, INK), (" (obligatoire), une ", False, MUT),
         ("Durée estimée", True, INK), (" (ou « Sans limite »), et un ", False, MUT), ("Mode", True, INK), (" :", False, MUT)],
        [("« Avec plan »", True, INK), (" prépare des angles de questions avant l'échange ; ", False, MUT),
         ("« Sans plan »", True, INK), (" démarre directement et structure après. Puis « Continuer ».", False, MUT)],
    ])
    step(doc, "2", "Plan", GREEN, "0E8F62", [
        [("Vérifiez et ajustez l'", False, MUT), ("arborescence", True, INK),
         (" (sujets et sous-sujets). Vous pouvez demander une modification en langage naturel puis « ", False, MUT),
         ("Appliquer", True, INK), (" ».", False, MUT)],
        [("Quand le plan vous convient : « ", False, MUT), ("Valider le plan", True, INK), (" ».", False, MUT)],
    ])
    figure(doc, "Capture — étape « Plan » : arborescence éditable et consigne en langage naturel.",
           "guide-capture-plan.png")
    step(doc, "3", "Capture (l'échange)", GREEN, "0E8F62", [
        [("« ", False, MUT), ("Lire", True, INK), (" » lit la question à voix haute ; répondez en parlant (« ", False, MUT),
         ("Parler", True, INK), (" ») ou au clavier. La position (sujet / sous-sujet) reste affichée en haut.", False, MUT)],
        [("À gauche : ", False, MUT), ("Sujets de capture", True, INK), (" et ", False, MUT), ("Questions IA", True, INK),
         (" (relances proposées par l'assistant).", False, MUT)],
        [("« ", False, MUT), ("Terminer la section", True, INK), (" » pour avancer ; « ", False, MUT),
         ("Terminer la capture", True, INK), (" » pour clore l'échange.", False, MUT)],
    ])
    figure(doc, "Capture — l'échange : question lue, réponse dictée, relances « Questions IA ».",
           "guide-capture-session.png")
    step(doc, "4", "Rapport", GREEN, "0E8F62", [
        [("Relisez la fiche structurée. « ", False, MUT), ("Éditer", True, INK),
         (" » pour la modifier ; « ", False, MUT), ("Plus d'actions", True, INK),
         (" » pour enregistrer, exporter ou régénérer.", False, MUT)],
        [("Traitez les ", False, MUT), ("Questions ouvertes", True, INK), (" à droite : « ", False, MUT),
         ("Répondre", True, INK), (" », « ", False, MUT), ("Laisser ouverte", True, INK), (" » ou « ", False, MUT),
         ("Invalider / Supprimer", True, INK), (" ». Une consigne courte + « ", False, MUT),
         ("Appliquer la consigne", True, INK), (" » fait corriger le texte.", False, MUT)],
        [("« ", False, MUT), ("Reprendre plus tard", True, INK), (" » sauvegarde l'état, sinon « ", False, MUT),
         ("Continuer vers publication", True, INK), (" ».", False, MUT)],
    ])
    figure(doc, "Capture — étape « Rapport » : fiche éditable et traitement des questions ouvertes.",
           "guide-capture-review.png")
    step(doc, "5", "Publier", GREEN, "0E8F62", [
        [("Confirmez le ", False, MUT), ("Titre final", True, INK), (", la ", False, MUT),
         ("Catégorie", True, INK), (" et la ", False, MUT), ("Destination", True, INK),
         (" (collection de connaissances). Le panneau ", False, MUT), ("Validation finale", True, INK),
         (" récapitule ce qui sera publié.", False, MUT)],
        [("« ", False, MUT), ("Publier dans la base de connaissances", True, INK),
         (" » : la fiche est découpée, indexée et disponible dans la Recherche. Message « ", False, MUT),
         ("Fiche publiée", True, INK), (" » à la fin.", False, MUT)],
    ])
    figure(doc, "Capture — étape « Publier » : titre, catégorie, destination et validation finale.",
           "guide-capture-publish.png")

    # ============================================ 4. Bon à savoir
    h1(doc, "4.  Bon à savoir")
    bullet(doc, [("Réponses sourcées — ", True, INK), ("la Recherche cite toujours ses documents ; vérifiez via les renvois [1][2].", False, MUT)])
    bullet(doc, [("Pas d'invention — ", True, INK), ("si la réponse n'est pas dans les sources, l'assistant vous le dit.", False, MUT)])
    bullet(doc, [("Boucle vertueuse — ", True, INK), ("une fiche publiée par la Capture devient immédiatement interrogeable dans la Recherche.", False, MUT)])
    bullet(doc, [("Vos corrections comptent — ", True, INK), ("une réponse corrigée est publiée immédiatement comme connaissance experte et priorisée dans les recherches suivantes.", False, MUT)])
    bullet(doc, [("Voix — ", True, INK), ("pendant la Capture vous pouvez écouter les questions et dicter vos réponses.", False, MUT)])
    bullet(doc, [("Reprendre plus tard — ", True, INK), ("une capture peut être mise en pause et reprise sans rien perdre.", False, MUT)])
    callout(doc, "En cas de doute",
            "Rapprochez-vous de votre administrateur d'espace de travail : il peut ajuster les sources, "
            "les collections et les accès. Vous, vous restez concentré sur vos questions et votre savoir.",
            NAVY, "0B2545", TINT["NAVY"])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("saved:", os.path.relpath(OUT, HERE))
    return OUT


if __name__ == "__main__":
    build()
