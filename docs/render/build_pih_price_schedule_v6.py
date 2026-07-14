# -*- coding: utf-8 -*-
"""PIH Price Schedule v5 — client-negotiated budget (validated 2026-07-08).

Aligned EXACTLY on the negotiated pricing table:
  Engagement initial ≈ USD 10k (months 1–3) · Wave 2 ≈ USD 26k (months 3–6)
  Wave 3 ≈ USD 14k (months 6–8) · Employee Services ≈ USD 10k (with Waves 2–3)
  Total build: Option 1 ≈ 50k · Option 2 ≈ 60k · ≈ 8 months
  Run: Option 1 USD 25k/yr · Option 2 USD 80k/yr (covers both scopes)

Day build-up mirrors the micro plan in 'Build Effort Detail - v1'
(same wave day budgets: 50 / 150 / 82 / 67 CoE days + 10 Paris oversight days).

Sheets: Summary · Rate Card · Build Detail · Run & Options · TCO Comparison
Output: docs/pih/Service Help Desk - Price Schedule - v6 Datategy.xlsx
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "pih", "Service Help Desk - Price Schedule - v6 Datategy.xlsx")

NAVY = "0B2545"; PANEL = "F3F5F8"; AMBER_BG = "FFF8E7"; GREEN_BG = "E2F2EC"; WHITE = "FFFFFF"
H_FONT = Font(name="Arial", bold=True, color=WHITE, size=10)
B_FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
TITLE = Font(name="Arial", bold=True, size=14, color=NAVY)
NOTE = Font(name="Arial", italic=True, size=9, color="596371")
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_P = PatternFill("solid", fgColor=PANEL)
FILL_A = PatternFill("solid", fgColor=AMBER_BG)
FILL_G = PatternFill("solid", fgColor=GREEN_BG)
THIN = Border(*[Side(style="thin", color="D7DCE3")] * 4)
MONEY = "#,##0"


def style_row(ws, row, cols, header=False, fill=None):
    for c in cols:
        cell = ws.cell(row=row, column=c)
        cell.font = H_FONT if header else B_FONT
        cell.border = THIN
        if header:
            cell.fill = FILL_H
        elif fill:
            cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


wb = openpyxl.Workbook()

# ================================================================ RATE CARD
rc = wb.active
rc.title = "Rate Card"
rc["A1"] = "Rate Card — indicative day rates (USD, excl. taxes)"
rc["A1"].font = TITLE
rc["A2"] = ("To be confirmed before contract. Build delivered end-to-end by Datategy's AI & Data Center of "
            "Excellence; Paris roles intervene as a thin architecture & QA oversight slice (10 days across "
            "the programme). On-site presence in Qatar: travel & accommodation at cost; premium to be agreed.")
rc["A2"].font = NOTE
for j, h in enumerate(["Role", "Location / model", "Day rate (USD)"], 1):
    rc.cell(row=4, column=j, value=h)
style_row(rc, 4, range(1, 4), header=True)

RATES = [
    ("Senior Solution Architect (oversight)", "France · remote/hybrid", 850),
    ("QA / Governance Lead (oversight)", "France · remote/hybrid", 500),
    ("CoE Agentic Tech Lead", "AI & Data Center of Excellence", 149),
    ("CoE AI Engineer", "AI & Data Center of Excellence", 149),
    ("CoE Integration Engineer (connectors)", "AI & Data Center of Excellence", 149),
    ("CoE QA / Evaluation Engineer", "AI & Data Center of Excellence", 149),
    ("CoE Delivery Lead", "AI & Data Center of Excellence", 149),
]
RATE_ROW = {}
for i, (role, loc, rate) in enumerate(RATES):
    r = 5 + i
    rc.cell(row=r, column=1, value=role)
    rc.cell(row=r, column=2, value=loc)
    rc.cell(row=r, column=3, value=rate).number_format = MONEY
    style_row(rc, r, range(1, 4), fill=FILL_P if i % 2 else None)
    RATE_ROW[role] = r
rc.column_dimensions["A"].width = 44
rc.column_dimensions["B"].width = 30
rc.column_dimensions["C"].width = 16

# ================================================================ BUILD DETAIL
bd = wb.create_sheet("Build Detail")
bd["A1"] = "Build — day-based build-up (USD excl. taxes) · negotiated, fixed price per wave"
bd["A1"].font = TITLE
bd["A2"] = ("Staged commitment: only the Engagement initial (≈ USD 10k) is engaged at signature; Waves 2–3 "
            "are pre-priced and engaged wave by wave on measured results. The per-use-case micro plan behind "
            "these day budgets is in the separate 'Build Effort Detail' annex. Compression levers: CoE-led "
            "delivery, pattern reuse across the 39 use cases, AI-assisted delivery.")
bd["A2"].font = NOTE

PHASES = [
    ("Engagement initial — Phase 0-lite + Pilot on the 4 foundational patterns (months 1–3)",
     "Foundation-lite (environments, security model, SDP + Entra ID + Teams connectivity, knowledge "
     "ingestion, golden sets, shared subflows) + pilot: deflection live, proactive detectors, identity "
     "write in parallel-run/simulate, one approval flow E2E.",
     [("Senior Solution Architect (oversight)", 3),
      ("CoE Agentic Tech Lead", 8), ("CoE AI Engineer", 25),
      ("CoE Integration Engineer (connectors)", 10), ("CoE QA / Evaluation Engineer", 4),
      ("CoE Delivery Lead", 3)]),
    ("Wave 2 — Migration 39/39 & privileged writes (months 3–6) · engaged on pilot results",
     "Full migration of the 39 use cases (pattern reuse), privileged writes staged live under mandates, "
     "Egnyte / SolarWinds ARM / firewall / Oracle / SuccessFactors connectors, 39/39 parallel-run "
     "evidence, extension to additional companies",
     [("Senior Solution Architect (oversight)", 3), ("QA / Governance Lead (oversight)", 2),
      ("CoE Agentic Tech Lead", 15), ("CoE AI Engineer", 75),
      ("CoE Integration Engineer (connectors)", 28), ("CoE QA / Evaluation Engineer", 22),
      ("CoE Delivery Lead", 10)]),
    ("Wave 3 — Scale & proactive (months 6–8) · engaged on Wave 2 results",
     "Multi-entity / multi-country rollout, extended RCA & trend analytics, optional Arabic UX, "
     "Datategy Academy handover, final acceptance",
     [("Senior Solution Architect (oversight)", 2),
      ("CoE Agentic Tech Lead", 8), ("CoE AI Engineer", 40),
      ("CoE QA / Evaluation Engineer", 20), ("CoE Delivery Lead", 14)]),
    ("Employee Services module (HR deflection & self-service) — Option 2 only",
     "Employee Q&A / self-service over HR policies & services: same channels, knowledge machinery and "
     "SuccessFactors connector reused; HR systems remain the systems of record",
     [("CoE Agentic Tech Lead", 4), ("CoE AI Engineer", 40),
      ("CoE QA / Evaluation Engineer", 17), ("CoE Delivery Lead", 6)]),
]

row = 4
SUBTOTALS = []
for title, scope, lines in PHASES:
    bd.cell(row=row, column=1, value=title).font = BOLD
    for c in range(1, 6):
        bd.cell(row=row, column=c).fill = FILL_P
        bd.cell(row=row, column=c).border = THIN
    bd.cell(row=row, column=2, value=scope).font = NOTE
    bd.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    row += 1
    for j, h in enumerate(["Role", "", "Days", "Rate (USD)", "Amount (USD)"], 1):
        if h:
            bd.cell(row=row, column=j, value=h)
    style_row(bd, row, range(1, 6), header=True)
    row += 1
    first = row
    for role, days in lines:
        bd.cell(row=row, column=1, value=role)
        bd.cell(row=row, column=3, value=days)
        bd.cell(row=row, column=4, value=f"='Rate Card'!C{RATE_ROW[role]}").number_format = MONEY
        bd.cell(row=row, column=5, value=f"=C{row}*D{row}").number_format = MONEY
        style_row(bd, row, range(1, 6))
        row += 1
    bd.cell(row=row, column=1, value="Subtotal").font = BOLD
    bd.cell(row=row, column=3, value=f"=SUM(C{first}:C{row - 1})").font = BOLD
    bd.cell(row=row, column=5, value=f"=SUM(E{first}:E{row - 1})").font = BOLD
    bd.cell(row=row, column=5).number_format = MONEY
    for c in range(1, 6):
        bd.cell(row=row, column=c).fill = FILL_A
        bd.cell(row=row, column=c).border = THIN
    SUBTOTALS.append(f"E{row}")
    row += 2

it_refs = SUBTOTALS[:3]
hr_ref = SUBTOTALS[3]
ENTRY_CELL = f"'Build Detail'!{SUBTOTALS[0]}"
bd.cell(row=row, column=1, value="TOTAL BUILD — OPTION 1 (IT scope, full path)").font = Font(name="Arial", bold=True, size=11, color=NAVY)
bd.cell(row=row, column=5, value="=" + "+".join(it_refs)).number_format = MONEY
bd.cell(row=row, column=5).font = Font(name="Arial", bold=True, size=11, color=NAVY)
for c in range(1, 6):
    bd.cell(row=row, column=c).fill = FILL_P
    bd.cell(row=row, column=c).border = THIN
OPT1_CELL = f"'Build Detail'!E{row}"
row += 1
bd.cell(row=row, column=1, value="TOTAL BUILD — OPTION 2 (IT + Employee Services, full path) · RECOMMENDED").font = Font(name="Arial", bold=True, size=11, color=NAVY)
bd.cell(row=row, column=5, value="=" + "+".join(it_refs + [hr_ref])).number_format = MONEY
bd.cell(row=row, column=5).font = Font(name="Arial", bold=True, size=11, color=NAVY)
for c in range(1, 6):
    bd.cell(row=row, column=c).fill = FILL_G
    bd.cell(row=row, column=c).border = THIN
OPT2_CELL = f"'Build Detail'!E{row}"
bd.column_dimensions["A"].width = 46
bd.column_dimensions["B"].width = 10
bd.column_dimensions["C"].width = 10
bd.column_dimensions["D"].width = 14
bd.column_dimensions["E"].width = 16

# ================================================================ RUN & OPTIONS
ro = wb.create_sheet("Run & Options")
ro["A1"] = "Run & Options (USD, excl. taxes)"
ro["A1"].font = TITLE
row = 3
ro.cell(row=row, column=1, value="Annual Run — Platform licence & editor support").font = BOLD
row += 1
ro.cell(row=row, column=1, value="Option 1 — IT scope")
ro.cell(row=row, column=2, value=25000).number_format = MONEY
style_row(ro, row, range(1, 3), fill=FILL_A)
RUN1_CELL = f"'Run & Options'!B{row}"
row += 1
ro.cell(row=row, column=1, value="Option 2 — IT + Employee Services (covers both scopes)")
ro.cell(row=row, column=2, value=80000).number_format = MONEY
style_row(ro, row, range(1, 3), fill=FILL_G)
RUN2_CELL = f"'Run & Options'!B{row}"
row += 1
for inc in [
    "Includes: platform licence (perimeter of the selected option), quarterly product releases, security "
    "patches, level-3 editor support (business hours, Qatar-time friendly), golden-set regression on every release.",
    "Option 2: the run fee covers the full IT + Employee Services perimeter — editor support is "
    "platform-level; day-to-day operations are internalised in PIH's team (Datategy Academy enablement).",
    "Excludes: infrastructure (PIH Azure tenant or on-premise, under PIH billing), day-to-day AgentOps, new "
    "use cases / connectors / entities (wave model or rate card).",
    "Perimeter bound to the systems, domains and companies listed in the offer; material extensions "
    "re-priced through the same wave model.",
]:
    ro.cell(row=row, column=1, value=inc).font = NOTE
    ro.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1

row += 1
ro.cell(row=row, column=1, value="Options (indicative)").font = BOLD
row += 1
for j, h in enumerate(["Option", "Unit", "Indicative price (USD)", "Notes"], 1):
    ro.cell(row=row, column=j, value=h)
style_row(ro, row, range(1, 5), header=True)
row += 1
OPTIONS = [
    ("Managed AgentOps (outsourced operations)", "per agent / month", "50",
     "Monitoring, exception management, prompt regression, reporting. Alternative to internalised model."),
    ("Additional use case — Simple", "fixed", "2,500 – 5,000", "Single system, existing pattern (~8–15 CoE days)."),
    ("Additional use case — Medium", "fixed", "6,000 – 12,000", "Multi-step / two systems, pattern adaptation."),
    ("Additional use case — Complex", "fixed", "15,000 – 30,000", "Multi-system, privileged writes, new pattern."),
    ("New connector (catalogue component)", "fixed", "4,000 – 10,000", "10–20 days on the connector framework, reusable."),
    ("Arabic user-experience package", "fixed", "15,000 – 25,000", "Arabic-capable model serving, KB coverage, UX validation."),
    ("Extended support (24/7)", "annual uplift", "on quote", "P1 coverage outside business hours."),
    ("Additional entity / company onboarding", "per batch", "wave model", "Scoped and priced as a mini-wave."),
]
for i, (opt, unit, price, note) in enumerate(OPTIONS):
    ro.cell(row=row, column=1, value=opt)
    ro.cell(row=row, column=2, value=unit)
    ro.cell(row=row, column=3, value=price)
    ro.cell(row=row, column=4, value=note).font = NOTE
    style_row(ro, row, range(1, 5), fill=FILL_P if i % 2 else None)
    ro.cell(row=row, column=4).font = NOTE
    row += 1
ro.column_dimensions["A"].width = 48
ro.column_dimensions["B"].width = 16
ro.column_dimensions["C"].width = 20
ro.column_dimensions["D"].width = 60

# ================================================================ TCO COMPARISON (INTERNAL)
tc = wb.create_sheet("TCO Comparison")
tc["A1"] = "TCO comparison — anonymised (USD, excl. taxes)"
tc["A1"].font = TITLE
tc["A2"] = ("INTERNAL CALIBRATION — baseline values below are informal estimates of the incumbent's licences. "
            "This workbook is NOT sent to the client. Never quote the incumbent's figures in writing.")
tc["A2"].font = Font(name="Arial", italic=True, bold=True, size=9, color="B45309")

tc["A4"] = "Baseline inputs (editable)"; tc["A4"].font = BOLD
tc["A5"] = "Current conversational platform — IT licence / year"
tc["B5"] = 38000; tc["B5"].number_format = MONEY; tc["B5"].fill = FILL_A; tc["B5"].border = THIN
tc["A6"] = "Current conversational platform — HR budget / year"
tc["B6"] = 70000; tc["B6"].number_format = MONEY; tc["B6"].fill = FILL_A; tc["B6"].border = THIN
tc["C6"] = "← confirmed via sales channel (~70k/yr, 2026-07-08)"; tc["C6"].font = NOTE

row = 8
for j, h in enumerate(["Scenario", "Year 1", "3 years", "5 years"], 1):
    tc.cell(row=row, column=j, value=h)
style_row(tc, row, range(1, 5), header=True)
row += 1
SCEN = [
    ("Current platform — IT only (rented)", "=$B$5*{n}", None),
    ("Current platform — IT + HR (rented)", "=($B$5+$B$6)*{n}", None),
    ("Option 1 — IT scope, full path (build + run)", f"={OPT1_CELL}+{RUN1_CELL}*{{n}}", None),
    ("Option 2 — IT + Employee Services, full path (build + run)", f"={OPT2_CELL}+{RUN2_CELL}*{{n}}", FILL_G),
]
opt2_row = None
combined_row = None
for i, (label, fmla, fill) in enumerate(SCEN):
    tc.cell(row=row, column=1, value=label).font = BOLD if fill else B_FONT
    for k, n in enumerate([1, 3, 5], 2):
        c = tc.cell(row=row, column=k, value=fmla.format(n=n))
        c.number_format = MONEY
    style_row(tc, row, range(1, 5), fill=fill if fill else (FILL_P if i % 2 else None))
    if label.startswith("Current platform — IT + HR"):
        combined_row = row
    if label.startswith("Option 2"):
        opt2_row = row
    row += 1
tc.cell(row=row, column=1, value="Delta — Option 2 vs current IT + HR (negative = above incumbent)").font = BOLD
for k in (2, 3, 4):
    col = chr(64 + k)
    c = tc.cell(row=row, column=k, value=f"={col}{combined_row}-{col}{opt2_row}")
    c.number_format = MONEY
    c.font = BOLD
style_row(tc, row, range(1, 5), fill=FILL_A)
row += 2
tc.cell(row=row, column=1,
        value="Reading (internal): with the incumbent's HR budget at ~70k/yr (confirmed), the baseline is "
              "~108k/yr IT+HR. Option 2 (60k build + 80k/yr run) crosses BELOW the incumbent at year 3 "
              "(299.8k vs 324k) and widens thereafter (5yr: 460k vs 540k). Year 1–2 is above baseline "
              "(build amortisation): the client-facing claim must be '3-year TCO below current combined "
              "spend', never a year-1 claim. Option 1's annual run (25k) stays below the IT-only licence "
              "(38k); its cumulative build+run crosses below at ~year 4.").font = Font(
        name="Arial", italic=True, bold=True, size=9, color="B45309")
tc.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
tc.column_dimensions["A"].width = 62
for col in ("B", "C", "D"):
    tc.column_dimensions[col].width = 14

# ================================================================ SUMMARY (first tab)
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "PIH Service Help Desk — Price Schedule (v6, client-negotiated)"
sm["A1"].font = TITLE
sm["A2"] = ("Currency USD, exclusive of taxes. Fixed price per wave as negotiated. Build delivered by "
            "Datategy's AI & Data Center of Excellence under senior Paris architecture & QA oversight. "
            "Per-use-case micro plan in the 'Build Effort Detail' annex. Aligned with the Technical & "
            "Commercial offer. Validity: 90 days.")
sm["A2"].font = NOTE

row = 4
for j, h in enumerate(["Component", "Option 1 — IT scope", "Option 2 — IT + Employee Services (RECOMMENDED)", "Notes"], 1):
    sm.cell(row=row, column=j, value=h)
style_row(sm, row, range(1, 5), header=True)
row += 1
LINES = [
    ("Engagement initial — Phase 0-lite + Pilot (4 patterns) · months 1–3", f"={ENTRY_CELL}", f"={ENTRY_CELL}",
     "Only amount engaged at signature"),
    ("Wave 2 — Migration 39/39, privileged writes, connectors · months 3–6", f"={SUBTOTALS[1].replace('E', chr(39) + 'Build Detail' + chr(39) + '!E')}", f"={SUBTOTALS[1].replace('E', chr(39) + 'Build Detail' + chr(39) + '!E')}",
     "Engaged on pilot results"),
    ("Wave 3 — Scale & proactive · months 6–8", f"={SUBTOTALS[2].replace('E', chr(39) + 'Build Detail' + chr(39) + '!E')}", f"={SUBTOTALS[2].replace('E', chr(39) + 'Build Detail' + chr(39) + '!E')}",
     "Engaged on Wave 2 results"),
    ("Employee Services module · with Waves 2–3", "—", f"={hr_ref.replace('E', chr(39) + 'Build Detail' + chr(39) + '!E')}",
     "HR deflection & self-service"),
    ("Total build — full path (one-time, staged) · ≈ 8 months", f"={OPT1_CELL}", f"={OPT2_CELL}",
     "≈ 50k / ≈ 60k as negotiated"),
    ("Annual Run — licence & editor support", f"={RUN1_CELL}", f"={RUN2_CELL}",
     "Option 2 run covers both scopes; from pilot go-live"),
    ("Indicative 3-year TCO (build + 3× run)", f"={OPT1_CELL}+3*{RUN1_CELL}", f"={OPT2_CELL}+3*{RUN2_CELL}",
     "RFP §9.3 requirement; excludes PIH-side infrastructure (under PIH billing)"),
]
for i, (comp, v1, v2, note) in enumerate(LINES):
    sm.cell(row=row, column=1, value=comp).font = BOLD
    a = sm.cell(row=row, column=2, value=v1)
    if v1 != "—":
        a.number_format = MONEY
    b = sm.cell(row=row, column=3, value=v2); b.number_format = MONEY; b.font = BOLD
    sm.cell(row=row, column=4, value=note).font = NOTE
    style_row(sm, row, range(1, 5), fill=FILL_P if i % 2 else None)
    sm.cell(row=row, column=3).fill = FILL_G
    row += 1
sm.cell(row=row, column=1, value="Options (both variants)").font = B_FONT
sm.cell(row=row, column=4, value="Managed AgentOps, additional use cases, Arabic UX, 24/7 — see 'Run & Options'").font = NOTE
style_row(sm, row, range(1, 5))
sm.column_dimensions["A"].width = 52
sm.column_dimensions["B"].width = 20
sm.column_dimensions["C"].width = 34
sm.column_dimensions["D"].width = 46

wb.save(OUT)

# ---- expected values
rates = {role: rate for role, _, rate in RATES}
subs = []
for title, _, lines in PHASES:
    sub = sum(d * rates[r] for r, d in lines)
    days = sum(d for _, d in lines)
    subs.append(sub)
    print(f"{title[:62]:62s} {days:4d}d  ${sub:>8,.0f}")
opt1 = sum(subs[:3]); opt2 = opt1 + subs[3]
print(f"{'TOTAL OPTION 1 (IT, full path)':62s}       ${opt1:>8,.0f}  (target ≈ 50k)")
print(f"{'TOTAL OPTION 2 (IT + Employee Services)':62s}       ${opt2:>8,.0f}  (target ≈ 60k)")
print(f"3-yr TCO: opt1 ${opt1+3*25000:,.0f} · opt2 ${opt2+3*80000:,.0f} · baseline IT+HR (38k+70k)×3 = $324,000")
print("saved:", OUT)
