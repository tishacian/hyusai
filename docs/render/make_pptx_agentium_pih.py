#!/usr/bin/env python3
"""
Generate a short English intro deck: Agentium alignment with PIH requirements.

Reuses visual helpers from the Translation Suite render pipeline.
Output: out/AGENTIUM_PIH_ALIGNMENT_EN.pptx
"""

import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

RENDER_ROOT = os.path.abspath(
    os.path.join(
        os.path.dirname(__file__),
        "../../../project-mt/OM/generic_code/docs/render",
    )
)
sys.path.insert(0, RENDER_ROOT)

from make_pptx import (  # noqa: E402
    SLIDE_W,
    SLIDE_H,
    BLACK,
    WHITE,
    GREY,
    DATATEGY_LOGO,
    add_textbox,
    add_rect,
    add_arrow,
    add_title,
    add_callout,
    hex_to_rgb,
)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out", "AGENTIUM_PIH_ALIGNMENT_EN.pptx")
SCREENSHOTS = os.path.join(HERE, "assets", "screenshots")

# Agentium palette
AGENTIUM_PRIMARY = "#1B3A5C"
AGENTIUM_ACCENT = "#4F46E5"
AGENTIUM_LIGHT = "#EEF2FF"
AGENTIUM_GREEN = "#059669"
AGENTIUM_AMBER = "#D97706"

PILLARS = [
    ("Sovereignty", AGENTIUM_PRIMARY),
    ("Robustness", AGENTIUM_GREEN),
    ("Security", "#7C3AED"),
    ("Observability & Scalability", AGENTIUM_AMBER),
    ("Token Performance", AGENTIUM_ACCENT),
]


def shot(name: str) -> str | None:
    path = os.path.join(SCREENSHOTS, name)
    return path if os.path.isfile(path) else None


def add_screenshot_frame(slide, image_path, x, y, w, h, caption: str = ""):
    """Place a screenshot with a subtle border and optional caption."""
    add_rect(slide, x, y, w, h, fill=WHITE, line=hex_to_rgb("#CCCCCC"), line_w=1.0)
    slide.shapes.add_picture(image_path, Inches(x + 0.06), Inches(y + 0.06),
                             width=Inches(w - 0.12), height=Inches(h - (0.42 if caption else 0.12)))
    if caption:
        add_textbox(
            slide, x + 0.06, y + h - 0.34, w - 0.12, 0.28,
            caption,
            font="Helvetica", size=9, color=GREY,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )


def add_agentium_header(slide, section: str = ""):
    accent = hex_to_rgb(AGENTIUM_PRIMARY)
    add_rect(slide, 0, 0, SLIDE_W, 0.18, fill=accent)
    if os.path.isfile(DATATEGY_LOGO):
        slide.shapes.add_picture(
            DATATEGY_LOGO, Inches(0.35), Inches(0.30), height=Inches(0.55)
        )
    add_textbox(
        slide,
        SLIDE_W - 3.2,
        0.32,
        2.85,
        0.50,
        "Agentium",
        font="Helvetica",
        size=20,
        bold=True,
        color=accent,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    if section:
        badge_w = 6.0
        badge_x = (SLIDE_W - badge_w) / 2
        add_rect(
            slide,
            badge_x,
            0.32,
            badge_w,
            0.50,
            fill=hex_to_rgb(AGENTIUM_LIGHT),
            line=accent,
            line_w=1.0,
        )
        add_textbox(
            slide,
            badge_x,
            0.33,
            badge_w,
            0.48,
            section,
            font="Helvetica",
            size=13,
            bold=True,
            color=accent,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def add_footer(slide, page_num: int, total: int):
    add_rect(slide, 0, SLIDE_H - 0.05, SLIDE_W, 0.05, fill=hex_to_rgb("#CCCCCC"))
    add_textbox(
        slide,
        0.35,
        SLIDE_H - 0.42,
        9.0,
        0.32,
        "Datategy — Agentium — PIH alignment intro",
        font="Helvetica",
        size=9,
        color=GREY,
    )
    add_textbox(
        slide,
        SLIDE_W - 2.2,
        SLIDE_H - 0.42,
        1.85,
        0.32,
        f"Page {page_num} / {total}",
        font="Helvetica",
        size=9,
        color=GREY,
        align=PP_ALIGN.RIGHT,
    )


def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SLIDE_W, 0.55, fill=hex_to_rgb(AGENTIUM_PRIMARY))
    if os.path.isfile(DATATEGY_LOGO):
        s.shapes.add_picture(DATATEGY_LOGO, Inches(2.8), Inches(1.15), height=Inches(1.2))

    add_textbox(
        s,
        0.5,
        2.55,
        SLIDE_W - 1.0,
        0.85,
        "Agentium",
        font="Helvetica",
        size=44,
        bold=True,
        color=BLACK,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        0.5,
        3.35,
        SLIDE_W - 1.0,
        0.65,
        "Enterprise alignment with PIH requirements",
        font="Helvetica",
        size=24,
        color=hex_to_rgb(AGENTIUM_PRIMARY),
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        0.5,
        4.05,
        SLIDE_W - 1.0,
        0.55,
        "Intro deck before live demo",
        font="Helvetica",
        size=16,
        color=GREY,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        0.5,
        4.75,
        SLIDE_W - 1.0,
        1.0,
        "Sovereignty  ·  Robustness  ·  Security\n"
        "Observability & Scalability  ·  Token Performance",
        font="Helvetica",
        size=14,
        color=hex_to_rgb("#444444"),
        align=PP_ALIGN.CENTER,
        line_spacing=1.35,
    )
    add_textbox(
        s,
        0.5,
        SLIDE_H - 0.85,
        SLIDE_W - 1.0,
        0.35,
        "Datategy — Confidential",
        font="Helvetica",
        size=11,
        color=GREY,
        align=PP_ALIGN.CENTER,
    )
    return s


def slide_agenda(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "Agenda")
    add_title(
        s,
        "What this deck covers",
        "How Agentium's operating model maps to PIH's five enterprise pillars",
    )
    items = [
        (
            "1.",
            "Agentium in one picture",
            "An OS for intelligent systems — not a chatbot builder.",
        ),
        (
            "2.",
            "PIH requirements → Agentium capabilities",
            "Five pillars, one alignment matrix.",
        ),
        (
            "3.",
            "Deep dive per pillar",
            "Sovereignty, robustness, security, observability, token performance.",
        ),
        (
            "4.",
            "Live demo",
            "System → Run → Evaluation → Decision → Action.",
        ),
    ]
    y = 2.05
    for num, title, body in items:
        add_rect(
            s,
            0.55,
            y,
            SLIDE_W - 1.1,
            1.15,
            fill=hex_to_rgb("#FAFAFA"),
            line=hex_to_rgb("#DDDDDD"),
        )
        add_textbox(
            s,
            0.75,
            y + 0.15,
            0.55,
            0.45,
            num,
            font="Helvetica",
            size=22,
            bold=True,
            color=hex_to_rgb(AGENTIUM_ACCENT),
        )
        add_textbox(
            s,
            1.35,
            y + 0.10,
            SLIDE_W - 2.0,
            0.45,
            title,
            font="Helvetica",
            size=16,
            bold=True,
            color=BLACK,
        )
        add_textbox(
            s,
            1.35,
            y + 0.55,
            SLIDE_W - 2.0,
            0.50,
            body,
            font="Helvetica",
            size=13,
            color=hex_to_rgb("#333333"),
        )
        y += 1.28
    return s


def slide_mental_model(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "Mental model")
    add_title(
        s,
        "Agentium — operating system for intelligent systems",
        "Objectives in, governed and measurable outcomes out",
    )

    loop = [
        ("Objective", "Business intent\ndeclared on System"),
        ("System", "Capability + Flow\n+ Skills + Context"),
        ("Run", "Stateful execution\ninstance with ledger"),
        ("Impact", "Cost · latency ·\nconfidence · value"),
        ("Optimize", "Hypervisor +\nControl Plane loop"),
    ]
    n = len(loop)
    margin = 0.45
    box_w = 2.15
    gap = (SLIDE_W - 2 * margin - n * box_w) / (n - 1)
    box_h = 1.55
    box_y = 2.15
    x = margin
    for i, (head, body) in enumerate(loop):
        add_rect(
            s,
            x,
            box_y,
            box_w,
            box_h,
            fill=hex_to_rgb(AGENTIUM_LIGHT),
            line=hex_to_rgb(AGENTIUM_ACCENT),
            line_w=1.3,
        )
        add_textbox(
            s,
            x,
            box_y + 0.10,
            box_w,
            0.35,
            head,
            font="Helvetica",
            size=13,
            bold=True,
            color=hex_to_rgb(AGENTIUM_PRIMARY),
            align=PP_ALIGN.CENTER,
        )
        add_textbox(
            s,
            x + 0.08,
            box_y + 0.48,
            box_w - 0.16,
            box_h - 0.55,
            body,
            font="Helvetica",
            size=11,
            color=BLACK,
            align=PP_ALIGN.CENTER,
            line_spacing=1.25,
        )
        if i < n - 1:
            add_arrow(
                s,
                x + box_w + gap / 2 - 0.15,
                box_y + box_h / 2 - 0.15,
                0.30,
                0.30,
                color=hex_to_rgb("#666666"),
            )
        x += box_w + gap

    add_callout(
        s,
        0.45,
        4.05,
        SLIDE_W - 0.9,
        1.55,
        [
            (
                "Canonical chain (shipped): ",
                "System → Run → Evaluation → Decision → Action. "
                "Replay, feedback, and canonical answers close the loop.",
            ),
            (
                "Positioning: ",
                "You don't build isolated agents — you compose intelligence into "
                "systems that produce measurable, governed value.",
            ),
        ],
        accent_hex=AGENTIUM_PRIMARY,
    )
    return s


def slide_product_overview(prs):
    """Full-width product screenshot — agentium.papai.ai live capture."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "Product")
    add_title(
        s,
        "Agentium in production — agentium.papai.ai",
        "Hypervisor cockpit: portfolio view, impact metrics, proactive recommendations",
    )
    img = shot("01_hypervisor.png")
    if img:
        add_screenshot_frame(
            s, img, 0.45, 2.05, SLIDE_W - 0.9, 4.35,
            caption="Live capture — Agentium Showcase workspace",
        )
    else:
        add_callout(s, 0.45, 2.5, SLIDE_W - 0.9, 2.0,
                    [("Note: ", "Run capture-pih-screenshots.mjs to refresh product screenshots.")],
                    accent_hex=AGENTIUM_ACCENT)
    return s


def slide_screenshot_gallery(prs):
    """2×3 grid of key UI surfaces captured from production."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "UI gallery")
    add_title(
        s,
        "Key surfaces — captured from agentium.papai.ai",
        "Each screen maps to a PIH pillar in the following slides",
    )
    tiles = [
        ("04_systems.png", "Systems\n(agent identity)"),
        ("08_run_detail.png", "Run detail\n(stateful execution)"),
        ("03_governance_audit.png", "Audit log\n(governance)"),
        ("05_observability.png", "Observability\n(SLO metrics)"),
        ("02_steering.png", "Control Plane\n(model guardrails)"),
        ("06_chat.png", "Chat + stream\n(token orchestration)"),
    ]
    cols, rows = 3, 2
    margin_x, margin_y = 0.40, 2.05
    gap_x, gap_y = 0.18, 0.22
    tile_w = (SLIDE_W - 2 * margin_x - (cols - 1) * gap_x) / cols
    tile_h = (SLIDE_H - margin_y - 0.55 - rows * gap_y) / rows - 0.05
    for idx, (fname, label) in enumerate(tiles):
        r, c = divmod(idx, cols)
        x = margin_x + c * (tile_w + gap_x)
        y = margin_y + r * (tile_h + gap_y)
        img = shot(fname)
        if img:
            add_screenshot_frame(s, img, x, y, tile_w, tile_h, caption=label.replace("\n", " · "))
        else:
            add_rect(s, x, y, tile_w, tile_h, fill=hex_to_rgb("#F5F5F5"), line=hex_to_rgb("#CCC"))
            add_textbox(s, x, y + tile_h / 2 - 0.2, tile_w, 0.4, label,
                        font="Helvetica", size=10, color=GREY, align=PP_ALIGN.CENTER)
    return s


def slide_pih_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "PIH requirements")
    add_title(
        s,
        "PIH enterprise pillars",
        "What regulated autonomous-agent platforms must guarantee",
    )

    rows = [
        (
            "Sovereignty",
            "Deploy and operate without mandatory external dependency; "
            "data and models under customer control.",
        ),
        (
            "Robustness",
            "Replay, resubmission, and stateful execution — failures are "
            "recoverable, not silent.",
        ),
        (
            "Security",
            "Model guardrails, distinct agent identity, strict RBAC for "
            "autonomous agents.",
        ),
        (
            "Observability & Scalability",
            "Full auditability of agent tool-calling; enterprise-grade scaling.",
        ),
        (
            "Token Performance",
            "Efficient, deterministic sovereign model orchestration.",
        ),
    ]
    y = 2.0
    for i, (pillar, desc) in enumerate(rows):
        color = PILLARS[i][1]
        add_rect(s, 0.45, y, 2.35, 0.72, fill=hex_to_rgb(color))
        add_textbox(
            s,
            0.45,
            y + 0.08,
            2.35,
            0.56,
            pillar,
            font="Helvetica",
            size=12,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            line_spacing=1.15,
        )
        add_rect(
            s,
            2.90,
            y,
            SLIDE_W - 3.35,
            0.72,
            fill=WHITE,
            line=hex_to_rgb("#CCCCCC"),
        )
        add_textbox(
            s,
            3.05,
            y + 0.08,
            SLIDE_W - 3.65,
            0.56,
            desc,
            font="Helvetica",
            size=11,
            color=BLACK,
            anchor=MSO_ANCHOR.MIDDLE,
            line_spacing=1.20,
        )
        y += 0.82

    add_callout(
        s,
        0.45,
        SLIDE_H - 1.05,
        SLIDE_W - 0.9,
        0.60,
        [
            (
                "Next slides: ",
                "For each pillar — PIH ask → Agentium answer → what you will see in the demo.",
            ),
        ],
        accent_hex=AGENTIUM_ACCENT,
    )
    return s


def slide_alignment_matrix(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "Alignment matrix")
    add_title(
        s,
        "PIH requirements → Agentium capabilities",
        "At-a-glance mapping before the live walkthrough",
    )

    rows = [
        (
            "Sovereignty",
            "On-prem / hybrid / sovereign cloud",
            "Ollama & Azure providers, Qdrant, MinIO — no cloud lock-in",
            "✅ Shipped",
        ),
        (
            "Robustness",
            "Replay, resubmission, stateful runs",
            "Run lineage, /runs/{id}/replay, skill_invocations ledger, DAG HITL resume",
            "✅ Shipped",
        ),
        (
            "Security",
            "Guardrails, agent identity, RBAC",
            "ControlPolicy, claim_audit_v1, Keycloak workspaces, audit on every mutation",
            "✅ Shipped",
        ),
        (
            "Observability",
            "Tool-call auditability, scale",
            "/audit, /observability, SSE stream, per-skill cost & latency",
            "✅ Shipped",
        ),
        (
            "Token performance",
            "Deterministic model orchestration",
            "allowed_models routing, Flow/DAG engine, token_delta streaming, budget caps",
            "✅ Shipped",
        ),
    ]

    cols = [2.0, 2.55, 5.55, 1.35]
    headers = ["PIH pillar", "Requirement", "Agentium capability", "Status"]
    x0 = 0.35
    y = 2.0
    h_head = 0.42
    h_row = 0.78
    accent = hex_to_rgb(AGENTIUM_PRIMARY)

    cur_x = x0
    for i, hw in enumerate(headers):
        add_rect(s, cur_x, y, cols[i], h_head, fill=accent)
        add_textbox(
            s,
            cur_x,
            y,
            cols[i],
            h_head,
            hw,
            font="Helvetica",
            size=10,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        cur_x += cols[i]
    y += h_head

    for row in rows:
        cur_x = x0
        for i, cell in enumerate(row):
            bg = hex_to_rgb(AGENTIUM_LIGHT) if i == 0 else WHITE
            add_rect(
                s,
                cur_x,
                y,
                cols[i],
                h_row,
                fill=bg,
                line=hex_to_rgb("#CCCCCC"),
            )
            bold = i in (0, 3)
            color = hex_to_rgb(AGENTIUM_GREEN) if i == 3 else BLACK
            add_textbox(
                s,
                cur_x + 0.06,
                y + 0.06,
                cols[i] - 0.12,
                h_row - 0.12,
                cell,
                font="Helvetica",
                size=9 if i == 2 else 10,
                bold=bold,
                color=color,
                anchor=MSO_ANCHOR.MIDDLE,
                line_spacing=1.15,
            )
            cur_x += cols[i]
        y += h_row

    return s


def _pillar_slide(
    prs,
    pillar: str,
    color_hex: str,
    pih_ask: str,
    bullets: list,
    demo: str,
    *,
    screenshot: str | None = None,
    screenshot_caption: str = "",
    show_boundary: bool = True,
):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, pillar)
    add_title(s, f"{pillar} — PIH alignment", pih_ask)

    left_w = 5.15 if screenshot else (SLIDE_W - 1.0) / 2 - 0.12
    content_h = 4.15 if screenshot else 3.55
    content_y = 2.05 if screenshot else 2.55

    add_rect(
        s,
        0.45,
        content_y,
        left_w,
        content_h,
        fill=hex_to_rgb("#FAFAFA"),
        line=hex_to_rgb(color_hex),
        line_w=1.2,
    )
    add_textbox(
        s,
        0.60,
        content_y + 0.10,
        left_w - 0.30,
        0.35,
        "How Agentium delivers",
        font="Helvetica",
        size=13,
        bold=True,
        color=hex_to_rgb(color_hex),
    )
    body = "\n".join(f"•  {b}" for b in bullets)
    add_textbox(
        s,
        0.60,
        content_y + 0.48,
        left_w - 0.30,
        content_h - 0.58,
        body,
        font="Helvetica",
        size=10 if screenshot else 11,
        color=BLACK,
        line_spacing=1.30,
    )

    if screenshot and os.path.isfile(screenshot):
        right_x = 0.45 + left_w + 0.20
        right_w = SLIDE_W - right_x - 0.45
        cap = screenshot_caption or demo.replace("\n", " · ")
        add_screenshot_frame(
            s, screenshot, right_x, content_y, right_w, content_h, caption=cap,
        )
    elif not screenshot:
        right_x = 0.45 + left_w + 0.24
        add_rect(
            s,
            right_x,
            content_y,
            left_w,
            content_h,
            fill=hex_to_rgb(AGENTIUM_LIGHT),
            line=hex_to_rgb(AGENTIUM_ACCENT),
            line_w=1.2,
        )
        add_textbox(
            s,
            right_x + 0.15,
            content_y + 0.10,
            left_w - 0.30,
            0.35,
            "In the demo",
            font="Helvetica",
            size=13,
            bold=True,
            color=hex_to_rgb(AGENTIUM_ACCENT),
        )
        add_textbox(
            s,
            right_x + 0.15,
            content_y + 0.48,
            left_w - 0.30,
            content_h - 0.58,
            demo,
            font="Helvetica",
            size=11,
            color=BLACK,
            line_spacing=1.35,
        )

    if show_boundary:
        add_callout(
            s,
            0.45,
            SLIDE_H - 0.88,
            SLIDE_W - 0.9,
            0.48,
            [
                (
                    "Honest boundary: ",
                    "Advanced custom RBAC roles and checkpoint resume tokens are on the "
                    "roadmap; core governance, replay, and audit paths are shipped today.",
                ),
            ],
            accent_hex=color_hex,
        )
    return s


def slide_sovereignty(prs):
    return _pillar_slide(
        prs,
        "Sovereignty",
        AGENTIUM_PRIMARY,
        "Operate intelligent systems under customer control — on-prem, hybrid, or sovereign cloud.",
        [
            "Deploy fully on-prem: local LLM (Ollama / vLLM), Qdrant vectors, MinIO storage.",
            "Hybrid mode: keep data local while optionally routing to approved cloud models.",
            "Workspace isolation — multi-tenant boundaries tested in production paths.",
            "No mandatory external SaaS: every skill, run, and policy stays in your perimeter.",
            "Same architecture philosophy as papAI — convergence-ready enterprise stack.",
        ],
        "Workspace-scoped Systems with local model providers.",
        screenshot=shot("04_systems.png"),
        screenshot_caption="Systems — distinct agent identity per workspace",
        show_boundary=False,
    )


def slide_robustness(prs):
    return _pillar_slide(
        prs,
        "Robustness",
        AGENTIUM_GREEN,
        "Replay, resubmission, and stateful execution — autonomous agents must recover, not fail silently.",
        [
            "Every Run is a durable execution record with status, input_ref, outcome, lineage.",
            "POST /runs/{id}/replay — resubmit with overrides (model, prompt, RAG mode, tokens).",
            "parent_run_id + replay_overrides preserve full audit lineage.",
            "skill_invocations[] ledger = per-step checkpoint (cost, latency, I/O).",
            "DAG engine: retry with backoff, HITL pause/resume, parallel branches.",
            "Evaluation → Decision → replay closes the quality loop without manual re-entry.",
        ],
        "Inspect skill invocation trail and replay lineage.",
        screenshot=shot("08_run_detail.png"),
        screenshot_caption="Run detail — skill ledger, cost, latency, lineage",
        show_boundary=False,
    )


def slide_security(prs):
    return _pillar_slide(
        prs,
        "Security",
        "#7C3AED",
        "Model guardrails, distinct agent identity, and strict RBAC for autonomous agents.",
        [
            "ControlPolicy: allowed_models, max_cost, max_latency, mandatory HITL thresholds.",
            "claim_audit_v1 — faithfulness guardrails on generated outputs.",
            "Each System = distinct agent identity (objective, capability, flow, policies).",
            "Keycloak / OIDC authentication; workspace members with built-in roles.",
            "Context permissions JSON + audit log on every mutation (actor, target, diff).",
            "AdaptivePolicy gates autonomous actions — no silent policy bypass.",
        ],
        "Governance → Audit log + Access & Roles.",
        screenshot=shot("03_governance_audit.png"),
        screenshot_caption="Audit log — actor, action, target, full traceability",
        show_boundary=False,
    )


def slide_observability(prs):
    return _pillar_slide(
        prs,
        "Observability & Scalability",
        AGENTIUM_AMBER,
        "Full auditability of agent tool-calling and enterprise-grade scaling.",
        [
            "Structured /audit stream — every mutation, upload, and policy change logged.",
            "skill_invocations persist full tool-call ledger: inputs, outputs, status, metrics.",
            "SSE /runs/{id}/stream — token_delta + structural node events in real time.",
            "/observability surfaces SLO: cost, latency, success rate, live telemetry.",
            "Async execution via background workers; Postgres-backed durable state.",
            "Horizontal scaling path: Celery runtime alignment for batch & event workloads.",
        ],
        "Observability dashboard + Hypervisor portfolio metrics.",
        screenshot=shot("05_observability.png"),
        screenshot_caption="Observability — SLO metrics and live telemetry",
        show_boundary=False,
    )


def slide_token_performance(prs):
    return _pillar_slide(
        prs,
        "Token Performance",
        AGENTIUM_ACCENT,
        "Efficient and deterministic sovereign model orchestration.",
        [
            "Control Plane routes to allowed_models only — no rogue model selection.",
            "Flow / DAG engine executes skills in declared order (sequential, parallel, gated).",
            "Per-skill cost_per_call + Run-level budget caps prevent token sprawl.",
            "token_delta SSE streaming — efficient UX without polling overhead.",
            "Deterministic paths: canonical answers bypass LLM when approved response exists.",
            "Impact preview (<300ms) before applying policy levers — predictable trade-offs.",
        ],
        "Steering Control Plane + chat token stream.",
        screenshot=shot("02_steering.png"),
        screenshot_caption="Control Plane — allowed models, cost & latency guardrails",
        show_boundary=True,
    )


def slide_rbac_and_chat(prs):
    """Supplementary slide: RBAC + chat streaming."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "Security & execution")
    add_title(
        s,
        "RBAC perimeter + live agent execution",
        "Access control before execution · token stream during run",
    )
    half = (SLIDE_W - 1.0) / 2 - 0.08
    y, h = 2.05, 4.35
    img_access = shot("07_access_roles.png")
    img_chat = shot("06_chat.png")
    if img_access:
        add_screenshot_frame(
            s, img_access, 0.45, y, half, h,
            caption="Access & Roles — workspace RBAC (Keycloak / OIDC)",
        )
    if img_chat:
        add_screenshot_frame(
            s, img_chat, 0.45 + half + 0.16, y, half, h,
            caption="Chat — grounded answers with streaming token orchestration",
        )
    return s


def slide_demo_bridge(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_agentium_header(s, "Live demo")
    add_title(
        s,
        "From requirements to proof — demo flow",
        "Five minutes to connect PIH pillars to what you will see on screen",
    )

    steps = [
        ("1", "Login & workspace", "RBAC, sovereign perimeter"),
        ("2", "System & Control Plane", "Agent identity + model guardrails"),
        ("3", "Run + stream", "Stateful execution + token orchestration"),
        ("4", "Skill ledger & audit", "Full tool-call traceability"),
        ("5", "Evaluation → Replay", "Robustness loop in action"),
    ]
    n = len(steps)
    margin = 0.55
    box_w = 2.25
    gap = (SLIDE_W - 2 * margin - n * box_w) / (n - 1)
    box_h = 2.35
    y = 2.20
    x = margin
    for i, (num, title, sub) in enumerate(steps):
        add_rect(
            s,
            x,
            y,
            box_w,
            box_h,
            fill=hex_to_rgb(AGENTIUM_LIGHT),
            line=hex_to_rgb(AGENTIUM_ACCENT),
            line_w=1.3,
        )
        sh = s.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + box_w / 2 - 0.28),
            Inches(y + 0.20),
            Inches(0.56),
            Inches(0.56),
        )
        sh.shadow.inherit = False
        sh.fill.solid()
        sh.fill.fore_color.rgb = hex_to_rgb(AGENTIUM_ACCENT)
        sh.line.fill.background()
        add_textbox(
            s,
            x + box_w / 2 - 0.28,
            y + 0.28,
            0.56,
            0.40,
            num,
            font="Helvetica",
            size=18,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        add_textbox(
            s,
            x + 0.10,
            y + 0.90,
            box_w - 0.20,
            0.55,
            title,
            font="Helvetica",
            size=12,
            bold=True,
            color=hex_to_rgb(AGENTIUM_PRIMARY),
            align=PP_ALIGN.CENTER,
        )
        add_textbox(
            s,
            x + 0.10,
            y + 1.50,
            box_w - 0.20,
            0.70,
            sub,
            font="Helvetica",
            size=10,
            color=hex_to_rgb("#444444"),
            align=PP_ALIGN.CENTER,
            line_spacing=1.25,
        )
        if i < n - 1:
            add_arrow(
                s,
                x + box_w + gap / 2 - 0.15,
                y + box_h / 2 - 0.15,
                0.30,
                0.30,
                color=hex_to_rgb("#999999"),
            )
        x += box_w + gap

    add_callout(
        s,
        0.45,
        4.85,
        SLIDE_W - 0.9,
        1.05,
        [
            (
                "Key message: ",
                "Agentium is not an agent playground — it is the control layer that makes "
                "autonomous systems sovereign, recoverable, secure, observable, and token-efficient.",
            ),
            (
                "Let's demo. ",
                "Questions welcome throughout — we will map each screen back to the PIH pillars.",
            ),
        ],
        accent_hex=AGENTIUM_PRIMARY,
    )
    return s


def slide_closing(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SLIDE_W, 0.55, fill=hex_to_rgb(AGENTIUM_PRIMARY))
    if os.path.isfile(DATATEGY_LOGO):
        s.shapes.add_picture(DATATEGY_LOGO, Inches(2.8), Inches(1.35), height=Inches(1.1))

    add_textbox(
        s,
        0.5,
        2.75,
        SLIDE_W - 1.0,
        0.75,
        "Ready for the live demo",
        font="Helvetica",
        size=36,
        bold=True,
        color=BLACK,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        0.5,
        3.55,
        SLIDE_W - 1.0,
        0.55,
        "Agentium — governed intelligent systems at enterprise scale",
        font="Helvetica",
        size=18,
        color=hex_to_rgb(AGENTIUM_PRIMARY),
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        0.5,
        4.25,
        SLIDE_W - 1.0,
        0.90,
        "Sovereign  ·  Robust  ·  Secure  ·  Observable  ·  Token-efficient",
        font="Helvetica",
        size=14,
        color=GREY,
        align=PP_ALIGN.CENTER,
        line_spacing=1.30,
    )
    add_textbox(
        s,
        0.5,
        SLIDE_H - 0.85,
        SLIDE_W - 1.0,
        0.35,
        "Datategy — Confidential",
        font="Helvetica",
        size=11,
        color=GREY,
        align=PP_ALIGN.CENTER,
    )
    return s


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide_cover(prs)
    slide_agenda(prs)
    slide_mental_model(prs)
    slide_product_overview(prs)
    slide_pih_overview(prs)
    slide_alignment_matrix(prs)
    slide_screenshot_gallery(prs)
    slide_sovereignty(prs)
    slide_robustness(prs)
    slide_security(prs)
    slide_observability(prs)
    slide_token_performance(prs)
    slide_rbac_and_chat(prs)
    slide_demo_bridge(prs)
    slide_closing(prs)

    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        if i == 1 or i == total:
            continue
        add_footer(slide, i, total)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(f"PIH alignment PPTX written: {OUT}")
    print(f"   {total} slides")


if __name__ == "__main__":
    build()
