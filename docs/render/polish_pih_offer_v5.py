# -*- coding: utf-8 -*-
"""Polish the PIH offer v4 -> v5: pure visual/formatting pass, zero wording change.

Root cause of the "wall of text" effect: paragraphs inserted by the iterate_pih_offer*.py
scripts (via python-docx `doc.add_paragraph()`/new `<w:p>` + `.style` assignment) never
carried the *direct* paragraph formatting (space before/after, line spacing, hanging
indent) that the original Google-Docs-exported paragraphs have on their own `<w:pPr>`.
The "normal" style itself defines no direct formatting either, so those paragraphs fall
back to bare defaults (0pt spacing, single line height, no indent) while the untouched
original paragraphs already look correct (12pt before/after, 1.15 line spacing; list
items additionally carry a hanging indent).

Fixes applied (formatting only — no text is changed, removed or reworded):
  1. Every body ("normal"-style) paragraph missing direct spacing gets the same spacing
     as the rest of the document: 12pt before, 1.15 line spacing, and 12pt after (0pt
     after + hanging indent for "•"-prefixed bullet paragraphs, matching the existing
     bullet convention already used elsewhere in the doc).
  2. A handful of already-correctly-indented list items that were written as plain
     one-per-line sentences (no bullet glyph at all — e.g. "Authentication handling",
     "Agent management.") get a leading "•  " marker, or a "1. " / "2. " ... numbered
     marker for the two sequential lists (the Teams walkthrough example, the factory
     lifecycle steps), matching the bullet convention used everywhere else in the offer.
     This only prepends a typographic marker; no word is added, removed or changed.

Input : docs/pih/Service Help Desk - Technical and Commercial offer - v4 Datategy.docx
Output: docs/pih/Service Help Desk - Technical and Commercial offer - v5 Datategy.docx
"""
import os
from docx import Document
from docx.shared import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
SRC = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v4 Datategy.docx")
OUT = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v5 Datategy.docx")

SPACE = Emu(152400)      # 12pt, matches the original document's own paragraph spacing
ZERO = Emu(0)
LEFT_INDENT = Emu(457200)      # matches the original document's own bullet hanging indent
FIRST_LINE = Emu(-228600)

doc = Document(SRC)

# capture full text before edits, to verify at the end that no word was touched
BEFORE_TEXT = [p.text for p in doc.paragraphs]

# ============================================================ 1. universal spacing fix
fixed = 0
for p in doc.paragraphs:
    if not p.style or p.style.name != "normal":
        continue
    if not p.text.strip():
        continue
    pf = p.paragraph_format
    if pf.space_before is not None:
        continue  # already explicitly formatted (original content) — leave untouched
    is_bullet = p.text.lstrip().startswith("•")
    pf.space_before = SPACE
    pf.line_spacing = 1.15
    if is_bullet:
        pf.space_after = ZERO
        pf.left_indent = LEFT_INDENT
        pf.first_line_indent = FIRST_LINE
    else:
        pf.space_after = SPACE
    fixed += 1
print(f"spacing fixed on {fixed} paragraphs")


# ============================================================ 2. marker glyphs on plain list items
def add_prefix(idx, expected_start, prefix):
    p = doc.paragraphs[idx]
    actual = p.text.strip()
    if not actual.startswith(expected_start):
        raise SystemExit(f"MISMATCH at paragraph {idx}: expected {expected_start!r}, got {actual[:60]!r}")
    if actual.startswith("•") or actual[0].isdigit():
        return  # already marked
    if p.runs:
        p.runs[0].text = prefix + p.runs[0].text
    else:
        p.add_run(prefix)


BULLET_RANGES = [
    (80, 87, ["Authentication handling", "API call abstraction", "Error handling", "Retry logic",
              "Logging", "Permission validation", "Response normalization", "Security controls"]),
    (91, 99, ["Agent management.", "Use-case configuration.", "Connector status monitoring.",
              "Workflow execution monitoring.", "Logs and audit trail visibility.",
              "Manual override or retry where authorized.", "Agent activation/deactivation.",
              "Parameter management.", "Basic operational dashboards."]),
    (104, 114, ["Agent definitions.", "Use-case configurations.", "Connector configuration metadata.",
                "Execution logs.", "Workflow states.", "Approval states.", "Audit records.",
                "Error records.", "User interaction metadata.", "API transaction history.",
                "Operational parameters."]),
    (119, 128, ["Intent detection.", "Request classification.", "User question answering.",
                "Knowledge search.", "Ticket summarization.", "Suggested resolution.",
                "Automated response generation.", "Agent reasoning.", "Multi-step workflow guidance.",
                "Translation or multilingual support if needed."]),
    (130, 136, ["Data protection.", "Prompt validation.", "Output validation.",
                "Restricted action boundaries.", "Hallucination reduction.",
                "Human-in-the-loop controls for sensitive actions.", "Logging of AI decisions."]),
    (140, 149, ["Ticket creation.", "Ticket enrichment.", "Ticket categorization.", "Ticket assignment.",
                "Ticket status tracking.", "SLA visibility.", "Ticket closure recommendations.",
                "User communication history.", "Linking automation actions to ticket records.",
                "Maintaining auditability of automated actions."]),
    (212, 218, ["Azure AD / Entra ID", "Microsoft Teams", "Email / Exchange / Microsoft 365", "Egnyte",
                "SolarWinds ARM", "Palo Alto / Fortinet Firewalls", "ManageEngine ServiceDesk Plus"]),
]
n_bullet = 0
for start, end, expected in BULLET_RANGES:
    for offset, idx in enumerate(range(start, end + 1)):
        add_prefix(idx, expected[offset], "•  ")
        n_bullet += 1

NUMBERED_RANGES = [
    (152, 163, ["User submits a request through Microsoft Teams.", "Agentium receives the request.",
                "The relevant AI agent is triggered.",
                "The run engine opens a stateful, checkpointed run and analyzes the request.",
                "The LLM identifies the user intent.",
                "The backend checks policy, authorization, and required approvals.",
                "If needed, a ticket is created or updated in ManageEngine.",
                "The backend calls the relevant connector, such as Azure AD or email.",
                "The action is executed or routed for approval.",
                "Logs and execution status are stored in PostgreSQL.",
                "The user receives a response through Teams.",
                "The ITSM record remains traceable in ManageEngine."]),
    (232, 244, ["detailed design", "technical specification", "connector mapping", "security validation",
                "development", "unit testing", "integration testing", "user acceptance testing",
                "deployment preparation", "production release", "monitoring", "stabilization",
                "Documentation"]),
]
n_numbered = 0
for start, end, expected in NUMBERED_RANGES:
    for offset, idx in enumerate(range(start, end + 1)):
        add_prefix(idx, expected[offset], f"{offset + 1}.  ")
        n_numbered += 1

print(f"bullet markers added: {n_bullet}, numbered markers added: {n_numbered}")

doc.save(OUT)

# ============================================================ sanity: no wording lost
d2 = Document(OUT)
AFTER_TEXT = [p.text for p in d2.paragraphs]


def strip_marker(t):
    t = t.strip()
    if t.startswith("•"):
        t = t.lstrip("•").strip()
    parts = t.split(".", 1)
    if len(parts) == 2 and parts[0].strip().isdigit():
        t = parts[1].strip()
    return t


assert len(BEFORE_TEXT) == len(AFTER_TEXT), "paragraph count changed!"
mismatches = 0
for i, (b, a) in enumerate(zip(BEFORE_TEXT, AFTER_TEXT)):
    if strip_marker(b) != strip_marker(a):
        mismatches += 1
        print(f"WORDING CHANGED at {i}: {b[:60]!r} -> {a[:60]!r}")
print(f"wording mismatches: {mismatches} (must be 0)")
print("saved:", OUT)
