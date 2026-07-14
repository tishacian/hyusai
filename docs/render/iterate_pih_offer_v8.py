# -*- coding: utf-8 -*-
"""Iterate on the PIH offer v7 -> v8: reinstate the 3-year TCO (RFP §6.4 / §9.3).

The negotiated table had dropped the TCO row, but the RFP explicitly requires a
3-year TCO estimate ("Must include a 3-year Total Cost of Ownership (TCO)
estimate", §9.3). Reinstated:
  - a TCO row in the pricing summary table: Option 1 ≈ USD 125k (50k build +
    3×25k run) · Option 2 ≈ USD 300k (60k build + 3×80k run);
  - the Option 2 prose claim, defensible with the confirmed incumbent baseline
    (~108k/yr IT+HR): 3-year TCO below current combined spend.

The added table row copies the styling of an existing row (deepcopy of the
<w:tr>) to avoid the unstyled-row bug seen previously with table.add_row().

Input : docs/pih/Service Help Desk - Technical and Commercial offer - v7 Datategy.docx
Output: docs/pih/Service Help Desk - Technical and Commercial offer - v8 Datategy.docx
"""
import copy
import os
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
SRC = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v7 Datategy.docx")
OUT = os.path.join(PIH, "Service Help Desk - Technical and Commercial offer - v8 Datategy.docx")

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

doc = Document(SRC)


def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise SystemExit(f"ANCHOR NOT FOUND: {prefix!r}")


def set_text(p, text, bold_lead=None):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


# ---- prose: Option 2 TCO claim
p_runfee = find("The Option 2 run fee is USD 80k per year")
set_text(p_runfee,
    "The Option 2 run fee is USD 80k per year and covers the full IT + Employee Services perimeter: "
    "platform licence, editor support and golden-set regression across both scopes, with day-to-day "
    "operations internalised in PIH's team for IT and employee services alike. Option 2 remains the "
    "recommended path: one front door for employees across IT and HR questions, the full "
    "replacement of the current conversational stack, a 3-year total cost of ownership "
    "(approximately USD 300k, build plus three years of run) designed to remain below the "
    "organisation's current combined conversational-platform spend across IT and employee "
    "services — and the conversion of rented workflows into an owned, governed, agentic platform "
    "whose returned analyst-hours are measured live in the value cockpit.")

# ---- table: add TCO row styled like an existing row
tbl = None
for t in doc.tables:
    if t.cell(0, 0).text.strip() == "Component":
        tbl = t
        break
if tbl is None:
    raise SystemExit("PRICING TABLE NOT FOUND")

# find the "Total build" row (4 distinct cells, right styling) to clone
src_row = None
for r in tbl.rows:
    if r.cells[0].text.strip().startswith("Total build"):
        src_row = r
        break
if src_row is None:
    raise SystemExit("TOTAL BUILD ROW NOT FOUND")

new_tr = copy.deepcopy(src_row._tr)
tbl._tbl.append(new_tr)
new_row = tbl.rows[-1]

VALUES = [
    "Indicative 3-year TCO (build + 3× run)",
    "≈ USD 125k",
    "≈ USD 300k",
    "—",
]
for cell, val in zip(new_row.cells, VALUES):
    # keep first paragraph & run formatting; replace text only
    para = cell.paragraphs[0]
    runs = para.runs
    if runs:
        runs[0].text = val
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        para.add_run(val)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)

doc.save(OUT)

# sanity
d2 = Document(OUT)
text = "\n".join(p.text for p in d2.paragraphs)
tbl2 = None
for t in d2.tables:
    if t.cell(0, 0).text.strip() == "Component":
        tbl2 = t
        break
last = [c.text.strip() for c in tbl2.rows[-1].cells]
print("last table row:", last)
print("prose TCO claim:", text.count("3-year total cost of ownership"))
print("USD 300k occurrences:", (text + " ".join(last)).count("USD 300k"))
print("saved:", OUT)
