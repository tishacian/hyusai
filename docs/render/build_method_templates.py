# -*- coding: utf-8 -*-
"""Datategy methodology templates (branded .docx, reusable, prospect-shareable).

Generates three official-process templates from the PIH working docs, genericised:
  1. Business Requirements                  -> Datategy-Template-Business-Requirements.docx
  2. Target Operating Model (TOM)           -> Datategy-Template-Target-Operating-Model.docx
  3. QA Acceptance Record (PV de recette)   -> Datategy-Template-QA-Acceptance-PV.docx
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "logo_datategy.png")
OUTDIR = os.path.join(HERE, "out")

def C(h):
    return RGBColor.from_string(h)

NAVY = C("0B2545"); BLUE = C("2D6CDF"); GREEN = C("0E8F62"); AMBER = C("B45309")
PURPLE = C("6D28D9"); INK = C("151A23"); MUT = C("596371"); WHITE = C("FFFFFF")
LINE = "D7DCE3"; PANEL = "F3F5F8"; FIELD = "FFF8E7"; TINT = "E5EDFB"
BODY = "Arial"


def _runfmt(run, size, color, bold=False, italic=False, mono=False):
    run.font.name = "Consolas" if mono else BODY
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = color


def _shade(cell, hexfill):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(sh)


def _margins(cell, t=60, b=60, l=110, r=90):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for tag, v in (("top", t), ("bottom", b), ("start", l), ("end", r), ("left", l), ("right", r)):
        e = OxmlElement(f"w:{tag}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def _borders(table, color=LINE, sz=4):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color); b.append(e)
    table._tbl.tblPr.append(b)


def _pbottom(p, color="C9D2DE", sz=6):
    pPr = p._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "4")
    e.set(qn("w:color"), color); pb.append(e); pPr.append(pb)


def _left_accent(cell, hexcolor, sz=30):
    tcPr = cell._tc.get_or_add_tcPr(); bd = OxmlElement("w:tcBorders")
    e = OxmlElement("w:left"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0"); e.set(qn("w:color"), hexcolor); bd.append(e); tcPr.append(bd)


def _pagefield(run, code="PAGE"):
    r = run._r
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "begin"); r.append(e)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code; r.append(it)
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); r.append(sep)
    t = OxmlElement("w:t"); t.text = "1"; r.append(t)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); r.append(end)


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY; r.font.name = BODY
    _pbottom(p); return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = BLUE; r.font.name = BODY
    return p


def para(doc, text, size=10.5, color=INK, bold=False, italic=False, before=3, after=4):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.14
    _runfmt(p.add_run(text), size, color, bold=bold, italic=italic); return p


def runs(doc, parts, before=3, after=4, size=10.5):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.14
    for text, bold, color in parts:
        _runfmt(p.add_run(text), size, color, bold=bold)
    return p


def bullet(doc, parts, size=10.5):
    p = doc.add_paragraph(style="List Bullet"); pf = p.paragraph_format
    pf.space_before = Pt(1); pf.space_after = Pt(1); pf.line_spacing = 1.12
    if isinstance(parts, str):
        parts = [(parts, False, INK)]
    for text, bold, color in parts:
        _runfmt(p.add_run(text), size, color, bold=bold)
    return p


def checklist(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(); pf = p.paragraph_format
        pf.space_before = Pt(1); pf.space_after = Pt(1); pf.line_spacing = 1.12
        pf.left_indent = Cm(0.4)
        _runfmt(p.add_run("☐  "), size + 1, MUT)
        if isinstance(it, tuple):
            head, body = it
            _runfmt(p.add_run(head + " — "), size, INK, bold=True)
            _runfmt(p.add_run(body), size, MUT)
        else:
            _runfmt(p.add_run(it), size, INK)


def field(doc, label, hint):
    """A fill-in line: Label: < hint to complete >."""
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(2); pf.space_after = Pt(2)
    _runfmt(p.add_run(label + " : "), 10.5, NAVY, bold=True)
    _runfmt(p.add_run("‹ " + hint + " ›"), 10, MUT, italic=True)


def guidance(doc, text):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0); cell.width = Cm(17); _shade(cell, PANEL)
    _left_accent(cell, "2D6CDF", sz=26); _margins(cell, t=70, b=70, l=160, r=130)
    cell.text = ""; p = cell.paragraphs[0]
    _runfmt(p.add_run("Guidance — "), 9, BLUE, bold=True, mono=True)
    _runfmt(p.add_run(text), 9.5, MUT, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def table(doc, headers, rows, widths=None, blank_rows=0, header_fill="0B2545"):
    n = len(headers)
    t = doc.add_table(rows=1, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.allow_autofit = False; _borders(t, LINE, 4)
    for j, htxt in enumerate(headers):
        c = t.cell(0, j); _shade(c, header_fill); _margins(c)
        c.text = ""; _runfmt(c.paragraphs[0].add_run(htxt), 9, WHITE, bold=True)
    all_rows = list(rows) + [[""] * n for _ in range(blank_rows)]
    for i, row in enumerate(all_rows):
        r = t.add_row(); fill = "FFFFFF" if i % 2 == 0 else PANEL
        for j, val in enumerate(row):
            c = r.cells[j]; _shade(c, fill if val else FIELD); _margins(c)
            c.text = ""; pp = c.paragraphs[0]; pp.paragraph_format.line_spacing = 1.1
            pp.paragraph_format.space_before = Pt(1); pp.paragraph_format.space_after = Pt(1)
            _runfmt(pp.add_run(val), 9, INK if val else MUT, bold=(j == 0 and bool(val)))
    if widths:
        for r in t.rows:
            for j, wv in enumerate(widths):
                r.cells[j].width = Cm(wv)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def signoff(doc, roles, decision_col=False):
    headers = ["Role", "Name"] + (["Decision"] if decision_col else []) + ["Date", "Signature"]
    rows = [[r, "", *(["" ] if decision_col else []), "", ""] for r in roles]
    widths = ([5.2, 4.2, 3.0, 2.3, 2.3] if decision_col else [5.6, 4.6, 3.4, 3.4])
    table(doc, headers, rows, widths=widths, header_fill="0B2545")


# ---------------------------------------------------------------- doc shell
def new_doc(short_title):
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = BODY; normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK; normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.14
    for lvl, sz, col in (("Heading 1", 15, NAVY), ("Heading 2", 12, BLUE)):
        st = doc.styles[lvl]; st.font.name = BODY; st.font.size = Pt(sz); st.font.bold = True
        st.font.color.rgb = col; st.paragraph_format.space_before = Pt(12 if lvl == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(4); st.paragraph_format.keep_with_next = True
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(1.9)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)
    sec.different_first_page_header_footer = True
    # header: small logo left + short title right
    hdr = sec.header; hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hp.add_run().add_picture(LOGO, height=Cm(0.6))
    hp.add_run("    ")
    r = hp.add_run(short_title); _runfmt(r, 8.5, MUT, italic=True)
    _pbottom(hp, LINE, 4)
    # footer: methodology label left, page right (borderless table)
    f = sec.footer; f.is_linked_to_previous = False
    ft = f.add_table(rows=1, cols=2, width=Cm(17)); ft.allow_autofit = False
    lc = ft.cell(0, 0); rc = ft.cell(0, 1); lc.width = Cm(12.5); rc.width = Cm(4.5)
    _runfmt(lc.paragraphs[0].add_run("Datategy · Agentium Delivery Methodology · Template"), 8, MUT)
    rp = rc.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _runfmt(rp.add_run("Page "), 8, MUT); _pagefield(rp.add_run())
    _empty = f.paragraphs[0]; _empty._element.getparent().remove(_empty._element)
    return doc


def cover(doc, kicker, title, subtitle, ref):
    for _ in range(2):
        doc.add_paragraph()
    cov = doc.add_paragraph(); cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cov.add_run().add_picture(LOGO, width=Cm(4.6))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(26)
    _runfmt(p.add_run(kicker), 11, BLUE, bold=True, mono=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10)
    _runfmt(p.add_run(title), 30, NAVY, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runfmt(p.add_run(subtitle), 13, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    # document-control table
    t = doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.allow_autofit = False; _borders(t, LINE, 4)
    ctrl = [("Template", title), ("Reference", ref), ("Version", "v1.0  ‹ adapt ›"),
            ("Owner", "Datategy — Delivery"), ("Prepared for", "‹ Client / Prospect ›"),
            ("Methodology", "ISO/IEC 42001 & 12792 (SC 42) · structured business-analysis practice · Two-Layer Agent-Ready Requirements"),
            ("Date", "‹ dd/mm/yyyy ›"), ("Classification", "Shared — methodology template")]
    for i, (a, b) in enumerate(ctrl):
        r = t.add_row(); fill = "FFFFFF" if i % 2 == 0 else PANEL
        for j, (val, col, bold, w) in enumerate([(a, NAVY, True, 5.0), (b, MUT, False, 12.0)]):
            c = r.cells[j]; _shade(c, fill); _margins(c); c.width = Cm(w)
            c.text = ""; _runfmt(c.paragraphs[0].add_run(val), 9.5, col, bold=bold)
    note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(16)
    _runfmt(note.add_run("Aligned with ISO/IEC 42001 (AI management systems) and ISO/IEC 12792 (transparency of AI systems), "
                         "ISO/IEC JTC 1/SC 42, and recognised business-analysis requirements practice. "), 9, MUT, italic=True)
    _runfmt(note.add_run("Datategy is a member of the AFNOR AI standardisation committee (CN IA — the French national "
                         "mirror of ISO/IEC JTC 1/SC 42) that contributes to drafting these standards."), 9, NAVY, italic=True, bold=True)
    doc.add_page_break()


def how_to_use(doc, text):
    h2(doc, "How to use this template")
    guidance(doc, text)
    para(doc, "Fields shown as ‹ … › are to be completed. Amber cells are blank rows to fill. "
              "Keep the structure — it maps the business need to an Agentium System (Skill → Capability → "
              "System → Suite), so architects and engineers can build directly from it.",
         size=9.5, color=MUT, italic=True)


def save(doc, name):
    os.makedirs(OUTDIR, exist_ok=True)
    out = os.path.join(OUTDIR, name); doc.save(out)
    print("saved:", os.path.relpath(out, HERE))


# ================================================================ 1. BUSINESS REQUIREMENTS
def build_business_requirements():
    doc = new_doc("Business Requirements — Template")
    cover(doc, "DELIVERY METHODOLOGY · TEMPLATE", "Business Requirements",
          "Capturing the business need for an AI system delivered on Agentium",
          "DTG-METH-BR-v1.0")
    how_to_use(doc, "Capture what the business needs and why — outcomes, functional & non-functional "
                    "requirements, data, KPIs and constraints. Capture enough to design; do not over-specify "
                    "(the architect derives the solution).")

    h1(doc, "1.  Context & objective")
    field(doc, "Client / entity", "name, sector, size")
    field(doc, "Sponsor & business owner", "who owns the outcome")
    field(doc, "System / agent in scope", "the job this AI system should take off the team's hands")
    field(doc, "Current situation", "how it is done today, pain points, volumes")

    h1(doc, "2.  Primary persona")
    guidance(doc, "The person the agent serves. Trust posture is decisive for adoption — capture what would "
                  "make this persona trust (or distrust) the agent.")
    field(doc, "Persona & volume", "role · how many people")
    field(doc, "Current tools", "systems used today")
    field(doc, "Job-to-be-done", "‹ When [situation], I want [capability], so I can [outcome] ›")
    field(doc, "AI literacy", "low / medium / high")
    field(doc, "Trust posture", "what makes them trust / distrust the agent (e.g. every claim sourced; no invented context)")

    h1(doc, "3.  Process: As-Is → To-Be")
    guidance(doc, "Capture today's steps (with time & pain) and the target flow showing which steps the agent "
                  "performs and where the human stays in the loop. Default: the agent acts autonomously; the human "
                  "integrates as quality review / exception handling.")
    para(doc, "As-Is", bold=True, color=NAVY, after=1)
    table(doc, ["Step", "Actor", "Systems", "Time", "Pain point"],
          [["", "", "", "", ""]], widths=(5.0, 2.6, 3.4, 2.0, 4.0), blank_rows=4)
    para(doc, "To-Be (agent + human)", bold=True, color=NAVY, after=1)
    table(doc, ["Step", "Performed by (agent / human)", "Systems", "Time"],
          [["", "", "", ""]], widths=(7.2, 3.6, 3.2, 3.0), blank_rows=4)

    h1(doc, "4.  Business outcomes")
    para(doc, "The measurable results the solution must deliver (the “why”).")
    table(doc, ["ID", "Business outcome", "Why it matters", "Target signal"],
          [["BO-1", "", "", ""]], widths=(1.6, 6.0, 5.4, 4.0), blank_rows=4)

    h1(doc, "5.  Functional business requirements")
    para(doc, "What the business needs the system to do (capability level, not technical design).")
    table(doc, ["ID", "Requirement", "Priority (M/S/C/W)", "Maps to capability"],
          [["FR-1", "", "", ""]], widths=(1.6, 7.6, 3.4, 4.4), blank_rows=5)

    h1(doc, "6.  Non-functional requirements")
    para(doc, "Confirm and specify the qualities the system must meet (the Agentium differentiators).")
    table(doc, ["Quality", "Required?", "Target / notes"],
          [["Sovereignty / data residency", "", "models & data inside the client boundary"],
           ["Security & identity", "", "RBAC + ABAC, distinct agent identity, guardrails"],
           ["Robustness", "", "stateful runs, checkpoints, replay, resubmission"],
           ["Observability & audit", "", "tool-call ledger, tamper-evident audit, decision trail"],
           ["Scalability", "", "multi-entity, concurrency, segmentation"],
           ["Performance & cost", "", "latency/SLA, deterministic token/compute budget"],
           ["Availability & support", "", "uptime, incident response, hypercare"],
           ["Maintainability / extensibility", "", "API-first, client can extend, no lock-in"]],
          widths=(5.4, 2.6, 9.0))

    h1(doc, "7.  Data & knowledge sources")
    para(doc, "Where the truth lives — systems of record, knowledge corpora (for RAG), and how to reach them.")
    table(doc, ["#", "Source", "Type", "Where it lives / access", "Owner", "Sensitivity"],
          [["S1", "", "", "", "", ""]], widths=(1.2, 3.6, 2.6, 4.6, 2.6, 2.4), blank_rows=5)
    field(doc, "Access method", "API / REST · MCP · direct DB · files / SFTP · events")
    field(doc, "Data location & residency", "data lake type (object store / warehouse / SQL / search); residency; PII")
    field(doc, "Compliance / regulatory", "applicable rules; human-oversight requirements")

    h1(doc, "8.  Decisions the agent will make")
    guidance(doc, "List each decision the agent makes, its inputs, possible outcomes, the confidence threshold "
                  "and what it does when uncertain, and when it escalates. This maps directly to an Agentium "
                  "Decision + Control Policy.")
    table(doc, ["ID", "Decision", "Inputs", "Possible outcomes", "Threshold / default-if-uncertain", "Escalation"],
          [["D1", "", "", "", "", ""]], widths=(1.2, 3.2, 3.4, 3.0, 3.6, 2.6), blank_rows=4)

    h1(doc, "9.  Business rules & guardrails")
    guidance(doc, "Two explicit lists: IF/THEN rules the agent must follow, and the “never-do” boundaries. "
                  "These become Agentium Control Policies / mandate — enforced by construction, and tested by "
                  "refusal cases in the golden set (see the QA template).")
    para(doc, "Business rules (IF / THEN)", bold=True, color=NAVY, after=1)
    table(doc, ["ID", "Rule (IF … THEN …)", "Source / rationale"],
          [["R1", "", ""]], widths=(1.2, 11.4, 4.4), blank_rows=4)
    para(doc, "Never do this", bold=True, color=NAVY, after=1)
    table(doc, ["ID", "The agent must never…"],
          [["N1", ""]], widths=(1.2, 15.8), blank_rows=4)

    h1(doc, "10.  Success metrics (KPIs)")
    table(doc, ["KPI", "Definition", "Baseline", "Target"],
          [["", "", "", ""]], widths=(4.6, 6.4, 3.0, 3.0), blank_rows=4)

    h1(doc, "11.  Constraints & assumptions")
    table(doc, ["#", "Constraint / assumption"], [["1", ""]], widths=(1.4, 15.6), blank_rows=4)

    h1(doc, "12.  Out of scope")
    table(doc, ["#", "Explicitly out of scope (the agent will not do this, even if asked)"],
          [["1", ""]], widths=(1.4, 15.6), blank_rows=3)

    h1(doc, "13.  Traceability")
    para(doc, "Each business outcome → functional requirement → decision/rule → operating-model element (TOM) → "
              "Agentium System / Capability / Skills, and → a golden case in the QA template. This guarantees every "
              "requirement is buildable, governable and testable.")

    h1(doc, "14.  Requirements acceptance")
    signoff(doc, ["Business owner", "Process SME", "Client IT", "Datategy — Delivery"])
    save(doc, "Datategy-Template-Business-Requirements.docx")


# ================================================================ 2. TARGET OPERATING MODEL
def build_tom():
    doc = new_doc("Target Operating Model — Template")
    cover(doc, "DELIVERY METHODOLOGY · TEMPLATE", "Target Operating Model",
          "How the AI-augmented function operates with Agentium",
          "DTG-METH-TOM-v1.0")
    how_to_use(doc, "Describe how the organisation will operate once the AI system is live: service, "
                    "human + agent value streams, roles, technology, data, governance, service management, "
                    "sourcing, finance and transition. The TOM bridges business requirements and the solution.")
    para(doc, "Positioning: Business Requirements (why/what) → Target Operating Model (how) → Solution "
              "(Agentium Systems / Capabilities / Skills).", italic=True, color=MUT)

    h1(doc, "1.  Service model & channels")
    guidance(doc, "What services, to whom, via which channels; self-service vs assisted; support tiers.")
    field(doc, "Services & tiers", "L0 self-service → L1/L2 automated → L3 human")
    field(doc, "Channels", "Teams / web / mobile / email / …")

    h1(doc, "2.  Value streams (human + agent)")
    guidance(doc, "The end-to-end flows and where the human integrates. The agent is expected to act "
                  "autonomously by default; the human integrates as feedback/quality and exception handling.")
    table(doc, ["Value stream", "Trigger", "What the agent does", "Human integration (HITL)", "KPI"],
          [["", "", "", "", ""]], widths=(3.4, 2.4, 4.4, 4.0, 2.8), blank_rows=4)

    h1(doc, "3.  Human + agent operating model")
    guidance(doc, "Default: agents act autonomously. Qualify how the human plugs in — (i) quality feedback "
                  "(review/score/correct → improvement loop), (ii) rare sign-off on genuinely sensitive actions "
                  "— and the human availability that feeds it.")
    table(doc, ["Agent / system", "Action autonomy (default)", "HITL integration", "Human availability / cadence"],
          [["", "Autonomous", "", ""]], widths=(4.0, 3.6, 5.0, 4.4), blank_rows=3)

    h1(doc, "4.  Organisation, roles & RACI")
    guidance(doc, "Operating roles (Executive, Operator, Governance, Builder + IAM/data/ML) and who is "
                  "Responsible/Accountable/Consulted/Informed. Principle: Datategy builds & enables; client operates & owns.")
    table(doc, ["Activity", "Datategy", "Client IT", "Business", "Governance"],
          [["Build & configure", "R/A", "C", "C", "I"],
           ["Operate the desk", "C", "R/A", "I", "I"],
           ["Approve sensitive actions", "I", "C", "C", "R/A"],
           ["", "", "", "", ""]], widths=(5.0, 3.0, 3.0, 3.0, 3.0), blank_rows=2)

    h1(doc, "5.  Technology & platform")
    field(doc, "Platform", "Agentium — orchestrate / build / evaluate")
    field(doc, "Deployment", "on-prem / private cloud / air-gap")
    field(doc, "Systems & Suites in scope", "the verticalised systems delivered")

    h1(doc, "6.  Data & knowledge operating model")
    field(doc, "Knowledge capture & curation", "how the KB stays current; owners")
    field(doc, "Data residency & isolation", "boundary, per-entity isolation")

    h1(doc, "7.  Governance, risk & control")
    checklist(doc, [("AI management system", "governed & audited per ISO/IEC 42001 (SC 42); transparency per ISO/IEC 12792"),
                    ("Mandates & control policies", "allowed actions, caps, scope, duration, revocation"),
                    ("Identity & access", "RBAC + ABAC, distinct agent identity"),
                    ("Audit", "tool-calling ledger, tamper-evident, exportable"),
                    ("Quality gates", "evaluation + review queue before any change (dress-rehearsal)"),
                    ("Regulatory alignment", "applicable rules; human oversight")])

    h1(doc, "8.  Service management & SLAs")
    table(doc, ["Service / process", "SLA / OLA", "Metric"], [["", "", ""]], widths=(7.0, 5.0, 5.0), blank_rows=4)

    h1(doc, "9.  Sourcing, locations & Center of Excellence")
    field(doc, "Delivery model", "editor (Datategy) + ESN/integrator partner; client CoE / AI Academy")
    field(doc, "Locations / multi-entity", "geographies, entity segmentation, chargeback")

    h1(doc, "10.  Performance management & financials")
    field(doc, "Value steering", "KPI/ROI cockpit (Hypervisor); per-capability cost/quality")
    field(doc, "Financial model", "3-yr TCO; per-entity chargeback; token/compute budget")

    h1(doc, "11.  Transition roadmap")
    table(doc, ["Phase", "Objective", "Milestone", "Exit criteria"],
          [["1 — Mobilise & POC", "", "", ""],
           ["2 — Migrate & enhance", "", "", ""],
           ["3 — Govern & secure", "", "", ""],
           ["4 — Scale & handover", "", "", ""]], widths=(4.0, 5.0, 4.0, 4.0))

    h1(doc, "12.  Operating-model acceptance")
    signoff(doc, ["Business owner", "Client IT", "Governance", "Datategy — Delivery"])
    save(doc, "Datategy-Template-Target-Operating-Model.docx")


# ================================================================ 3. QA ACCEPTANCE (PV de recette)
def build_qa_pv():
    doc = new_doc("QA Acceptance Record (PV) — Template")
    cover(doc, "DELIVERY METHODOLOGY · TEMPLATE", "QA Acceptance Record",
          "Procès-verbal de recette — qualifying an Agentium system for go-live",
          "DTG-METH-QA-v1.0")
    how_to_use(doc, "A QA uses this to qualify (pass/fail, sign-off) a delivered Agentium system. A delivery "
                    "is qualified only when every must-pass gate (★) is green, there are 0 open Critical/High "
                    "defects, and the sovereignty egress check is clean.")

    h1(doc, "1.  System under test")
    field(doc, "System / agent", "name & version")
    field(doc, "Environment", "staging / sandbox; build / commit")
    field(doc, "QA owner & date", "name · dd/mm/yyyy")

    h1(doc, "2.  Entry criteria")
    checklist(doc, ["System manifest (systems, flows, skills, capabilities, connectors, mandates, KB) delivered",
                    "Spec & traceability (use cases → flows; acceptance criteria) available",
                    "Isolated staging/sandbox with synthetic identities & sandbox connectors",
                    "Test data + golden evaluation set provided",
                    "Runbooks, API reference, rollback/cutover plan delivered",
                    "Build/version pinned; migrations applied"])

    h1(doc, "3.  Test execution safety")
    checklist(doc, ["Staging/sandbox only — no production mutation",
                    "Privileged writes tested in simulate/dry-run first; one safe write on sandbox only",
                    "Sandbox credentials; no production secrets in test config",
                    "Synthetic PII only; test data purged after run"])

    h1(doc, "4.  Qualification gates")
    para(doc, "Status: P (pass) · F (fail) · N/A. Attach evidence (run ID, audit export, golden report, screenshot). "
              "★ = must-pass.")
    gates = [
        ("A ★", "Scope & functional coverage", "every spec'd feature has a test + evidence; outcomes met"),
        ("B ★", "Flow correctness & branches", "main + rejected/missed/failed; HITL gates; idempotency"),
        ("C", "Skills & tool-calling", "correct outputs; typed errors; bounded retries"),
        ("D ★", "Conversational / RAG quality", "grounded & cited; hallucination below threshold; golden ≥ target"),
        ("E ★", "Robustness", "resume/checkpoint; replay reproduces or discrepancy report; resubmission"),
        ("F ★", "Security & identity", "RBAC/ABAC enforced; agent identity; multi-entity isolation"),
        ("G ★", "Governance, mandates & audit", "privileged actions ledgered; out-of-mandate blocked; audit export"),
        ("H", "Evaluation & quality gates", "auto-scoring; thresholds; review queue; drift/rollback"),
        ("I", "Human-in-the-loop", "review/feedback integration; exception handling; corrections reinjected"),
        ("J ★", "Connectors & integration", "auth + happy path + error/rate-limit + mapping per connector"),
        ("K", "Data & knowledge", "ingestion; KB isolation; residency; facts queryable"),
        ("L", "Performance, capacity & cost", "latency vs SLA; token budget; load/concurrency; QoS"),
        ("M ★", "Observability & analytics", "KPI dashboards reconcile (MTTR, deflection, ROI, SLA)"),
        ("N ★", "Sovereignty & deployment", "on-prem/air-gap; egress clean (no phone-home); residency"),
        ("O ★", "Migration acceptance", "legacy use cases at parity; cutover & rollback rehearsed"),
        ("P ★", "UAT / business acceptance", "outcomes validated vs requirements & KPIs"),
        ("Q", "Documentation & enablement", "runbooks, API docs, knowledge transfer / CoE"),
        ("R ★", "Go-live readiness & hypercare", "go/no-go, rollback tested, hypercare plan"),
    ]
    table(doc, ["Gate", "Dimension", "What it confirms", "Status", "Evidence"],
          [[g, d, w, "", ""] for g, d, w in gates],
          widths=(1.5, 4.0, 6.5, 1.7, 3.3))

    h1(doc, "5.  Golden set")
    guidance(doc, "Representative cases the system is qualified against — and re-run for regression. Cover four "
                  "kinds: routine (the common case, must be flawless), refusal / adversarial (the agent must "
                  "refuse — e.g. an out-of-mandate or forbidden request), escalation (must route to a human), "
                  "and ambiguity (must abstain below the confidence threshold). Each refusal/never-do rule in the "
                  "Business Requirements should have a golden case here.")
    table(doc, ["ID", "Input / scenario", "Expected behaviour", "Why this case matters"],
          [["GS1", "", "", ""]], widths=(1.3, 5.0, 7.1, 3.6), blank_rows=6)

    h1(doc, "6.  Defect register")
    table(doc, ["ID", "Severity", "Gate", "Description", "Status", "Evidence"],
          [["", "", "", "", "", ""]], widths=(1.3, 2.2, 1.5, 6.5, 2.2, 3.3), blank_rows=5)

    h1(doc, "7.  Roll-up & exit criteria")
    table(doc, ["Metric", "Value", "Target"],
          [["Features qualified", "", "100% of scope"],
           ["Tests / execution evidence", "", "all gates evidenced"],
           ["Open Critical / High defects", "", "0"],
           ["Golden set", "", "≥ target"],
           ["Sovereignty egress", "", "clean"]], widths=(7.0, 5.0, 5.0))
    checklist(doc, ["All must-pass (★) gates green",
                    "Scope coverage 100% (brownfield: legacy use cases at parity)",
                    "0 open Critical/High defects",
                    "Golden ≥ target · replay reproducible · isolation verified · egress clean",
                    "KPI dashboards reconcile · UAT signed"])

    h1(doc, "8.  Acceptance decision (PV de recette)")
    para(doc, "Decision per signatory: Accept · Accept with reservations · Reject. Qualification is granted "
              "only when all parties accept.")
    signoff(doc, ["QA Lead", "Governance Officer", "Business Owner", "Datategy — Delivery", "Client IT Owner"],
            decision_col=True)
    save(doc, "Datategy-Template-QA-Acceptance-PV.docx")


if __name__ == "__main__":
    build_business_requirements()
    build_tom()
    build_qa_pv()
