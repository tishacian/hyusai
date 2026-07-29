# -*- coding: utf-8 -*-
"""Align PIH AI Factory RFI Response v3 with Annex B v4.

Output: docs/pih/RFI-PIH-AI-Factory-Response-v4.docx

v4 replaces the dispersed recurring model with one post-build Enterprise
Platform & Managed Operations Baseline: USD 100k/year for 50 live agents,
then AgentOps overage by risk class.
"""
import copy
import os

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
PIH = os.path.join(os.path.dirname(HERE), "pih")
SRC = os.path.join(PIH, "RFI-PIH-AI-Factory-Response-v3.docx")
OUT = os.path.join(PIH, "RFI-PIH-AI-Factory-Response-v4.docx")

doc = Document(SRC)


def find_contains(text):
    for p in doc.paragraphs:
        if text in p.text:
            return p
    raise SystemExit(f"ANCHOR NOT FOUND: {text!r}")


def set_text(p, text, bold_lead=None):
    for run in list(p.runs):
        run._r.getparent().remove(run._r)
    if bold_lead:
        run = p.add_run(bold_lead)
        run.bold = True
    p.add_run(text)


def set_cell(cell, value):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        run._r.getparent().remove(run._r)
    p.add_run(value)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def commercial_table():
    for table in doc.tables:
        if table.cell(0, 0).text.strip() == "Component":
            return table
    raise SystemExit("COMMERCIAL TABLE NOT FOUND")


# ---------------------------------------------------------------- §3.2 commercial principle
p = find_contains("Commercial principle")
set_text(
    p,
    "PIH buys delivered agents, not per-user licences or opaque execution meters. Each build price is "
    "all-inclusive for the scope agreed in its design card: Agentium runtime during the Factory programme, "
    "agreed knowledge scope, testing, security, deployment, hypercare and one knowledge-transfer session. "
    "Catalogue prices apply to reuse instances — new agents configured from delivered and validated "
    "patterns, skills and connectors — while first-of-type assets are identified and priced before build.",
    bold_lead="Commercial principle — ",
)

# ---------------------------------------------------------------- §3.2 table
tbl = commercial_table()
rows = [
    (
        "Agent build — catalogue unit price (reuse instance)",
        "Per accepted agent, by S/M/C tier; all-inclusive front-office unit",
        "Simple ≈ $3.5–5k · Medium ≈ $9–13k · Complex ≈ $20–30k. Runtime, agreed knowledge scope, "
        "testing, security, deployment, hypercare and transfer included during the Factory programme.",
    ),
    (
        "First-of-type premium",
        "First agent creating/materially extending a pattern, skill or connector",
        "+50% / +50–75% / +75–100% applied to the approved design-card catalogue price before "
        "volume reduction. Non-degressive and scoped before build.",
    ),
    (
        "Volume reduction",
        "Accepted reuse instances only",
        "−10% only on incremental units 101–250; −20% only on incremental units 251+. No retroactive "
        "rebate; first-of-type agents are excluded.",
    ),
    (
        "Enterprise Platform & Managed Operations Baseline — post-build",
        "USD 100k/year for the first 50 live agents",
        "Starts after final programme acceptance, or 90 days after the last accepted delivery if no "
        "successor factory plan is active. Includes platform continuity, updates, security patches, "
        "compatibility, golden-set regression, N3 support, standard AgentOps and knowledge service.",
    ),
    (
        "AgentOps overage — post-build",
        "Per accepted live agent above the first 50, per month",
        "Low $50 · Medium $60 · High $70. High-impact scope includes reinforced drift watch and "
        "quarterly adversarial/red-team regression.",
    ),
    (
        "T&M day rates",
        "Pilot staff augmentation, discovery, change request or scope outside catalogue pricing",
        "AI/Data engineer $149 · Developer $149 · QA/Eval $149 · CoE Director $349 · Solution "
        "Architect/Expert $440 per person/day. Not added to an all-inclusive unit-priced agent.",
    ),
    (
        "Azure / SaaS inference, where PIH approved",
        "Datategy fee",
        "$0. Azure/OpenAI consumption remains contracted and billed directly between PIH and its cloud "
        "provider; Datategy adds no markup.",
    ),
    (
        "Private / sovereign inference",
        "PIH-controlled GPU capacity",
        "$0 Datategy infrastructure fee. Hosted in PIH-controlled Azure Qatar, private-cloud or "
        "on-premise capacity; model sizing remains a PIH infrastructure decision.",
    ),
]
while len(tbl.rows) - 1 < len(rows):
    tbl._tbl.append(copy.deepcopy(tbl.rows[-1]._tr))
while len(tbl.rows) - 1 > len(rows):
    tbl._tbl.remove(tbl.rows[-1]._tr)
for idx, values in enumerate(rows, 1):
    for cell, value in zip(tbl.rows[idx].cells, values):
        set_cell(cell, value)

# ---------------------------------------------------------------- §3.3 inference
p = find_contains("Commercial structure & trajectory")
set_text(
    p,
    "Where PIH approves Azure/SaaS endpoints, Datategy charges $0 and PIH pays its cloud provider "
    "directly at the applicable tariff. Private/sovereign inference uses PIH-controlled GPU capacity. "
    "We target a −30–50% reduction in cost per execution over 12 months through caching, right-sized "
    "routing and prompt optimisation, measured against a jointly agreed baseline and reviewed quarterly. "
    "This is a jointly managed optimisation target, not a guaranteed reduction.",
    bold_lead="Commercial structure & trajectory (Q5A.2) — ",
)

# ---------------------------------------------------------------- §3.4 payment
p = find_contains("Cash profile")
set_text(
    p,
    "we request a mobilisation payment at signature covering Factory onboarding and internal capacity "
    "reservation; the pilot/staff-augmentation phase is invoiced monthly; and industrial delivery is "
    "paid through short monthly or wave milestones rather than a deferred 200-agent step.",
    bold_lead="Cash profile — ",
)

# ---------------------------------------------------------------- §5 factory economics
p = find_contains("Measurement method (Q2A.1)")
set_text(
    p,
    "we propose a 30% pattern-reuse / 15% component-standardisation measurement method, subject to "
    "PIH clarification of the baseline, calculation and contractual status. The proposed baseline is "
    "first-of-type cost per pattern; reporting is through the Factory ledger; conditions include stable "
    "pattern taxonomy and PIH data/environment access.",
    bold_lead="Measurement method (Q2A.1) — ",
)

# ---------------------------------------------------------------- §13 AgentOps
p = find_contains("We can deliver AgentOps as build-partner")
set_text(
    p,
    "After the Factory build programme, the Enterprise Platform & Managed Operations Baseline covers "
    "the first 50 live agents. PIH may operate agents itself or appoint another operator; delivered "
    "flows, design cards, evaluation evidence and documented connector mappings are handover-ready. "
    "PIH retains a non-exclusive right to execute accepted delivered versions on PIH-controlled "
    "infrastructure. This does not transfer Agentium product IP, future releases or source code; "
    "updates, compatibility, regression and N3 support remain part of the post-build continuity regime.",
)

# ---------------------------------------------------------------- terminology and header
for section in doc.sections:
    for p in section.header.paragraphs:
        for run in p.runs:
            if "DRAFT v3 — internal work in progress" in run.text:
                run.text = run.text.replace("DRAFT v3", "DRAFT v4")
for p in doc.paragraphs:
    for run in p.runs:
        if "DRAFT v3 — internal" in run.text:
            run.text = run.text.replace("DRAFT v3", "DRAFT v4")

doc.save(OUT)

# ---------------------------------------------------------------- verification
out = Document(OUT)
text = "\n".join(p.text for p in out.paragraphs)
table_text = "\n".join(cell.text for table in out.tables for row in table.rows for cell in row.cells)
checks = {
    "baseline 100k present": "USD 100k/year for the first 50 live agents" in table_text,
    "overage present": "Low $50 · Medium $60 · High $70" in table_text,
    "SaaS zero fee present": "$0. Azure/OpenAI consumption remains" in table_text,
    "old Platform Subscription row removed": "Platform Subscription — post-programme" not in table_text,
    "old knowledge price row removed": "Knowledge capacity (Document Center)" not in table_text,
    "no consumption tax removed": "no consumption tax" not in text,
    "v4 heading": "DRAFT v4" in text,
}
for name, passed in checks.items():
    print(("OK" if passed else "FAIL"), name)
print("saved:", OUT)
