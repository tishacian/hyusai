"""Build editable documents and PostgreSQL fixtures; never connects to a DB.

Run with the backend Python environment, which provides PyMuPDF and python-docx.
All inputs are in fixture_spec.json; only files in this fixture directory change.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf as fitz
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
LABEL = "DÉMONSTRATION SYNTHÉTIQUE - aucun client ni remboursement réel"
BUILD_TIME = datetime(2026, 10, 1, tzinfo=timezone.utc)

DDL = """-- Luma Maison, fixture v1; never loads measured time or ROI.
-- Run with an operator account on the intended demo DB, not the agent reader.
-- Inserts preserve existing rows; no DROP, DELETE or state reset.
BEGIN;
CREATE SCHEMA IF NOT EXISTS showcase_ecommerce;

CREATE TABLE IF NOT EXISTS showcase_ecommerce.customers (
  customer_id TEXT PRIMARY KEY, display_name TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS showcase_ecommerce.orders (
  order_id TEXT PRIMARY KEY,
  customer_id TEXT NOT NULL REFERENCES showcase_ecommerce.customers(customer_id),
  ordered_at TIMESTAMPTZ NOT NULL, paid_amount NUMERIC(12,2) NOT NULL CHECK (paid_amount >= 0),
  currency TEXT NOT NULL CHECK (currency = 'EUR'), shipping_postcode TEXT NOT NULL,
  payment_status TEXT NOT NULL CHECK (payment_status IN ('paid','unpaid'))
);
CREATE TABLE IF NOT EXISTS showcase_ecommerce.order_items (
  item_id TEXT PRIMARY KEY, order_id TEXT NOT NULL REFERENCES showcase_ecommerce.orders(order_id),
  product_name TEXT NOT NULL, quantity INTEGER NOT NULL CHECK (quantity > 0),
  unit_price NUMERIC(12,2) NOT NULL CHECK (unit_price >= 0)
);
CREATE TABLE IF NOT EXISTS showcase_ecommerce.shipments (
  shipment_id TEXT PRIMARY KEY, order_id TEXT NOT NULL REFERENCES showcase_ecommerce.orders(order_id),
  tracking_id TEXT NOT NULL UNIQUE, carrier TEXT NOT NULL,
  status TEXT NOT NULL CHECK (status IN ('delivered','lost','in_transit')), status_at TIMESTAMPTZ NOT NULL
);
CREATE TABLE IF NOT EXISTS showcase_ecommerce.refunds (
  refund_id TEXT PRIMARY KEY, order_id TEXT NOT NULL REFERENCES showcase_ecommerce.orders(order_id),
  amount NUMERIC(12,2) NOT NULL CHECK (amount > 0), currency TEXT NOT NULL CHECK (currency = 'EUR'),
  status TEXT NOT NULL CHECK (status IN ('prepared','approved','executed','failed')),
  executed_at TIMESTAMPTZ,
  CHECK ((status = 'executed' AND executed_at IS NOT NULL) OR (status <> 'executed' AND executed_at IS NULL))
);
CREATE TABLE IF NOT EXISTS showcase_ecommerce.claims (
  claim_id TEXT PRIMARY KEY, order_id TEXT NOT NULL REFERENCES showcase_ecommerce.orders(order_id),
  reason TEXT NOT NULL, state TEXT NOT NULL CHECK (state IN ('to_investigate','waiting_information','awaiting_decision','resolved')),
  opened_at TIMESTAMPTZ NOT NULL
);
CREATE TABLE IF NOT EXISTS showcase_ecommerce.document_refs (
  document_key TEXT PRIMARY KEY, source_filename TEXT NOT NULL, document_type TEXT NOT NULL,
  version TEXT NOT NULL, effective_from DATE NOT NULL, is_current BOOLEAN NOT NULL,
  order_id TEXT REFERENCES showcase_ecommerce.orders(order_id),
  sha256 TEXT NOT NULL CHECK (sha256 ~ '^[0-9a-f]{64}$'),
  knowledge_source_id TEXT
);
COMMENT ON SCHEMA showcase_ecommerce IS 'Synthetic Luma Maison demo; no measured business gains';
COMMENT ON COLUMN showcase_ecommerce.document_refs.knowledge_source_id IS
  'NULL until a source is actually ingested in the matching Agentium workspace';
"""


def validate(spec: dict) -> None:
    assert spec["evidence_kind"] == "synthetic_demo"
    assert spec["workspace_slug"] == "agentium-showcase"
    customers = {row["customer_id"] for row in spec["customers"]}
    orders = {row["order_id"]: row for row in spec["orders"]}
    assert len(orders) == len(spec["orders"])
    assert all(row["customer_id"] in customers for row in orders.values())
    for table in ("order_items", "shipments", "refunds", "claims"):
        assert all(row["order_id"] in orders for row in spec[table])
    for order_id, order in orders.items():
        total = sum(Decimal(row["unit_price"]) * row["quantity"]
                    for row in spec["order_items"] if row["order_id"] == order_id)
        assert total == Decimal(order["paid_amount"])
        refunds = sum(Decimal(row["amount"]) for row in spec["refunds"]
                      if row["order_id"] == order_id and row["status"] == "executed")
        assert refunds <= total
    documents = {row["reference"]: row for row in spec["documents"]}
    assert len(documents) == len(spec["documents"])
    assert len({row["filename"] for row in documents.values()}) == len(documents)
    for row in documents.values():
        assert Path(row["filename"]).name == row["filename"]
        assert re.fullmatch(r"[a-z0-9][a-z0-9-]*", row["reference"])
        assert row["order_id"] is None or row["order_id"] in orders
    assert orders["LM-1042"]["shipping_postcode"] == "75011"
    assert "75012" in " ".join(documents["delivery-lm1042"]["paragraphs"])
    assert not documents["refund-policy-v1-archived"]["current"]
    assert documents["refund-policy-v2"]["current"]
    assert [row["order_id"] for row in spec["refunds"]] == ["LM-1044"]


def source_text(row: dict) -> str:
    status = "En vigueur / pièce courante" if row["current"] else "ARCHIVE"
    return (f"# {row['title']}\n\n{LABEL}\n\n"
            f"Version : {row['version']} | Date : {row['effective_from']} | {status}\n\n"
            + "\n\n".join(row["paragraphs"]) + "\n")


def write_pdf(row: dict, path: Path) -> None:
    with fitz.open() as pdf:
        page = pdf.new_page(width=595, height=842)
        page.draw_rect(fitz.Rect(0, 0, 595, 68), color=None, fill=(0.04, 0.34, 0.40))
        page.insert_text((42, 32), "LUMA MAISON", fontsize=18, fontname="hebo", color=(1, 1, 1))
        page.insert_text((42, 52), "Document de démonstration", fontsize=10, color=(1, 1, 1))
        result = page.insert_textbox(fitz.Rect(42, 90, 553, 151), row["title"],
                                    fontsize=17, fontname="hebo", color=(0.1, 0.16, 0.20))
        assert result >= 0, f"Title does not fit: {row['reference']}"
        state = "Courant" if row["current"] else "ARCHIVE"
        page.insert_text((42, 162), f"Version {row['version']} | {row['effective_from']} | {state}",
                         fontsize=10, color=(0.3, 0.35, 0.4))
        result = page.insert_textbox(fitz.Rect(42, 189, 553, 765),
                                    "\n\n".join(row["paragraphs"]), fontsize=11,
                                    lineheight=1.35, color=(0.13, 0.17, 0.20))
        assert result >= 0, f"Body does not fit: {row['reference']}"
        page.insert_text((42, 799), LABEL, fontsize=8, color=(0.35, 0.4, 0.45))
        page.insert_text((540, 799), "1", fontsize=8)
        pdf.set_metadata({"title": row["title"], "author": "Luma Maison - demonstration",
                          "subject": "synthetic_demo", "creationDate": "D:20261001000000Z",
                          "modDate": "D:20261001000000Z"})
        pdf.save(path, garbage=4, deflate=True, no_new_id=True)


def write_docx(row: dict, path: Path) -> None:
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("LUMA MAISON", level=1)
    document.add_heading(row["title"], level=2)
    document.add_paragraph(LABEL)
    document.add_paragraph(f"Version {row['version']} | {row['effective_from']}")
    for paragraph in row["paragraphs"]:
        document.add_paragraph(paragraph)
    document.sections[0].footer.paragraphs[0].text = LABEL
    props = document.core_properties
    props.title, props.author, props.subject = row["title"], "Luma Maison - demonstration", "synthetic_demo"
    props.created, props.modified = BUILD_TIME, BUILD_TIME
    memory = io.BytesIO()
    document.save(memory)
    # Canonical ZIP timestamps make the manifest stable across regenerations.
    with ZipFile(io.BytesIO(memory.getvalue())) as archive, ZipFile(path, "w", ZIP_DEFLATED) as canonical:
        for item in archive.infolist():
            item.date_time = (2026, 10, 1, 0, 0, 0)
            canonical.writestr(item, archive.read(item.filename))


def sql_value(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def insert_rows(table: str, rows: list[dict]) -> str:
    if not rows:
        return ""
    columns = list(rows[0])
    assert all(list(row) == columns for row in rows)
    values = ",\n".join("  (" + ", ".join(sql_value(row[col]) for col in columns) + ")" for row in rows)
    return (f"\nINSERT INTO showcase_ecommerce.{table} ({', '.join(columns)}) VALUES\n"
            + values + f"\nON CONFLICT ({columns[0]}) DO NOTHING;\n")


def main() -> None:
    spec = json.loads((ROOT / "fixture_spec.json").read_text())
    validate(spec)
    for folder in ("documents", "source-texts", "postgres"):
        (ROOT / folder).mkdir(exist_ok=True)
    manifest = {"fixture_version": spec["fixture_version"], "workspace_slug": spec["workspace_slug"],
                "evidence_kind": spec["evidence_kind"], "documents": []}
    refs = []
    for row in spec["documents"]:
        text = source_text(row)
        (ROOT / "source-texts" / f"{row['reference']}.md").write_text(text, encoding="utf-8")
        path = ROOT / "documents" / row["filename"]
        if path.suffix == ".pdf":
            write_pdf(row, path)
        elif path.suffix == ".docx":
            write_docx(row, path)
        else:
            assert path.suffix == ".md"
            path.write_text(text, encoding="utf-8")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        manifest["documents"].append({"reference": row["reference"], "filename": row["filename"],
                                      "type": row["type"], "version": row["version"],
                                      "effective_from": row["effective_from"], "current": row["current"],
                                      "order_id": row["order_id"], "sha256": digest})
        refs.append({"document_key": row["reference"], "source_filename": row["filename"],
                     "document_type": row["type"], "version": row["version"],
                     "effective_from": row["effective_from"], "is_current": row["current"],
                     "order_id": row["order_id"], "sha256": digest, "knowledge_source_id": None})
    sql = DDL
    for table in ("customers", "orders", "order_items", "shipments", "refunds", "claims"):
        sql += insert_rows(table, spec[table])
    sql += insert_rows("document_refs", refs) + "\nCOMMIT;\n"
    (ROOT / "postgres" / "seed.sql").write_text(sql, encoding="utf-8")
    (ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Prepared {len(refs)} documents and PostgreSQL fixture; no database connection attempted.")


if __name__ == "__main__":
    main()
