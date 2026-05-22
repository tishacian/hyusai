"""Spreadsheet parser for Excel/CSV workbooks.

The parser emits several views of the same table artifact:
- row chunks for backwards-compatible RAG
- schema chunks so retrieval can see sheet/table shape
- cell/table facts with coordinates and labels
- semantic sentences for natural-language lookup

The goal is not to infer business truth. It exposes workbook evidence in a
more queryable form and keeps uncertainty visible.
"""
from __future__ import annotations

import asyncio
import csv
import re
from pathlib import Path
from typing import Any

from app.services.document_parser.base import BaseDocumentParser, DocumentType, ParsedDocument


_EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
_DELIMITED_EXTENSIONS = {".csv", ".tsv"}
_SUPPORTED_EXTENSIONS = _EXCEL_EXTENSIONS | _DELIMITED_EXTENSIONS
_MAX_ROWS_PER_CHUNK = 80
_MAX_SHEET_ROWS = 5_000
_MAX_SHEET_COLS = 80
_MAX_FACT_LINES_PER_CHUNK = 60
_MAX_FACT_CHUNKS_PER_SHEET = 800
_MAX_STRUCTURED_FACTS = 2_000
_MAX_CHUNK_CHARS = 6_000


class SpreadsheetParser(BaseDocumentParser):
    """Parse spreadsheets into structured, RAG-friendly table artifacts."""

    async def parse(self, file_path: str, **kwargs) -> ParsedDocument:
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".xls":
            raise ValueError("Legacy .xls spreadsheets are not supported for Knowledge ingestion yet")
        if suffix not in _SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported spreadsheet extension: {path.suffix}")

        metadata = await self.extract_metadata(file_path)
        rows_per_chunk = _safe_positive_int(kwargs.get("spreadsheet_rows_per_chunk"), _MAX_ROWS_PER_CHUNK)
        max_sheet_rows = _safe_positive_int(kwargs.get("spreadsheet_max_rows"), _MAX_SHEET_ROWS)
        max_sheet_cols = _safe_positive_int(kwargs.get("spreadsheet_max_cols"), _MAX_SHEET_COLS)

        loop = asyncio.get_event_loop()
        if suffix in _DELIMITED_EXTENSIONS:
            parsed = await loop.run_in_executor(
                None,
                lambda: self._parse_delimited(
                    path,
                    rows_per_chunk=rows_per_chunk,
                    max_sheet_rows=max_sheet_rows,
                    max_sheet_cols=max_sheet_cols,
                ),
            )
            document_type = DocumentType.CSV
            encoding = "utf-8"
        else:
            parsed = await loop.run_in_executor(
                None,
                lambda: self._parse_workbook(
                    path,
                    rows_per_chunk=rows_per_chunk,
                    max_sheet_rows=max_sheet_rows,
                    max_sheet_cols=max_sheet_cols,
                ),
            )
            document_type = DocumentType.SPREADSHEET
            encoding = "binary"

        metadata.update(
            {
                "source_path": file_path,
                "sheet_count": len(parsed["sheets"]),
                "spreadsheet_truncated": parsed["truncated"],
                "table_artifact_version": "table_artifact_v1",
            }
        )
        raw_content = "\n\n".join(chunk["content"] for chunk in parsed["chunks"])

        return ParsedDocument(
            id=self._generate_id(file_path),
            filename=path.name,
            file_path=file_path,
            document_type=document_type,
            raw_content=raw_content,
            structured_content={
                "sheet_count": len(parsed["sheets"]),
                "sheets": parsed["sheets"],
                "chunks_count": len(parsed["chunks"]),
                "truncated": parsed["truncated"],
                "table_artifacts": parsed["table_artifacts"],
            },
            metadata=metadata,
            chunks=parsed["chunks"],
            images=[],
            tables=parsed["tables"],
            code_blocks=[],
            entities=[],
            relationships=[],
            language="unknown",
            encoding=encoding,
            file_size=metadata["file_size"],
            created_at=metadata["created_at"],
            modified_at=metadata["modified_at"],
        )

    def supports(self, file_path: str) -> bool:
        return Path(file_path).suffix.lower() in _SUPPORTED_EXTENSIONS

    def _parse_workbook(
        self,
        path: Path,
        *,
        rows_per_chunk: int,
        max_sheet_rows: int,
        max_sheet_cols: int,
    ) -> dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Excel parsing requires openpyxl to ingest modern spreadsheet files") from exc

        value_workbook = load_workbook(path, read_only=True, data_only=True)
        formula_workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheets_payload: list[dict[str, Any]] = []
            tables: list[dict[str, Any]] = []
            all_rows: list[tuple[str, dict[str, Any], dict[str, Any]]] = []
            truncated = False

            for worksheet in value_workbook.worksheets:
                formula_ws = formula_workbook[worksheet.title] if worksheet.title in formula_workbook.sheetnames else None
                rows = self._extract_rows(
                    worksheet,
                    formula_ws=formula_ws,
                    max_sheet_rows=max_sheet_rows,
                    max_sheet_cols=max_sheet_cols,
                )
                truncated = truncated or bool(rows["truncated"])
                sheet_info = _sheet_info(
                    worksheet,
                    row_count=len(rows["rows"]),
                    truncated=rows["truncated"],
                    max_col_seen=rows["max_col"],
                )
                sheets_payload.append(sheet_info)
                all_rows.append((worksheet.title, rows, sheet_info))
                if rows["rows"]:
                    tables.append(
                        {
                            "sheet_name": worksheet.title,
                            "table_region_id": _table_region_id(worksheet.title, 1),
                            "rows": rows["rows"][:50],
                            "truncated": rows["truncated"] or len(rows["rows"]) > 50,
                        }
                    )

            chunks, table_artifacts = self._build_chunks_and_artifacts(
                path=path,
                sheets_payload=sheets_payload,
                all_rows=all_rows,
                rows_per_chunk=rows_per_chunk,
            )
            return {
                "chunks": chunks,
                "sheets": sheets_payload,
                "tables": tables,
                "truncated": truncated,
                "table_artifacts": table_artifacts,
            }
        finally:
            value_workbook.close()
            formula_workbook.close()

    def _parse_delimited(
        self,
        path: Path,
        *,
        rows_per_chunk: int,
        max_sheet_rows: int,
        max_sheet_cols: int,
    ) -> dict[str, Any]:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        rows: list[dict[str, Any]] = []
        max_col_seen = 1
        truncated = False
        with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            for row_index, row in enumerate(reader, start=1):
                if row_index > max_sheet_rows:
                    truncated = True
                    break
                cells = []
                for col_index, value in enumerate(row[:max_sheet_cols], start=1):
                    clean = _format_cell(value)
                    if clean == "":
                        continue
                    max_col_seen = max(max_col_seen, col_index)
                    cells.append(
                        {
                            "column_index": col_index,
                            "column": _column_letter(col_index),
                            "value": clean,
                            "cell_ref": f"{_column_letter(col_index)}{row_index}",
                        }
                    )
                if len(row) > max_sheet_cols:
                    truncated = True
                if cells:
                    rows.append({"row_index": row_index, "cells": cells})

        sheet_name = path.stem or "Sheet1"
        sheet_info = {
            "name": sheet_name,
            "rows": len(rows),
            "max_row": len(rows),
            "max_column": max_col_seen,
            "max_column_seen": max_col_seen,
            "hidden": False,
            "merged_ranges": [],
            "truncated": truncated,
        }
        rows_payload = {"rows": rows, "max_col": max_col_seen, "truncated": truncated}
        chunks, table_artifacts = self._build_chunks_and_artifacts(
            path=path,
            sheets_payload=[sheet_info],
            all_rows=[(sheet_name, rows_payload, sheet_info)],
            rows_per_chunk=rows_per_chunk,
        )
        return {
            "chunks": chunks,
            "sheets": [sheet_info],
            "tables": [
                {
                    "sheet_name": sheet_name,
                    "table_region_id": _table_region_id(sheet_name, 1),
                    "rows": rows[:50],
                    "truncated": truncated or len(rows) > 50,
                }
            ]
            if rows
            else [],
            "truncated": truncated,
            "table_artifacts": table_artifacts,
        }

    def _build_chunks_and_artifacts(
        self,
        *,
        path: Path,
        sheets_payload: list[dict[str, Any]],
        all_rows: list[tuple[str, dict[str, Any], dict[str, Any]]],
        rows_per_chunk: int,
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        chunks: list[dict[str, Any]] = []
        table_artifacts = {
            "version": "table_artifact_v1",
            "workbook_schema": {
                "filename": path.name,
                "sheet_count": len(sheets_payload),
                "sheets": sheets_payload,
            },
            "cell_facts": [],
            "table_facts": [],
            "semantic_sentences": [],
            "interpretation_notes": [],
        }
        current_pos = 0

        for sheet_name, rows, sheet_info in all_rows:
            schema_content = self._render_schema_chunk(path.name, sheet_info)
            chunks.append(
                _chunk(
                    content=schema_content,
                    start_char=current_pos,
                    chunk_index=len(chunks),
                    source_path=str(path),
                    semantic_type="spreadsheet_schema",
                    sheet_name=sheet_name,
                    table_region_id=_table_region_id(sheet_name, 1),
                    cell_range=f"A1:{_column_letter(max(1, rows['max_col']))}{max(1, sheet_info['max_row'] or 1)}",
                )
            )
            current_pos += len(schema_content) + 2

            for batch in _batched(rows["rows"], rows_per_chunk):
                content = self._render_row_chunk(sheet_name, batch)
                row_start = batch[0]["row_index"]
                row_end = batch[-1]["row_index"]
                cell_range = f"A{row_start}:{_column_letter(rows['max_col'])}{row_end}"
                chunks.append(
                    _chunk(
                        content=content,
                        start_char=current_pos,
                        chunk_index=len(chunks),
                        source_path=str(path),
                        semantic_type="spreadsheet_row",
                        sheet_name=sheet_name,
                        row_start=row_start,
                        row_end=row_end,
                        cell_range=cell_range,
                        table_region_id=_table_region_id(sheet_name, 1),
                    )
                )
                current_pos += len(content) + 2

            facts = self._build_fact_rows(sheet_name, rows["rows"])
            if len(facts["cell_facts"]) + len(table_artifacts["cell_facts"]) <= _MAX_STRUCTURED_FACTS:
                table_artifacts["cell_facts"].extend(facts["cell_facts"])
            if len(facts["table_facts"]) + len(table_artifacts["table_facts"]) <= _MAX_STRUCTURED_FACTS:
                table_artifacts["table_facts"].extend(facts["table_facts"])
            if len(facts["semantic_sentences"]) + len(table_artifacts["semantic_sentences"]) <= _MAX_STRUCTURED_FACTS:
                table_artifacts["semantic_sentences"].extend(facts["semantic_sentences"])
            table_artifacts["interpretation_notes"].extend(facts["interpretation_notes"])

            for fact in facts["chunks"][:_MAX_FACT_CHUNKS_PER_SHEET]:
                content = fact["content"]
                chunks.append(
                    _chunk(
                        content=content,
                        start_char=current_pos,
                        chunk_index=len(chunks),
                        source_path=str(path),
                        semantic_type=fact["semantic_type"],
                        sheet_name=sheet_name,
                        row_start=fact.get("row_index"),
                        row_end=fact.get("row_index"),
                        cell_ref=fact.get("cell_ref"),
                        cell_range=fact.get("cell_range"),
                        row_label=fact.get("row_label"),
                        column_header=fact.get("column_header"),
                        unit=fact.get("unit"),
                        table_region_id=_table_region_id(sheet_name, 1),
                        interpretation_note=fact.get("interpretation_note"),
                    )
                )
                current_pos += len(content) + 2
            if len(facts["chunks"]) > _MAX_FACT_CHUNKS_PER_SHEET:
                note = (
                    f'Spreadsheet interpretation note: sheet="{sheet_name}" '
                    f"fact chunks truncated at {_MAX_FACT_CHUNKS_PER_SHEET} for ingestion safety."
                )
                table_artifacts["interpretation_notes"].append(note)

        return chunks, table_artifacts

    def _extract_rows(
        self,
        worksheet: Any,
        *,
        formula_ws: Any | None = None,
        max_sheet_rows: int,
        max_sheet_cols: int,
    ) -> dict[str, Any]:
        rows: list[dict[str, Any]] = []
        max_col_seen = 1
        truncated = False
        formula_rows = (
            formula_ws.iter_rows(max_row=max_sheet_rows + 1, max_col=max_sheet_cols, values_only=False)
            if formula_ws is not None
            else None
        )
        value_rows = worksheet.iter_rows(max_row=max_sheet_rows + 1, max_col=max_sheet_cols, values_only=True)

        for row_index, row in enumerate(value_rows, start=1):
            if row_index > max_sheet_rows:
                truncated = True
                break
            formula_row = next(formula_rows, None) if formula_rows is not None else None
            cells = []
            for col_index, value in enumerate(row, start=1):
                formula = _formula_for_cell(formula_row, col_index)
                clean = _format_cell(value)
                if clean == "" and formula:
                    clean = formula
                if clean == "":
                    continue
                max_col_seen = max(max_col_seen, col_index)
                cell_ref = f"{_column_letter(col_index)}{row_index}"
                cell = {
                    "column_index": col_index,
                    "column": _column_letter(col_index),
                    "value": clean,
                    "cell_ref": cell_ref,
                }
                if formula and formula != clean:
                    cell["formula"] = formula
                cells.append(cell)
            if cells:
                rows.append({"row_index": row_index, "cells": cells})
        if worksheet.max_column and worksheet.max_column > max_sheet_cols:
            truncated = True
        return {"rows": rows, "max_col": max_col_seen, "truncated": truncated}

    def _render_schema_chunk(self, filename: str, sheet_info: dict[str, Any]) -> str:
        hidden = "hidden" if sheet_info.get("hidden") else "visible"
        merged = ", ".join(sheet_info.get("merged_ranges") or []) or "none"
        return (
            f"Spreadsheet schema: workbook={filename} sheet={sheet_info['name']} "
            f"visibility={hidden} rows={sheet_info['max_row']} columns={sheet_info['max_column']} "
            f"merged_ranges={merged}"
        )

    def _render_row_chunk(self, sheet_name: str, rows: list[dict[str, Any]]) -> str:
        lines = [f"Spreadsheet sheet: {sheet_name}"]
        for row in rows:
            cells = row["cells"]
            rendered_cells = []
            for cell in cells:
                formula = f" formula={cell['formula']}" if cell.get("formula") else ""
                rendered_cells.append(f"{cell['cell_ref']}={cell['value']}{formula}")
            line = f"Row {row['row_index']}: " + " | ".join(rendered_cells)
            if len(cells) >= 2:
                label = cells[0]["value"]
                value = cells[1]["value"]
                if label:
                    line += f" | {label} = {value}"
            lines.append(line)
        return "\n".join(lines)

    def _build_fact_rows(
        self,
        sheet_name: str,
        rows: list[dict[str, Any]],
    ) -> dict[str, Any]:
        facts: dict[str, list[Any]] = {
            "chunks": [],
            "cell_facts": [],
            "table_facts": [],
            "semantic_sentences": [],
            "interpretation_notes": [],
        }
        active_headers: dict[int, str] = {}
        header_row_index: int | None = None
        table_region_id = _table_region_id(sheet_name, 1)

        for row in rows:
            row_index = row["row_index"]
            cells = row["cells"]
            row_label_cell = _first_text_cell(cells)
            row_label = row_label_cell["value"] if row_label_cell else None

            if _looks_like_header_row(cells):
                active_headers = {cell["column_index"]: cell["value"] for cell in cells}
                header_row_index = row_index
                header_content = (
                    f'Spreadsheet schema: sheet="{sheet_name}" header_row={row_index} | '
                    + " | ".join(f'{cell["cell_ref"]}="{cell["value"]}"' for cell in cells)
                )
                facts["chunks"].append(
                    {
                        "semantic_type": "spreadsheet_schema",
                        "row_index": row_index,
                        "cell_range": _row_cell_range(cells),
                        "content": header_content,
                    }
                )

            if len(cells) >= 2 and row_label_cell:
                value_cell = next(
                    (cell for cell in cells if cell["column_index"] != row_label_cell["column_index"]),
                    cells[1],
                )
                unit = _extract_unit(row_label)
                fact = {
                    "sheet_name": sheet_name,
                    "table_region_id": table_region_id,
                    "row_index": row_index,
                    "row_label": row_label,
                    "label_cell": row_label_cell["cell_ref"],
                    "value_cell": value_cell["cell_ref"],
                    "cell_ref": value_cell["cell_ref"],
                    "value": value_cell["value"],
                    "unit": unit,
                }
                facts["cell_facts"].append(fact)
                content = (
                    f'Spreadsheet cell fact: sheet="{sheet_name}" row={row_index} '
                    f'label="{row_label}" value="{value_cell["value"]}" '
                    f'label_cell={row_label_cell["cell_ref"]} value_cell={value_cell["cell_ref"]} '
                    f'cell={value_cell["cell_ref"]} row_label="{row_label}"'
                    + (f' unit="{unit}"' if unit else "")
                    + f" | {row_label} = {value_cell['value']}"
                )
                sentence = (
                    f'In spreadsheet sheet "{sheet_name}", label "{row_label}" has value '
                    f'"{value_cell["value"]}" at cell {value_cell["cell_ref"]}. '
                    f'En feuille "{sheet_name}", le libellé "{row_label}" vaut "{value_cell["value"]}".'
                )
                if unit:
                    sentence += f" Unit: {unit}."
                else:
                    sentence += " Unit: not explicit in the source cell."
                facts["semantic_sentences"].append(sentence)
                facts["chunks"].append(
                    {
                        "semantic_type": "spreadsheet_cell_fact",
                        "row_index": row_index,
                        "cell_ref": value_cell["cell_ref"],
                        "cell_range": f'{row_label_cell["cell_ref"]}:{value_cell["cell_ref"]}',
                        "row_label": row_label,
                        "unit": unit,
                        "content": content,
                    }
                )
                facts["chunks"].append(
                    {
                        "semantic_type": "spreadsheet_semantic_sentence",
                        "row_index": row_index,
                        "cell_ref": value_cell["cell_ref"],
                        "cell_range": f'{row_label_cell["cell_ref"]}:{value_cell["cell_ref"]}',
                        "row_label": row_label,
                        "unit": unit,
                        "content": f"Spreadsheet semantic sentence: {sentence}",
                    }
                )

            interpreted_cells = []
            for cell in cells:
                header = active_headers.get(cell["column_index"])
                if row_label and row_label_cell and cell["column_index"] == row_label_cell["column_index"]:
                    continue
                if not row_label and not header:
                    continue
                if header_row_index is not None and row_index == header_row_index:
                    continue

                unit = _extract_unit(row_label or header or "")
                fact = {
                    "sheet_name": sheet_name,
                    "table_region_id": table_region_id,
                    "row_index": row_index,
                    "cell_ref": cell["cell_ref"],
                    "value": cell["value"],
                    "row_label": row_label,
                    "column_header": header,
                    "unit": unit,
                }
                if cell.get("formula"):
                    fact["formula"] = cell["formula"]
                facts["table_facts"].append(fact)
                parts = [
                    f'sheet="{sheet_name}"',
                    f"row={row_index}",
                    f'cell={cell["cell_ref"]}',
                    f'value="{cell["value"]}"',
                ]
                if row_label:
                    parts.append(f'row_label="{row_label}"')
                    parts.append(f'metric="{row_label}"')
                    parts.append(f'parameter="{row_label}"')
                if header and header != cell["value"]:
                    parts.append(f'column_header="{header}"')
                    parts.append(f'trial_or_sample="{header}"')
                if unit:
                    parts.append(f'unit="{unit}"')
                if cell.get("formula"):
                    parts.append(f'formula="{cell["formula"]}"')
                interpreted_cells.append(" ".join(parts))
                facts["chunks"].append(
                    {
                        "semantic_type": "spreadsheet_table_fact",
                        "row_index": row_index,
                        "cell_ref": cell["cell_ref"],
                        "cell_range": cell["cell_ref"],
                        "row_label": row_label,
                        "column_header": header,
                        "unit": unit,
                        "content": "Spreadsheet table fact: " + " ".join(parts),
                    }
                )

            if interpreted_cells and len(interpreted_cells) > 1:
                facts["chunks"].append(
                    {
                        "semantic_type": "spreadsheet_table_fact",
                        "row_index": row_index,
                        "cell_range": _row_cell_range(cells),
                        "row_label": row_label,
                        "content": (
                            f'Spreadsheet table fact row summary: sheet="{sheet_name}" row={row_index} | '
                            + " | ".join(interpreted_cells[:24])
                        ),
                    }
                )

        if facts["cell_facts"]:
            facts["interpretation_notes"].append(
                "Label-value facts are local workbook evidence; they are not a global reference unless a Knowledge Guide explicitly says so."
            )
        return facts


def _chunk(content: str, start_char: int, chunk_index: int, **metadata: Any) -> dict[str, Any]:
    if len(content) > _MAX_CHUNK_CHARS:
        content = (
            content[: _MAX_CHUNK_CHARS - 160].rstrip()
            + "\nSpreadsheet interpretation note: chunk text truncated for embedding/upsert safety; "
            "cell-level facts remain indexed separately when available."
        )
    return {
        "content": content,
        "start_char": start_char,
        "end_char": start_char + len(content),
        "chunk_index": chunk_index,
        **{key: value for key, value in metadata.items() if value is not None},
    }


def _sheet_info(worksheet: Any, *, row_count: int, truncated: bool, max_col_seen: int) -> dict[str, Any]:
    merged_ranges: list[str] = []
    try:
        merged_ranges = [str(item) for item in getattr(worksheet, "merged_cells", []).ranges]
    except Exception:
        merged_ranges = []
    return {
        "name": worksheet.title,
        "rows": row_count,
        "max_row": worksheet.max_row,
        "max_column": worksheet.max_column,
        "max_column_seen": max_col_seen,
        "hidden": getattr(worksheet, "sheet_state", "visible") != "visible",
        "merged_ranges": merged_ranges,
        "truncated": truncated,
    }


def _formula_for_cell(formula_row: Any | None, col_index: int) -> str | None:
    if formula_row is None or col_index - 1 >= len(formula_row):
        return None
    try:
        value = formula_row[col_index - 1].value
    except Exception:
        return None
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\n", " ").strip()


def _first_text_cell(cells: list[dict[str, Any]]) -> dict[str, Any] | None:
    for cell in cells:
        if _looks_like_text_label(cell["value"]):
            return cell
    return cells[0] if cells else None


def _looks_like_text_label(value: str) -> bool:
    return any(char.isalpha() for char in value)


def _looks_like_header_row(cells: list[dict[str, Any]]) -> bool:
    if len(cells) < 2:
        return False
    joined = " ".join(cell["value"].lower() for cell in cells)
    header_markers = (
        "trial",
        "trials",
        "sample",
        "echantillon",
        "échantillon",
        "moy",
        "avg",
        "target",
        "min",
        "max",
        "n°",
        "no ",
        "test",
        "date",
        "customer",
        "client",
    )
    if any(marker in joined for marker in header_markers):
        return True
    return False


def _extract_unit(label: str | None) -> str | None:
    if not label:
        return None
    match = re.search(r"\(([^)]+)\)", label)
    if match:
        return match.group(1).strip()
    for unit in ("g/m²", "g/m2", "mm", "m/min", "mbar", "bar", "%", "N/50 mm", "daN"):
        if unit.lower() in label.lower():
            return unit
    return None


def _row_cell_range(cells: list[dict[str, Any]]) -> str | None:
    if not cells:
        return None
    return f"{cells[0]['cell_ref']}:{cells[-1]['cell_ref']}"


def _safe_positive_int(value: Any, default: int) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return default
    return parsed if parsed > 0 else default


def _batched(rows: list[dict[str, Any]], size: int) -> list[list[dict[str, Any]]]:
    return [rows[index : index + size] for index in range(0, len(rows), size)]


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters or "A"


def _table_region_id(sheet_name: str, index: int) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", sheet_name.lower()).strip("-") or "sheet"
    return f"{slug}-{index}"
