"""Tabular datasets: ingest, profile, and hand frames to the transform nodes.

Ownership split
---------------
The API never reads dataset bytes. It stages an upload into the ObjectStore and
queues :func:`ingest_dataset`; the worker parses, writes the canonical Parquet,
and folds everything a surface needs back into the row — ``schema_json``,
``preview_json`` and ``stats_json``. Rendering a dataset (list, detail, node
preview, feature picker) is therefore a single indexed SELECT.

Profiling
---------
:func:`profile_frame` is the one place that turns a frame into that read model.
Numeric columns get min/max/mean/std plus histogram bins; low-cardinality and
text columns get their top values; every column gets nulls and distincts. It
runs on the frame already in memory at ingest or transform time, so the cost is
one extra pass and the dataset detail page gets column histograms for free.

Storage
-------
Keys are ObjectStore-relative, never absolute paths: the API container and the
workers reach the same bytes through different backends (a local directory in
dev, MinIO in production). Workers materialize Parquet into a scratch file
before handing it to polars or duckdb, which keeps duckdb free of the httpfs
extension (no runtime download on a locked-down VM).
"""

from __future__ import annotations

import math
import re
import tempfile
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

from sqlalchemy.orm import Session as DBSession

from app.core.config import settings
from app.core.logging import get_logger
from app.models.tabular import TabularDataset
from app.services.object_store import get_object_store

logger = get_logger(__name__)

DATASET_INGEST_TASK = "agentium.dataset_ingest"

# The ingest steps a row reports through ``status_detail``, in order. They are
# codes rather than sentences because two locales poll the same row; see
# :func:`mark_step`. The frontend mirrors this tuple to render the whole
# check-list, so appending a step here is a contract change for both. A step may
# carry the row count it is working on as ``profiling:8412``.
INGEST_STEPS: tuple[str, ...] = ("queued", "reading", "profiling", "writing")

# Canonical column kinds. The UI maps these to icons and the training plane uses
# them to split numeric from categorical features, so they must stay stable and
# storage-agnostic (never a raw polars/duckdb dtype string).
KIND_INTEGER = "integer"
KIND_FLOAT = "float"
KIND_BOOLEAN = "boolean"
KIND_DATETIME = "datetime"
KIND_STRING = "string"
KIND_OTHER = "other"
NUMERIC_KINDS = frozenset({KIND_INTEGER, KIND_FLOAT})

_SUPPORTED_SUFFIXES = {
    ".csv": "csv",
    ".tsv": "tsv",
    ".txt": "csv",
    ".parquet": "parquet",
    ".json": "json",
    ".jsonl": "ndjson",
    ".ndjson": "ndjson",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
}
_SLUG_RE = re.compile(r"[^a-z0-9]+")


@dataclass(slots=True)
class TabularError(Exception):
    """Business error carrying the canonical ``{code, message}`` payload."""

    code: str
    message: str
    status_code: int = 422
    details: dict[str, Any] | None = None

    def payload(self) -> dict[str, Any]:
        return {
            "error": self.code.lower(),
            "code": self.code,
            "message": self.message,
            **dict(self.details or {}),
        }

    def __str__(self) -> str:  # pragma: no cover - repr convenience
        return f"{self.code}: {self.message}"


# ---------------------------------------------------------------------------
# Naming and keys
# ---------------------------------------------------------------------------


def slugify(value: str, *, fallback: str = "dataset") -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    ascii_only = normalized.encode("ascii", "ignore").decode("ascii").lower()
    slug = _SLUG_RE.sub("-", ascii_only).strip("-")
    return (slug or fallback)[:180]


def dataset_prefix(workspace_id: str, dataset_id: str) -> str:
    store = get_object_store()
    return store.key("workspaces", workspace_id, "tabular", "datasets", dataset_id)


def _parquet_key(workspace_id: str, dataset_id: str) -> str:
    return f"{dataset_prefix(workspace_id, dataset_id)}/data.parquet"


def _upload_key(workspace_id: str, dataset_id: str, filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    safe = f"source{suffix}" if suffix in _SUPPORTED_SUFFIXES else "source.bin"
    return f"{dataset_prefix(workspace_id, dataset_id)}/{safe}"


def detect_format(filename: str, content_type: str | None = None) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in _SUPPORTED_SUFFIXES:
        return _SUPPORTED_SUFFIXES[suffix]
    lowered = (content_type or "").lower()
    if "csv" in lowered:
        return "csv"
    if "parquet" in lowered:
        return "parquet"
    if "json" in lowered:
        return "json"
    if "sheet" in lowered or "excel" in lowered:
        return "xlsx"
    raise TabularError(
        code="DATASET_FORMAT_UNSUPPORTED",
        message=(
            "Supported files: CSV, TSV, Parquet, JSON/NDJSON and XLSX."
        ),
        details={"filename": str(filename or "")[:200]},
    )


def next_version(db: DBSession, *, workspace_id: str, slug: str) -> int:
    rows = (
        db.query(TabularDataset.version)
        .filter(
            TabularDataset.workspace_id == workspace_id,
            TabularDataset.slug == slug,
        )
        .all()
    )
    return max((int(row[0] or 0) for row in rows), default=0) + 1


# ---------------------------------------------------------------------------
# Profiling
# ---------------------------------------------------------------------------


def _column_kind(dtype: Any) -> str:
    import polars as pl

    if dtype == pl.Boolean:
        return KIND_BOOLEAN
    if dtype.is_integer():
        return KIND_INTEGER
    if dtype.is_float() or dtype.is_decimal():
        return KIND_FLOAT
    if dtype.is_temporal():
        return KIND_DATETIME
    if dtype in (pl.String, pl.Categorical, pl.Enum):
        return KIND_STRING
    return KIND_OTHER


def _json_scalar(value: Any) -> Any:
    """Make one cell JSON-safe without lying about it.

    Postgres JSON columns reject NaN/Infinity, and dates must reach the UI as
    ISO strings rather than repr(). Anything exotic degrades to ``str``.
    """

    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if isinstance(value, (datetime,)):
        return value.isoformat()
    if isinstance(value, str):
        return value[:2000]
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:  # noqa: BLE001 - fall through to str()
            pass
    return str(value)[:2000]


def _numeric_stats(series: Any, kind: str) -> dict[str, Any]:
    stats: dict[str, Any] = {}
    for label, getter in (
        ("min", series.min),
        ("max", series.max),
        ("mean", series.mean),
        ("std", series.std),
    ):
        try:
            stats[label] = _json_scalar(getter())
        except Exception:  # noqa: BLE001 - a stat is never worth failing on
            stats[label] = None
    bins = int(settings.tabular_histogram_bins)
    try:
        clean = series.drop_nulls().drop_nans() if kind == KIND_FLOAT else series.drop_nulls()
        if clean.len() > 1 and stats.get("min") != stats.get("max"):
            hist = clean.hist(bin_count=bins, include_category=False)
            counts_col = "count" if "count" in hist.columns else hist.columns[-1]
            edge_col = next(
                (c for c in hist.columns if c in ("breakpoint", "break_point")),
                hist.columns[0],
            )
            stats["histogram"] = [
                {
                    "upper": _json_scalar(row[edge_col]),
                    "count": int(row[counts_col] or 0),
                }
                for row in hist.iter_rows(named=True)
            ]
    except Exception:  # noqa: BLE001 - histogram is decoration, never a blocker
        logger.debug("tabular: histogram skipped", column=series.name)
    return stats


def _top_values(series: Any) -> list[dict[str, Any]]:
    limit = int(settings.tabular_top_values)
    try:
        counts = series.value_counts(sort=True).head(limit)
        value_col = series.name if series.name in counts.columns else counts.columns[0]
        count_col = "count" if "count" in counts.columns else counts.columns[-1]
        return [
            {
                "value": _json_scalar(row[value_col]),
                "count": int(row[count_col] or 0),
            }
            for row in counts.iter_rows(named=True)
        ]
    except Exception:  # noqa: BLE001 - decoration again
        return []


def profile_frame(frame: Any) -> dict[str, Any]:
    """Turn a polars frame into the read model every tabular surface renders."""

    preview_rows = int(settings.tabular_preview_rows)
    row_count = int(frame.height)
    schema: list[dict[str, Any]] = []
    stats: dict[str, Any] = {}

    for name, dtype in zip(frame.columns, frame.dtypes, strict=True):
        kind = _column_kind(dtype)
        series = frame.get_column(name)
        null_count = int(series.null_count())
        try:
            distinct = int(series.n_unique())
        except Exception:  # noqa: BLE001 - nested dtypes cannot be hashed
            distinct = 0
        schema.append(
            {
                "name": name,
                "kind": kind,
                "dtype": str(dtype),
                "nullable": null_count > 0,
            }
        )
        column_stats: dict[str, Any] = {
            "kind": kind,
            "nulls": null_count,
            "null_ratio": round(null_count / row_count, 4) if row_count else 0.0,
            "distinct": distinct,
        }
        if kind in NUMERIC_KINDS:
            column_stats.update(_numeric_stats(series, kind))
        elif kind == KIND_DATETIME:
            column_stats["min"] = _json_scalar(series.min())
            column_stats["max"] = _json_scalar(series.max())
        elif kind in (KIND_STRING, KIND_BOOLEAN):
            column_stats["top_values"] = _top_values(series)
        stats[name] = column_stats

    preview = [
        {key: _json_scalar(value) for key, value in row.items()}
        for row in frame.head(preview_rows).iter_rows(named=True)
    ]
    return {
        "row_count": row_count,
        "column_count": int(frame.width),
        "schema": schema,
        "preview": preview,
        "stats": stats,
    }


# ---------------------------------------------------------------------------
# Reading and writing frames
# ---------------------------------------------------------------------------


def _read_source(path: Path, fmt: str) -> Any:
    import polars as pl

    if fmt == "parquet":
        return pl.read_parquet(path)
    if fmt in ("csv", "tsv"):
        separator = "\t" if fmt == "tsv" else ","
        return pl.read_csv(
            path,
            separator=separator,
            try_parse_dates=True,
            infer_schema_length=10_000,
            ignore_errors=False,
            null_values=["", "NA", "N/A", "null", "NULL", "nan"],
        )
    if fmt == "ndjson":
        return pl.read_ndjson(path)
    if fmt == "json":
        return pl.read_json(path)
    if fmt == "xlsx":
        # openpyxl is a hard dependency of the document pipeline, so reading
        # sheets this way avoids polars' optional excel engines entirely.
        from openpyxl import load_workbook

        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise TabularError(
                code="DATASET_EMPTY",
                message="The uploaded sheet has no rows.",
            ) from exc
        columns = [
            str(value) if value not in (None, "") else f"column_{index + 1}"
            for index, value in enumerate(header)
        ]
        records = [dict(zip(columns, row, strict=False)) for row in rows]
        workbook.close()
        if not records:
            raise TabularError(
                code="DATASET_EMPTY", message="The uploaded sheet has no data rows."
            )
        return pl.DataFrame(records, infer_schema_length=None)
    raise TabularError(
        code="DATASET_FORMAT_UNSUPPORTED",
        message=f"Unsupported dataset format: {fmt}.",
    )


def materialize(dataset: TabularDataset, destination: Path) -> Path:
    """Copy a dataset's Parquet out of the ObjectStore into a scratch file."""

    if not dataset.storage_key:
        raise TabularError(
            code="DATASET_NOT_READY",
            message="The dataset has no materialized Parquet yet.",
            status_code=409,
        )
    store = get_object_store()
    return store.copy_to_local(dataset.storage_key, destination)


def read_frame(dataset: TabularDataset, *, columns: list[str] | None = None) -> Any:
    """Load a dataset as a polars frame (worker side)."""

    import polars as pl

    with tempfile.TemporaryDirectory(prefix="agentium-tabular-") as tmp:
        local = materialize(dataset, Path(tmp) / "data.parquet")
        return pl.read_parquet(local, columns=columns)


def read_page(
    dataset: TabularDataset, *, offset: int = 0, limit: int | None = None
) -> dict[str, Any]:
    """A window of a dataset's rows, read from the Parquet on demand.

    The fifty rows stored on the row at ingest are enough for a page to open
    without a second request, and they are the wrong answer to "show me more":
    every dataset here is immutable, so row 8 400 of an 8 412-row dataset is a
    fact that exists and simply is not in that cache. This reads it.

    Scanned and sliced rather than read and trimmed, so the cost is the page and
    not the file — polars pushes the slice into the Parquet reader, which skips
    whole row groups. The row count still comes off the database row rather than
    from counting here: it was computed at ingest and has not changed since.
    """

    import polars as pl

    total = int(dataset.row_count or 0)
    start = max(0, int(offset))
    size = int(limit if limit is not None else settings.tabular_preview_rows)
    size = max(1, min(size, int(settings.tabular_preview_page_max)))
    if total and start >= total:
        # Past the end is an empty page, not a refusal: a reader who paged one
        # step too far should see "no more rows", not an error dialog.
        return {"rows": [], "offset": start, "limit": size, "total": total}
    with tempfile.TemporaryDirectory(prefix="agentium-tabular-") as tmp:
        local = materialize(dataset, Path(tmp) / "data.parquet")
        frame = pl.scan_parquet(local).slice(start, size).collect()
    rows = [
        {key: _json_scalar(value) for key, value in row.items()}
        for row in frame.iter_rows(named=True)
    ]
    return {"rows": rows, "offset": start, "limit": size, "total": total}


def _write_frame(frame: Any, *, workspace_id: str, dataset_id: str) -> tuple[str, int]:
    store = get_object_store()
    key = _parquet_key(workspace_id, dataset_id)
    with tempfile.TemporaryDirectory(prefix="agentium-tabular-") as tmp:
        local = Path(tmp) / "data.parquet"
        frame.write_parquet(local, compression="zstd")
        size = int(local.stat().st_size)
        store.write_bytes(key, local.read_bytes())
    return key, size


# ---------------------------------------------------------------------------
# Lifecycle
# ---------------------------------------------------------------------------


def create_upload(
    db: DBSession,
    *,
    workspace_id: str,
    name: str,
    filename: str,
    content_type: str | None,
    payload: bytes,
    description: str | None = None,
    created_by: str | None = None,
) -> TabularDataset:
    """Stage raw upload bytes and return the ``pending`` row to ingest."""

    if not settings.tabular_data_enabled:
        raise TabularError(
            code="TABULAR_DISABLED",
            message="The data plane is not enabled on this deployment.",
            status_code=409,
        )
    if not payload:
        raise TabularError(code="DATASET_EMPTY", message="The uploaded file is empty.")
    limit = int(settings.tabular_upload_max_bytes)
    if len(payload) > limit:
        raise TabularError(
            code="DATASET_TOO_LARGE",
            message=f"The file exceeds the {limit // 1024 // 1024} MB upload limit.",
            details={"size_bytes": len(payload), "limit_bytes": limit},
        )
    fmt = detect_format(filename, content_type)
    label = (name or Path(filename or "").stem or "dataset").strip()[:200]
    slug = slugify(label)
    dataset = TabularDataset(
        id=str(uuid4()),
        workspace_id=workspace_id,
        name=label,
        slug=slug,
        version=next_version(db, workspace_id=workspace_id, slug=slug),
        description=(description or None),
        source="upload",
        status="pending",
        status_detail=INGEST_STEPS[0],
        original_filename=str(filename or "")[:400],
        content_type=(content_type or None),
        size_bytes=len(payload),
        created_by=created_by,
        parent_ids=[],
        lineage_json={"format": fmt},
    )
    dataset.upload_key = _upload_key(workspace_id, dataset.id, filename or "")
    get_object_store().write_bytes(dataset.upload_key, payload)
    db.add(dataset)
    db.flush()
    return dataset


def mark_step(
    db: DBSession,
    dataset: TabularDataset,
    step: str,
    *,
    rows: int | None = None,
) -> None:
    """Publish which ingest step the worker is on, as a code the UI translates.

    A code and not a sentence: this row is polled by a French and an English
    surface, and a worker that wrote "Reading the uploaded file" would put an
    English string on both. The codes are :data:`INGEST_STEPS`, in the order the
    worker passes through them, so a client can render the whole list and mark
    how far it got instead of showing one line at a time.

    ``rows`` rides along as ``profiling:8412`` once the parse has produced a
    height. It is the difference between a progress line and a fact: "profiling"
    is what a spinner says, "8 412 rows" is what tells the person watching that
    their file arrived whole. It stays a number rather than a formatted string
    because the thousands separator belongs to the reader's locale.
    """

    detail = str(step)
    if rows is not None:
        detail = f"{detail}:{int(rows)}"
    dataset.status_detail = detail[:300]
    dataset.updated_at = datetime.utcnow()
    db.commit()


def ingest_dataset(dataset_id: str) -> dict[str, Any]:
    """Parse a staged upload into the canonical Parquet + profile (worker side)."""

    from app.db.base import SessionLocal

    started = time.monotonic()
    with SessionLocal() as db:
        dataset = (
            db.query(TabularDataset).filter(TabularDataset.id == dataset_id).first()
        )
        if dataset is None:
            return {"id": dataset_id, "status": "missing"}
        if dataset.status in ("ready", "deleted"):
            return {"id": dataset_id, "status": dataset.status}
        if not dataset.upload_key:
            _fail(db, dataset, "dataset_no_upload")
            return {"id": dataset_id, "status": "failed"}

        dataset.status = "ingesting"
        mark_step(db, dataset, "reading")
        fmt = str((dataset.lineage_json or {}).get("format") or "csv")
        try:
            with tempfile.TemporaryDirectory(prefix="agentium-ingest-") as tmp:
                local = get_object_store().copy_to_local(
                    dataset.upload_key, Path(tmp) / f"source.{fmt}"
                )
                frame = _read_source(local, fmt)
                if frame.height == 0:
                    raise TabularError(
                        code="DATASET_EMPTY", message="The file has no data rows."
                    )
                max_columns = int(settings.tabular_max_columns)
                if frame.width > max_columns:
                    raise TabularError(
                        code="DATASET_TOO_WIDE",
                        message=f"The file has more than {max_columns} columns.",
                    )
                mark_step(db, dataset, "profiling", rows=frame.height)
                profile = profile_frame(frame)
                mark_step(db, dataset, "writing", rows=frame.height)
                key, size = _write_frame(
                    frame,
                    workspace_id=dataset.workspace_id,
                    dataset_id=dataset.id,
                )
        except TabularError as exc:
            _fail(db, dataset, exc.code.lower())
            return {"id": dataset_id, "status": "failed", "code": exc.code}
        except Exception as exc:  # noqa: BLE001 - surface the parser message
            _fail(db, dataset, f"ingest_failed: {exc}"[:2000])
            logger.warning("tabular: ingest failed", dataset_id=dataset_id, error=str(exc))
            return {"id": dataset_id, "status": "failed"}

        _apply_profile(dataset, profile)
        dataset.storage_key = key
        dataset.size_bytes = size
        dataset.status = "ready"
        dataset.status_detail = None
        dataset.error = None
        dataset.ingested_at = datetime.utcnow()
        dataset.ingest_duration_ms = (time.monotonic() - started) * 1000
        db.commit()
        logger.info(
            "tabular: ingested",
            dataset_id=dataset.id,
            rows=dataset.row_count,
            columns=dataset.column_count,
            size_bytes=size,
        )
        return {
            "id": dataset.id,
            "status": "ready",
            "rows": int(dataset.row_count or 0),
            "columns": int(dataset.column_count or 0),
        }


def _apply_profile(dataset: TabularDataset, profile: dict[str, Any]) -> None:
    dataset.row_count = int(profile["row_count"])
    dataset.column_count = int(profile["column_count"])
    dataset.schema_json = profile["schema"]
    dataset.preview_json = profile["preview"]
    dataset.stats_json = profile["stats"]


def _fail(db: DBSession, dataset: TabularDataset, error: str) -> None:
    dataset.status = "failed"
    dataset.status_detail = None
    dataset.error = error
    dataset.updated_at = datetime.utcnow()
    db.commit()


def register_frame(
    db: DBSession,
    *,
    workspace_id: str,
    name: str,
    frame: Any,
    source: str = "transform",
    produced_by: str | None = None,
    parent_ids: Iterable[str] | None = None,
    run_id: str | None = None,
    node_id: str | None = None,
    description: str | None = None,
    created_by: str | None = None,
    lineage: dict[str, Any] | None = None,
) -> TabularDataset:
    """Persist an in-memory frame as a ready dataset (transform/score output)."""

    parents = [str(p) for p in (parent_ids or []) if p]
    max_rows = int(settings.tabular_transform_max_rows)
    if frame.height > max_rows:
        raise TabularError(
            code="DATASET_TOO_LARGE",
            message=f"The result exceeds {max_rows:,} rows.",
            details={"row_count": int(frame.height)},
        )
    label = (name or "result").strip()[:200] or "result"
    slug = slugify(label)
    dataset = TabularDataset(
        id=str(uuid4()),
        workspace_id=workspace_id,
        name=label,
        slug=slug,
        version=next_version(db, workspace_id=workspace_id, slug=slug),
        description=(description or None),
        source=source,
        status="ingesting",
        produced_by=produced_by,
        parent_ids=parents,
        run_id=run_id,
        node_id=node_id,
        created_by=created_by,
        lineage_json=dict(lineage or {}),
    )
    db.add(dataset)
    db.flush()

    profile = profile_frame(frame)
    key, size = _write_frame(
        frame, workspace_id=workspace_id, dataset_id=dataset.id
    )
    _apply_profile(dataset, profile)
    dataset.storage_key = key
    dataset.size_bytes = size
    dataset.status = "ready"
    dataset.status_detail = None
    dataset.ingested_at = datetime.utcnow()
    db.flush()
    return dataset


def soft_delete(db: DBSession, dataset: TabularDataset) -> TabularDataset:
    """Retire a dataset.

    The production object policy is append-only (versioned bucket, no delete
    grant), so the row is the authority and byte reclamation belongs to the
    storage lifecycle. Best-effort removal still runs for local deployments.
    """

    dataset.status = "deleted"
    dataset.status_detail = None
    dataset.updated_at = datetime.utcnow()
    store = get_object_store()
    if store.backend == "local":
        try:
            store.delete_prefix(dataset_prefix(dataset.workspace_id, dataset.id))
        except Exception:  # noqa: BLE001 - row flip stays authoritative
            logger.debug("tabular: local delete skipped", dataset_id=dataset.id)
    db.commit()
    return dataset


# ---------------------------------------------------------------------------
# Projections
# ---------------------------------------------------------------------------


def serialize_dataset(
    dataset: TabularDataset, *, include_preview: bool = False
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "id": dataset.id,
        "name": dataset.name,
        "slug": dataset.slug,
        "version": int(dataset.version or 1),
        "description": dataset.description,
        "source": dataset.source,
        "status": dataset.status,
        "status_detail": dataset.status_detail,
        "error": dataset.error,
        "row_count": int(dataset.row_count) if dataset.row_count is not None else None,
        "column_count": (
            int(dataset.column_count) if dataset.column_count is not None else None
        ),
        "size_bytes": int(dataset.size_bytes) if dataset.size_bytes is not None else None,
        "schema": list(dataset.schema_json or []),
        "original_filename": dataset.original_filename,
        "run_id": dataset.run_id,
        "node_id": dataset.node_id,
        "parent_ids": list(dataset.parent_ids or []),
        "produced_by": dataset.produced_by,
        "created_at": dataset.created_at.isoformat() if dataset.created_at else None,
        "updated_at": dataset.updated_at.isoformat() if dataset.updated_at else None,
        "ingested_at": dataset.ingested_at.isoformat() if dataset.ingested_at else None,
        "ingest_duration_ms": dataset.ingest_duration_ms,
    }
    if include_preview:
        payload["preview"] = list(dataset.preview_json or [])
        payload["stats"] = dict(dataset.stats_json or {})
        payload["lineage"] = _public_lineage(dataset)
    return payload


# Which of a producer's lineage notes the detail surface is allowed to render.
# A curated list rather than the whole bag: the block also holds the authored
# statement, and a dataset page is not where a Flow's program is published.
_PUBLIC_LINEAGE_KEYS = (
    "engine",
    "model",
    "added_columns",
    "output_model",
    "models",
    "tests_total",
    "tests_failed",
    "code_lines",
    "duration_ms",
)


def _public_lineage(dataset: TabularDataset) -> dict[str, Any]:
    """How this dataset was produced, as the detail page tells the story.

    The interesting entry is ``model``: a scored dataset knows which model
    version wrote its extra columns, which is what lets those columns carry a
    "scored by Churn v3" badge instead of appearing out of nowhere.
    """

    raw = dataset.lineage_json if isinstance(dataset.lineage_json, dict) else {}
    return {key: raw[key] for key in _PUBLIC_LINEAGE_KEYS if key in raw}


def dataset_reference(dataset: TabularDataset) -> dict[str, Any]:
    """The envelope shape datasets travel as inside the DAG.

    By reference, never by value: nodes downstream resolve the id, and the
    canvas badges read ``rows``/``columns`` straight off this block.
    """

    return {
        "dataset_id": dataset.id,
        "name": dataset.name,
        "slug": dataset.slug,
        "version": int(dataset.version or 1),
        "rows": int(dataset.row_count or 0),
        "columns": int(dataset.column_count or 0),
        "schema": list(dataset.schema_json or []),
    }


def get_dataset(
    db: DBSession, *, dataset_id: str, workspace_id: str, ready_only: bool = False
) -> TabularDataset:
    dataset = (
        db.query(TabularDataset)
        .filter(
            TabularDataset.id == dataset_id,
            TabularDataset.workspace_id == workspace_id,
        )
        .first()
    )
    if dataset is None or dataset.status == "deleted":
        raise TabularError(
            code="DATASET_NOT_FOUND",
            message="The dataset does not exist in this workspace.",
            status_code=404,
        )
    if ready_only and dataset.status != "ready":
        raise TabularError(
            code="DATASET_NOT_READY",
            message="The dataset is still being prepared.",
            status_code=409,
            details={"status": dataset.status},
        )
    return dataset


def resolve_dataset_ref(
    db: DBSession, *, workspace_id: str, ref: Any, ready_only: bool = True
) -> TabularDataset:
    """Accept an id, a ``{dataset_id: ...}`` envelope or a slug and resolve it.

    Both ``slug`` and ``dataset_slug`` are read: the first is how an upstream
    envelope names its own lineage, the second is how a node pins one. They mean
    the same thing — follow the latest ready version of that lineage — and a
    resolver that honoured only one silently ignored half the pins ever written.
    """

    dataset_id: str | None = None
    slug: str | None = None
    if isinstance(ref, str):
        dataset_id = ref.strip() or None
    elif isinstance(ref, dict):
        raw_id = ref.get("dataset_id") or ref.get("id")
        dataset_id = str(raw_id).strip() if raw_id else None
        raw_slug = ref.get("slug") or ref.get("dataset_slug")
        slug = str(raw_slug).strip() if raw_slug else None
    if dataset_id:
        try:
            return get_dataset(
                db,
                dataset_id=dataset_id,
                workspace_id=workspace_id,
                ready_only=ready_only,
            )
        except TabularError:
            if not slug:
                raise
    if slug:
        latest = (
            db.query(TabularDataset)
            .filter(
                TabularDataset.workspace_id == workspace_id,
                TabularDataset.slug == slug,
                TabularDataset.status == "ready",
            )
            .order_by(TabularDataset.version.desc())
            .first()
        )
        if latest is not None:
            return latest
    raise TabularError(
        code="DATASET_NOT_FOUND",
        message="No dataset matched the reference.",
        status_code=404,
    )


__all__ = [
    "DATASET_INGEST_TASK",
    "INGEST_STEPS",
    "KIND_BOOLEAN",
    "KIND_DATETIME",
    "KIND_FLOAT",
    "KIND_INTEGER",
    "KIND_OTHER",
    "KIND_STRING",
    "NUMERIC_KINDS",
    "TabularError",
    "create_upload",
    "dataset_prefix",
    "dataset_reference",
    "detect_format",
    "get_dataset",
    "ingest_dataset",
    "mark_step",
    "materialize",
    "next_version",
    "profile_frame",
    "read_frame",
    "register_frame",
    "resolve_dataset_ref",
    "serialize_dataset",
    "slugify",
    "soft_delete",
]
