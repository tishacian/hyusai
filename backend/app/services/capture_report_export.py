"""Branded Andritz export (PDF / DOCX) for capture / FSE intervention reports.

On-demand rendering from an existing proposal + session plan. Does not alter
publication, indexing, or object-store flows — the markdown fiche remains the
canonical KB artefact; PDF/DOCX are presentation derivatives.
"""
from __future__ import annotations

import html
import io
import re
from pathlib import Path
from typing import Any, Dict, Literal, Mapping, Optional, Tuple

from app.core.logging import get_logger
from app.models.expert_capture import ExpertCaptureSession, KnowledgeUpdateProposal
from app.services.capture_templates import (
    header_fields_from_plan,
    publication_filename_from_header,
    publication_title_from_header,
)

logger = get_logger(__name__)

ExportFormat = Literal["pdf", "docx"]

RESOURCES_DIR = Path(__file__).resolve().parents[1] / "resources" / "andritz"
ANDRITZ_LOGO = RESOURCES_DIR / "andritz.png"

# Andritz institutional blue (close to historical EX70 covers).
_ANDRITZ_BLUE = "#003366"
_ANDRITZ_ACCENT = "#0055A4"


def export_urls_for_proposal(proposal_id: str) -> Dict[str, str]:
    base = f"/api/v1/knowledge-capture/proposals/{proposal_id}/export"
    return {
        "pdf_url": f"{base}?format=pdf",
        "docx_url": f"{base}?format=docx",
    }


def build_capture_report_export(
    *,
    proposal: KnowledgeUpdateProposal,
    session: Optional[ExpertCaptureSession],
    fmt: ExportFormat,
) -> Tuple[bytes, str, str]:
    """Return ``(bytes, media_type, filename)`` for the requested format."""
    payload = proposal.proposal if isinstance(proposal.proposal, Mapping) else {}
    markdown = str(
        payload.get("report_markdown")
        or (payload.get("recommended_ingestion") or {}).get("content")
        or ""
    ).strip()
    if not markdown:
        raise ValueError("Report markdown is empty — finalize the capture before exporting")

    plan = (session.plan if session and isinstance(session.plan, Mapping) else {}) or {}
    header = header_fields_from_plan(plan)
    title = publication_title_from_header(
        header,
        fallback=str(payload.get("title") or (session.title if session else "") or "Rapport FSE"),
    )
    doc_ref = _doc_ref_from_plan(plan)
    type_label = _type_label_from_plan(plan)

    stem = publication_filename_from_header(header, fallback_session_id=proposal.session_id)
    if stem.endswith(".md"):
        stem = stem[:-3]

    if fmt == "pdf":
        content = _build_pdf(
            markdown=markdown,
            title=title,
            header=header,
            doc_ref=doc_ref,
            type_label=type_label,
        )
        return content, "application/pdf", f"{stem}.pdf"

    if fmt == "docx":
        content = _build_docx(
            markdown=markdown,
            title=title,
            header=header,
            doc_ref=doc_ref,
            type_label=type_label,
        )
        return (
            content,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{stem}.docx",
        )

    raise ValueError(f"Unsupported export format: {fmt}")


def _doc_ref_from_plan(plan: Mapping[str, Any]) -> str:
    snap = plan.get("capture_template") if isinstance(plan.get("capture_template"), Mapping) else {}
    return str(snap.get("doc_ref") or "").strip()


def _type_label_from_plan(plan: Mapping[str, Any]) -> str:
    snap = plan.get("capture_template") if isinstance(plan.get("capture_template"), Mapping) else {}
    return str(
        snap.get("intervention_type_label")
        or plan.get("intervention_type")
        or ""
    ).strip()


def _cartouche_rows(header: Mapping[str, Any]) -> list[tuple[str, str]]:
    from app.services.capture_report_templates import _fse_display
    from app.services.capture_templates import _DISTRIBUTION_OPTIONS

    return [
        ("Client", _fse_display(header.get("customer"))),
        ("Pays", _fse_display(header.get("country"))),
        ("Site / machine", _fse_display(header.get("site_or_machine"))),
        ("Référence", _fse_display(header.get("reference"))),
        ("Participants", _fse_display(header.get("participants"))),
        ("Émis par", _fse_display(header.get("issued_by"))),
        ("Date", _fse_display(header.get("intervention_date"))),
        ("Semaine", _fse_display(header.get("week"), empty=_fse_display(header.get("intervention_date")))),
        (
            "Diffusion",
            _fse_display(header.get("distribution"), options=list(_DISTRIBUTION_OPTIONS)),
        ),
    ]


def _markdown_to_html_body(markdown_text: str) -> str:
    try:
        import markdown as md  # type: ignore

        return md.markdown(markdown_text, extensions=["extra", "tables", "toc"])
    except Exception as exc:  # noqa: BLE001
        logger.info("capture_report_export.markdown_lib_missing", error=str(exc))
        escaped = html.escape(markdown_text)
        return f"<pre>{escaped}</pre>"


def _branded_html(
    *,
    markdown: str,
    title: str,
    header: Mapping[str, Any],
    doc_ref: str,
    type_label: str,
) -> str:
    logo_uri = ""
    if ANDRITZ_LOGO.is_file():
        logo_uri = ANDRITZ_LOGO.resolve().as_uri()

    cartouche_cells = "".join(
        f"<tr><th>{html.escape(label)}</th><td>{html.escape(value)}</td></tr>"
        for label, value in _cartouche_rows(header)
    )
    meta_bits = [bit for bit in (type_label, doc_ref) if bit]
    meta_line = " · ".join(meta_bits)
    body = _markdown_to_html_body(markdown)
    logo_block = (
        f'<img class="logo" src="{logo_uri}" alt="ANDRITZ" />' if logo_uri else '<div class="logo-text">ANDRITZ</div>'
    )

    return f"""<!doctype html>
<html lang="fr">
<head>
<meta charset="utf-8" />
<title>{html.escape(title)}</title>
<style>
  @page {{
    size: A4;
    margin: 16mm 14mm 18mm 14mm;
    @bottom-center {{
      content: "ANDRITZ · {html.escape(doc_ref or 'Rapport intervention')} · page " counter(page);
      font-size: 8pt;
      color: #667788;
    }}
  }}
  body {{
    font-family: "Helvetica Neue", Helvetica, Arial, sans-serif;
    color: #1a2332;
    font-size: 10.5pt;
    line-height: 1.45;
  }}
  .banner {{
    display: flex;
    align-items: center;
    justify-content: space-between;
    border-bottom: 3px solid {_ANDRITZ_BLUE};
    padding-bottom: 6mm;
    margin-bottom: 6mm;
  }}
  .logo {{ height: 14mm; }}
  .logo-text {{
    font-weight: 800;
    font-size: 18pt;
    letter-spacing: 0.12em;
    color: {_ANDRITZ_BLUE};
  }}
  .banner-meta {{
    text-align: right;
    font-size: 9pt;
    color: #445566;
  }}
  h1 {{
    color: {_ANDRITZ_BLUE};
    font-size: 16pt;
    margin: 0 0 3mm;
  }}
  h2 {{
    color: {_ANDRITZ_ACCENT};
    font-size: 12pt;
    margin-top: 7mm;
    border-bottom: 1px solid #c5d0dc;
    padding-bottom: 1.5mm;
  }}
  h3 {{ color: {_ANDRITZ_BLUE}; font-size: 11pt; margin-top: 5mm; }}
  table.cartouche {{
    border-collapse: collapse;
    width: 100%;
    margin: 3mm 0 6mm;
    font-size: 9.5pt;
  }}
  table.cartouche th {{
    width: 32%;
    background: #e8eef5;
    color: {_ANDRITZ_BLUE};
    text-align: left;
    padding: 3px 6px;
    border: 1px solid #b8c6d6;
  }}
  table.cartouche td {{
    padding: 3px 6px;
    border: 1px solid #b8c6d6;
  }}
  table {{
    border-collapse: collapse;
    width: 100%;
    margin: 3mm 0;
    font-size: 9.5pt;
  }}
  th, td {{ border: 1px solid #c3cbd6; padding: 3px 6px; text-align: left; }}
  th {{ background: #eaf1fb; color: {_ANDRITZ_BLUE}; }}
  ul, ol {{ margin: 2mm 0 2mm 5mm; }}
  p {{ margin: 2mm 0; }}
</style>
</head>
<body>
  <div class="banner">
    {logo_block}
    <div class="banner-meta">
      <div style="font-weight:700; color:{_ANDRITZ_BLUE};">Field Service Excellence</div>
      <div>{html.escape(meta_line)}</div>
    </div>
  </div>
  <h1>{html.escape(title)}</h1>
  <table class="cartouche">{cartouche_cells}</table>
  {body}
</body>
</html>
"""


def _build_pdf(
    *,
    markdown: str,
    title: str,
    header: Mapping[str, Any],
    doc_ref: str,
    type_label: str,
) -> bytes:
    html_doc = _branded_html(
        markdown=markdown,
        title=title,
        header=header,
        doc_ref=doc_ref,
        type_label=type_label,
    )
    try:
        from weasyprint import HTML  # type: ignore

        base = str(RESOURCES_DIR if RESOURCES_DIR.is_dir() else Path(__file__).resolve().parents[3])
        return HTML(string=html_doc, base_url=base).write_pdf()
    except Exception as exc:  # noqa: BLE001
        logger.info("capture_report_export.weasyprint_unavailable", error=str(exc))
        return _stub_pdf_bytes(title=title)


def _stub_pdf_bytes(*, title: str) -> bytes:
    """Minimal valid PDF when WeasyPrint system libs are unavailable."""
    safe = re.sub(r"[()]", "", title)[:80] or "Rapport FSE"
    notice = "Rendu de secours (weasyprint indisponible)."
    stream = f"BT /F1 14 Tf 50 760 Td ({safe}) Tj ET BT /F1 10 Tf 50 740 Td ({notice}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    parts = ["%PDF-1.4\n"]
    offsets: list[int] = []
    cursor = len(parts[0])
    for idx, obj in enumerate(objects, start=1):
        offsets.append(cursor)
        chunk = f"{idx} 0 obj\n{obj}\nendobj\n"
        parts.append(chunk)
        cursor += len(chunk)
    xref_pos = cursor
    xref = [f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"]
    for off in offsets:
        xref.append(f"{off:010d} 00000 n \n")
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    )
    parts.extend(xref)
    parts.append(trailer)
    return "".join(parts).encode("latin-1", errors="replace")


def _build_docx(
    *,
    markdown: str,
    title: str,
    header: Mapping[str, Any],
    doc_ref: str,
    type_label: str,
) -> bytes:
    from docx import Document  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Inches, Pt, RGBColor  # type: ignore

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Banner
    if ANDRITZ_LOGO.is_file():
        try:
            doc.add_picture(str(ANDRITZ_LOGO), width=Inches(1.6))
        except Exception as exc:  # noqa: BLE001
            logger.info("capture_report_export.logo_embed_failed", error=str(exc))
            p = doc.add_paragraph()
            run = p.add_run("ANDRITZ")
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    else:
        p = doc.add_paragraph()
        run = p.add_run("ANDRITZ")
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta_run = meta.add_run(
        " · ".join(bit for bit in ("Field Service Excellence", type_label, doc_ref) if bit)
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x44, 0x55, 0x66)

    heading = doc.add_heading(title, level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    # Cartouche table
    rows = _cartouche_rows(header)
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for idx, (label, value) in enumerate(rows):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value

    doc.add_paragraph()
    _append_markdown_to_docx(doc, markdown)

    footer = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    footer.text = f"ANDRITZ · {doc_ref or 'Rapport intervention'}"
    if footer.runs:
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = RGBColor(0x66, 0x77, 0x88)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _append_markdown_to_docx(doc: Any, markdown_text: str) -> None:
    """Lightweight markdown → DOCX (headings, lists, paragraphs, pipe tables)."""
    from docx.shared import Pt, RGBColor  # type: ignore

    lines = markdown_text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1].strip()):
            table_lines = [stripped]
            i += 1
            # skip separator
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_pipe_table(doc, table_lines)
            continue
        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            h = doc.add_heading(text, level=min(level, 3))
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x00, 0x55, 0xA4)
            i += 1
            continue
        if stripped.startswith(("- ", "* ")):
            while i < len(lines) and lines[i].strip().startswith(("- ", "* ")):
                item = re.sub(r"^[-*]\s+", "", lines[i].strip())
                doc.add_paragraph(item, style="List Bullet")
                i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                item = re.sub(r"^\d+\.\s+", "", lines[i].strip())
                doc.add_paragraph(item, style="List Number")
                i += 1
            continue
        para = doc.add_paragraph(_strip_inline_md(stripped))
        for run in para.runs:
            run.font.size = Pt(10)
        i += 1


def _add_pipe_table(doc: Any, table_lines: list[str]) -> None:
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            table.rows[r_idx].cells[c_idx].text = row[c_idx] if c_idx < len(row) else ""


def _strip_inline_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text
