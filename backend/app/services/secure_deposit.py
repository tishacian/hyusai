"""Secure Deposit service.

This module backs both the public web drop portal and the future SFTP
subsystem. A deposit link is workspace-bound, password-protected, and writes
only to staging until an authorized Agentium user promotes a file to Knowledge.
"""
from __future__ import annotations

import base64
import hashlib
import hmac
import json
import mimetypes
import os
import re
import secrets
import shutil
import tempfile
import zipfile
from datetime import datetime, timedelta
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Optional
from urllib.parse import quote

import jwt
from fastapi import HTTPException, UploadFile, status
from sqlalchemy.orm import Session as DBSession

from app.core.config import settings
from app.core.settings_manager import get_resolved_settings
from app.models.knowledge_collection import KnowledgeCollection, WorkerJob
from app.models.secure_deposit import DepositAccessLink, DepositFile
from app.models.user import User
from app.models.workspace import Workspace
from app.services.audit_logger import emit_audit_event
from app.services.knowledge_collections import (
    collection_inventory,
    create_or_get_collection,
    create_worker_job,
    document_manifest_key,
    original_key,
    update_collection_status,
)
from app.services.knowledge_collections import (
    serialize_job as serialize_worker_job,
)
from app.services.object_store import get_object_store
from app.services.rag.document_service import DocumentService
from app.services.rag.project_references import (
    LEGACY_PROJECT_REFERENCE_RE,
    andritz_project_scheme_active,
    derive_project_reference,
    using_workspace_project_scheme,
)
from app.services.rag.vector_store_config import resolve_vector_db_type
from app.services.worker_dispatch import dispatch_worker_job

try:  # pragma: no cover - exercised when dependency is installed.
    from argon2 import PasswordHasher
    from argon2.exceptions import VerifyMismatchError

    _ARGON2 = PasswordHasher()
except Exception:  # noqa: BLE001 - dev/test fallback when deps are stale.
    PasswordHasher = None  # type: ignore
    VerifyMismatchError = Exception  # type: ignore
    _ARGON2 = None


_PBKDF2_PREFIX = "pbkdf2_sha256"
_TOKEN_ALGORITHM = "HS256"
SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY = "release_a_canary_v1"
SFTP_RELEASE_A_CANARY_ACCESS_ID_PREFIX = "ra1_"
_UPLOAD_CHUNK_BYTES = 1024 * 1024
_TEXT_PREVIEW_BYTES = 1024 * 1024
_STRUCTURED_PREVIEW_MAX_BYTES = 25 * 1024 * 1024
_INLINE_PREVIEW_MAX_BYTES = 25 * 1024 * 1024
# PDFs (e.g. Andritz carde manuals) routinely exceed the 25 MB image cap; the
# viewer streams the bytes lazily as a blob, so a generous cap is fine.
_PDF_PREVIEW_MAX_BYTES = 256 * 1024 * 1024
_TEXT_EXTENSIONS = {
    "csv",
    "json",
    "log",
    "md",
    "rst",
    "txt",
    "xml",
    "yaml",
    "yml",
}
_HTML_EXTENSIONS = {"html", "htm"}
_SPREADSHEET_EXTENSIONS = {"xlsx", "xlsm", "xltx", "xltm"}
_LEGACY_SPREADSHEET_EXTENSIONS = {"xls"}
_DOCX_EXTENSIONS = {"docx"}
_IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "bmp", "tif", "tiff", "webp"}
_DOCX_PREVIEW_MAX_BYTES = 25 * 1024 * 1024
_DOCX_PREVIEW_MAX_CHARS = 200_000
_ARCHIVE_BROWSER_MAX_ENTRIES = 5_000
_ARCHIVE_PROMOTION_EXTENSIONS = {"csv", "html", "htm", "md", "pdf", "txt"} | _DOCX_EXTENSIONS | _IMAGE_EXTENSIONS
_WORKER_PROMOTION_EXTENSIONS = (
    _ARCHIVE_PROMOTION_EXTENSIONS
    | _SPREADSHEET_EXTENSIONS
    | _TEXT_EXTENSIONS
    | _DOCX_EXTENSIONS
    | _IMAGE_EXTENSIONS
    | {"log", "markdown", "rst"}
)
_BULK_PROMOTION_MAX_FILES = 50
_BULK_PROMOTION_MAX_DOCUMENTS = 200
def enabled_workspace_slugs() -> set[str]:
    return {
        item.strip().lower()
        for item in (settings.secure_deposit_enabled_workspace_slugs or "").split(",")
        if item.strip()
    }


def is_workspace_enabled(workspace: Workspace) -> bool:
    from app.services.workspace_features import feature_enabled

    return feature_enabled(
        workspace,
        "secure_deposit",
        csv_fallback=settings.secure_deposit_enabled_workspace_slugs,
    )


def default_allowed_extensions() -> list[str]:
    return [
        item.strip().lower().lstrip(".")
        for item in (settings.secure_deposit_allowed_extensions or "").split(",")
        if item.strip()
    ]


def generate_access_id() -> str:
    # The ordinary generator never enters the reserved canary namespace.
    while True:
        access_id = secrets.token_urlsafe(14).replace("_", "-")
        if not access_id.startswith(SFTP_RELEASE_A_CANARY_ACCESS_ID_PREFIX):
            return access_id


def generate_release_a_canary_access_id() -> str:
    return SFTP_RELEASE_A_CANARY_ACCESS_ID_PREFIX + secrets.token_urlsafe(14).replace(
        "_", "-"
    )


def is_release_a_sftp_canary_access_id(access_id: object) -> bool:
    return (
        isinstance(access_id, str)
        and access_id.startswith(SFTP_RELEASE_A_CANARY_ACCESS_ID_PREFIX)
        and len(access_id) > len(SFTP_RELEASE_A_CANARY_ACCESS_ID_PREFIX)
    )


def generate_deposit_password() -> str:
    return secrets.token_urlsafe(18)


def hash_password(password: str) -> str:
    if _ARGON2 is not None:
        return _ARGON2.hash(password)
    salt = os.urandom(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 260_000)
    return f"{_PBKDF2_PREFIX}${base64.b64encode(salt).decode()}${base64.b64encode(digest).decode()}"


def verify_password(password_hash: str, password: str) -> bool:
    if password_hash.startswith(f"{_PBKDF2_PREFIX}$"):
        try:
            _, salt_raw, digest_raw = password_hash.split("$", 2)
            salt = base64.b64decode(salt_raw.encode())
            expected = base64.b64decode(digest_raw.encode())
            actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 260_000)
            return hmac.compare_digest(actual, expected)
        except Exception:  # noqa: BLE001
            return False
    if _ARGON2 is None:
        return False
    try:
        return bool(_ARGON2.verify(password_hash, password))
    except VerifyMismatchError:
        return False
    except Exception:  # noqa: BLE001
        return False


def safe_filename(filename: str | None) -> str:
    name = PurePosixPath(str(filename or "upload").replace("\\", "/")).name
    name = re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip(" .")
    return (name or "upload")[:180]


def safe_relative_path(path: str | None) -> str:
    """Sanitize a user-provided relative path while preserving folders."""

    raw = str(path or "upload").replace("\\", "/").strip()
    clean_parts: list[str] = []
    for part in PurePosixPath(raw).parts:
        if part in {"", "/", ".", ".."}:
            continue
        safe_part = re.sub(r"[^A-Za-z0-9._ -]+", "_", part).strip(" .")
        if safe_part:
            clean_parts.append(safe_part[:96])
    if not clean_parts:
        clean_parts = [safe_filename(path)]
    relative = "/".join(clean_parts)
    if len(relative) <= 240:
        return relative
    tail = safe_filename(clean_parts[-1])
    budget = max(1, 240 - len(tail) - 1)
    prefix = "/".join(clean_parts[:-1])[:budget].rstrip("/")
    return f"{prefix}/{tail}" if prefix else tail[:240]


def _storage_root() -> Path:
    root = Path(settings.secure_deposit_storage_dir).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _storage_key(*parts: object) -> str:
    path = PurePosixPath("/".join(str(part).strip("/") for part in parts if str(part).strip("/")))
    clean_parts = [part for part in path.parts if part not in ("", "/", ".")]
    if any(part == ".." for part in clean_parts):
        raise ValueError("Invalid secure deposit storage key")
    return "/".join(clean_parts)


def _storage_path(key: str) -> Path:
    root = _storage_root()
    path = (root / _storage_key(key)).resolve()
    if root not in path.parents and path != root:
        raise ValueError("Secure deposit storage key escapes root")
    return path


def _write_staged_bytes(key: str, content: bytes) -> None:
    path = _storage_path(key)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)


async def _write_staged_upload(key: str, upload: UploadFile, max_bytes: int) -> tuple[int, str]:
    path = _storage_path(key)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=".upload-", suffix=".part", dir=str(path.parent))
    tmp_path = Path(tmp_name)
    size = 0
    digest = hashlib.sha256()
    try:
        with os.fdopen(fd, "wb") as handle:
            while True:
                chunk = await upload.read(_UPLOAD_CHUNK_BYTES)
                if not chunk:
                    break
                size += len(chunk)
                if size > max_bytes:
                    raise HTTPException(status_code=413, detail="File is too large")
                digest.update(chunk)
                handle.write(chunk)
        tmp_path.replace(path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
    return size, digest.hexdigest()


def _copy_staged_to_local(key: str, destination: Path) -> Path:
    source = _storage_path(key)
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


def staged_file_path(file: DepositFile) -> Path:
    path = _storage_path(file.object_key)
    if not path.exists() or not path.is_file():
        raise HTTPException(status_code=409, detail=f"Staged file is missing: {file.filename}")
    return path


def staged_file_media_type(file: DepositFile) -> str:
    guessed, _ = mimetypes.guess_type(file.filename or "")
    if guessed:
        return guessed
    return file.content_type or "application/octet-stream"


def staged_file_download_name(file: DepositFile) -> str:
    return safe_filename(PurePosixPath(file.filename or "deposit-file").name)


def _hash_local_file(path: Path) -> tuple[int, str]:
    size = 0
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(_UPLOAD_CHUNK_BYTES)
            if not chunk:
                break
            size += len(chunk)
            digest.update(chunk)
    return size, digest.hexdigest()


def _cell_preview(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value)[:240]


def _spreadsheet_preview(
    path: Path, *, sheet_name: str | None = None, cell_range: str | None = None,
) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    bounds = None
    if cell_range:
        if not sheet_name or not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}(?::[A-Za-z]{1,3}[1-9][0-9]{0,6})?", cell_range):
            raise HTTPException(status_code=422, detail="A cell range requires a sheet and valid cell coordinates")
        bounds = range_boundaries(cell_range.upper())
        c1, r1, c2, r2 = bounds
        if not (1 <= c1 <= c2 <= 16384 and 1 <= r1 <= r2 <= 1048576):
            raise HTTPException(status_code=422, detail="Cell range is outside spreadsheet limits")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise HTTPException(status_code=404, detail="Referenced worksheet not found")
        sheet = workbook[sheet_name if sheet_name is not None else workbook.sheetnames[0]]
        max_row, max_col = sheet.max_row or 1, sheet.max_column or 1
        if bounds and (bounds[2] > max_col or bounds[3] > max_row):
            raise HTTPException(status_code=404, detail="Referenced cells are outside the source worksheet")
        row_start, col_start, column_limit = 1, 1, 12
        if bounds:
            width, height = bounds[2] - bounds[0] + 1, bounds[3] - bounds[1] + 1
            # Prefer the cited cells over context; keep large selections bounded.
            column_limit = min(40, max(12, width + 2))
            row_start = max(1, bounds[1] - min(2, max(0, 40 - height)))
            col_start = max(1, bounds[0] - min(2, max(0, column_limit - width)))
        row_end = min(max_row, row_start + 39)
        col_end = min(max_col, col_start + column_limit - 1)
        rows = [
            [_cell_preview(value) for value in row]
            for row in sheet.iter_rows(min_row=row_start, max_row=row_end,
                                       min_col=col_start, max_col=col_end, values_only=True)
        ]
        return {
            "kind": "spreadsheet",
            "sheet_name": sheet.title,
            "rows": rows,
            "row_start": row_start,
            "column_start": col_start,
            "columns": [get_column_letter(col) for col in range(col_start, col_end + 1)],
            "cell_range": cell_range.upper() if cell_range else None,
            "selection": ({"row_start": bounds[1], "row_end": bounds[3],
                           "column_start": bounds[0], "column_end": bounds[2]} if bounds else None),
            "selection_truncated": bool(bounds and (bounds[3] > row_end or bounds[2] > col_end)),
            "truncated": row_start > 1 or col_start > 1 or row_end < max_row or col_end < max_col,
        }
    finally:
        workbook.close()


def _docx_preview(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    blocks: list[str] = []
    truncated = False

    def append_block(text: str) -> None:
        nonlocal truncated
        clean = text.strip()
        if not clean or truncated:
            return
        current = sum(len(block) for block in blocks) + max(0, len(blocks) - 1) * 2
        remaining = _DOCX_PREVIEW_MAX_CHARS - current
        if remaining <= 0:
            truncated = True
            return
        if len(clean) > remaining:
            blocks.append(clean[:remaining].rstrip())
            truncated = True
        else:
            blocks.append(clean)

    for paragraph in document.paragraphs:
        append_block(paragraph.text)

    for table in document.tables:
        rows: list[str] = []
        for row in table.rows[:40]:
            cells = [_cell_preview(cell.text.replace("\n", " ")) for cell in row.cells[:12]]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            append_block("\n".join(rows))
        if truncated:
            break

    content = "\n\n".join(blocks).strip()
    return {
        "kind": "text",
        "source_kind": "docx",
        "content": content or "No textual content found in this DOCX.",
        "truncated": truncated,
    }


def build_file_preview(
    path: Path,
    *,
    filename: str,
    media_type: str,
    size_bytes: int,
    download_url: str,
    sheet_name: str | None = None,
    cell_range: str | None = None,
) -> dict[str, Any]:
    """Render an inline preview payload for a local file.

    Shared by the Secure Deposit staging queue and the Knowledge collection
    Sources browser so both surfaces produce the same ``kind`` contract
    (``text`` | ``html`` | ``spreadsheet`` | ``image`` | ``pdf`` | ``binary``).
    """
    ext = extension_for(filename)
    size = int(size_bytes or 0)
    base: dict[str, Any] = {
        "filename": filename,
        "content_type": media_type,
        "size_bytes": size,
        "download_url": download_url,
    }

    if ext in _SPREADSHEET_EXTENSIONS and size <= _STRUCTURED_PREVIEW_MAX_BYTES:
        return {**base, **_spreadsheet_preview(path, sheet_name=sheet_name, cell_range=cell_range)}

    if ext in _DOCX_EXTENSIONS and size <= _DOCX_PREVIEW_MAX_BYTES:
        return {**base, **_docx_preview(path)}

    if ext in _HTML_EXTENSIONS or media_type == "text/html":
        if size > _TEXT_PREVIEW_BYTES:
            return {**base, "kind": "binary", "reason": "html_preview_too_large"}
        data = path.read_bytes()
        truncated = len(data) > _TEXT_PREVIEW_BYTES
        data = data[:_TEXT_PREVIEW_BYTES]
        content = _decode_preview_text(data)
        return {**base, "kind": "html", "content": content, "truncated": truncated}

    if ext in _TEXT_EXTENSIONS or media_type.startswith("text/"):
        if size > _TEXT_PREVIEW_BYTES:
            return {**base, "kind": "binary", "reason": "text_preview_too_large"}
        data = path.read_bytes()
        truncated = len(data) > _TEXT_PREVIEW_BYTES
        data = data[:_TEXT_PREVIEW_BYTES]
        content = _decode_preview_text(data)
        return {**base, "kind": "text", "content": content, "truncated": truncated}

    if media_type == "application/pdf" and size <= _PDF_PREVIEW_MAX_BYTES:
        return {**base, "kind": "pdf"}

    if media_type.startswith("image/") and size <= _INLINE_PREVIEW_MAX_BYTES:
        return {**base, "kind": "image"}

    return {**base, "kind": "binary"}


def _decode_preview_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def preview_needs_file_bytes(filename: str, media_type: str, size_bytes: int) -> bool:
    """Whether :func:`build_file_preview` reads the file content for this input.

    PDF / image / binary previews only need the size and media type, so callers
    backed by remote object storage (MinIO/S3) can avoid downloading the original
    bytes just to build the preview envelope. Kept in lock-step with the reading
    branches of :func:`build_file_preview`.
    """
    ext = extension_for(filename)
    size = int(size_bytes or 0)
    if ext in _SPREADSHEET_EXTENSIONS and size <= _STRUCTURED_PREVIEW_MAX_BYTES:
        return True
    if ext in _DOCX_EXTENSIONS and size <= _DOCX_PREVIEW_MAX_BYTES:
        return True
    if (ext in _TEXT_EXTENSIONS or media_type.startswith("text/")) and size <= _TEXT_PREVIEW_BYTES:
        return True
    return False


def preview_deposit_file(file: DepositFile) -> dict[str, Any]:
    return build_file_preview(
        staged_file_path(file),
        filename=file.filename,
        media_type=staged_file_media_type(file),
        size_bytes=int(file.size_bytes or 0),
        download_url=f"/api/v1/sftp/deposits/{file.id}/download",
    )


def _normalize_archive_browser_path(path: str | None) -> str:
    raw = str(path or "").replace("\\", "/").strip("/")
    if not raw:
        return ""
    pure = PurePosixPath(raw)
    if pure.is_absolute() or any(part in {"", ".", ".."} for part in pure.parts):
        raise HTTPException(status_code=422, detail="Unsafe ZIP path")
    if any(":" in part for part in pure.parts):
        raise HTTPException(status_code=422, detail="Unsafe ZIP path")
    return pure.as_posix()


def _assert_zip_deposit(file: DepositFile) -> Path:
    if extension_for(file.filename or "") != "zip":
        raise HTTPException(status_code=415, detail="Archive browsing is only available for ZIP files")
    return staged_file_path(file)


def _open_zip_archive(path: Path) -> zipfile.ZipFile:
    try:
        return zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=422, detail="Archive is not a valid ZIP file") from exc


def _archive_member_previewable(filename: str, media_type: str, size_bytes: int) -> bool:
    ext = extension_for(filename)
    if ext in _SPREADSHEET_EXTENSIONS:
        return size_bytes <= _STRUCTURED_PREVIEW_MAX_BYTES
    if ext in _DOCX_EXTENSIONS:
        return size_bytes <= _DOCX_PREVIEW_MAX_BYTES
    if ext in _TEXT_EXTENSIONS or media_type.startswith("text/"):
        return size_bytes <= _TEXT_PREVIEW_BYTES
    return size_bytes <= _INLINE_PREVIEW_MAX_BYTES and (media_type.startswith("image/") or media_type == "application/pdf")


def _archive_entry_payload(path: str, *, kind: str, size_bytes: int = 0, compressed_size: int = 0) -> dict[str, Any]:
    name = PurePosixPath(path).name or path
    extension = extension_for(name) if kind == "file" else ""
    media_type = mimetypes.guess_type(name)[0] or "application/octet-stream"
    return {
        "kind": kind,
        "name": name,
        "path": path,
        "extension": extension,
        "content_type": media_type if kind == "file" else None,
        "size_bytes": int(size_bytes or 0),
        "compressed_size_bytes": int(compressed_size or 0),
        "previewable": kind == "file" and _archive_member_previewable(name, media_type, int(size_bytes or 0)),
    }


def list_deposit_zip_archive(file: DepositFile, *, path: str | None = None, max_entries: int = _ARCHIVE_BROWSER_MAX_ENTRIES) -> dict[str, Any]:
    archive_path = _assert_zip_deposit(file)
    current_path = _normalize_archive_browser_path(path)
    prefix = f"{current_path}/" if current_path else ""
    folders: dict[str, dict[str, Any]] = {}
    files: list[dict[str, Any]] = []
    total_files = 0
    total_size = 0
    truncated = False

    with _open_zip_archive(archive_path) as archive:
        for info in archive.infolist():
            try:
                member_path = _validated_archive_member_path(info)
            except HTTPException:
                raise
            if member_path is None:
                continue
            member = member_path.as_posix()
            total_files += 1
            total_size += int(info.file_size or 0)
            if current_path and not member.startswith(prefix):
                continue
            remainder = member[len(prefix) :] if prefix else member
            if not remainder:
                continue
            head, *rest = remainder.split("/")
            if rest:
                folder_path = f"{prefix}{head}" if prefix else head
                folder = folders.get(folder_path)
                if folder is None:
                    folder = _archive_entry_payload(folder_path, kind="folder")
                    folder["count"] = 0
                    folder["size_bytes"] = 0
                    folders[folder_path] = folder
                folder["count"] += 1
                folder["size_bytes"] += int(info.file_size or 0)
                continue
            if len(files) + len(folders) >= max_entries:
                truncated = True
                continue
            if info.flag_bits & 0x1:
                files.append({**_archive_entry_payload(member, kind="file", size_bytes=info.file_size, compressed_size=info.compress_size), "encrypted": True, "previewable": False})
            else:
                files.append(_archive_entry_payload(member, kind="file", size_bytes=info.file_size, compressed_size=info.compress_size))

    if total_files <= 0:
        raise HTTPException(status_code=422, detail="Archive contains no files")
    folder_items = list(folders.values())
    if len(folder_items) + len(files) > max_entries:
        truncated = True
        folder_items = folder_items[:max_entries]
        files = files[: max(0, max_entries - len(folder_items))]
    items = sorted(folder_items, key=lambda item: item["name"].lower()) + sorted(files, key=lambda item: item["name"].lower())
    return {
        "file_id": file.id,
        "filename": file.filename,
        "path": current_path,
        "items": items,
        "truncated": truncated,
        "max_entries": max_entries,
        "total_files": total_files,
        "total_size_bytes": total_size,
    }


def _find_zip_member(archive: zipfile.ZipFile, member_path: str) -> zipfile.ZipInfo:
    wanted = _normalize_archive_browser_path(member_path)
    if not wanted:
        raise HTTPException(status_code=422, detail="ZIP member path is required")
    for info in archive.infolist():
        candidate = _validated_archive_member_path(info)
        if candidate is not None and candidate.as_posix() == wanted:
            if info.flag_bits & 0x1:
                raise HTTPException(status_code=422, detail="Encrypted ZIP member is not supported")
            return info
    raise HTTPException(status_code=404, detail="ZIP member not found")


def _extract_zip_info_to_temp(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> tuple[Path, str]:
    safe_name = safe_filename(PurePosixPath(info.filename).name)
    suffix = PurePosixPath(safe_name).suffix
    fd, tmp_name = tempfile.mkstemp(prefix=".zip-member-", suffix=suffix)
    tmp_path = Path(tmp_name)
    try:
        with os.fdopen(fd, "wb") as handle, archive.open(info) as source:
            shutil.copyfileobj(source, handle, length=_UPLOAD_CHUNK_BYTES)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
    return tmp_path, safe_name


def extract_deposit_zip_member_to_temp(file: DepositFile, *, member_path: str) -> tuple[Path, zipfile.ZipInfo, str]:
    archive_path = _assert_zip_deposit(file)
    with _open_zip_archive(archive_path) as archive:
        info = _find_zip_member(archive, member_path)
        tmp_path, safe_name = _extract_zip_info_to_temp(archive, info)
        return tmp_path, info, safe_name


def _zip_member_download_url(file: DepositFile, member_path: str) -> str:
    return f"/api/v1/sftp/deposits/{file.id}/archive/member/download?path={quote(member_path, safe='')}"


def preview_deposit_zip_member(file: DepositFile, *, member_path: str) -> dict[str, Any]:
    normalized = _normalize_archive_browser_path(member_path)
    archive_path = _assert_zip_deposit(file)
    with _open_zip_archive(archive_path) as archive:
        info = _find_zip_member(archive, normalized)
        safe_name = safe_filename(PurePosixPath(info.filename).name)
        media_type = mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
        size = int(info.file_size or 0)
        download_url = _zip_member_download_url(file, normalized)
        if not preview_needs_file_bytes(safe_name, media_type, size):
            return build_file_preview(Path(safe_name), filename=safe_name, media_type=media_type, size_bytes=size, download_url=download_url)
        temp_path, _safe_name = _extract_zip_info_to_temp(archive, info)
        try:
            return build_file_preview(temp_path, filename=safe_name, media_type=media_type, size_bytes=size, download_url=download_url)
        finally:
            temp_path.unlink(missing_ok=True)


def _archive_component(value: str | None, fallback: str) -> str:
    component = safe_filename(value or fallback).replace(" ", "_")
    return component[:96] or fallback


def _unique_archive_name(path: str, used: set[str]) -> str:
    if path not in used:
        used.add(path)
        return path
    stem, dot, suffix = path.rpartition(".")
    base = stem if dot else path
    ext = f".{suffix}" if dot else ""
    index = 2
    while f"{base}-{index}{ext}" in used:
        index += 1
    unique = f"{base}-{index}{ext}"
    used.add(unique)
    return unique


_ARCHIVE_DOCUMENT_NAME_MAX_CHARS = 220


def _bounded_archive_document_name(name: str) -> str:
    if len(name) <= _ARCHIVE_DOCUMENT_NAME_MAX_CHARS:
        return name

    suffix = PurePosixPath(name).suffix
    if len(suffix) > 16:
        suffix = ""
    stem = name[: -len(suffix)] if suffix else name
    digest = hashlib.sha1(name.encode("utf-8")).hexdigest()[:12]
    marker = f"__{digest}__"
    budget = max(8, _ARCHIVE_DOCUMENT_NAME_MAX_CHARS - len(suffix) - len(marker))
    head_budget = max(4, budget * 2 // 3)
    tail_budget = max(4, budget - head_budget)
    head = stem[:head_budget].rstrip(" ._-")
    tail = stem[-tail_budget:].lstrip(" ._-")
    bounded = f"{head}{marker}{tail}{suffix}"
    return bounded[:_ARCHIVE_DOCUMENT_NAME_MAX_CHARS]


def _validated_archive_member_path(info: zipfile.ZipInfo) -> PurePosixPath | None:
    raw_name = str(info.filename or "").replace("\\", "/").strip()
    if not raw_name or info.is_dir():
        return None
    path = PurePosixPath(raw_name)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise HTTPException(status_code=422, detail=f"Unsafe ZIP member path: {info.filename}")
    if any(":" in part for part in path.parts):
        raise HTTPException(status_code=422, detail=f"Unsafe ZIP member path: {info.filename}")
    return path


def archive_document_namespace(deposit_filename: str | None) -> str:
    """Stable per-archive prefix used to namespace flattened document names.

    Two archives that share an internal member path (e.g. ``OPERATING_MANUAL/
    page1.pdf``) would otherwise flatten to the *same* document name and clobber
    each other in a shared collection. Prefixing with a namespace derived from
    the deposit's source path (its immediate folder + archive stem, e.g.
    ``Notices_Techniques_SPL/A/Manual_ASY100.zip`` -> ``A__Manual_ASY100``)
    keeps cross-archive / cross-prefix documents distinct.

    A short digest of the *full* deposit path is appended when the readable
    prefix would be too long, so the namespace stays globally unique without
    growing object keys unbounded.
    """
    raw = str(deposit_filename or "").replace("\\", "/").strip("/")
    if not raw:
        return ""
    pure = PurePosixPath(raw)
    stem = pure.name[: -len(pure.suffix)] if pure.suffix else pure.name
    parent = pure.parent.name  # immediate folder letter (A/B/C/D/...)
    namespace = "__".join(part for part in (parent, stem) if part) or stem
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", namespace).strip("_")
    if len(safe) > 80:
        digest = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:10]
        safe = f"{safe[:60]}_{digest}"
    return safe or hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12]


def _archive_document_name(
    member_path: PurePosixPath,
    used: set[str],
    *,
    namespace: str | None = None,
) -> str:
    safe_path = safe_relative_path(member_path.as_posix())
    flattened = safe_path.replace("/", "__")
    if namespace:
        flattened = f"{namespace}__{flattened}"
    return _unique_archive_name(_bounded_archive_document_name(flattened), used)


def _extract_andritz_project_reference(
    *values: str | None,
    scheme: str | None = None,
) -> dict[str, str]:
    """Compatibility wrapper around the canonical source-aware resolver."""

    return derive_project_reference(*values, scheme=scheme)


def _extract_machine_reference(
    *values: str | None,
    exclude: str | None = None,
    scheme: str | None = None,
) -> dict[str, str]:
    """First series reference (e.g. BBA120) distinct from the project code.

    Conservative on purpose: archive paths usually carry the machine/series as
    a second alphanum reference next to the project code; anything fuzzier
    belongs in a knowledge guide, not in payload metadata.  The series grammar
    is Andritz-only; other workspaces fail closed.
    """
    if not andritz_project_scheme_active(scheme):
        return {}
    for value in values:
        for match in LEGACY_PROJECT_REFERENCE_RE.finditer(str(value or "")):
            reference = f"{match.group(1).upper()}{match.group(2)}{(match.group(3) or '').upper()}"
            if exclude and reference == exclude:
                continue
            return {"machine": reference}
    return {}


def _classify_archive_source_family(path: str, extension: str | None = None) -> str:
    haystack = path.replace("_", " ").replace("-", " ").lower()
    ext = (extension or PurePosixPath(path).suffix.lstrip(".")).lower()
    if any(token in haystack for token in ("spare part", "parts list", "spareparts")):
        return "spare_parts_list"
    if any(token in haystack for token in ("commissioning", "check list", "checklist")):
        return "commissioning"
    if any(token in haystack for token in ("safety", "declaration of conformity", "certif", "conformity")):
        return "safety"
    if any(token in haystack for token in ("annex", "annexe", "annexes")):
        return "annex"
    if any(token in haystack for token in ("maintenance", "service manual", "manuel de service", "cleaning", "nettoyage")):
        return "maintenance"
    if any(token in haystack for token in ("operating manual", "operator manual", "user manual", "users manual", "user's manual", "manual")):
        return "operating_manual"
    if ext in {"html", "htm"}:
        return "html_manual"
    return "unknown"


def _archive_document_metadata(
    *,
    deposit_filename: str | None,
    archive_path: str | None,
    document_name: str,
    extension: str | None,
    source_deposit_file_id: str | None = None,
    scheme: str | None = None,
) -> dict[str, Any]:
    archive_name = PurePosixPath(str(deposit_filename or "")).name if deposit_filename else None
    metadata: dict[str, Any] = {
        "source_origin": "secure_deposit",
        "source_family": _classify_archive_source_family(" ".join([archive_path or "", document_name]), extension),
        "archive_name": archive_name,
        "source_deposit_path": deposit_filename,
        "source_deposit_file_id": source_deposit_file_id,
        "inner_document_path": archive_path,
    }
    metadata.update(
        derive_project_reference(
            deposit_filename,
            archive_path,
            document_name,
            scheme=scheme,
        )
    )
    if archive_path:
        metadata.update(
            _extract_machine_reference(
                archive_path,
                document_name,
                exclude=str(metadata.get("project_code") or "") or None,
                scheme=scheme,
            )
        )
    return {key: value for key, value in metadata.items() if value is not None}


def _read_supported_archive_documents(
    path: Path,
    *,
    deposit_filename: str | None = None,
    source_deposit_file_id: str | None = None,
    max_files: int | None = None,
    max_uncompressed_bytes: int | None = None,
    on_limit: str = "error",
    document_namespace: str | None = None,
    used_names: set[str] | None = None,
    include_content: bool = True,
    document_callback: Callable[[dict[str, Any]], None] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Extract supported documents from a staged ZIP.

    ``document_namespace`` (opt-in) prefixes each flattened document name with a
    per-archive namespace so identical internal paths in different archives do
    not collide on a shared collection. ``used_names`` lets a caller share the
    de-dup set *across* archives within a single wave (rather than only within
    one ZIP); when omitted a fresh per-call set preserves the legacy behaviour.
    """
    limit_files = max(1, int(max_files or settings.secure_deposit_archive_promotion_max_files or 50))
    limit_bytes = int(max_uncompressed_bytes) if max_uncompressed_bytes else None
    documents: list[dict[str, Any]] = []
    used_names = used_names if used_names is not None else set()
    truncated_files = 0
    skipped_uncompressed_bytes = 0
    skipped_read_error_count = 0
    skipped_read_errors: list[dict[str, str]] = []
    accepted_files = 0
    accepted_uncompressed_bytes = 0

    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=422, detail="Archive is not a valid ZIP file") from exc

    with archive:
        for info in archive.infolist():
            member_path = _validated_archive_member_path(info)
            if member_path is None:
                continue
            ext = extension_for(member_path.name)
            if ext not in _ARCHIVE_PROMOTION_EXTENSIONS:
                continue
            if info.flag_bits & 0x1:
                raise HTTPException(status_code=422, detail=f"Encrypted ZIP member is not supported: {info.filename}")
            file_size = int(info.file_size or 0)
            if accepted_files >= limit_files:
                if on_limit == "truncate":
                    truncated_files += 1
                    continue
                raise HTTPException(
                    status_code=413,
                    detail=f"Archive contains more than {limit_files} supported documents",
                )
            if limit_bytes is not None and accepted_uncompressed_bytes + file_size > limit_bytes:
                if on_limit == "truncate":
                    truncated_files += 1
                    skipped_uncompressed_bytes += file_size
                    continue
                raise HTTPException(
                    status_code=413,
                    detail=f"Archive uncompressed payload exceeds {limit_bytes} bytes",
                )
            document_name = _archive_document_name(member_path, used_names, namespace=document_namespace)
            archive_path = member_path.as_posix()
            metadata = _archive_document_metadata(
                deposit_filename=deposit_filename or path.name,
                archive_path=archive_path,
                document_name=document_name,
                extension=ext,
                source_deposit_file_id=source_deposit_file_id,
            )
            if document_namespace:
                # Record the namespace + the legacy (un-namespaced) flattened
                # name so collisions can be detected and legacy lookups mapped.
                metadata["document_namespace"] = document_namespace
                metadata["legacy_document_name"] = safe_relative_path(member_path.as_posix()).replace("/", "__")
            document = {
                "archive_path": archive_path,
                "filename": document_name,
                "extension": ext,
                "size_bytes": file_size,
                "metadata": metadata,
            }
            if include_content:
                try:
                    document["content"] = archive.read(info)
                except (zipfile.BadZipFile, RuntimeError, OSError, EOFError) as exc:
                    if on_limit == "truncate":
                        truncated_files += 1
                        skipped_read_error_count += 1
                        if len(skipped_read_errors) < 50:
                            skipped_read_errors.append({"archive_path": archive_path, "reason": str(exc)[:200]})
                        continue
                    raise HTTPException(
                        status_code=422,
                        detail=f"ZIP member could not be read: {info.filename}",
                    ) from exc
            if document_callback is not None:
                document_callback(document)
            else:
                documents.append(document)
            accepted_files += 1
            accepted_uncompressed_bytes += file_size

    if accepted_files <= 0:
        raise HTTPException(status_code=422, detail="Archive contains no supported documents")
    stats = {
        "document_count": accepted_files,
        "uncompressed_bytes": accepted_uncompressed_bytes,
        "truncated_files": truncated_files,
        "skipped_uncompressed_bytes": skipped_uncompressed_bytes,
        "skipped_read_error_count": skipped_read_error_count,
        "skipped_read_errors": skipped_read_errors,
        "max_files": limit_files,
        "max_uncompressed_bytes": limit_bytes,
    }
    return documents, stats


def build_deposit_archive(
    files: list[DepositFile],
    *,
    links_by_id: dict[str, DepositAccessLink],
    workspace_slug: str,
) -> tuple[Path, str]:
    if not files:
        raise HTTPException(status_code=404, detail="No staged files to archive")

    stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    archive_filename = f"{_archive_component(workspace_slug, 'workspace')}-secure-deposit-{stamp}.zip"
    archive_dir = _storage_root() / "_archives"
    archive_dir.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=".secure-deposit-", suffix=".zip.part", dir=str(archive_dir))
    os.close(fd)
    archive_path = Path(tmp_name)

    manifest: list[dict[str, Any]] = []
    used_names: set[str] = set()
    try:
        with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_STORED, allowZip64=True) as archive:
            for file in files:
                source = _storage_path(file.object_key)
                if not source.exists():
                    raise HTTPException(status_code=409, detail=f"Staged file is missing: {file.filename}")
                link = links_by_id.get(file.access_link_id)
                link_label = _archive_component(link.label if link else None, "deposit-link")
                link_access = _archive_component(link.access_id if link else file.access_link_id, file.access_link_id)
                filename = safe_relative_path(file.filename)
                archive_name = _unique_archive_name(
                    f"{link_label}-{link_access}/{file.status}/{filename}",
                    used_names,
                )
                archive.write(source, archive_name)
                manifest.append(
                    {
                        "archive_path": archive_name,
                        "file_id": file.id,
                        "filename": file.filename,
                        "status": file.status,
                        "size_bytes": file.size_bytes,
                        "sha256": file.sha256,
                        "uploaded_at": file.uploaded_at.isoformat() if file.uploaded_at else None,
                        "access_link_id": file.access_link_id,
                        "access_id": link.access_id if link else None,
                        "deposit_link_label": link.label if link else None,
                    }
                )
            archive.writestr(
                "_manifest.json",
                json.dumps(
                    {
                        "workspace": workspace_slug,
                        "generated_at": datetime.utcnow().isoformat() + "Z",
                        "file_count": len(manifest),
                        "files": manifest,
                    },
                    indent=2,
                    sort_keys=True,
                ),
            )
    except Exception:
        archive_path.unlink(missing_ok=True)
        raise

    final_path = archive_dir / archive_filename
    if final_path.exists():
        final_path = archive_dir / f"{archive_filename.removesuffix('.zip')}-{secrets.token_hex(4)}.zip"
        archive_filename = final_path.name
    archive_path.replace(final_path)
    return final_path, archive_filename


def extension_for(filename: str) -> str:
    if "." not in filename:
        return ""
    return filename.rsplit(".", 1)[1].lower()


def public_url(access_id: str) -> str:
    base = (settings.secure_deposit_public_base_url or "").rstrip("/")
    return f"{base}/deposit/{access_id}" if base else f"/deposit/{access_id}"


def _session_secret() -> str:
    return (
        settings.secure_deposit_session_secret
        or settings.keycloak_client_secret
        or "agentium-dev-secure-deposit-secret"
    )


def issue_session_token(link: DepositAccessLink) -> tuple[str, datetime]:
    expires_at = datetime.utcnow() + timedelta(seconds=settings.secure_deposit_session_ttl_seconds)
    token = jwt.encode(
        {
            "typ": "deposit_session",
            "sub": link.access_id,
            "link_id": link.id,
            "workspace_id": link.workspace_id,
            "exp": expires_at,
            "iat": datetime.utcnow(),
            "jti": secrets.token_urlsafe(10),
        },
        _session_secret(),
        algorithm=_TOKEN_ALGORITHM,
    )
    return token, expires_at


def verify_session_token(token: str, access_id: str) -> dict[str, Any]:
    try:
        payload = jwt.decode(token, _session_secret(), algorithms=[_TOKEN_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Deposit session expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid deposit session")
    if payload.get("typ") != "deposit_session" or payload.get("sub") != access_id:
        raise HTTPException(status_code=401, detail="Invalid deposit session")
    return payload


def serialize_link(link: DepositAccessLink, *, reveal_password: str | None = None) -> dict[str, Any]:
    return {
        "id": link.id,
        "workspace_id": link.workspace_id,
        "created_by_user_id": link.created_by_user_id,
        "label": link.label,
        "access_id": link.access_id,
        "public_url": public_url(link.access_id),
        "status": link.status,
        "expires_at": link.expires_at.isoformat() if link.expires_at else None,
        "max_file_size_mb": link.max_file_size_mb,
        "allowed_extensions": link.allowed_extensions or [],
        "sftp_auth_audit_profile": (
            SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY
            if is_release_a_sftp_canary_access_id(link.access_id)
            else None
        ),
        "created_at": link.created_at.isoformat() if link.created_at else None,
        "updated_at": link.updated_at.isoformat() if link.updated_at else None,
        "generated_password": reveal_password,
    }


def serialize_public_link(link: DepositAccessLink) -> dict[str, Any]:
    return {
        "label": link.label,
        "access_id": link.access_id,
        "public_url": public_url(link.access_id),
        "status": link.status,
        "expires_at": link.expires_at.isoformat() if link.expires_at else None,
        "max_file_size_mb": link.max_file_size_mb,
        "allowed_extensions": link.allowed_extensions or [],
    }


def serialize_file(file: DepositFile) -> dict[str, Any]:
    return {
        "id": file.id,
        "access_link_id": file.access_link_id,
        "filename": file.filename,
        "content_type": file.content_type,
        "size_bytes": file.size_bytes,
        "sha256": file.sha256,
        "status": file.status,
        "uploaded_at": file.uploaded_at.isoformat() if file.uploaded_at else None,
        "promoted_at": file.promoted_at.isoformat() if file.promoted_at else None,
        "promoted_collection_slug": file.promoted_collection_slug,
        "worker_job_id": file.worker_job_id,
        "promotion_result": file.promotion_result,
        "rejection_reason": file.rejection_reason,
    }


def serialize_public_file(file: DepositFile) -> dict[str, Any]:
    return {
        "id": file.id,
        "filename": file.filename,
        "content_type": file.content_type,
        "size_bytes": file.size_bytes,
        "sha256": file.sha256,
        "status": file.status,
        "uploaded_at": file.uploaded_at.isoformat() if file.uploaded_at else None,
        "promoted_at": file.promoted_at.isoformat() if file.promoted_at else None,
    }


def _target_collection_slug(workspace: Workspace, collection_slug: str | None) -> str:
    default_collection_slug = f"{workspace.slug}-secure-deposit"
    return (collection_slug or default_collection_slug).strip() or default_collection_slug


def _find_collection(db: DBSession, *, workspace: Workspace, collection_ref: str) -> KnowledgeCollection | None:
    ref = collection_ref.strip()
    if not ref:
        return None
    return (
        db.query(KnowledgeCollection)
        .filter(
            KnowledgeCollection.workspace_id == workspace.id,
            (
                (KnowledgeCollection.id == ref)
                | (KnowledgeCollection.slug == ref)
                | (KnowledgeCollection.name == ref)
            ),
        )
        .first()
    )


def _collection_indexing_snapshot(
    db: DBSession,
    *,
    workspace: Workspace,
    collection_slug: str,
) -> dict[str, Any]:
    collection = _find_collection(db, workspace=workspace, collection_ref=collection_slug)
    if not collection:
        return {
            "slug": collection_slug,
            "exists": False,
            "status": "missing",
            "source_count": 0,
            "document_count": 0,
            "chunk_count": 0,
            "zero_chunk_sources": 0,
            "error_sources": 0,
            "job_counts": {},
            "latest_job": None,
            "jobs": [],
        }

    inventory = collection_inventory(db, collection=collection, include_sources=False)
    jobs = (
        db.query(WorkerJob)
        .filter(WorkerJob.workspace_id == workspace.id, WorkerJob.collection_id == collection.id)
        .order_by(WorkerJob.updated_at.desc(), WorkerJob.created_at.desc())
        .limit(8)
        .all()
    )
    job_counts: dict[str, int] = {}
    for job in jobs:
        job_counts[job.status] = job_counts.get(job.status, 0) + 1
    serialized_jobs = [serialize_worker_job(job) for job in jobs]
    return {
        "id": collection.id,
        "slug": collection.slug or collection_slug,
        "name": collection.name,
        "exists": True,
        "status": collection.status,
        "source_count": inventory.get("source_count", 0),
        "document_count": inventory.get("document_count", 0),
        "chunk_count": inventory.get("chunk_count", 0),
        "zero_chunk_sources": inventory.get("zero_chunk_sources", 0),
        "error_sources": inventory.get("error_sources", 0),
        "job_counts": job_counts,
        "latest_job": serialized_jobs[0] if serialized_jobs else None,
        "jobs": serialized_jobs,
    }


def _archive_assist_summary(path: Path, *, deposit_filename: str | None) -> tuple[dict[str, Any] | None, str | None]:
    try:
        documents, stats = _read_supported_archive_documents(
            path,
            deposit_filename=deposit_filename,
            include_content=False,
            on_limit="truncate",
        )
    except HTTPException as exc:
        detail = exc.detail if isinstance(exc.detail, str) else "ZIP archive could not be analyzed"
        return None, str(detail)
    except (OSError, RuntimeError, zipfile.BadZipFile) as exc:
        return None, str(exc)
    return {
        "supported_document_count": len(documents),
        "supported_extensions": sorted({str(doc.get("extension") or "") for doc in documents if doc.get("extension")}),
        "uncompressed_bytes": stats.get("uncompressed_bytes", 0),
        "truncated_files": stats.get("truncated_files", 0),
        "max_files": stats.get("max_files"),
    }, None


def _recommend_received_file(file: DepositFile) -> tuple[str, str, str, bool, dict[str, Any] | None]:
    filename = file.filename or "upload"
    ext = extension_for(filename)
    source_path = staged_file_path(file)
    archive_summary: dict[str, Any] | None = None

    if not source_path.exists():
        return (
            "unsupported",
            "Missing staged file",
            "The database row exists, but the raw staged file is missing on disk.",
            False,
            None,
        )
    if (file.size_bytes or 0) <= 0:
        return ("unsupported", "Empty file", "Empty files are kept in staging and are not indexable.", False, None)
    if ext in _LEGACY_SPREADSHEET_EXTENSIONS:
        return (
            "unsupported",
            "Unsupported",
            "Legacy .xls spreadsheets are not supported for Knowledge promotion yet.",
            False,
            None,
        )
    if ext == "zip":
        archive_summary, archive_error = _archive_assist_summary(source_path, deposit_filename=filename)
        if archive_error:
            if "no supported" in archive_error.lower():
                return (
                    "unsupported",
                    "Unsupported",
                    "ZIP contains no currently supported Knowledge document.",
                    False,
                    archive_summary,
                )
            return (
                "inspect_archive",
                "Inspect archive",
                f"ZIP received, but the archive needs inspection before promotion: {archive_error}",
                False,
                archive_summary,
            )
        supported_count = int((archive_summary or {}).get("supported_document_count") or 0)
        truncated_count = int((archive_summary or {}).get("truncated_files") or 0)
        if supported_count <= 0:
            return (
                "unsupported",
                "Unsupported",
                "ZIP contains no currently supported Knowledge document.",
                False,
                archive_summary,
            )
        if supported_count > 10 or truncated_count > 0:
            return (
                "inspect_archive",
                "Inspect archive",
                f"ZIP contains {supported_count} supported document(s); browse it before batch promotion.",
                False,
                archive_summary,
            )
        return (
            "promote_now",
            "Promote now",
            f"ZIP contains {supported_count} supported document(s) and is eligible for Knowledge ingestion.",
            True,
            archive_summary,
        )
    if ext in _SPREADSHEET_EXTENSIONS and not zipfile.is_zipfile(source_path):
        return (
            "unsupported",
            "Unsupported",
            "Office spreadsheet container is invalid or incomplete.",
            False,
            None,
        )
    if ext not in _WORKER_PROMOTION_EXTENSIONS:
        return (
            "unsupported",
            "Unsupported",
            f".{ext or 'unknown'} is not currently supported for Knowledge promotion.",
            False,
            None,
        )
    return (
        "promote_now",
        "Promote now",
        "Received file is supported and ready for Knowledge ingestion.",
        True,
        None,
    )


def build_indexing_assist_snapshot(
    db: DBSession,
    *,
    workspace: Workspace,
    file_ids: list[str],
    collection_slug: str | None,
) -> dict[str, Any]:
    safe_file_ids = [str(file_id).strip() for file_id in file_ids if str(file_id).strip()]
    if len(safe_file_ids) > 200:
        raise HTTPException(status_code=422, detail="Indexing assist is limited to 200 files")

    target_slug = _target_collection_slug(workspace, collection_slug)
    rows = (
        db.query(DepositFile)
        .filter(DepositFile.workspace_id == workspace.id, DepositFile.id.in_(safe_file_ids))
        .all()
    )
    rows_by_id = {row.id: row for row in rows}
    recommendations: list[dict[str, Any]] = []
    recommended_file_ids: list[str] = []
    summary = {
        "total_files": len(safe_file_ids),
        "found_files": len(rows),
        "promote_now_count": 0,
        "inspect_archive_count": 0,
        "unsupported_count": 0,
        "already_promoted_count": 0,
        "needs_target_count": 0,
        "recommended_count": 0,
        "recommended_bytes": 0,
        "zip_count": 0,
        "missing_count": max(0, len(safe_file_ids) - len(rows)),
    }

    for file_id in safe_file_ids:
        file = rows_by_id.get(file_id)
        if not file:
            summary["unsupported_count"] += 1
            recommendations.append(
                {
                    "file_id": file_id,
                    "filename": "",
                    "status": "missing",
                    "extension": "",
                    "recommendation": "unsupported",
                    "label": "Missing",
                    "reason": "File is not visible in this workspace.",
                    "eligible_for_batch": False,
                    "size_bytes": 0,
                    "archive": None,
                }
            )
            continue

        ext = extension_for(file.filename or "")
        if ext == "zip":
            summary["zip_count"] += 1
        if not target_slug:
            recommendation, label, reason, eligible, archive_summary = (
                "needs_target",
                "Target missing",
                "Choose a target Knowledge collection before promotion.",
                False,
                None,
            )
        elif file.status == "promoted":
            recommendation, label, reason, eligible, archive_summary = (
                "already_promoted",
                "Promoted",
                f"Already promoted to {file.promoted_collection_slug or 'Knowledge'}.",
                False,
                None,
            )
        elif file.status != "received":
            recommendation, label, reason, eligible, archive_summary = (
                "unsupported",
                "Unsupported",
                f"File status is {file.status}; only received files can be promoted.",
                False,
                None,
            )
        else:
            recommendation, label, reason, eligible, archive_summary = _recommend_received_file(file)

        if recommendation == "promote_now":
            summary["promote_now_count"] += 1
        elif recommendation == "inspect_archive":
            summary["inspect_archive_count"] += 1
        elif recommendation == "already_promoted":
            summary["already_promoted_count"] += 1
        elif recommendation == "needs_target":
            summary["needs_target_count"] += 1
        else:
            summary["unsupported_count"] += 1

        if eligible:
            recommended_file_ids.append(file.id)
            summary["recommended_count"] += 1
            summary["recommended_bytes"] += int(file.size_bytes or 0)

        recommendations.append(
            {
                "file_id": file.id,
                "filename": file.filename,
                "status": file.status,
                "extension": ext,
                "recommendation": recommendation,
                "label": label,
                "reason": reason,
                "eligible_for_batch": eligible,
                "size_bytes": file.size_bytes or 0,
                "archive": archive_summary,
                "target_collection_slug": target_slug,
                "promoted_collection_slug": file.promoted_collection_slug,
                "worker_job_id": file.worker_job_id,
                "promotion_result": file.promotion_result,
            }
        )

    summary["recommended_file_ids"] = recommended_file_ids
    summary["target_collection_slug"] = target_slug
    return {
        "summary": summary,
        "recommendations": recommendations,
        "collection": _collection_indexing_snapshot(db, workspace=workspace, collection_slug=target_slug),
    }


def get_link_by_access_id(db: DBSession, access_id: str) -> DepositAccessLink:
    link = db.query(DepositAccessLink).filter(DepositAccessLink.access_id == access_id).first()
    if not link:
        raise HTTPException(status_code=404, detail="Deposit link not found")
    return link


def assert_link_usable(db: DBSession, link: DepositAccessLink) -> None:
    if link.status != "active":
        raise HTTPException(status_code=403, detail="Deposit link is not active")
    if link.expires_at and link.expires_at < datetime.utcnow():
        link.status = "expired"
        db.flush()
        raise HTTPException(status_code=403, detail="Deposit link has expired")


def create_link(
    db: DBSession,
    *,
    workspace: Workspace,
    user: User,
    label: str,
    expires_at: Optional[datetime],
    max_file_size_mb: Optional[int],
    allowed_extensions: Optional[list[str]],
    sftp_auth_audit_profile: str | None = None,
) -> tuple[DepositAccessLink, str]:
    if not is_workspace_enabled(workspace):
        raise HTTPException(status_code=403, detail="Secure Deposit is not enabled for this workspace")
    if sftp_auth_audit_profile not in {
        None,
        SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY,
    }:
        raise HTTPException(status_code=422, detail="Unsupported SFTP auth audit profile")

    password = generate_deposit_password()
    access_id_factory = (
        generate_release_a_canary_access_id
        if sftp_auth_audit_profile == SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY
        else generate_access_id
    )
    access_id = access_id_factory()
    while db.query(DepositAccessLink).filter(DepositAccessLink.access_id == access_id).first():
        access_id = access_id_factory()

    link = DepositAccessLink(
        workspace_id=workspace.id,
        created_by_user_id=user.id,
        label=(label or "External deposit")[:255],
        access_id=access_id,
        password_hash=hash_password(password),
        expires_at=expires_at,
        max_file_size_mb=max_file_size_mb or settings.secure_deposit_default_max_file_size_mb,
        allowed_extensions=(
            default_allowed_extensions()
            if allowed_extensions is None
            else [item.lower().lstrip(".") for item in allowed_extensions if item]
        ),
    )
    db.add(link)
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=workspace.id,
        event_type="deposit.link.created",
        actor=user.email or user.username or user.id,
        details={"access_id": link.access_id, "link_id": link.id, "label": link.label},
    )
    return link, password


def rotate_link_password(
    db: DBSession,
    *,
    link: DepositAccessLink,
    user: User,
) -> str:
    if is_release_a_sftp_canary_access_id(link.access_id):
        raise HTTPException(
            status_code=409,
            detail="SFTP canary links cannot rotate credentials",
        )
    password = generate_deposit_password()
    link.password_hash = hash_password(password)
    link.updated_at = datetime.utcnow()
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=link.workspace_id,
        event_type="deposit.link.password_rotated",
        actor=user.email or user.username or user.id,
        details={"access_id": link.access_id, "link_id": link.id},
    )
    return password


def revoke_link(db: DBSession, *, link: DepositAccessLink, user: User) -> None:
    link.status = "revoked"
    link.updated_at = datetime.utcnow()
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=link.workspace_id,
        event_type="deposit.link.revoked",
        actor=user.email or user.username or user.id,
        details={"access_id": link.access_id, "link_id": link.id},
    )


def authenticate_link(db: DBSession, *, access_id: str, password: str) -> tuple[DepositAccessLink, str, datetime]:
    link = get_link_by_access_id(db, access_id)
    if is_release_a_sftp_canary_access_id(link.access_id):
        raise HTTPException(
            status_code=403,
            detail="SFTP canary links do not support public sessions",
        )
    workspace = db.query(Workspace).filter(Workspace.id == link.workspace_id).first()
    if not workspace or not is_workspace_enabled(workspace):
        raise HTTPException(status_code=403, detail="Secure Deposit is not enabled")
    try:
        assert_link_usable(db, link)
    except HTTPException:
        emit_audit_event(
            workspace_id=link.workspace_id,
            event_type="deposit.auth.failed",
            actor=f"deposit:{access_id}",
            severity="warning",
            details={"access_id": access_id, "reason": "inactive_or_expired"},
        )
        raise

    if not verify_password(link.password_hash, password):
        emit_audit_event(
            workspace_id=link.workspace_id,
            event_type="deposit.auth.failed",
            actor=f"deposit:{access_id}",
            severity="warning",
            details={"access_id": access_id, "reason": "bad_password"},
        )
        raise HTTPException(status_code=401, detail="Invalid deposit password")

    token, expires_at = issue_session_token(link)
    emit_audit_event(
        db=db,
        workspace_id=link.workspace_id,
        event_type="deposit.auth.success",
        actor=f"deposit:{access_id}",
        details={"access_id": access_id, "link_id": link.id},
    )
    return link, token, expires_at


async def receive_file(
    db: DBSession,
    *,
    link: DepositAccessLink,
    upload: UploadFile,
) -> DepositFile:
    assert_link_usable(db, link)
    filename = safe_filename(upload.filename)
    allowed = [item.lower().lstrip(".") for item in (link.allowed_extensions or []) if item]
    ext = extension_for(filename)
    if allowed and ext not in allowed:
        emit_audit_event(
            workspace_id=link.workspace_id,
            event_type="deposit.file.rejected",
            actor=f"deposit:{link.access_id}",
            severity="warning",
            details={"access_id": link.access_id, "filename": filename, "reason": "extension_not_allowed"},
        )
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="File extension is not allowed")

    max_bytes = int(link.max_file_size_mb or settings.secure_deposit_default_max_file_size_mb) * 1024 * 1024
    file = DepositFile(
        workspace_id=link.workspace_id,
        access_link_id=link.id,
        filename=filename,
        content_type=upload.content_type,
        object_key="pending",
        size_bytes=0,
        sha256="",
        status="received",
    )
    db.add(file)
    db.flush()
    object_key = _storage_key(
        "workspaces",
        link.workspace_id,
        "secure-deposit",
        link.access_id,
        file.id,
        filename,
    )
    try:
        size_bytes, sha256 = await _write_staged_upload(object_key, upload, max_bytes)
    except HTTPException as exc:
        if exc.status_code == 413:
            emit_audit_event(
                workspace_id=link.workspace_id,
                event_type="deposit.file.rejected",
                actor=f"deposit:{link.access_id}",
                severity="warning",
                details={"access_id": link.access_id, "filename": filename, "reason": "file_too_large"},
            )
        db.delete(file)
        db.flush()
        raise
    except Exception:
        db.delete(file)
        db.flush()
        raise
    file.object_key = object_key
    file.size_bytes = size_bytes
    file.sha256 = sha256
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=link.workspace_id,
        event_type="deposit.file.received",
        actor=f"deposit:{link.access_id}",
        details={
            "access_id": link.access_id,
            "link_id": link.id,
            "file_id": file.id,
            "filename": filename,
            "size_bytes": file.size_bytes,
            "sha256": sha256,
        },
    )
    return file


def record_staged_file_from_path(
    db: DBSession,
    *,
    link: DepositAccessLink,
    source_path: Path,
    filename: str,
    content_type: str | None = None,
    actor: str | None = None,
    transport: str = "sftp",
) -> DepositFile:
    """Persist an already-written upload into the Secure Deposit staging store."""

    assert_link_usable(db, link)
    safe_name = safe_relative_path(filename)
    allowed = [item.lower().lstrip(".") for item in (link.allowed_extensions or []) if item]
    ext = extension_for(safe_name)
    audit_actor = actor or f"deposit:{link.access_id}"
    if allowed and ext not in allowed:
        emit_audit_event(
            workspace_id=link.workspace_id,
            event_type="deposit.file.rejected",
            actor=audit_actor,
            severity="warning",
            details={
                "access_id": link.access_id,
                "filename": safe_name,
                "reason": "extension_not_allowed",
                "transport": transport,
            },
        )
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="File extension is not allowed")

    if not source_path.exists() or not source_path.is_file():
        raise HTTPException(status_code=400, detail="Uploaded file is missing")

    size_bytes, sha256 = _hash_local_file(source_path)
    max_bytes = int(link.max_file_size_mb or settings.secure_deposit_default_max_file_size_mb) * 1024 * 1024
    if size_bytes > max_bytes:
        emit_audit_event(
            workspace_id=link.workspace_id,
            event_type="deposit.file.rejected",
            actor=audit_actor,
            severity="warning",
            details={
                "access_id": link.access_id,
                "filename": safe_name,
                "reason": "file_too_large",
                "transport": transport,
            },
        )
        raise HTTPException(status_code=413, detail="File is too large")

    file = DepositFile(
        workspace_id=link.workspace_id,
        access_link_id=link.id,
        filename=safe_name,
        content_type=content_type,
        object_key="pending",
        size_bytes=size_bytes,
        sha256=sha256,
        status="received",
    )
    db.add(file)
    db.flush()
    object_key = _storage_key(
        "workspaces",
        link.workspace_id,
        "secure-deposit",
        link.access_id,
        file.id,
        safe_name,
    )
    destination = _storage_path(object_key)
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(source_path), destination)
    file.object_key = object_key
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=link.workspace_id,
        event_type="deposit.file.received",
        actor=audit_actor,
        details={
            "access_id": link.access_id,
            "link_id": link.id,
            "file_id": file.id,
            "filename": safe_name,
            "size_bytes": file.size_bytes,
            "sha256": sha256,
            "transport": transport,
        },
    )
    # Live SFTP path: emit on staging close (not only reconciliation).
    if transport == "sftp":
        _emit_sftp_file_arrived_on_staging(
            db,
            workspace_id=link.workspace_id,
            file_id=file.id,
            filename=safe_name,
            access_id=link.access_id,
            size_bytes=int(file.size_bytes or 0),
            sha256=sha256,
        )
    return file


def _emit_sftp_file_arrived_on_staging(
    db: DBSession,
    *,
    workspace_id: str | None,
    file_id: str,
    filename: str,
    access_id: str,
    size_bytes: int,
    sha256: str,
) -> None:
    """Fire ``sftp.file_arrived`` when an SFTP upload is staged (flag-gated).

    Complements the reconciliation emission path so showcase live mode can
    react at connection close. Exception-safe: never breaks staging.
    """
    try:
        from app.services.run_engine import triggers

        if not triggers.is_event_triggers_enabled(workspace_id, db=db):
            return
        triggers.emit_sftp_file_arrived(
            db,
            workspace_id=workspace_id,
            payload={
                "workspace_id": workspace_id,
                "file_id": file_id,
                "filename": filename,
                "access_id": access_id,
                "size_bytes": size_bytes,
                "sha256": sha256,
                "source": "sftp_staging",
            },
        )
    except Exception:  # noqa: BLE001 — never break staging.
        pass


def _emit_deposit_promoted_event(
    db: DBSession,
    *,
    workspace_id: str,
    collection_slug: str | None,
    file_ids: list[str],
) -> None:
    """Fire the Phase 3 ``deposit.promoted`` event trigger (flag-gated, safe).

    Inert unless the global master switch OR the workspace opt-in is ON. Any
    failure — import, registry, dispatch — is swallowed so an event-trigger
    problem can NEVER break a Knowledge promotion. ``deposit.promoted`` is the
    only governance-permitted event that may feed a side-effecting downstream
    run (still HITL-gated); see docs/adr-flow-source-nodes.md §6.
    """
    try:
        from app.services.run_engine import triggers

        if not triggers.is_event_triggers_enabled(workspace_id, db=db):
            return
        triggers.emit_deposit_promoted(
            db,
            workspace_id=workspace_id,
            collection_slug=collection_slug,
            file_ids=file_ids,
        )
    except Exception:  # noqa: BLE001 — never break promotion.
        pass


async def promote_file_to_collection(
    db: DBSession,
    *,
    deposit_file: DepositFile,
    workspace: Workspace,
    user: User,
    collection_slug: str,
) -> DepositFile:
    with using_workspace_project_scheme(workspace):
        return await _promote_file_to_collection(
            db,
            deposit_file=deposit_file,
            workspace=workspace,
            user=user,
            collection_slug=collection_slug,
        )


async def _promote_file_to_collection(
    db: DBSession,
    *,
    deposit_file: DepositFile,
    workspace: Workspace,
    user: User,
    collection_slug: str,
) -> DepositFile:
    if deposit_file.workspace_id != workspace.id:
        raise HTTPException(status_code=404, detail="Deposit file not found")
    if deposit_file.status != "received":
        raise HTTPException(status_code=409, detail="Deposit file is not pending promotion")

    default_collection_slug = f"{workspace.slug}-secure-deposit"
    collection_slug = (collection_slug or default_collection_slug).strip() or default_collection_slug
    extension = extension_for(deposit_file.filename or "")
    if extension == "zip":
        promoted = _promote_archive_file_to_collection(
            db,
            deposit_file=deposit_file,
            workspace=workspace,
            user=user,
            collection_slug=collection_slug,
        )
        _emit_deposit_promoted_event(
            db,
            workspace_id=workspace.id,
            collection_slug=promoted.promoted_collection_slug or collection_slug,
            file_ids=[promoted.id],
        )
        return promoted
    if extension in _LEGACY_SPREADSHEET_EXTENSIONS:
        raise HTTPException(
            status_code=422,
            detail="Legacy .xls spreadsheets are not supported for Knowledge promotion yet",
        )
    if extension in _SPREADSHEET_EXTENSIONS:
        promoted = _promote_single_worker_file_to_collection(
            db,
            deposit_file=deposit_file,
            workspace=workspace,
            user=user,
            collection_slug=collection_slug,
            mode="spreadsheet",
        )
        _emit_deposit_promoted_event(
            db,
            workspace_id=workspace.id,
            collection_slug=promoted.promoted_collection_slug or collection_slug,
            file_ids=[promoted.id],
        )
        return promoted

    app_settings = get_resolved_settings(workspace_id=workspace.id)
    db_type = resolve_vector_db_type(app_settings)
    tmp_dir = tempfile.mkdtemp()
    try:
        tmp_path = Path(tmp_dir) / deposit_file.filename
        _copy_staged_to_local(deposit_file.object_key, tmp_path)
        doc_service = DocumentService(
            collection_name=collection_slug,
            vector_db_type=db_type,
            workspace_slug=workspace.slug,
        )
        document_metadata = _archive_document_metadata(
            deposit_filename=deposit_file.filename,
            archive_path=None,
            document_name=_single_document_name(deposit_file),
            extension=extension,
            source_deposit_file_id=deposit_file.id,
        )
        result = await doc_service.ingest_document(
            str(tmp_path),
            document_metadata=document_metadata,
        )
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    deposit_file.status = "promoted"
    deposit_file.promoted_at = datetime.utcnow()
    deposit_file.promoted_by_user_id = user.id
    deposit_file.promoted_collection_slug = collection_slug
    deposit_file.promotion_result = result
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=workspace.id,
        event_type="deposit.file.promoted",
        actor=user.email or user.username or user.id,
        details={
            "file_id": deposit_file.id,
            "filename": deposit_file.filename,
            "collection_slug": collection_slug,
            "result": result,
        },
    )
    _emit_deposit_promoted_event(
        db,
        workspace_id=workspace.id,
        collection_slug=collection_slug,
        file_ids=[deposit_file.id],
    )
    return deposit_file


def _single_document_name(deposit_file: DepositFile) -> str:
    raw_name = str(deposit_file.filename or "upload").replace("\\", "/")
    safe_path = safe_relative_path(raw_name)
    return safe_path.replace("/", "__")


def _promote_single_worker_file_to_collection(
    db: DBSession,
    *,
    deposit_file: DepositFile,
    workspace: Workspace,
    user: User,
    collection_slug: str,
    mode: str,
) -> DepositFile:
    source_path = staged_file_path(deposit_file)
    document_name = _single_document_name(deposit_file)
    collection = create_or_get_collection(
        db,
        workspace=workspace,
        name=collection_slug,
        description=f"Secure Deposit {mode} promotion from {deposit_file.filename}",
        created_by_user_id=user.id,
        slug=collection_slug,
    )
    store = get_object_store()
    existing_names = list(collection.document_names or [])
    if document_name not in existing_names:
        existing_names.append(document_name)
    store.write_bytes(original_key(collection, document_name), source_path.read_bytes())
    document_metadata = _archive_document_metadata(
        deposit_filename=deposit_file.filename,
        archive_path=None,
        document_name=document_name,
        extension=extension_for(deposit_file.filename or ""),
        source_deposit_file_id=deposit_file.id,
    )
    manifest_key = document_manifest_key(collection)
    document_manifest: dict[str, Any] = {}
    if store.exists(manifest_key):
        try:
            loaded_manifest = json.loads(store.read_bytes(manifest_key).decode("utf-8"))
            if isinstance(loaded_manifest, dict):
                document_manifest = loaded_manifest
        except Exception:
            document_manifest = {}
    document_manifest[document_name] = document_metadata
    store.write_text(manifest_key, json.dumps(document_manifest, ensure_ascii=True, indent=2, sort_keys=True))

    update_collection_status(
        db,
        collection.id,
        status="queued",
        document_names=existing_names,
        document_count=len(existing_names),
    )
    job = create_worker_job(
        db,
        workspace_id=workspace.id,
        collection_id=collection.id,
        kind="document_ingest_index",
    )
    db.commit()
    db.refresh(job)
    db.refresh(collection)

    celery_task_id = dispatch_worker_job(db, job)
    db.commit()
    db.refresh(job)
    db.refresh(collection)

    result = {
        "status": "queued",
        "indexing_status": "queued",
        "mode": mode,
        "collection_id": collection.id,
        "collection_slug": collection.slug,
        "job_id": job.id,
        "celery_task_id": job.celery_task_id or celery_task_id,
        mode: {
            "filename": deposit_file.filename,
            "document_name": document_name,
            "extension": extension_for(deposit_file.filename or ""),
            "size_bytes": deposit_file.size_bytes,
            "metadata": document_metadata,
        },
    }
    deposit_file.status = "promoted"
    deposit_file.promoted_at = datetime.utcnow()
    deposit_file.promoted_by_user_id = user.id
    deposit_file.promoted_collection_slug = collection.slug
    deposit_file.worker_job_id = job.id
    deposit_file.promotion_result = result
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=workspace.id,
        event_type="deposit.file.promoted",
        actor=user.email or user.username or user.id,
        details={
            "file_id": deposit_file.id,
            "filename": deposit_file.filename,
            "collection_slug": collection.slug,
            "worker_job_id": job.id,
            "mode": mode,
            "result": result,
        },
    )
    return deposit_file


def promote_files_to_collection_batch(
    db: DBSession,
    *,
    deposit_files: list[DepositFile],
    workspace: Workspace,
    user: User,
    collection_slug: str,
) -> dict[str, Any]:
    """Promote several Knowledge-supported deposit files with one worker job.

    Qdrant itself supports incremental upserts. The single-job constraint comes
    from Agentium's current canonical worker: ``document_ingest_index`` reloads
    ``collection.document_names`` and ingests that full collection snapshot.
    Dispatching one job per deposit file would therefore re-ingest the same
    growing collection repeatedly.
    """
    with using_workspace_project_scheme(workspace):
        return _promote_files_to_collection_batch(
            db,
            deposit_files=deposit_files,
            workspace=workspace,
            user=user,
            collection_slug=collection_slug,
        )


def _promote_files_to_collection_batch(
    db: DBSession,
    *,
    deposit_files: list[DepositFile],
    workspace: Workspace,
    user: User,
    collection_slug: str,
) -> dict[str, Any]:
    if len(deposit_files) > _BULK_PROMOTION_MAX_FILES:
        raise HTTPException(
            status_code=422,
            detail=f"Bulk promotion is limited to {_BULK_PROMOTION_MAX_FILES} files",
        )

    default_collection_slug = f"{workspace.slug}-secure-deposit"
    collection_slug = (collection_slug or default_collection_slug).strip() or default_collection_slug
    selected: list[dict[str, Any]] = []
    skipped: list[dict[str, str]] = []
    seen_ids: set[str] = set()
    document_total = 0

    for deposit_file in deposit_files:
        if deposit_file.id in seen_ids:
            continue
        seen_ids.add(deposit_file.id)
        reason = ""
        extension = extension_for(deposit_file.filename or "")
        source_path: Path | None = None
        if deposit_file.workspace_id != workspace.id:
            reason = "wrong_workspace"
        elif deposit_file.status != "received":
            reason = f"status_{deposit_file.status}"
        else:
            source_path = staged_file_path(deposit_file)
            if extension in _LEGACY_SPREADSHEET_EXTENSIONS:
                reason = "legacy_xls_unsupported"
            elif extension != "zip" and extension not in _WORKER_PROMOTION_EXTENSIONS:
                reason = "unsupported_for_knowledge_bulk"
            else:
                if not source_path.exists():
                    reason = "staged_file_missing"
                elif (deposit_file.size_bytes or 0) <= 0:
                    reason = "empty_file"
                elif extension in _SPREADSHEET_EXTENSIONS and not zipfile.is_zipfile(source_path):
                    reason = "invalid_office_spreadsheet"

        if reason:
            skipped.append({"file_id": deposit_file.id, "filename": deposit_file.filename or "", "reason": reason})
            continue

        if source_path is None:
            source_path = staged_file_path(deposit_file)

        if extension == "zip":
            try:
                archive_documents, _archive_stats = _read_supported_archive_documents(
                    source_path,
                    deposit_filename=deposit_file.filename,
                    source_deposit_file_id=deposit_file.id,
                )
            except HTTPException as exc:
                skipped.append(
                    {
                        "file_id": deposit_file.id,
                        "filename": deposit_file.filename or "",
                        "reason": f"archive_{exc.detail}",
                    }
                )
                continue
            document_total += len(archive_documents)
            selected.append(
                {
                    "deposit_file": deposit_file,
                    "mode": "archive",
                    "extension": extension,
                    "documents": archive_documents,
                }
            )
        else:
            document_total += 1
            selected.append(
                {
                    "deposit_file": deposit_file,
                    "mode": "document",
                    "extension": extension,
                    "documents": [
                        {
                            "filename": _single_document_name(deposit_file),
                            "extension": extension,
                            "size_bytes": int(deposit_file.size_bytes or 0),
                            "metadata": _archive_document_metadata(
                                deposit_filename=deposit_file.filename,
                                archive_path=None,
                                document_name=_single_document_name(deposit_file),
                                extension=extension,
                                source_deposit_file_id=deposit_file.id,
                            ),
                            "source_path": source_path,
                        }
                    ],
                }
            )

        if document_total > _BULK_PROMOTION_MAX_DOCUMENTS:
            raise HTTPException(
                status_code=413,
                detail=f"Bulk promotion is limited to {_BULK_PROMOTION_MAX_DOCUMENTS} extracted documents",
            )

    if not selected:
        raise HTTPException(status_code=422, detail={"message": "No promotable Knowledge documents found", "skipped": skipped})

    collection = create_or_get_collection(
        db,
        workspace=workspace,
        name=collection_slug,
        description=f"Secure Deposit bulk promotion ({len(selected)} files)",
        created_by_user_id=user.id,
        slug=collection_slug,
    )
    store = get_object_store()
    document_names = list(collection.document_names or [])
    document_name_set = set(document_names)
    manifest_key = document_manifest_key(collection)
    document_manifest: dict[str, Any] = {}
    if store.exists(manifest_key):
        try:
            loaded_manifest = json.loads(store.read_bytes(manifest_key).decode("utf-8"))
            if isinstance(loaded_manifest, dict):
                document_manifest = loaded_manifest
        except Exception:
            document_manifest = {}

    promoted_payload: list[dict[str, Any]] = []
    for item in selected:
        deposit_file: DepositFile = item["deposit_file"]
        item_documents: list[dict[str, Any]] = []
        for document in item["documents"]:
            document_name = str(document["filename"])
            if "content" in document:
                store.write_bytes(original_key(collection, document_name), document["content"])
            else:
                store.write_bytes(original_key(collection, document_name), Path(document["source_path"]).read_bytes())
            if document_name not in document_name_set:
                document_names.append(document_name)
                document_name_set.add(document_name)
            document_metadata = dict(document.get("metadata") or {})
            if document_metadata:
                document_manifest[document_name] = document_metadata
            item_documents.append(
                {
                    "document_name": document_name,
                    "archive_path": document.get("archive_path"),
                    "extension": document.get("extension") or item["extension"],
                    "size_bytes": int(document.get("size_bytes") or deposit_file.size_bytes or 0),
                    "metadata": document_metadata,
                }
            )
        promoted_payload.append(
            {
                "file_id": deposit_file.id,
                "filename": deposit_file.filename,
                "mode": item["mode"],
                "extension": item["extension"],
                "size_bytes": deposit_file.size_bytes,
                "documents": item_documents,
            }
        )
    store.write_text(manifest_key, json.dumps(document_manifest, ensure_ascii=True, indent=2, sort_keys=True))

    update_collection_status(
        db,
        collection.id,
        status="queued",
        document_names=document_names,
        document_count=len(document_names),
    )
    job = create_worker_job(
        db,
        workspace_id=workspace.id,
        collection_id=collection.id,
        kind="document_ingest_index",
    )
    db.commit()
    db.refresh(job)
    db.refresh(collection)

    celery_task_id = dispatch_worker_job(db, job)
    db.commit()
    db.refresh(job)
    db.refresh(collection)

    result = {
        "status": "queued",
        "indexing_status": "queued",
        "mode": "document_bulk",
        "collection_id": collection.id,
        "collection_slug": collection.slug,
        "job_id": job.id,
        "celery_task_id": job.celery_task_id or celery_task_id,
        "promoted_count": len(selected),
        "document_count": document_total,
        "skipped": skipped,
        "files": promoted_payload,
    }

    for item in selected:
        deposit_file: DepositFile = item["deposit_file"]
        file_payload = next((payload for payload in promoted_payload if payload["file_id"] == deposit_file.id), {})
        file_result = {
            **result,
            "file": file_payload,
        }
        deposit_file.status = "promoted"
        deposit_file.promoted_at = datetime.utcnow()
        deposit_file.promoted_by_user_id = user.id
        deposit_file.promoted_collection_slug = collection.slug
        deposit_file.worker_job_id = job.id
        deposit_file.promotion_result = file_result
        emit_audit_event(
            db=db,
            workspace_id=workspace.id,
            event_type="deposit.file.promoted",
            actor=user.email or user.username or user.id,
            details={
                "file_id": deposit_file.id,
                "filename": deposit_file.filename,
                "collection_slug": collection.slug,
                "worker_job_id": job.id,
                "mode": "document_bulk",
                "result": file_result,
            },
        )

    emit_audit_event(
        db=db,
        workspace_id=workspace.id,
        event_type="deposit.bulk_promote.created",
        actor=user.email or user.username or user.id,
        details={
            "collection_slug": collection.slug,
            "worker_job_id": job.id,
            "promoted_count": len(selected),
            "document_count": document_total,
            "skipped_count": len(skipped),
        },
    )
    db.flush()
    promoted_files = [item["deposit_file"] for item in selected]
    _emit_deposit_promoted_event(
        db,
        workspace_id=workspace.id,
        collection_slug=collection.slug,
        file_ids=[f.id for f in promoted_files],
    )
    return {
        "collection": collection,
        "job": job,
        "promoted_files": promoted_files,
        "skipped": skipped,
        "result": result,
    }


def promote_spreadsheet_files_to_collection(
    db: DBSession,
    *,
    deposit_files: list[DepositFile],
    workspace: Workspace,
    user: User,
    collection_slug: str,
) -> dict[str, Any]:
    """Backward-compatible wrapper for older callers/tests."""
    return promote_files_to_collection_batch(
        db,
        deposit_files=deposit_files,
        workspace=workspace,
        user=user,
        collection_slug=collection_slug,
    )


def _promote_archive_file_to_collection(
    db: DBSession,
    *,
    deposit_file: DepositFile,
    workspace: Workspace,
    user: User,
    collection_slug: str,
) -> DepositFile:
    archive_path = staged_file_path(deposit_file)
    documents, _archive_stats = _read_supported_archive_documents(
        archive_path,
        deposit_filename=deposit_file.filename,
        source_deposit_file_id=deposit_file.id,
    )

    collection = create_or_get_collection(
        db,
        workspace=workspace,
        name=collection_slug,
        description=f"Secure Deposit archive promotion from {deposit_file.filename}",
        created_by_user_id=user.id,
        slug=collection_slug,
    )
    store = get_object_store()
    existing_names = list(collection.document_names or [])
    extracted_manifest: list[dict[str, Any]] = []
    manifest_key = document_manifest_key(collection)
    document_manifest: dict[str, Any] = {}
    if store.exists(manifest_key):
        try:
            loaded_manifest = json.loads(store.read_bytes(manifest_key).decode("utf-8"))
            if isinstance(loaded_manifest, dict):
                document_manifest = loaded_manifest
        except Exception:
            document_manifest = {}
    for document in documents:
        filename = str(document["filename"])
        store.write_bytes(original_key(collection, filename), document["content"])
        if filename not in existing_names:
            existing_names.append(filename)
        document_metadata = dict(document.get("metadata") or {})
        if document_metadata:
            document_manifest[filename] = document_metadata
        extracted_manifest.append(
            {
                "archive_path": document["archive_path"],
                "filename": filename,
                "extension": document["extension"],
                "size_bytes": document["size_bytes"],
                "metadata": document_metadata,
            }
        )
    store.write_text(manifest_key, json.dumps(document_manifest, ensure_ascii=True, indent=2, sort_keys=True))

    update_collection_status(
        db,
        collection.id,
        status="queued",
        document_names=existing_names,
        document_count=len(existing_names),
    )
    job = create_worker_job(
        db,
        workspace_id=workspace.id,
        collection_id=collection.id,
        kind="document_ingest_index",
    )
    db.commit()
    db.refresh(job)
    db.refresh(collection)

    celery_task_id = dispatch_worker_job(db, job)
    db.commit()
    db.refresh(job)
    db.refresh(collection)

    result = {
        "status": "queued",
        "indexing_status": "queued",
        "mode": "archive",
        "collection_id": collection.id,
        "collection_slug": collection.slug,
        "job_id": job.id,
        "celery_task_id": job.celery_task_id or celery_task_id,
        "archive": {
            "filename": deposit_file.filename,
            "supported_extensions": sorted(_ARCHIVE_PROMOTION_EXTENSIONS),
            "max_files": int(settings.secure_deposit_archive_promotion_max_files or 50),
            "extracted_count": len(extracted_manifest),
            "documents": extracted_manifest,
        },
    }
    deposit_file.status = "promoted"
    deposit_file.promoted_at = datetime.utcnow()
    deposit_file.promoted_by_user_id = user.id
    deposit_file.promoted_collection_slug = collection.slug
    deposit_file.worker_job_id = job.id
    deposit_file.promotion_result = result
    db.flush()
    emit_audit_event(
        db=db,
        workspace_id=workspace.id,
        event_type="deposit.file.promoted",
        actor=user.email or user.username or user.id,
        details={
            "file_id": deposit_file.id,
            "filename": deposit_file.filename,
            "collection_slug": collection.slug,
            "worker_job_id": job.id,
            "archive_extracted_count": len(extracted_manifest),
            "result": result,
        },
    )
    db.commit()
    db.refresh(deposit_file)
    return deposit_file
