"""Deterministic line-item reconciliation toolkit.

Pure, platform-generic functions behind four canonical skills:

* :func:`extract_spreadsheet_table` — typed rows out of an ``.xlsx`` table;
* :func:`extract_pdf_text` / :func:`extract_invoice_fields` — labeled header
  fields, line items and totals out of a text PDF invoice;
* :func:`reconcile_line_items` — pairs two sets of line items by description
  and compares quantity/unit price against a tolerance;
* :func:`render_reconciliation_report` — a stable plain-text report (no
  timestamps, fixed ordering) so the same input always yields the same bytes.

No I/O and no tenant awareness here: byte resolution and workspace scoping
live in ``skills_registry.file_resolution``; these functions receive bytes or
already-extracted structures.
"""
from __future__ import annotations

import difflib
import re
from collections.abc import Mapping
from io import BytesIO
from typing import Any

# ---------------------------------------------------------------------------
# Spreadsheet extraction
# ---------------------------------------------------------------------------

_FORMULA_RE = re.compile(r"^=\s*([A-Z]{1,3}[0-9]+)\s*([*+/-])\s*([A-Z]{1,3}[0-9]+)\s*$")


def _detect_header_row(raw_rows: list[tuple[Any, ...]], max_column: int) -> int:
    """First 1-based row where most cells are non-empty strings.

    Title/banner rows carry one long string; a real header row labels most of
    the table's columns. We require string cells in at least half the sheet's
    columns (and at least two).
    """
    threshold = max(2, (max_column + 1) // 2)
    for index, row in enumerate(raw_rows, start=1):
        strings = sum(
            1 for cell in row if isinstance(cell, str) and cell.strip() and not str(cell).startswith("=")
        )
        if strings >= threshold:
            return index
    for index, row in enumerate(raw_rows, start=1):
        if sum(1 for cell in row if cell not in (None, "")) >= 2:
            return index
    return 1


def _cell_value(
    cached: Any,
    raw: Any,
    cached_grid: dict[str, Any],
) -> Any:
    """Prefer the cached (computed) value; evaluate simple formulas otherwise.

    ``data_only=True`` yields the last value the authoring tool computed. When
    no cache exists (e.g. a workbook written by openpyxl itself), a formula
    like ``=E5*F5`` survives as a string: we evaluate the binary arithmetic
    against the already-resolved referenced cells so quantity×price totals
    stay usable. Anything more exotic degrades to ``None`` rather than leaking
    a formula string into typed rows.
    """
    if cached is not None and not (isinstance(cached, str) and cached.startswith("=")):
        return cached
    if isinstance(raw, str) and raw.startswith("="):
        match = _FORMULA_RE.match(raw)
        if not match:
            return None
        left = cached_grid.get(match.group(1))
        right = cached_grid.get(match.group(3))
        if not isinstance(left, int | float) or not isinstance(right, int | float):
            return None
        operator = match.group(2)
        if operator == "*":
            return left * right
        if operator == "+":
            return left + right
        if operator == "-":
            return left - right
        return left / right if right else None
    return raw if cached is None else cached


def extract_spreadsheet_table(
    data: bytes,
    *,
    sheet: str | None = None,
    header_row: int | None = None,
    filters: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Extract the typed table body of one worksheet.

    The body stops at the first fully-empty row after the header, which keeps
    subtotal blocks and footnotes below the table out of the rows. ``filters``
    is an exact string match per column name.
    """
    from openpyxl import load_workbook

    cached_wb = load_workbook(BytesIO(data), data_only=True, read_only=False)
    raw_wb = load_workbook(BytesIO(data), data_only=False, read_only=False)
    if sheet:
        if sheet not in cached_wb.sheetnames:
            raise ValueError(
                f"spreadsheet_sheet_not_found: {sheet!r} (has {cached_wb.sheetnames})"
            )
        cached_ws, raw_ws = cached_wb[sheet], raw_wb[sheet]
    else:
        cached_ws, raw_ws = cached_wb.active, raw_wb.active

    cached_rows = [tuple(row) for row in cached_ws.iter_rows(values_only=True)]
    raw_rows = [tuple(row) for row in raw_ws.iter_rows(values_only=True)]
    if not cached_rows:
        return {"columns": [], "rows": [], "row_count": 0}

    cached_grid: dict[str, Any] = {}
    for row_index, row in enumerate(cached_rows, start=1):
        for col_index, value in enumerate(row):
            cached_grid[f"{_column_letter(col_index + 1)}{row_index}"] = value

    header_index = header_row or _detect_header_row(raw_rows, cached_ws.max_column)
    header = raw_rows[header_index - 1]
    columns = [
        str(cell).strip() if cell not in (None, "") else f"column_{position + 1}"
        for position, cell in enumerate(header)
    ]
    while columns and columns[-1].startswith("column_"):
        columns.pop()

    rows: list[dict[str, Any]] = []
    for row_index in range(header_index + 1, len(cached_rows) + 1):
        cached = cached_rows[row_index - 1]
        raw = raw_rows[row_index - 1]
        values = [
            _cell_value(
                cached[position] if position < len(cached) else None,
                raw[position] if position < len(raw) else None,
                cached_grid,
            )
            for position in range(len(columns))
        ]
        if all(value in (None, "") for value in values):
            break
        record = {
            column: value.strip() if isinstance(value, str) else value
            for column, value in zip(columns, values)
        }
        rows.append(record)

    for column, wanted in (filters or {}).items():
        rows = [
            record
            for record in rows
            if str(record.get(column, "")).strip() == str(wanted).strip()
        ]

    return {"columns": columns, "rows": rows, "row_count": len(rows)}


def _column_letter(index: int) -> str:
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


# ---------------------------------------------------------------------------
# Invoice PDF extraction
# ---------------------------------------------------------------------------

_MONEY = r"[\d,]+\.\d{2}"
_INLINE_ITEM_RE = re.compile(
    rf"^(\d+)\s+(.+?)\s+(\d+(?:\.\d+)?)\s+({_MONEY})\s+({_MONEY})$"
)


def extract_pdf_text(data: bytes) -> str:
    """Text of every page, using the same libraries as the document parser.

    Mirrors ``PDFParser``: pdfplumber preferred, PyPDF2 as fallback — both are
    declared backend dependencies, nothing new is introduced.
    """
    try:
        import pdfplumber

        with pdfplumber.open(BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except ImportError:
        import PyPDF2

        reader = PyPDF2.PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)


def _number(text: str) -> float:
    return float(text.replace(",", ""))


def _labeled(text: str, label: str) -> str | None:
    match = re.search(rf"^{label}\s*:\s*(.+)$", text, re.MULTILINE)
    return match.group(1).strip() if match else None


def _total(text: str, label: str) -> float | None:
    match = re.search(rf"{label}\b.*?({_MONEY})", text, re.DOTALL)
    return _number(match.group(1)) if match else None


def _vendor(lines: list[str]) -> str | None:
    for index, line in enumerate(lines):
        if line.strip().lower().startswith("from:"):
            inline = line.split(":", 1)[1].strip()
            if inline:
                return inline
            for candidate in lines[index + 1 :]:
                if candidate.strip():
                    return candidate.strip()
    return None


def _line_items(lines: list[str]) -> list[dict[str, Any]]:
    """Numbered item rows, resilient to both PDF text layouts.

    pdfplumber emits one physical line per row; PyPDF2 emits one token block
    per line (index / description / qty / unit price / line total each on
    their own line). Try the single-line shape first, then reassemble groups.
    """
    items: list[dict[str, Any]] = []
    for line in lines:
        match = _INLINE_ITEM_RE.match(line.strip())
        if match:
            items.append(
                {
                    "description": match.group(2).strip(),
                    "qty": _number(match.group(3)),
                    "unit_price": _number(match.group(4)),
                    "line_total": _number(match.group(5)),
                }
            )
    if items:
        return items

    money_re = re.compile(rf"^{_MONEY}$")
    expected_index = 1
    position = 0
    while position < len(lines):
        if lines[position].strip() != str(expected_index):
            position += 1
            continue
        description_parts: list[str] = []
        cursor = position + 1
        qty: float | None = None
        while cursor < len(lines):
            token = lines[cursor].strip()
            if re.fullmatch(r"\d+(?:\.\d+)?", token) and description_parts:
                qty = _number(token)
                cursor += 1
                break
            if token:
                description_parts.append(token)
            cursor += 1
        numbers: list[float] = []
        while cursor < len(lines) and len(numbers) < 2:
            token = lines[cursor].strip()
            if money_re.fullmatch(token):
                numbers.append(_number(token))
                cursor += 1
                continue
            if token:
                break
            cursor += 1
        if qty is None or len(numbers) != 2:
            break
        items.append(
            {
                "description": " ".join(description_parts),
                "qty": qty,
                "unit_price": numbers[0],
                "line_total": numbers[1],
            }
        )
        expected_index += 1
        position = cursor
    return items


def extract_invoice_fields(text: str) -> dict[str, Any]:
    """Parse labeled header fields, line items and totals from invoice text."""
    lines = text.splitlines()
    result: dict[str, Any] = {
        "invoice_number": _labeled(text, r"Invoice No\.?"),
        "invoice_date": _labeled(text, "Invoice Date"),
        "po_reference": _labeled(text, "PO Reference"),
        "due_date": _labeled(text, "Due Date"),
        "vendor": _vendor(lines),
        "line_items": _line_items(lines),
        "subtotal": _total(text, "Subtotal"),
        "vat": _total(text, "VAT"),
        "total_due": _total(text, "Total Due"),
    }
    currency = re.search(r"(?:Total Due|Subtotal)\s*\(([A-Z]{3})\)", text)
    if currency:
        result["currency"] = currency.group(1)
    return result


# ---------------------------------------------------------------------------
# Reconciliation
# ---------------------------------------------------------------------------

_CANONICAL_FIELDS = {
    "description": "description",
    "designation": "description",
    "item": "description",
    "line item description": "description",
    "qty": "qty",
    "quantity": "qty",
    "unit price": "unit_price",
    "unit_price": "unit_price",
    "price": "unit_price",
    "line total": "line_total",
    "line_total": "line_total",
    "total": "line_total",
    "total amount": "line_total",
    "amount": "line_total",
}


def _normalize_column(name: str) -> str:
    """Case-insensitive, parenthesized units stripped: ``Unit Price (QAR)`` → ``unit price``."""
    cleaned = re.sub(r"\([^)]*\)", " ", str(name))
    return re.sub(r"\s+", " ", cleaned).strip().lower()


def _canonical_line(
    line: Mapping[str, Any], field_map: Mapping[str, str] | None
) -> dict[str, Any]:
    explicit = {str(key): str(value) for key, value in (field_map or {}).items()}
    record: dict[str, Any] = {}
    for key, value in line.items():
        target = explicit.get(str(key)) or _CANONICAL_FIELDS.get(_normalize_column(str(key)))
        if target and target not in record:
            record[target] = value
    result = {"description": str(record.get("description") or "").strip()}
    for field in ("qty", "unit_price", "line_total"):
        raw = record.get(field)
        if isinstance(raw, str):
            raw = raw.replace(",", "").strip()
        try:
            result[field] = float(raw) if raw not in (None, "") else None
        except (TypeError, ValueError):
            result[field] = None
    return result


def _description_key(description: str) -> str:
    return re.sub(r"\s+", " ", description).strip().casefold()


def _variance_pct(po_value: float | None, invoice_value: float | None) -> float | None:
    if po_value in (None, 0) or invoice_value is None:
        return None
    return round((invoice_value - po_value) / po_value * 100.0, 2)


def _line_total(line: Mapping[str, Any]) -> float:
    if line.get("line_total") is not None:
        return float(line["line_total"])
    qty, price = line.get("qty"), line.get("unit_price")
    if qty is None or price is None:
        return 0.0
    return float(qty) * float(price)


def reconcile_line_items(
    po_lines: list[Mapping[str, Any]],
    invoice_lines: list[Mapping[str, Any]],
    *,
    po_reference: str | None = None,
    tolerance_pct: float = 2.0,
    po_field_map: Mapping[str, str] | None = None,
    invoice_field_map: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    """Pair lines by description and flag qty/price variances beyond tolerance.

    Pairing is exact on the normalized description first, then the best
    ``difflib.SequenceMatcher`` ratio above 0.85 among what remains. Output
    ordering is stable: PO lines in input order, then unmatched invoice lines
    in input order.
    """
    po_records = [_canonical_line(line, po_field_map) for line in po_lines]
    invoice_records = [_canonical_line(line, invoice_field_map) for line in invoice_lines]

    unmatched_invoice = list(range(len(invoice_records)))
    pairs: list[tuple[int, int | None]] = []
    for po_index, po_record in enumerate(po_records):
        key = _description_key(po_record["description"])
        match_index = next(
            (
                candidate
                for candidate in unmatched_invoice
                if _description_key(invoice_records[candidate]["description"]) == key
            ),
            None,
        )
        if match_index is None:
            best_ratio = 0.85
            for candidate in unmatched_invoice:
                ratio = difflib.SequenceMatcher(
                    None, key, _description_key(invoice_records[candidate]["description"])
                ).ratio()
                if ratio > best_ratio:
                    best_ratio, match_index = ratio, candidate
        if match_index is not None:
            unmatched_invoice.remove(match_index)
        pairs.append((po_index, match_index))

    lines: list[dict[str, Any]] = []
    for po_index, invoice_index in pairs:
        po_record = po_records[po_index]
        if invoice_index is None:
            lines.append(
                {
                    "description": po_record["description"],
                    "po_qty": po_record["qty"],
                    "invoice_qty": None,
                    "qty_variance_pct": None,
                    "po_unit_price": po_record["unit_price"],
                    "invoice_unit_price": None,
                    "price_variance_pct": None,
                    "status": "unmatched_po_line",
                }
            )
            continue
        invoice_record = invoice_records[invoice_index]
        qty_variance = _variance_pct(po_record["qty"], invoice_record["qty"])
        price_variance = _variance_pct(po_record["unit_price"], invoice_record["unit_price"])
        qty_off = qty_variance is not None and abs(qty_variance) > tolerance_pct
        price_off = price_variance is not None and abs(price_variance) > tolerance_pct
        if qty_off and price_off:
            status = "both_mismatch"
        elif qty_off:
            status = "qty_mismatch"
        elif price_off:
            status = "price_mismatch"
        else:
            status = "matched"
        lines.append(
            {
                "description": po_record["description"],
                "po_qty": po_record["qty"],
                "invoice_qty": invoice_record["qty"],
                "qty_variance_pct": qty_variance,
                "po_unit_price": po_record["unit_price"],
                "invoice_unit_price": invoice_record["unit_price"],
                "price_variance_pct": price_variance,
                "status": status,
            }
        )
    for invoice_index in unmatched_invoice:
        invoice_record = invoice_records[invoice_index]
        lines.append(
            {
                "description": invoice_record["description"],
                "po_qty": None,
                "invoice_qty": invoice_record["qty"],
                "qty_variance_pct": None,
                "po_unit_price": None,
                "invoice_unit_price": invoice_record["unit_price"],
                "price_variance_pct": None,
                "status": "unmatched_invoice_line",
            }
        )

    po_total = round(sum(_line_total(record) for record in po_records), 2)
    invoice_total = round(sum(_line_total(record) for record in invoice_records), 2)
    matched_count = sum(1 for line in lines if line["status"] == "matched")
    return {
        "po_reference": po_reference,
        "lines": lines,
        "matched_count": matched_count,
        "flagged_count": len(lines) - matched_count,
        "po_total": po_total,
        "invoice_total": invoice_total,
        "total_variance_pct": _variance_pct(po_total, invoice_total),
        "tolerance_pct": tolerance_pct,
    }


# ---------------------------------------------------------------------------
# Report rendering
# ---------------------------------------------------------------------------

def _cell(value: Any, width: int, *, money: bool = False) -> str:
    if value is None:
        text = "-"
    elif money:
        text = f"{float(value):,.2f}"
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    return text.rjust(width) if isinstance(value, int | float) else text.ljust(width)


def _signed_pct(value: Any) -> str:
    if value is None:
        return "-"
    return f"{value:+.2f}%"


def _flag_sentence(line: Mapping[str, Any]) -> str:
    description = line["description"]
    status = line["status"]
    if status == "unmatched_invoice_line":
        return f"- Investigate '{description}': billed on the invoice but absent from the PO."
    if status == "unmatched_po_line":
        return f"- Investigate '{description}': ordered on the PO but absent from the invoice."
    parts = []
    if status in ("qty_mismatch", "both_mismatch"):
        parts.append(
            f"quantity billed {_cell(line['invoice_qty'], 0).strip()} vs "
            f"{_cell(line['po_qty'], 0).strip()} ordered ({_signed_pct(line['qty_variance_pct'])})"
        )
    if status in ("price_mismatch", "both_mismatch"):
        parts.append(
            f"unit price billed {float(line['invoice_unit_price']):,.2f} vs "
            f"{float(line['po_unit_price']):,.2f} ordered ({_signed_pct(line['price_variance_pct'])})"
        )
    return f"- Review '{description}': " + " and ".join(parts) + "."


def render_reconciliation_report(
    reconciliation: Mapping[str, Any],
    verdict: str,
    invoice_meta: Mapping[str, Any] | None = None,
) -> str:
    """Deterministic plain-text report: no timestamps, stable line ordering."""
    meta = invoice_meta or {}
    lines_data = list(reconciliation.get("lines") or [])
    tolerance = reconciliation.get("tolerance_pct")

    widths = (34, 8, 8, 10, 10, 10, 11, 22)
    headers = (
        "Description", "PO Qty", "Inv Qty", "Qty Var", "PO Price", "Inv Price", "Price Var", "Status",
    )
    header_row = "  ".join(title.ljust(width) for title, width in zip(headers, widths))
    separator = "-" * len(header_row)

    out: list[str] = []
    out.append("PO / INVOICE RECONCILIATION REPORT")
    out.append("=" * len(out[0]))
    out.append(f"Verdict: {verdict}")
    identity = [
        f"Invoice: {meta['invoice_number']}" if meta.get("invoice_number") else None,
        f"Vendor: {meta['vendor']}" if meta.get("vendor") else None,
        f"Invoice date: {meta['invoice_date']}" if meta.get("invoice_date") else None,
        f"Total due: {float(meta['total_due']):,.2f}" if meta.get("total_due") is not None else None,
    ]
    identity_line = " | ".join(part for part in identity if part)
    if identity_line:
        out.append(identity_line)
    if reconciliation.get("po_reference"):
        out.append(f"PO Reference: {reconciliation['po_reference']}")
    if tolerance is not None:
        out.append(f"Tolerance: +/-{float(tolerance):.1f}%")
    out.append("")
    out.append(header_row)
    out.append(separator)
    for line in lines_data:
        flag = "" if line["status"] == "matched" else "  <-- FLAG"
        out.append(
            "  ".join(
                (
                    _cell(line["description"], widths[0]),
                    _cell(line["po_qty"], widths[1]),
                    _cell(line["invoice_qty"], widths[2]),
                    _signed_pct(line["qty_variance_pct"]).rjust(widths[3]),
                    _cell(line["po_unit_price"], widths[4], money=True),
                    _cell(line["invoice_unit_price"], widths[5], money=True),
                    _signed_pct(line["price_variance_pct"]).rjust(widths[6]),
                    _cell(line["status"], widths[7]),
                )
            ).rstrip()
            + flag
        )
    out.append(separator)
    po_total = reconciliation.get("po_total")
    invoice_total = reconciliation.get("invoice_total")
    total_variance = reconciliation.get("total_variance_pct")
    out.append(
        f"Totals: PO {po_total:,.2f} | Invoice {invoice_total:,.2f}"
        + (f" | Variance {_signed_pct(total_variance)}" if total_variance is not None else "")
    )
    out.append(
        f"Matched lines: {reconciliation.get('matched_count', 0)} | "
        f"Flagged lines: {reconciliation.get('flagged_count', 0)}"
    )
    flagged = [line for line in lines_data if line["status"] != "matched"]
    if flagged:
        out.append("")
        out.append("Recommendations:")
        out.extend(_flag_sentence(line) for line in flagged)
    return "\n".join(out) + "\n"
