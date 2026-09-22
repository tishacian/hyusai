from __future__ import annotations

import hashlib
import json
import stat
import zipfile
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException, UploadFile

from app.core.config import settings
from app.models.audit import AuditLog
from app.models.knowledge_collection import KnowledgeCollection, WorkerJob
from app.models.secure_deposit import DepositAccessLink, DepositFile
from app.models.user import User
from app.models.workspace import Workspace, WorkspaceMember
from app.services import secure_deposit as secure_deposit_service
from app.services import secure_deposit_sftp
from app.services.knowledge_collections import document_manifest_key
from app.services.object_store import get_object_store
from app.services.rag.project_references import derive_project_reference
from app.services.secure_deposit import (
    SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY,
    _archive_document_metadata,
    authenticate_link,
    build_deposit_archive,
    build_file_preview,
    create_link,
    preview_deposit_file,
    promote_file_to_collection,
    receive_file,
    record_staged_file_from_path,
    revoke_link,
    rotate_link_password,
    safe_filename,
    safe_relative_path,
    verify_session_token,
)
from app.services.secure_deposit_sftp import (
    _build_asyncssh_components,
    _is_root_path,
    _remote_dir_path,
    _remote_relative_path,
    _sftp_longname,
)


def _workspace_user(db_session):
    workspace = Workspace(
        id="ws-andritz",
        name="Andritz",
        slug="andritz",
        settings={"family": "andritz", "features": {"secure_deposit": True}},
    )
    user = User(id="user-1", username="thib", email="operator@example.test")
    db_session.add_all([workspace, user])
    db_session.flush()
    db_session.add(
        WorkspaceMember(
            user_id=user.id,
            workspace_id=workspace.id,
            role="owner",
            role_template="workspace_owner",
        )
    )
    db_session.flush()
    return workspace, user


class _FakeAsyncSSH:
    class SSHServer:
        pass

    class SFTPServer:
        def __init__(self, _chan):
            pass

    class SFTPAttrs:
        def __init__(self, **kwargs):
            self.__dict__.update(kwargs)

    class SFTPName:
        def __init__(self, *args, **kwargs):
            self.args = args
            self.kwargs = kwargs

    class SFTPError(Exception):
        pass

    class SFTPNoSuchFile(Exception):  # noqa: N818 - mirrors AsyncSSH API
        pass

    class SFTPFailure(Exception):  # noqa: N818 - mirrors AsyncSSH API
        pass

    class SFTPPermissionDenied(Exception):  # noqa: N818 - mirrors AsyncSSH API
        pass


def _ssh_auth_server(client_version: str):
    ssh_server, _sftp_server = _build_asyncssh_components(_FakeAsyncSSH)
    server = ssh_server()
    server.connection_made(
        SimpleNamespace(
            get_extra_info=lambda key: client_version if key == "client_version" else None
        )
    )
    return server


def test_deposit_link_password_is_one_time_and_session_scoped(db_session):
    workspace, user = _workspace_user(db_session)

    link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Supplier upload",
        expires_at=None,
        max_file_size_mb=20,
        allowed_extensions=["pdf"],
    )
    assert password not in link.password_hash

    authed_link, token, _ = authenticate_link(
        db_session,
        access_id=link.access_id,
        password=password,
    )
    payload = verify_session_token(token, link.access_id)

    assert authed_link.id == link.id
    assert payload["sub"] == link.access_id
    assert payload["workspace_id"] == workspace.id


def test_deposit_session_token_is_bound_to_access_id(db_session):
    workspace, user = _workspace_user(db_session)
    first_link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="First supplier upload",
        expires_at=None,
        max_file_size_mb=20,
        allowed_extensions=["pdf"],
    )
    second_link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Second supplier upload",
        expires_at=None,
        max_file_size_mb=20,
        allowed_extensions=["pdf"],
    )

    _authed_link, token, _ = authenticate_link(
        db_session,
        access_id=first_link.access_id,
        password=password,
    )

    with pytest.raises(HTTPException) as exc:
        verify_session_token(token, second_link.access_id)

    assert exc.value.status_code == 401


def test_deposit_link_rejects_bad_password(db_session):
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Supplier upload",
        expires_at=None,
        max_file_size_mb=None,
        allowed_extensions=None,
    )

    with pytest.raises(HTTPException) as exc:
        authenticate_link(db_session, access_id=link.access_id, password="wrong")

    assert exc.value.status_code == 401


def test_release_a_sftp_canary_success_audit_recovers_after_commit_before_response(
    db_session,
    monkeypatch,
):
    workspace, user = _workspace_user(db_session)
    link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Release A disposable SFTP proof",
        expires_at=None,
        max_file_size_mb=1,
        allowed_extensions=[],
        sftp_auth_audit_profile=SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY,
    )
    db_session.commit()
    server = _ssh_auth_server("SSH-2.0-AgentiumReleaseASFTPPositiveCanary")
    original = secure_deposit_sftp._emit_password_auth_audit
    crashed = False

    def _commit_then_crash(*args, **kwargs):
        nonlocal crashed
        event_id = original(*args, **kwargs)
        if kwargs.get("retry_safe_canary") and not crashed:
            crashed = True
            raise RuntimeError("simulated producer disconnect after audit commit")
        return event_id

    monkeypatch.setattr(
        secure_deposit_sftp,
        "_emit_password_auth_audit",
        _commit_then_crash,
    )

    # The first SSH decision is lost after the authoritative audit commit.
    assert server.validate_password(link.access_id, password) is False
    monkeypatch.setattr(
        secure_deposit_sftp,
        "_emit_password_auth_audit",
        original,
    )

    # Redelivery authenticates and reuses exactly the committed event.
    assert server.validate_password(link.access_id, password) is True
    db_session.expire_all()
    rows = (
        db_session.query(AuditLog)
        .filter(AuditLog.event_type == "deposit.sftp.auth.success")
        .all()
    )
    assert len(rows) == 1
    assert rows[0].id == secure_deposit_sftp._canary_audit_id(
        link,
        event_type="deposit.sftp.auth.success",
        reason=None,
    )
    assert rows[0].details == {"access_id": link.access_id, "link_id": link.id}


def test_release_a_sftp_canary_denial_audit_is_retry_safe_after_revocation(
    db_session,
    monkeypatch,
):
    workspace, user = _workspace_user(db_session)
    link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Release A disposable SFTP proof",
        expires_at=None,
        max_file_size_mb=1,
        allowed_extensions=[],
        sftp_auth_audit_profile=SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY,
    )
    db_session.commit()
    revoke_link(db_session, link=link, user=user)
    db_session.commit()
    server = _ssh_auth_server("SSH-2.0-AgentiumReleaseASFTPPositiveCanary")
    original = secure_deposit_sftp._emit_password_auth_audit
    crashed = False

    def _commit_then_crash(*args, **kwargs):
        nonlocal crashed
        event_id = original(*args, **kwargs)
        if kwargs.get("retry_safe_canary") and not crashed:
            crashed = True
            raise RuntimeError("simulated producer disconnect after denial audit commit")
        return event_id

    monkeypatch.setattr(
        secure_deposit_sftp,
        "_emit_password_auth_audit",
        _commit_then_crash,
    )
    assert server.validate_password(link.access_id, password) is False
    monkeypatch.setattr(
        secure_deposit_sftp,
        "_emit_password_auth_audit",
        original,
    )
    assert server.validate_password(link.access_id, password) is False

    db_session.expire_all()
    rows = (
        db_session.query(AuditLog)
        .filter(AuditLog.event_type == "deposit.sftp.auth.failed")
        .all()
    )
    assert len(rows) == 1
    assert rows[0].id == secure_deposit_sftp._canary_audit_id(
        link,
        event_type="deposit.sftp.auth.failed",
        reason="inactive_or_expired",
    )
    assert rows[0].details == {
        "access_id": link.access_id,
        "link_id": link.id,
        "reason": "inactive_or_expired",
    }


@pytest.mark.parametrize(
    ("profile", "client_version"),
    [
        (None, "SSH-2.0-AgentiumReleaseASFTPPositiveCanary"),
        (SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY, "SSH-2.0-OpenSSH_9.9"),
    ],
)
def test_sftp_business_authentication_is_never_deduplicated_by_one_marker_alone(
    db_session,
    profile,
    client_version,
):
    workspace, user = _workspace_user(db_session)
    link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Ordinary SFTP link",
        expires_at=None,
        max_file_size_mb=20,
        allowed_extensions=[],
        sftp_auth_audit_profile=profile,
    )
    db_session.commit()
    server = _ssh_auth_server(client_version)

    assert server.validate_password(link.access_id, password) is True
    assert server.validate_password(link.access_id, password) is True

    db_session.expire_all()
    rows = (
        db_session.query(AuditLog)
        .filter(AuditLog.event_type == "deposit.sftp.auth.success")
        .all()
    )
    assert len(rows) == 2
    assert rows[0].id != rows[1].id


def test_reserved_canary_access_namespace_requires_the_explicit_profile(
    db_session,
    monkeypatch,
):
    workspace, user = _workspace_user(db_session)
    generated = iter(("password-token", "ra1_attempted-prefix"))
    monkeypatch.setattr(
        secure_deposit_service.secrets,
        "token_urlsafe",
        lambda _size: next(generated),
    )

    link, _password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="ra1_ is only text in this ordinary label",
        expires_at=None,
        max_file_size_mb=20,
        allowed_extensions=[],
    )

    assert link.access_id == "ra1-attempted-prefix"
    assert not secure_deposit_service.is_release_a_sftp_canary_access_id(
        link.access_id
    )
    assert secure_deposit_service.serialize_link(link)[
        "sftp_auth_audit_profile"
    ] is None


def test_release_a_sftp_canary_bad_password_audits_remain_per_attempt(db_session):
    workspace, user = _workspace_user(db_session)
    link, _password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Release A disposable SFTP proof",
        expires_at=None,
        max_file_size_mb=1,
        allowed_extensions=[],
        sftp_auth_audit_profile=SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY,
    )
    db_session.commit()
    server = _ssh_auth_server("SSH-2.0-AgentiumReleaseASFTPPositiveCanary")

    assert server.validate_password(link.access_id, "incorrect-one") is False
    assert server.validate_password(link.access_id, "incorrect-two") is False

    db_session.expire_all()
    rows = (
        db_session.query(AuditLog)
        .filter(AuditLog.event_type == "deposit.sftp.auth.failed")
        .all()
    )
    assert len(rows) == 2
    assert {row.details["reason"] for row in rows} == {"bad_password"}


def test_release_a_sftp_canary_profile_is_sftp_only_and_credential_immutable(db_session):
    workspace, user = _workspace_user(db_session)
    link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Release A disposable SFTP proof",
        expires_at=None,
        max_file_size_mb=1,
        allowed_extensions=[],
        sftp_auth_audit_profile=SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY,
    )

    with pytest.raises(HTTPException) as public_error:
        authenticate_link(db_session, access_id=link.access_id, password=password)
    with pytest.raises(HTTPException) as rotation_error:
        rotate_link_password(db_session, link=link, user=user)

    assert public_error.value.status_code == 403
    assert rotation_error.value.status_code == 409


def test_release_a_sftp_canary_rejects_a_conflicting_deterministic_audit(db_session):
    workspace, user = _workspace_user(db_session)
    link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Release A disposable SFTP proof",
        expires_at=None,
        max_file_size_mb=1,
        allowed_extensions=[],
        sftp_auth_audit_profile=SFTP_AUTH_AUDIT_PROFILE_RELEASE_A_CANARY,
    )
    event_id = secure_deposit_sftp._canary_audit_id(
        link,
        event_type="deposit.sftp.auth.success",
        reason=None,
    )
    db_session.add(
        AuditLog(
            id=event_id,
            workspace_id=workspace.id,
            event_type="deposit.sftp.auth.success",
            actor=f"sftp:{link.access_id}",
            details={"access_id": link.access_id, "link_id": "different-link"},
        )
    )
    db_session.commit()
    server = _ssh_auth_server("SSH-2.0-AgentiumReleaseASFTPPositiveCanary")

    assert server.validate_password(link.access_id, password) is False
    db_session.expire_all()
    assert db_session.query(AuditLog).filter(AuditLog.id == event_id).count() == 1


def test_disabled_workspace_blocks_link_creation(db_session):
    workspace = Workspace(
        id="ws-disabled-deposit",
        name="Disabled deposit",
        slug="disabled-deposit",
        settings={"features": {"secure_deposit": False}},
    )
    user = User(id="user-disabled", username="disabled", email="disabled@example.test")
    db_session.add_all([workspace, user])
    db_session.flush()
    db_session.add(
        WorkspaceMember(
            user_id=user.id,
            workspace_id=workspace.id,
            role="owner",
            role_template="workspace_owner",
        )
    )
    db_session.flush()

    with pytest.raises(HTTPException) as exc:
        create_link(
            db_session,
            workspace=workspace,
            user=user,
            label="Should not be created",
            expires_at=None,
            max_file_size_mb=None,
            allowed_extensions=None,
        )

    assert exc.value.status_code == 403
    assert db_session.query(DepositAccessLink).count() == 0


def test_disabled_workspace_blocks_existing_link_auth(db_session):
    workspace, user = _workspace_user(db_session)
    link, password = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Supplier upload",
        expires_at=None,
        max_file_size_mb=None,
        allowed_extensions=None,
    )
    workspace.settings = {"features": {"secure_deposit": False}}
    db_session.add(workspace)
    db_session.commit()

    with pytest.raises(HTTPException) as exc:
        authenticate_link(db_session, access_id=link.access_id, password=password)

    assert exc.value.status_code == 403


def test_empty_allowed_extensions_means_any_file_type(db_session):
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Any upload",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )

    assert link.allowed_extensions == []
    assert link.max_file_size_mb == 30 * 1024


@pytest.mark.asyncio
async def test_allowed_extension_policy_is_case_insensitive_for_public_and_sftp_paths(
    db_session,
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="PDF only",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=["PDF"],
    )

    public_upload = UploadFile(filename="MANUAL.PDF", file=BytesIO(b"%PDF public upload"))
    public_row = await receive_file(db_session, link=link, upload=public_upload)

    staged = tmp_path / "sftp-upload.part"
    staged.write_bytes(b"%PDF staged upload")
    sftp_row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="Line A/manual.pdf",
        content_type="application/pdf",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    blocked = tmp_path / "blocked.part"
    blocked.write_bytes(b"MZ")
    with pytest.raises(HTTPException) as exc:
        record_staged_file_from_path(
            db_session,
            link=link,
            source_path=blocked,
            filename="Line A/manual.EXE",
            content_type="application/octet-stream",
            actor=f"sftp:{link.access_id}",
            transport="sftp",
        )

    assert public_row.filename == "MANUAL.PDF"
    assert sftp_row.filename == "Line A/manual.pdf"
    assert exc.value.status_code == 415
    assert blocked.exists()


@pytest.mark.asyncio
async def test_public_deposit_upload_accepts_supported_file_and_rejects_without_ghost_rows(
    db_session,
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="QA upload",
        expires_at=None,
        max_file_size_mb=1,
        allowed_extensions=["pdf"],
    )

    accepted = await receive_file(
        db_session,
        link=link,
        upload=UploadFile(filename="Manual.pdf", file=BytesIO(b"%PDF synthetic manual")),
    )

    with pytest.raises(HTTPException) as unsupported:
        await receive_file(
            db_session,
            link=link,
            upload=UploadFile(filename="loader.exe", file=BytesIO(b"MZ")),
        )
    with pytest.raises(HTTPException) as too_large:
        await receive_file(
            db_session,
            link=link,
            upload=UploadFile(filename="huge.pdf", file=BytesIO(b"x" * (1024 * 1024 + 1))),
        )

    rows = db_session.query(DepositFile).filter(DepositFile.access_link_id == link.id).all()
    staged_paths = list((tmp_path / "store").rglob("*"))

    assert accepted.filename == "Manual.pdf"
    assert accepted.status == "received"
    assert unsupported.value.status_code == 415
    assert too_large.value.status_code == 413
    assert [row.filename for row in rows] == ["Manual.pdf"]
    assert all(path.name != "huge.pdf" for path in staged_paths)
    assert all(not path.name.endswith(".part") for path in staged_paths)


def test_sftp_server_rejects_read_and_delete_operations():
    class PermissionDenied(Exception):
        pass

    class FakeAsyncSSH:
        class SSHServer:
            pass

        class SFTPServer:
            def __init__(self, _chan):
                pass

        class SFTPAttrs:
            def __init__(self, **kwargs):
                self.__dict__.update(kwargs)

        class SFTPName:
            def __init__(self, *args, **kwargs):
                self.args = args
                self.kwargs = kwargs

        class SFTPError(Exception):
            pass

        class SFTPNoSuchFile(Exception):
            pass

        class SFTPFailure(Exception):
            pass

        SFTPPermissionDenied = PermissionDenied

    _ssh_server, sftp_server = _build_asyncssh_components(FakeAsyncSSH)
    chan = SimpleNamespace(get_extra_info=lambda key: "access-1" if key == "username" else None)
    server = sftp_server(chan)

    with pytest.raises(PermissionDenied):
        server.open(b"existing.pdf", 0x00000001, None)
    with pytest.raises(PermissionDenied):
        server.remove(b"existing.pdf")


def test_record_staged_file_from_path_moves_sftp_upload(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="SFTP upload",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "upload.part"
    content = b"drive commissioning report"
    staged.write_bytes(content)

    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="../../commissioning.zip",
        content_type="application/octet-stream",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    assert row.filename == "commissioning.zip"
    assert row.size_bytes == len(content)
    assert row.sha256 == hashlib.sha256(content).hexdigest()
    assert row.object_key.endswith("/commissioning.zip")
    assert not staged.exists()
    assert (tmp_path / "store" / row.object_key).read_bytes() == content


def test_record_staged_file_from_path_preserves_safe_relative_path(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="SFTP directory upload",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "upload.part"
    content = b"nested drive report"
    staged.write_bytes(content)

    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="Line A/../Motor #1/Photos/Tms1.jpg",
        content_type="application/octet-stream",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    assert row.filename == "Line A/Motor _1/Photos/Tms1.jpg"
    assert row.object_key.endswith("/Line A/Motor _1/Photos/Tms1.jpg")
    assert (tmp_path / "store" / row.object_key).read_bytes() == content


def test_build_deposit_archive_contains_files_and_manifest(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="QA unrestricted",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    content = b"industrial context"
    object_key = f"workspaces/{workspace.id}/secure-deposit/{link.access_id}/file-1/manuals/drive/demo.txt"
    source = tmp_path / object_key
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_bytes(content)
    row = DepositFile(
        workspace_id=workspace.id,
        access_link_id=link.id,
        filename="manuals/drive/demo.txt",
        content_type="text/plain",
        object_key=object_key,
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        status="received",
    )
    db_session.add(row)
    db_session.flush()

    archive_path, archive_name = build_deposit_archive(
        [row],
        links_by_id={link.id: link},
        workspace_slug=workspace.slug,
    )
    try:
        assert archive_name.endswith(".zip")
        with zipfile.ZipFile(archive_path) as archive:
            names = archive.namelist()
            assert "_manifest.json" in names
            assert any(name.endswith("received/manuals/drive/demo.txt") for name in names)
            manifest = json.loads(archive.read("_manifest.json"))
            assert manifest["file_count"] == 1
            assert manifest["files"][0]["sha256"] == row.sha256
    finally:
        archive_path.unlink(missing_ok=True)


def test_preview_deposit_file_returns_text_content(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="QA unrestricted",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "upload.part"
    staged.write_text("# Maintenance note\nTorque setting: 42 Nm", encoding="utf-8")
    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="manuals/drive/note.md",
        content_type="text/markdown",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    preview = preview_deposit_file(row)

    assert preview["kind"] == "text"
    assert "Torque setting" in preview["content"]
    assert preview["download_url"].endswith(f"/{row.id}/download")


@pytest.mark.parametrize("cell_range, start, truncated", [
    ("C3:AP42", 3, False),
    ("C3:AX50", 3, True),
    (None, 1, False),
])
def test_spreadsheet_preview_prioritizes_selection_with_bounded_size(tmp_path, cell_range, start, truncated):
    from openpyxl import Workbook

    book = Workbook()
    book.active["C3"] = "Selection starts"
    book.active["AP42"] = "Selection ends"
    book.active["AX50"] = "Outside bounded window"
    source = tmp_path / "wide.xlsx"
    book.save(source)
    book.close()
    preview = build_file_preview(source, filename=source.name,
                                 media_type="application/octet-stream",
                                 size_bytes=source.stat().st_size, download_url="/original",
                                 sheet_name="Sheet", cell_range=cell_range)
    assert preview["row_start"] == preview["column_start"] == start
    assert len(preview["rows"]) == 40
    assert len(preview["columns"]) == (40 if cell_range else 12)
    assert preview["selection_truncated"] is truncated
    if cell_range:
        assert preview["rows"][0][0] == "Selection starts"
        assert preview["rows"][-1][-1] == "Selection ends"


def test_build_file_preview_large_pdf_renders_inline(tmp_path):
    # Carde manuals routinely exceed the 25 MB image cap (the real one was
    # ~125 MB). The PDF branch only needs size + media type, so we assert the
    # large-PDF path without materialising 130 MB of bytes on disk.
    preview = build_file_preview(
        tmp_path / "70060-Carde ELM001Y.pdf",
        filename="70060-Carde ELM001Y.pdf",
        media_type="application/pdf",
        size_bytes=130 * 1024 * 1024,
        download_url="/api/v1/documents/doc-1/raw",
    )

    assert preview["kind"] == "pdf"
    assert preview["size_bytes"] == 130 * 1024 * 1024


def test_build_file_preview_html_returns_html_content(tmp_path):
    source = tmp_path / "report.html"
    source.write_text(
        "<html><body><h1>Commissioning</h1><p>Torque 42 Nm</p></body></html>",
        encoding="utf-8",
    )

    preview = build_file_preview(
        source,
        filename="report.html",
        media_type="text/html",
        size_bytes=source.stat().st_size,
        download_url="/api/v1/documents/doc-2/raw",
    )

    assert preview["kind"] == "html"
    assert "Commissioning" in preview["content"]
    assert preview["truncated"] is False


def test_preview_deposit_file_returns_spreadsheet_rows(db_session, monkeypatch, tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="QA unrestricted",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "upload.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "AKK"
    sheet.append(["Reference", "Value"])
    sheet.append(["PULP80", 154.8])
    workbook.save(staged)

    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="AKK-WET-PAT-PULP80.xlsx",
        content_type="application/octet-stream",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    preview = preview_deposit_file(row)

    assert preview["kind"] == "spreadsheet"
    assert preview["sheet_name"] == "AKK"
    assert preview["rows"][0] == ["Reference", "Value"]
    assert preview["rows"][1] == ["PULP80", "154.8"]


@pytest.mark.asyncio
async def test_promote_spreadsheet_queues_collection_ingest(db_session, monkeypatch, tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "object_store_backend", "local")
    monkeypatch.setattr(settings, "object_store_base_path", str(tmp_path / "objects"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Excel pilot",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "GEOTEX-SPL-Y25.05.22-PIL.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Def strips"
    sheet.append(["A", 80])
    sheet.append(["B", 85])
    workbook.save(staged)
    staged_bytes = staged.read_bytes()
    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="1-NON-WOVENS/FRANCE/GEOTEX/2025-05-PIL-tests diff_rents filets/GEOTEX-SPL-Y25.05.22-PIL.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    def fake_dispatch(db, job):
        job.celery_task_id = "task-excel"
        return "task-excel"

    monkeypatch.setattr("app.services.secure_deposit.dispatch_worker_job", fake_dispatch)

    promoted = await promote_file_to_collection(
        db_session,
        deposit_file=row,
        workspace=workspace,
        user=user,
        collection_slug="andritz-non-wovens-france-excel-pilot",
    )

    collection = (
        db_session.query(KnowledgeCollection)
        .filter(KnowledgeCollection.slug == "andritz-non-wovens-france-excel-pilot")
        .one()
    )
    job = db_session.query(WorkerJob).filter(WorkerJob.id == promoted.worker_job_id).one()

    assert promoted.status == "promoted"
    assert promoted.promoted_collection_slug == "andritz-non-wovens-france-excel-pilot"
    assert promoted.promotion_result["mode"] == "spreadsheet"
    assert promoted.promotion_result["spreadsheet"]["extension"] == "xlsx"
    assert promoted.promotion_result["celery_task_id"] == "task-excel"
    assert collection.status == "queued"
    assert collection.document_count == 1
    assert job.status == "queued"
    assert job.celery_task_id == "task-excel"
    stored_name = collection.document_names[0]
    assert stored_name.endswith("GEOTEX-SPL-Y25.05.22-PIL.xlsx")
    assert (
        tmp_path
        / "objects"
            / collection.artifact_prefix
            / "original"
            / stored_name
        ).read_bytes() == staged_bytes


@pytest.mark.asyncio
async def test_promote_legacy_xls_is_rejected_without_text_fallback(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Legacy Excel",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "legacy.xls"
    staged.write_bytes(b"legacy binary")
    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="1-NON-WOVENS/FRANCE/legacy.xls",
        content_type="application/vnd.ms-excel",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    with pytest.raises(HTTPException) as exc:
        await promote_file_to_collection(
            db_session,
            deposit_file=row,
            workspace=workspace,
            user=user,
            collection_slug="andritz-non-wovens-france-excel-pilot",
        )

    db_session.refresh(row)
    assert exc.value.status_code == 422
    assert row.status == "received"
    assert db_session.query(KnowledgeCollection).count() == 0


@pytest.mark.asyncio
async def test_promote_already_promoted_file_is_rejected_without_duplicate_collection_work(
    db_session,
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Already promoted guard",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "manual.pdf"
    staged.write_bytes(b"%PDF already promoted")
    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="1-NON-WOVENS/FRANCE/manual.pdf",
        content_type="application/pdf",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )
    row.status = "promoted"
    row.promoted_collection_slug = "andritz-existing"
    row.promotion_result = {"status": "done", "collection_slug": "andritz-existing"}
    db_session.flush()

    with pytest.raises(HTTPException) as exc:
        await promote_file_to_collection(
            db_session,
            deposit_file=row,
            workspace=workspace,
            user=user,
            collection_slug="andritz-duplicate-target",
        )

    db_session.refresh(row)
    assert exc.value.status_code == 409
    assert row.status == "promoted"
    assert row.promoted_collection_slug == "andritz-existing"
    assert row.promotion_result == {"status": "done", "collection_slug": "andritz-existing"}
    assert db_session.query(KnowledgeCollection).count() == 0
    assert db_session.query(WorkerJob).count() == 0


def test_preview_deposit_file_returns_docx_text(db_session, monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="QA unrestricted",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "commissioning.docx"
    document = docx.Document()
    document.add_heading("Commissioning report", level=1)
    document.add_paragraph("Inspect the dryer pressure before restart.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Asset"
    table.cell(0, 1).text = "Action"
    table.cell(1, 0).text = "Dryer"
    table.cell(1, 1).text = "Verify pressure"
    document.save(staged)

    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="reports/commissioning.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    preview = preview_deposit_file(row)

    assert preview["kind"] == "text"
    assert preview["source_kind"] == "docx"
    assert "Commissioning report" in preview["content"]
    assert "Dryer | Verify pressure" in preview["content"]


def _staged_zip(
    db_session,
    tmp_path,
    *,
    workspace: Workspace,
    link,
    entries: dict[str, bytes],
    filename: str = "Manual_BBA120.zip",
) -> DepositFile:
    staged = tmp_path / filename
    with zipfile.ZipFile(staged, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename=filename,
        content_type="application/zip",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )


@pytest.mark.asyncio
async def test_promote_zip_queues_collection_ingest(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "object_store_backend", "local")
    monkeypatch.setattr(settings, "object_store_base_path", str(tmp_path / "objects"))
    monkeypatch.setattr(settings, "secure_deposit_archive_promotion_max_files", 50)
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Manual upload",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    row = _staged_zip(
        db_session,
        tmp_path,
        workspace=workspace,
        link=link,
        entries={
            "Manual_BBA120/Declaration/Declaration.pdf": b"%PDF declaration",
            "Manual_BBA120/Operator manual/Chapter 01.pdf": b"%PDF chapter",
            "Manual_BBA120/image.png": b"ocr image",
        },
    )

    def fake_dispatch(db, job):
        job.celery_task_id = "task-bba120"
        return "task-bba120"

    monkeypatch.setattr("app.services.secure_deposit.dispatch_worker_job", fake_dispatch)

    promoted = await promote_file_to_collection(
        db_session,
        deposit_file=row,
        workspace=workspace,
        user=user,
        collection_slug="andritz-manuals-bba120-pilot",
    )

    collection = (
        db_session.query(KnowledgeCollection)
        .filter(KnowledgeCollection.slug == "andritz-manuals-bba120-pilot")
        .one()
    )
    job = db_session.query(WorkerJob).filter(WorkerJob.id == promoted.worker_job_id).one()

    assert promoted.status == "promoted"
    assert promoted.promoted_collection_slug == "andritz-manuals-bba120-pilot"
    assert promoted.promotion_result["archive"]["extracted_count"] == 3
    assert promoted.promotion_result["celery_task_id"] == "task-bba120"
    assert collection.status == "queued"
    assert collection.document_count == 3
    assert job.status == "queued"
    assert job.celery_task_id == "task-bba120"
    stored_names = sorted(collection.document_names)
    assert stored_names == [
        "Manual_BBA120__Declaration__Declaration.pdf",
        "Manual_BBA120__Operator manual__Chapter 01.pdf",
        "Manual_BBA120__image.png",
    ]
    assert (
        tmp_path
        / "objects"
        / collection.artifact_prefix
        / "original"
        / "Manual_BBA120__Declaration__Declaration.pdf"
    ).read_bytes() == b"%PDF declaration"
    manifest = json.loads(get_object_store().read_bytes(document_manifest_key(collection)).decode("utf-8"))
    chapter_metadata = manifest["Manual_BBA120__Operator manual__Chapter 01.pdf"]
    assert chapter_metadata["project_code"] == "BBA120"
    assert chapter_metadata["initial_buyer_code"] == "BBA"
    assert chapter_metadata["project_position"] == "120"
    assert chapter_metadata["project_reference_kind"] == "andritz_project"
    assert chapter_metadata["source_family"] == "operating_manual"
    assert chapter_metadata["archive_name"] == "Manual_BBA120.zip"
    assert chapter_metadata["inner_document_path"] == "Manual_BBA120/Operator manual/Chapter 01.pdf"


@pytest.mark.asyncio
async def test_promote_invalid_zip_leaves_deposit_received(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "object_store_backend", "local")
    monkeypatch.setattr(settings, "object_store_base_path", str(tmp_path / "objects"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Manual upload",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    staged = tmp_path / "Manual_BHX100_revE.zip"
    staged.write_bytes(b"not a zip")
    row = record_staged_file_from_path(
        db_session,
        link=link,
        source_path=staged,
        filename="Manual_BHX100_revE.zip",
        content_type="application/zip",
        actor=f"sftp:{link.access_id}",
        transport="sftp",
    )

    with pytest.raises(HTTPException) as exc:
        await promote_file_to_collection(
            db_session,
            deposit_file=row,
            workspace=workspace,
            user=user,
            collection_slug="andritz-manuals-bba120-pilot",
        )

    db_session.refresh(row)
    assert exc.value.status_code == 422
    assert row.status == "received"
    assert db_session.query(KnowledgeCollection).count() == 0


@pytest.mark.asyncio
async def test_promote_zip_rejects_path_traversal(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "object_store_backend", "local")
    monkeypatch.setattr(settings, "object_store_base_path", str(tmp_path / "objects"))
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Manual upload",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    row = _staged_zip(
        db_session,
        tmp_path,
        workspace=workspace,
        link=link,
        entries={"../evil.pdf": b"%PDF evil"},
    )

    with pytest.raises(HTTPException) as exc:
        await promote_file_to_collection(
            db_session,
            deposit_file=row,
            workspace=workspace,
            user=user,
            collection_slug="andritz-manuals-bba120-pilot",
        )

    db_session.refresh(row)
    assert exc.value.status_code == 422
    assert row.status == "received"


@pytest.mark.asyncio
async def test_promote_zip_rejects_supported_file_count_over_limit(db_session, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "secure_deposit_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "object_store_backend", "local")
    monkeypatch.setattr(settings, "object_store_base_path", str(tmp_path / "objects"))
    monkeypatch.setattr(settings, "secure_deposit_archive_promotion_max_files", 1)
    workspace, user = _workspace_user(db_session)
    link, _ = create_link(
        db_session,
        workspace=workspace,
        user=user,
        label="Manual upload",
        expires_at=None,
        max_file_size_mb=30 * 1024,
        allowed_extensions=[],
    )
    row = _staged_zip(
        db_session,
        tmp_path,
        workspace=workspace,
        link=link,
        entries={
            "Manual_BBA120/Chapter 01.pdf": b"%PDF one",
            "Manual_BBA120/Chapter 02.pdf": b"%PDF two",
        },
    )

    with pytest.raises(HTTPException) as exc:
        await promote_file_to_collection(
            db_session,
            deposit_file=row,
            workspace=workspace,
            user=user,
            collection_slug="andritz-manuals-bba120-pilot",
        )

    db_session.refresh(row)
    assert exc.value.status_code == 413
    assert row.status == "received"


@pytest.mark.parametrize(
    "source, expected_code, expected_position",
    [
        # Trailing single-letter suffix (the previously-unresolvable codes).
        ("R__ELM001Y", "ELM001Y", "001"),
        ("R__RJT003Y", "RJT003Y", "003"),
        ("R__MUR002Y", "MUR002Y", "002"),
        # Two-letter suffix.
        ("N__NBD100ZH", "NBD100ZH", "100"),
        # Codes embedded in a flattened document name resolve from the prefix.
        ("R__ELM001Y__70060-Carde ELM001Y.pdf", "ELM001Y", "001"),
        # Backward compatibility: no suffix => exact historical output.
        ("Manual_AKK200/Operator manual/Chapter 01.pdf", "AKK200", "200"),
        ("Manual_BHX100_revE.zip", "BHX100", "100"),
        # Lower-case suffix is normalised to upper-case.
        ("elm001y", "ELM001Y", "001"),
    ],
)
def test_extract_andritz_project_reference_supports_letter_suffix(
    source, expected_code, expected_position
):
    reference = derive_project_reference(source, scheme="andritz")
    assert reference["project_code"] == expected_code
    assert reference["project_position"] == expected_position
    assert reference["initial_buyer_code"] == expected_code[:3]
    assert reference["project_reference_kind"] == "andritz_project"


@pytest.mark.parametrize(
    "source",
    [
        "KD724",  # only two leading letters
        "L10080",  # only one leading letter
        "ABC1234XYZ",  # 3+ trailing letters must not be swallowed
        "manuals/drive/note.md",  # no Andritz code at all
    ],
)
def test_extract_andritz_project_reference_rejects_non_codes(source):
    assert derive_project_reference(source, scheme="andritz") == {}


def test_direct_needlepunch_document_metadata_keeps_structural_project_identity():
    source_path = (
        "Notices_Techniques_Needlepunch/60000-69999/"
        "61035 - Customer/Manuals/TTN17829J.pdf"
    )

    metadata = _archive_document_metadata(
        deposit_filename=source_path,
        archive_path=None,
        document_name="Needlepunch__TTN17829J.pdf",
        extension="pdf",
        source_deposit_file_id="deposit-61035-manual",
        scheme="andritz",
    )

    assert metadata["project_code"] == "61035"
    assert metadata["project_code_scheme"] == "needlepunch_numeric5"
    assert metadata["business_scope"] == "needlepunch"
    assert metadata["project_range"] == "60000-69999"
    assert metadata["project_folder"] == "61035 - Customer"
    assert metadata["source_deposit_path"] == source_path
    assert metadata["source_deposit_file_id"] == "deposit-61035-manual"
    assert "initial_buyer_code" not in metadata
    assert "project_position" not in metadata


def test_direct_needlepunch_document_metadata_rejects_deeper_numeric_part_reference():
    metadata = _archive_document_metadata(
        deposit_filename=(
            "Notices_Techniques_Needlepunch/60000-69999/"
            "Manuals/61035/TTN17829J.pdf"
        ),
        archive_path=None,
        document_name="TTN17829J.pdf",
        extension="pdf",
        source_deposit_file_id="deposit-invalid-structure",
        scheme="andritz",
    )

    assert "project_code" not in metadata


def test_archive_metadata_does_not_invent_andritz_codes_without_scheme():
    metadata = _archive_document_metadata(
        deposit_filename="Manual_BBA120.zip",
        archive_path="Manual_BBA120/Operator manual/Chapter 01.pdf",
        document_name="Manual_BBA120__Operator manual__Chapter 01.pdf",
        extension="pdf",
    )

    assert "project_code" not in metadata
    assert "machine" not in metadata
    assert "project_reference_kind" not in metadata


def test_safe_filename_strips_paths_and_unsafe_characters():
    assert safe_filename("../../secret report?.pdf") == "secret report_.pdf"
    assert safe_filename("..\\..\\motor#1.xlsx") == "motor_1.xlsx"


def test_safe_relative_path_preserves_folders_without_traversal():
    assert safe_relative_path("../../line A/motor#1/photo?.jpg") == "line A/motor_1/photo_.jpg"
    assert safe_relative_path("..\\manuals\\2026\\drive.zip") == "manuals/2026/drive.zip"


def test_sftp_upload_root_alias_maps_to_deposit_root():
    assert _is_root_path("/upload")
    assert _remote_dir_path("/upload") == ""
    assert _remote_relative_path("/upload/manuals/drive.zip") == "manuals/drive.zip"


def test_sftp_posix_root_aliases_map_to_deposit_root():
    """asyncssh composes ``stat(".")`` into ``/./`` and sends ``/.`` for the
    session root.  These must resolve to the deposit root, not to a literal
    ``upload`` directory (which made the canary fail with SFTPNoSuchFile)."""
    for alias in ("/.", "./", "/./", "/"):
        assert _is_root_path(alias), alias
        assert _remote_dir_path(alias) == "", alias


def test_sftp_dotdot_is_clamped_to_deposit_root():
    """A leading ``..`` must not traverse above the virtual deposit root."""
    assert _is_root_path("..")
    assert _is_root_path("/..")
    assert _remote_relative_path("upload/../manuals/drive.zip") == "manuals/drive.zip"
    assert _remote_relative_path("a/b/../c") == "a/c"


def test_sftp_longname_marks_virtual_directories_for_filezilla():
    attrs = SimpleNamespace(
        permissions=stat.S_IFDIR | 0o755,
        size=0,
        mtime=1_747_000_000,
    )

    longname = _sftp_longname("1-NON-WOVENS", attrs).decode("utf-8")

    assert longname.startswith("drwxr-xr-x ")
    assert longname.endswith(" 1-NON-WOVENS")
