#!/usr/bin/env python3
"""Append 4 SENTINEL-CI S3 Phase 2.1 surface slides to the presenter deck."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SRC = REPO_ROOT / "docs" / "sentinel-ci-demo-presenter-deck-2026-05-25.pptx"
DEFAULT_DST = DEFAULT_SRC

BG = RGBColor(0x0A, 0x11, 0x18)
ACCENT = RGBColor(0x0F, 0x6F, 0x3F)
TEXT = RGBColor(0xE8, 0xED, 0xF2)
DIM = RGBColor(0x9A, 0xA8, 0xB8)

SLIDES = [
    (
        "S3.V21.1 — Onglet Securite (rail principal)",
        "Rail : Cockpit · Carte · Securite · Reputation · Agenda · Presse · Arbitrages\n"
        "Vue Securite : posture dual-axis + theatre Sahel + rumeur inline\n"
        "Barre sticky Conseil Defense 15h00\n"
        "Plan B : rail gauche → Securite",
    ),
    (
        "S3.V21.2 — Security Monitor plein ecran",
        "Route : /hypervisor/mission-room/securite/monitor\n"
        "Carte Sahel · military-air + border-tension\n"
        "Alertes ADS-B · fil rumeur · signaux sociaux\n"
        "Badge CACHE BASELINE · advisory only\n"
        "Plan B : Securite → Ouvrir Security Monitor",
    ),
    (
        "S3.V21.3 — Reputation rail + Veille sociale",
        "Reputation promue dans le rail (drill 2+/1-)\n"
        "Page Veille sociale : /veille-sociale\n"
        "Filtres bucket · sentiment · tri · export CSV advisory\n"
        "Plan B : URL veille-sociale ou AYA pulsation sociale",
    ),
    (
        "S3.V21.4 — Documents security-briefs",
        "Collection sentinel-ci-security-briefs dans scope vigie\n"
        "Onglet Documents du shell Securite\n"
        "Briefs Sahel · Conseil Defense · dossier rumeur · ADS-B\n"
        "Plan B : Securite → onglet Documents",
    ),
]


def _style_title(shape) -> None:
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = shape.text if hasattr(shape, "text") else ""
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = ACCENT


def add_slide(prs: Presentation, title: str, body: str) -> None:
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    if slide.shapes.title:
        slide.shapes.title.text = title
        for p in slide.shapes.title.text_frame.paragraphs:
            for run in p.runs:
                run.font.color.rgb = ACCENT
                run.font.size = Pt(24)
                run.font.bold = True
    body_shape = slide.placeholders[1] if len(slide.placeholders) > 1 else None
    if body_shape is not None:
        tf = body_shape.text_frame
        tf.clear()
        for idx, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            p.level = 0
            for run in p.runs:
                run.font.size = Pt(16)
                run.font.color.rgb = TEXT if idx == 0 else DIM


def main() -> None:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_DST
    prs = Presentation(str(src))
    for title, body in SLIDES:
        add_slide(prs, title, body)
    prs.save(str(dst))
    print(f"Added {len(SLIDES)} v2.1 slides -> {dst} ({len(prs.slides)} total)")


if __name__ == "__main__":
    main()
