#!/usr/bin/env python3
"""Build the SENTINEL-CI Scénario 3 presenter trame card as a minimal Word
document.

Same format as ``build_trame_card_docx.py`` (sober green accent, presenter
friendly, no emoji, no dev jargon). Six steps S3.1 → S3.6 covering the dual
axis security posture scenario, each with verbatim phrase, AYA expected reply,
what appears on screen and a Plan B click path.

Usage:
    python3 scripts/build_trame_card_s3_docx.py [output.docx]

Default output: docs/sentinel-ci-demo-trame-s3-presenter-card-2026-05-25.docx
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = REPO_ROOT / "docs" / "sentinel-ci-demo-trame-s3-presenter-card-2026-05-25.docx"

FONT_FAMILY = "Calibri"
BODY_PT = 10.5
SMALL_PT = 9.5
TITLE_PT = 18
H2_PT = 13
H3_PT = 11.5

ACCENT = RGBColor(0x0F, 0x6F, 0x3F)
DIM = RGBColor(0x55, 0x60, 0x6E)
TEXT = RGBColor(0x12, 0x18, 0x22)


PROD_URL = "https://agentium.papai.ai/hypervisor/mission-room/cockpit?workspace=sentinel-ci"
PROFIL = "Profil vigie_executive · Bouton « Parler à AYA » (barre AYA en bas du cockpit)"

PREAMBLE = [
    "Le Scénario 3 s'enchaîne après la Décision Option B (S2.7). Cible : posture sécuritaire dual-axis avant le Conseil Défense restreint de 15h00.",
    "Lire chaque phrase mot pour mot, sans filler poli en début de tour.",
    "Si AYA ne réagit pas en 2 secondes, copier la phrase telle quelle dans le Quick Panel (Cmd / Ctrl + K).",
    "Si toujours rien, suivre le Plan B clic indiqué pour l'étape — même effet à l'écran.",
    "Tout est advisory only et démo-safe : ADS-B advisory, comptes citoyens pseudonymisés, OSINT publique.",
]

SCENARIO_3 = (
    "Scénario 3 — Posture sécuritaire dual-axis",
    [
        {
            "id": "S3.1",
            "title": "Posture sécuritaire du jour",
            "say": "AYA, montre-moi la posture sécuritaire du jour.",
            "reply": "« Posture dual-axis : intérieur vigilance, extérieur Sahel élevée. Prochain Conseil Défense restreint à 15h. »",
            "screen": "Cockpit, bloc « Posture sécuritaire » mis en avant avec deux cartes Intérieur / Extérieur et la pastille « 15h00 — Conseil Défense restreint ».",
            "plan_b": "Sidebar gauche → chip « Posture sécuritaire » dans la barre de statut VP (ouvre directement le bloc).",
        },
        {
            "id": "S3.2",
            "title": "Pulsation sociale Abidjan",
            "say": "AYA, montre la pulsation sociale à Abidjan.",
            "reply": "« 18 tweets sur la dernière heure : 5 officiels, 8 citoyens pseudonymisés, 5 signaux rumeur frontière. Snapshot demo-safe. »",
            "screen": "Vue Stratégie, carte avec couche social-geo activée (points colorés sentiment) et drawer de droite affichant la liste des tweets groupés (officiels / citoyens / rumeur).",
            "plan_b": "Carte Stratégie → bouton couche « Pulsation sociale » → drawer presse rouvert sur la pulsation.",
        },
        {
            "id": "S3.3",
            "title": "Trace de la rumeur frontière",
            "say": "AYA, d'où vient la rumeur frontière Nord ?",
            "reply": "« Origine : tweet citoyen pseudonyme à 11h42 secteur Bouna. Repris sur Telegram à 12h08, blog régional à 12h48. Démentis officiels FANCI et Préfecture Nord publiés à 13h45. »",
            "screen": "Carte zoomée sur le Nord, couche border-tension orange, drawer « Trace OSINT » avec la chaîne tweet → telegram → blog → démentis. Proposition AYA « Rédiger un communiqué ».",
            "plan_b": "Drawer Brouillons → « Dossier rumeur frontière Nord » (la trace OSINT s'ouvre identique).",
        },
        {
            "id": "S3.4",
            "title": "Mouvements de troupes Sahel",
            "say": "AYA, montre les mouvements de troupes au Sahel.",
            "reply": "« Snapshot ADS-B advisory : 10 traces sur l'axe Bamako, Ouagadougou, Niamey. 3 zones de surveillance OSINT, 2 bases CEDEAO en alerte standard. Lecture advisory only. »",
            "screen": "Carte avec couches military-air (triangles) et border-tension actives, drawer « Snapshot ADS-B » listant les traces et zones de surveillance.",
            "plan_b": "Carte Stratégie → bouton couche « Trafic militaire » → drawer presse rouvert sur le snapshot ADS-B.",
        },
        {
            "id": "S3.5",
            "title": "Drill réputation 2 positifs / 1 critique",
            "say": "AYA, montre le drill de réputation 2 positifs et 1 critique.",
            "reply": "« Score 72 sur 100, plus 4 points. Deux signaux positifs (Jeune Afrique, Fraternité Matin) compensent une critique de L'Inter sur le budget défense. »",
            "screen": "Vue Réputation, score 72 affiché, trois cartes drill : 2 vertes + 1 orange, lien direct vers l'article L'Inter.",
            "plan_b": "Sidebar → onglet Réputation → ancre « Drill 2+/1- » en bas de page.",
        },
        {
            "id": "S3.6",
            "title": "Communiqué sécurité (préparation)",
            "say": "AYA, prépare un communiqué de sécurité sur la rumeur frontière Nord.",
            "reply": "« Brouillon prêt : démenti officiel cité, posture FANCI rappelée, advisory only. Validation requise avant publication. »",
            "screen": "Drawer brouillon communiqué (subject, body), boutons Valider / Modifier / Annuler, mention « Validation advisory requise ».",
            "plan_b": "Drawer Brouillons → « Nouveau communiqué » → template « Démenti rumeur Nord ».",
        },
    ],
)

SECTIONS = [SCENARIO_3]

FOOTER_NOTES = [
    "Lancer ``predemo_reset.sh`` à H-30 pour réinitialiser ``last_focus``, ``awaiting``, ``current_meeting`` et réactiver le pack ``sentinel_ci_aya_security_v1``.",
    "Hard reload (Cmd / Ctrl + Shift + R) avant la démo pour purger tout état résiduel.",
    "Pas de wake-word côté client : bouton « Parler à AYA » ou Quick Panel.",
    "Tout signal présenté est demo-safe : ADS-B advisory only, handles citoyens pseudonymisés ``@citoyen_***`` / ``@rumeur_***``, score critique = article L'Inter déjà publié (S1).",
    "Garde-fou rédactionnel : si la voix échoue, présenter en mode Quick Panel texte (même résolveur, même effet à l'écran).",
]


def _set_run(run, *, size_pt=BODY_PT, bold=False, italic=False, color=TEXT):
    run.font.name = FONT_FAMILY
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(qn(f"w:{attr}"), FONT_FAMILY)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _tighten(paragraph, *, before=0, after=2, line=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _add_label_value(document, label, value, *, value_bold=False, value_italic=False,
                     value_color=TEXT):
    p = document.add_paragraph()
    _tighten(p, before=0, after=1, line=1.15)
    lab = p.add_run(f"{label} ")
    _set_run(lab, size_pt=BODY_PT, bold=True, color=DIM)
    val = p.add_run(value)
    _set_run(val, size_pt=BODY_PT, bold=value_bold, italic=value_italic, color=value_color)
    return p


def _add_horizontal_rule(document):
    p = document.add_paragraph()
    _tighten(p, before=2, after=4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFC6D1")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _set_margins(section, *, top=1.6, bottom=1.6, left=1.8, right=1.8):
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)


def _style_heading(document, style_name, size_pt, color=ACCENT, bold=True):
    style = document.styles[style_name]
    style.font.name = FONT_FAMILY
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def build_document(output_path: Path) -> Path:
    document = Document()

    base_style = document.styles["Normal"]
    base_style.font.name = FONT_FAMILY
    base_style.font.size = Pt(BODY_PT)
    base_style.font.color.rgb = TEXT

    _style_heading(document, "Title", TITLE_PT, color=ACCENT, bold=True)
    _style_heading(document, "Heading 1", TITLE_PT, color=ACCENT, bold=True)
    _style_heading(document, "Heading 2", H2_PT, color=ACCENT, bold=True)
    _style_heading(document, "Heading 3", H3_PT, color=TEXT, bold=True)

    for section in document.sections:
        _set_margins(section)

    title = document.add_heading("SENTINEL-CI — Carte de trame VP, Scénario 3 · 25 mai 2026", level=1)
    _tighten(title, before=0, after=4)

    sub = document.add_paragraph()
    _tighten(sub, before=0, after=2)
    r1 = sub.add_run("URL prod : ")
    _set_run(r1, size_pt=SMALL_PT, bold=True, color=DIM)
    r2 = sub.add_run(PROD_URL)
    _set_run(r2, size_pt=SMALL_PT, color=TEXT)

    sub2 = document.add_paragraph()
    _tighten(sub2, before=0, after=6)
    r3 = sub2.add_run(PROFIL)
    _set_run(r3, size_pt=SMALL_PT, italic=True, color=DIM)

    pre_h = document.add_paragraph()
    _tighten(pre_h, before=0, after=2)
    rh = pre_h.add_run("Mode d'emploi (5 lignes)")
    _set_run(rh, size_pt=BODY_PT, bold=True, color=ACCENT)

    for line in PREAMBLE:
        bullet = document.add_paragraph(style="List Bullet")
        _tighten(bullet, before=0, after=1)
        run = bullet.add_run(line)
        _set_run(run, size_pt=BODY_PT, color=TEXT)

    _add_horizontal_rule(document)

    for idx, (section_title, steps) in enumerate(SECTIONS):
        h2 = document.add_heading(section_title, level=2)
        _tighten(h2, before=6 if idx > 0 else 0, after=3)

        for step in steps:
            h3 = document.add_heading(f"{step['id']} — {step['title']}", level=3)
            _tighten(h3, before=4, after=2)

            say_p = document.add_paragraph()
            _tighten(say_p, before=0, after=1)
            lab = say_p.add_run("À dire : ")
            _set_run(lab, size_pt=BODY_PT, bold=True, color=DIM)
            quote = say_p.add_run(f"« {step['say']} »")
            _set_run(quote, size_pt=BODY_PT, bold=True, color=TEXT)

            _add_label_value(document, "AYA répond :", step["reply"], value_italic=True)
            _add_label_value(document, "À l'écran :", step["screen"])
            _add_label_value(document, "Plan B clic :", step["plan_b"], value_color=DIM)

    _add_horizontal_rule(document)

    foot_h = document.add_paragraph()
    _tighten(foot_h, before=4, after=2)
    rf = foot_h.add_run("Garde-fous présentateur")
    _set_run(rf, size_pt=BODY_PT, bold=True, color=ACCENT)

    for line in FOOTER_NOTES:
        bullet = document.add_paragraph(style="List Bullet")
        _tighten(bullet, before=0, after=1)
        run = bullet.add_run(line)
        _set_run(run, size_pt=SMALL_PT, color=TEXT)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def main(argv: list[str]) -> int:
    out = Path(argv[1]).resolve() if len(argv) > 1 else DEFAULT_OUTPUT
    written = build_document(out)
    print(f"Wrote {written}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
