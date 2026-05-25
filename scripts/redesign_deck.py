#!/usr/bin/env python3
"""Redesign SENTINEL-CI demo pptx with a dark cockpit visual identity.

Applies an Agentium / SENTINEL-CI design pass on top of a pandoc-generated
PowerPoint deck while preserving all content (text, screenshots, tables):

- Dark cockpit background (#0a1118) on every slide
- Title placeholders in Agentium green (#3ee68a), large bold
- Body text in light slate (#e2e8f0), captions dimmed
- Tables: dark zebra rows, accent-green header, badge tokens colorised
- Top accent bar (1.7 mm) and bottom footer with page number on every slide
- Title slide: oversize accent title, italic subtitle, "Confidentiel" footer
- Section header slides: centred accent title, no body, darker background
- PASS / WARN / FAIL / AYA only tokens recoloured where they stand alone

Usage:
    python3 scripts/redesign_deck.py <input.pptx> [output.pptx]
"""

from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ---------------------------------------------------------------------------
# Palette — Agentium / SENTINEL-CI cockpit dark theme
# ---------------------------------------------------------------------------
BG_DARK = RGBColor(0x0A, 0x11, 0x18)
BG_DARKER = RGBColor(0x06, 0x0B, 0x12)
BG_TABLE_ROW_A = RGBColor(0x0D, 0x15, 0x1E)
BG_TABLE_ROW_B = RGBColor(0x12, 0x1C, 0x28)
BG_TABLE_HEADER = RGBColor(0x12, 0x2A, 0x1F)

ACCENT_GREEN = RGBColor(0x3E, 0xE6, 0x8A)
ACCENT_ORANGE = RGBColor(0xF5, 0x9E, 0x0B)
ACCENT_RED = RGBColor(0xEF, 0x44, 0x44)
ACCENT_VIOLET = RGBColor(0xA7, 0x8B, 0xFA)

TEXT_PRIMARY = RGBColor(0xE2, 0xE8, 0xF0)
TEXT_SECONDARY = RGBColor(0x94, 0xA3, 0xB8)
TEXT_DIM = RGBColor(0x64, 0x74, 0x8B)

FONT_FAMILY = "Calibri"

FOOTER_LEFT = "SENTINEL-CI · Démo Vice Premier Ministre · 25 mai 2026"

BADGES = (
    ("[PASS]", ACCENT_GREEN, True),
    ("PASS", ACCENT_GREEN, True),
    ("[WARN]", ACCENT_ORANGE, True),
    ("WARN", ACCENT_ORANGE, True),
    ("[FAIL]", ACCENT_RED, True),
    ("FAIL", ACCENT_RED, True),
    ("[CRITICAL]", ACCENT_RED, True),
    ("[AYA only]", ACCENT_VIOLET, True),
    ("AYA only", ACCENT_VIOLET, True),
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_slide_background(slide, color: RGBColor) -> None:
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def style_run(run, *, color=None, bold=None, italic=None, size=None, font=None) -> None:
    if font is not None:
        run.font.name = font
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def recolor_badges_in_paragraph(paragraph) -> None:
    """Recolour PASS / WARN / FAIL / AYA only tokens when they are the run's text."""
    for run in paragraph.runs:
        token = run.text.strip()
        if not token:
            continue
        for marker, color, bold in BADGES:
            if token == marker or token.strip("·") == marker:
                style_run(run, color=color, bold=bold)
                break


def style_text_frame_body(tf, *, color=TEXT_PRIMARY, size: float = 14.0) -> None:
    tf.word_wrap = True
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            if run.font.size is None:
                run.font.size = Pt(size)
            run.font.color.rgb = color
        recolor_badges_in_paragraph(paragraph)


def style_title_shape(shape, *, color: RGBColor = ACCENT_GREEN, size: float = 24.0,
                      center: bool = False, anchor_middle: bool = False) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    if anchor_middle:
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in tf.paragraphs:
        if center:
            paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            run.font.bold = True
            run.font.size = Pt(size)
            run.font.color.rgb = color


def style_caption_shape(shape, *, color: RGBColor = TEXT_SECONDARY, size: float = 11.0,
                        italic: bool = True) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            run.font.size = Pt(size)
            run.font.italic = italic
            run.font.color.rgb = color


def _clear_table_style(table) -> None:
    """Drop the inherited tblStyle reference so cell-level fill wins."""
    tbl_pr = table._tbl.find(qn("a:tblPr"))
    if tbl_pr is None:
        return
    tbl_pr.set("firstRow", "0")
    tbl_pr.set("bandRow", "0")
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is not None:
        tbl_pr.remove(style_id)


def _cell_border(cell, color: RGBColor) -> None:
    """Set a thin border on all four sides of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnT", "a:lnB", "a:lnL", "a:lnR"):
        existing = tcPr.find(qn(tag))
        if existing is not None:
            tcPr.remove(existing)
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", "6350")  # 0.5pt
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        solid = etree.SubElement(ln, qn("a:solidFill"))
        srgb = etree.SubElement(solid, qn("a:srgbClr"))
        srgb.set("val", "{:02X}{:02X}{:02X}".format(color[0], color[1], color[2]))
        prst = etree.SubElement(ln, qn("a:prstDash"))
        prst.set("val", "solid")


def style_table(table) -> None:
    _clear_table_style(table)
    rows = list(table.rows)
    for r_idx, row in enumerate(rows):
        is_header = r_idx == 0
        for cell in row.cells:
            if is_header:
                bg = BG_TABLE_HEADER
                fg = ACCENT_GREEN
                bold = True
            else:
                bg = BG_TABLE_ROW_A if (r_idx % 2) else BG_TABLE_ROW_B
                fg = TEXT_PRIMARY
                bold = False
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.margin_left = Emu(60000)
            cell.margin_right = Emu(60000)
            cell.margin_top = Emu(30000)
            cell.margin_bottom = Emu(30000)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT_FAMILY
                    if run.font.size is None:
                        run.font.size = Pt(11)
                    run.font.color.rgb = fg
                    if bold:
                        run.font.bold = True
                recolor_badges_in_paragraph(paragraph)
            _cell_border(cell, RGBColor(0x1F, 0x29, 0x37))


def add_top_accent_bar(slide, slide_w: int, color: RGBColor = ACCENT_GREEN,
                       height_emu: int = 60000) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_w, Emu(height_emu))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def _add_textbox(slide, left, top, width, height, text, *, size=9.0,
                 color=TEXT_DIM, align=PP_ALIGN.LEFT, italic=False, bold=False) -> None:
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.margin_left = Emu(20000)
    tf.margin_right = Emu(20000)
    tf.word_wrap = True
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_FAMILY
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.italic = italic
    run.font.bold = bold


def add_footer(slide, slide_w: int, slide_h: int, page_num: int, total: int) -> None:
    tb_h = Emu(260000)
    bottom_margin = Emu(80000)
    top = slide_h - tb_h - bottom_margin
    left_margin = Emu(180000)
    left_w = Emu(int(slide_w * 0.55))
    right_w = Emu(int(slide_w * 0.35))
    right_left = slide_w - right_w - left_margin

    _add_textbox(slide, left_margin, top, left_w, tb_h, FOOTER_LEFT,
                 size=9, color=TEXT_DIM, align=PP_ALIGN.LEFT)
    _add_textbox(slide, right_left, top, right_w, tb_h,
                 f"{page_num} / {total}",
                 size=9, color=TEXT_SECONDARY, align=PP_ALIGN.RIGHT)


def is_section_header(slide) -> bool:
    return slide.slide_layout.name == "Section Header"


# ---------------------------------------------------------------------------
# Per-slide processors
# ---------------------------------------------------------------------------
def process_title_slide(slide, slide_w: int, slide_h: int) -> None:
    set_slide_background(slide, BG_DARKER)

    add_top_accent_bar(slide, slide_w, color=ACCENT_GREEN, height_emu=120000)

    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.is_placeholder:
            continue
        ph = shape.placeholder_format
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
        if ph.idx == 0:
            style_title_shape(shape, color=ACCENT_GREEN, size=44,
                              center=True, anchor_middle=True)
        else:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT_FAMILY
                    run.font.size = Pt(20)
                    run.font.italic = True
                    run.font.bold = False
                    run.font.color.rgb = TEXT_PRIMARY

    _add_textbox(
        slide,
        Emu(180000),
        slide_h - Emu(420000),
        slide_w - Emu(360000),
        Emu(260000),
        "Confidentiel · Démo interne · Datategy / Plateforme Agentium",
        size=10,
        color=TEXT_DIM,
        align=PP_ALIGN.CENTER,
        italic=True,
    )


def process_section_header(slide, idx: int, total: int, slide_w: int, slide_h: int) -> None:
    set_slide_background(slide, BG_DARKER)

    add_top_accent_bar(slide, slide_w, color=ACCENT_GREEN, height_emu=120000)

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        ph = shape.placeholder_format if shape.is_placeholder else None
        if ph is not None and ph.idx == 0:
            style_title_shape(shape, color=ACCENT_GREEN, size=36,
                              center=True, anchor_middle=True)
        else:
            style_text_frame_body(shape.text_frame, color=TEXT_SECONDARY, size=14)
            for paragraph in shape.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER

    _add_textbox(
        slide,
        Emu(180000),
        slide_h - Emu(420000),
        Emu(int(slide_w * 0.5)),
        Emu(260000),
        FOOTER_LEFT,
        size=9,
        color=TEXT_DIM,
        align=PP_ALIGN.LEFT,
    )
    right_w = Emu(int(slide_w * 0.35))
    _add_textbox(
        slide,
        slide_w - right_w - Emu(180000),
        slide_h - Emu(420000),
        right_w,
        Emu(260000),
        f"{idx + 1} / {total}",
        size=9,
        color=TEXT_SECONDARY,
        align=PP_ALIGN.RIGHT,
    )


def fix_pandoc_legend_overlap(slide, slide_w: int, slide_h: int) -> "Optional[object]":
    """Reposition the extra `Content Placeholder 2` (blockquote / *Légende*) when it
    sits underneath a Picture on a "Content with Caption" pandoc layout.

    Pandoc maps the markdown blockquote that follows an image to a hidden second
    content placeholder that spans the entire right column. With the dark theme
    its light text becomes visible and visually overlaps the screenshot. Move it
    into a tidy bottom strip beneath the picture + caption, dimmed.
    """
    picture = None
    caption_textbox = None
    extra_content = None
    for shape in slide.shapes:
        if shape.shape_type == 13:  # PICTURE
            picture = shape
        elif shape.has_text_frame and shape.is_placeholder:
            if shape.name == "TextBox 3":
                caption_textbox = shape
            elif shape.name == "Content Placeholder 2":
                extra_content = shape
    if picture is None or extra_content is None:
        return None

    pic_left = picture.left
    pic_width = picture.width
    new_pic_top = Emu(420000)
    new_pic_height = Emu(2350000)
    picture.top = new_pic_top
    picture.height = new_pic_height

    caption_top = new_pic_top + new_pic_height + Emu(60000)
    caption_height = Emu(280000)
    if caption_textbox is not None:
        caption_textbox.left = pic_left
        caption_textbox.top = caption_top
        caption_textbox.width = pic_width
        caption_textbox.height = caption_height

    extra_top = caption_top + caption_height + Emu(80000)
    extra_height = Emu(slide_h - extra_top - 480000)
    extra_content.left = pic_left
    extra_content.top = extra_top
    extra_content.width = pic_width
    extra_content.height = extra_height
    extra_content.text_frame.word_wrap = True
    for paragraph in extra_content.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            run.font.size = Pt(9)
            run.font.italic = True
            run.font.color.rgb = TEXT_DIM
    return extra_content


def process_content_slide(slide, idx: int, total: int, slide_w: int, slide_h: int) -> None:
    set_slide_background(slide, BG_DARK)
    add_top_accent_bar(slide, slide_w)

    legend_shape = fix_pandoc_legend_overlap(slide, slide_w, slide_h)

    for shape in list(slide.shapes):
        if shape is legend_shape:
            continue  # already styled by the overlap fix
        if shape.has_text_frame:
            ph = shape.placeholder_format if shape.is_placeholder else None
            is_title = ph is not None and ph.idx == 0
            if is_title:
                style_title_shape(shape, color=ACCENT_GREEN, size=22)
            elif shape.name.startswith("TextBox") and ph is not None and ph.idx not in (0, 1):
                style_caption_shape(shape, color=TEXT_SECONDARY, size=11, italic=True)
            else:
                style_text_frame_body(shape.text_frame, color=TEXT_PRIMARY, size=13)
        if shape.has_table:
            style_table(shape.table)

    add_footer(slide, slide_w, slide_h, page_num=idx + 1, total=total)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main(src: str, dst: str) -> None:
    prs = Presentation(src)
    total = len(prs.slides)
    slide_w = prs.slide_width
    slide_h = prs.slide_height

    for idx, slide in enumerate(prs.slides):
        if idx == 0:
            process_title_slide(slide, slide_w, slide_h)
        elif is_section_header(slide):
            process_section_header(slide, idx, total, slide_w, slide_h)
        else:
            process_content_slide(slide, idx, total, slide_w, slide_h)

    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    prs.save(dst)
    print(f"Saved {dst} ({total} slides redesigned)")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    src_path = sys.argv[1]
    dst_path = sys.argv[2] if len(sys.argv) > 2 else src_path
    main(src_path, dst_path)
